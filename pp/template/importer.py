"""Bounded-memory OOXML importer for synthetic workbook templates."""

from __future__ import annotations

import hashlib
import json
import logging
import math
import posixpath
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
import re
import shutil
import sqlite3
import tempfile
import zipfile
import xml.etree.ElementTree as ET
import defusedxml.ElementTree as SafeET

from openpyxl.formula.translate import Translator
from openpyxl.styles.numbers import BUILTIN_FORMATS
from openpyxl.utils.cell import get_column_letter, range_boundaries
from openpyxl.utils.escape import unescape as unescape_ooxml

from .colors import Palette, indexed_palette, theme_colors
from .drawings import (DRAWING_A, DRAWING_REL, DRAWING_XDR, VML_REL, anchor_box, anchors, drawing_objects,
                       related_parts, shape_look, sheet_drawing_parts)
from .dashboard import DashboardUnsupported, read_dashboard_metadata
from .store import publish, replace_snapshot, staging_directory


LOGGER = logging.getLogger(__name__)


NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
RNS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
REL_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}"
CELL_REF = re.compile(r"([A-Z]+)([1-9][0-9]*)")
IMPORT_VERSION = 17
NOTE_COUNT = 10_000
NOTE_TEXT_BYTES = 64 * 1024
TABLE_COUNT = 1_000
TABLE_PART_BYTES = 2 * 1024 * 1024
MAX_SHEET_ROWS = 1_048_576
MAX_SHEET_COLS = 16_384
CF_RULE_TYPES = frozenset({"expression", "cellIs", "containsText", "duplicateValues", "beginsWith", "containsErrors"})
CF_BORDER_STYLES = {
    "thin": 1, "hair": 2, "dotted": 3, "dashed": 4,
    "dashDot": 5, "dashDotDot": 6, "double": 7, "medium": 8,
    "mediumDashed": 9, "mediumDashDot": 10, "mediumDashDotDot": 11,
    "slantDashDot": 12, "thick": 13,
}


class TemplateImportError(Exception):
    """A source archive could not be safely imported."""


@dataclass(frozen=True)
class StyleDefinition:
    index: int
    node: ET.Element
    base: ET.Element | None


@dataclass(frozen=True)
class CellContext:
    database: sqlite3.Connection
    sheet: int
    unlocked: set[int]


@dataclass
class NoteBudget:
    count: int = 0


@dataclass(frozen=True)
class SheetContext:
    archive: zipfile.ZipFile
    database: sqlite3.Connection
    unlocked: set[int]
    dxf_styles: list[dict]
    destination: Path
    note_budget: NoteBudget
    table_styles: dict[str, dict]
    palette: Palette


@dataclass(frozen=True)
class SheetState:
    rows: dict
    metadata: dict
    cells: CellContext


@dataclass(frozen=True)
class FormulaSync:
    destination: Path
    database: sqlite3.Connection
    metadata: list[dict]


@dataclass(frozen=True)
class CapturedSource:
    path: Path
    digest: str
    name: str


def _source_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _capture_source(source: Path, parent: Path) -> CapturedSource:
    parent.mkdir(parents=True, exist_ok=True)
    handle = tempfile.NamedTemporaryFile(prefix=".captured-", dir=parent, delete=False)
    temporary = Path(handle.name)
    try:
        with source.open("rb") as reader, handle:
            shutil.copyfileobj(reader, handle, length=1024 * 1024)
        return CapturedSource(temporary, _source_hash(temporary), source.name)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise


def _write_json(path: Path, value: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, separators=(",", ":"), sort_keys=True), encoding="utf-8")


def _preserve_source(source: Path, destination: Path) -> None:
    target = destination / "source.xlsm"
    if target.exists():
        return
    destination.mkdir(parents=True, exist_ok=True)
    handle = tempfile.NamedTemporaryFile(prefix=".source-", dir=destination, delete=False)
    temporary = Path(handle.name)
    try:
        with source.open("rb") as reader, handle:
            shutil.copyfileobj(reader, handle, length=1024 * 1024)
        temporary.replace(target)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise


def _cell_ref(ref: str) -> tuple[int, int]:
    match = CELL_REF.fullmatch(ref.replace("$", ""))
    if match is None:
        raise TemplateImportError("invalid cell coordinate")
    col = 0
    for letter in match.group(1):
        col = col * 26 + ord(letter) - 64
    return int(match.group(2)) - 1, col - 1


def _bounds(ref: str) -> list[int]:
    start, _, end = ref.partition(":")
    row1, col1 = _cell_ref(start)
    row2, col2 = _cell_ref(end or start)
    return [min(row1, row2), min(col1, col2), max(row1, row2), max(col1, col2)]


def _font_name(font: ET.Element | None) -> str | None:
    if font is None:
        return None
    name = font.find(f"{NS}name")
    return name.get("val") if name is not None else None


def _component(items: list[ET.Element], index: int) -> ET.Element | None:
    return items[index] if 0 <= index < len(items) else None


def _component_id(xf: ET.Element, base: ET.Element | None, name: str) -> int:
    apply_name = "applyNumberFormat" if name == "numFmt" else f"apply{name[0].upper()}{name[1:]}"
    apply = xf.get(apply_name)
    source = base if apply in {"0", "false"} else xf
    return int(source.get(f"{name}Id", "0")) if source is not None else 0


def _locked(xf: ET.Element | None, inherited: bool = True) -> bool:
    if xf is None or xf.get("applyProtection") in {"0", "false"}:
        return inherited
    protection = xf.find(f"{NS}protection")
    if protection is None:
        return inherited
    return protection.get("locked", "1") not in {"0", "false"}


def _style_record(style: StyleDefinition, root: ET.Element, palette: Palette) -> dict:
    index, xf, base = style.index, style.node, style.base
    fonts = list(root.findall(f"{NS}fonts/{NS}font"))
    fills = list(root.findall(f"{NS}fills/{NS}fill"))
    borders = list(root.findall(f"{NS}borders/{NS}border"))
    font = _component(fonts, _component_id(xf, base, "font"))
    fill = _component(fills, _component_id(xf, base, "fill"))
    border = _component(borders, _component_id(xf, base, "border"))
    alignment = None if xf.get("applyAlignment") in {"0", "false"} else xf.find(f"{NS}alignment")
    if alignment is None and base is not None:
        alignment = base.find(f"{NS}alignment")
    num_id = _component_id(xf, base, "numFmt")
    custom = {int(item.get("numFmtId", "0")): item.get("formatCode", "General")
              for item in root.findall(f"{NS}numFmts/{NS}numFmt")}
    def side(name):
        node = border.find(f"{NS}{name}") if border is not None else None
        if node is None or node.get("style") in {None, "none"}:
            return None
        return {"style": node.get("style"), "color": palette.cell(node.find(f"{NS}color")) or "#000000"}
    pattern = fill.find(f"{NS}patternFill") if fill is not None else None
    # shortcut: hatch patterns (gray125, darkGrid...) show no fill; draw them if a template uses one
    solid = pattern is not None and pattern.get("patternType") == "solid"
    return {
        "id": index, "nf": custom.get(num_id, BUILTIN_FORMATS.get(num_id, "General")),
        "bold": font is not None and font.find(f"{NS}b") is not None,
        "italic": font is not None and font.find(f"{NS}i") is not None,
        "font_name": _font_name(font),
        "color": palette.cell(font.find(f"{NS}color") if font is not None else None),
        "fill": palette.cell(pattern.find(f"{NS}fgColor")) if solid else None,
        "align": alignment.get("horizontal") if alignment is not None else None,
        "valign": alignment.get("vertical") if alignment is not None else None,
        "wrap": alignment is not None and alignment.get("wrapText") in {"1", "true"},
        "font_size": float(font.find(f"{NS}sz").get("val")) if font is not None and font.find(f"{NS}sz") is not None else None,
        "locked": _locked(xf, _locked(base)),
        "border": {"t": side("top"), "r": side("right"), "b": side("bottom"), "l": side("left")},
    }


def _style_data(archive: zipfile.ZipFile) -> tuple[list[dict], set[int]]:
    if "xl/styles.xml" not in archive.namelist():
        return [_default_style()], set()
    root = SafeET.fromstring(archive.read("xl/styles.xml"))
    palette = Palette(theme_colors(archive), indexed_palette(root))
    bases = list(root.findall(f"{NS}cellStyleXfs/{NS}xf"))
    styles, unlocked = [], set()
    for index, xf in enumerate(root.findall(f"{NS}cellXfs/{NS}xf")):
        parent = int(xf.get("xfId", "0"))
        base = bases[parent] if parent < len(bases) else None
        styles.append(_style_record(StyleDefinition(index, xf, base), root, palette))
        if not _locked(xf, _locked(base)):
            unlocked.add(index)
    return styles or [_default_style()], unlocked


def _dxf_data(archive: zipfile.ZipFile) -> list[dict]:
    if "xl/styles.xml" not in archive.namelist():
        return []
    root = SafeET.fromstring(archive.read("xl/styles.xml"))
    palette = Palette(theme_colors(archive), indexed_palette(root))
    return [_dxf_style(node, palette) for node in root.findall(f"{NS}dxfs/{NS}dxf")]


def _dxf_style(node: ET.Element, palette: Palette) -> dict:
    style, font, fill = {}, node.find(f"{NS}font"), node.find(f"{NS}fill")
    if font is not None:
        bold, color = font.find(f"{NS}b"), palette.cell(font.find(f"{NS}color"))
        if bold is not None:
            style["bold"] = bold.get("val", "1").lower() not in {"0", "false"}
        if color is not None:
            style["color"] = color
    if fill is not None:
        pattern = fill.find(f"{NS}patternFill")
        if pattern is not None and pattern.get("patternType") == "none":
            style["fill"] = None
        elif pattern is not None:
            color = palette.cell(pattern.find(f"{NS}bgColor"))
            color = color or palette.cell(pattern.find(f"{NS}fgColor"))
            if color is not None:
                style["fill"] = color
    number = node.find(f"{NS}numFmt")
    if number is not None and number.get("formatCode") is not None:
        style["nf"] = number.get("formatCode")
    border = node.find(f"{NS}border")
    if border is not None:
        edges = {}
        for source, target in (("top", "t"), ("right", "r"), ("bottom", "b"), ("left", "l")):
            side = border.find(f"{NS}{source}")
            if side is None:
                continue
            border_style = CF_BORDER_STYLES.get(side.get("style"))
            if border_style is None:
                continue
            color = palette.cell(side.find(f"{NS}color"))
            edge = {"s": border_style}
            if color is not None:
                edge["cl"] = {"rgb": color}
            edges[target] = edge
        if edges:
            style["border"] = edges
    return style


def _table_dxf_style(node: ET.Element, palette: Palette) -> dict:
    style = _dxf_style(node, palette)
    border = node.find(f"{NS}border")
    if border is None:
        return style
    edges = dict(style.get("border", {}))
    for name in ("vertical", "horizontal"):
        side = border.find(f"{NS}{name}")
        border_style = None if side is None else CF_BORDER_STYLES.get(side.get("style"))
        if border_style is None:
            continue
        edge = {"s": border_style}
        if color := palette.cell(side.find(f"{NS}color")):
            edge["cl"] = {"rgb": color}
        edges[name] = edge
    if edges:
        style["border"] = edges
    return style


def _table_style_elements(node, dxfs, palette, built_in=False):
    output = {}
    for element in node.findall(f"{NS}tableStyleElement"):
        try:
            dxf_id = int(element.get("dxfId", "-1")) - (1 if built_in else 0)
            size = max(1, int(element.get("size", "1")))
        except ValueError:
            continue
        if not 0 <= dxf_id < len(dxfs):
            continue
        output[element.get("type", "")] = {
            "size": size,
            "style": _table_dxf_style(dxfs[dxf_id], palette),
        }
    return output


def _custom_table_styles(archive, palette):
    if "xl/styles.xml" not in archive.namelist():
        return {}
    root = SafeET.fromstring(archive.read("xl/styles.xml"))
    dxfs = list(root.findall(f"{NS}dxfs/{NS}dxf"))
    return {
        node.get("name", ""): _table_style_elements(node, dxfs, palette)
        for node in root.findall(f"{NS}tableStyles/{NS}tableStyle")
        if node.get("table", "1") not in {"0", "false"}
    }


@lru_cache(maxsize=1)
def _preset_table_style_nodes():
    path = Path(__file__).parents[1] / "fidelity" / "preset_table_styles.xml"
    root = SafeET.parse(path).getroot()
    return {node.tag: node for node in root}


def _preset_table_style(name, palette):
    node = _preset_table_style_nodes().get(name)
    if node is None:
        return {}
    normalized = SafeET.fromstring(ET.tostring(node, encoding="utf-8"))
    for item in normalized.iter():
        item.tag = f"{NS}{item.tag}"
    dxfs = list(normalized.findall(f"{NS}dxfs/{NS}dxf"))
    style = normalized.find(f"{NS}tableStyles/{NS}tableStyle")
    return {} if style is None else _table_style_elements(style, dxfs, palette, True)


def _default_style() -> dict:
    return {"id": 0, "nf": "General", "bold": False, "italic": False, "color": None,
            "fill": None, "align": None, "valign": None, "wrap": False, "font_size": None,
            "font_name": None, "locked": True,
            "border": {"t": None, "r": None, "b": None, "l": None}}


def _workbook_data(archive: zipfile.ZipFile) -> tuple[list[dict], list[dict]]:
    workbook = SafeET.fromstring(archive.read("xl/workbook.xml"))
    rels = SafeET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    targets = {item.get("Id"): item.get("Target") for item in rels}
    view = workbook.find(f"{NS}bookViews/{NS}workbookView")
    active_tab = int(view.get("activeTab", "0")) if view is not None and view.get("activeTab") else 0
    sheets = []
    for index, node in enumerate(workbook.findall(f"{NS}sheets/{NS}sheet")):
        target = targets[node.get(f"{RNS}id")].lstrip("/")
        state = node.get("state", "visible")
        sheets.append({"index": index, "name": node.get("name", ""),
                       "part": target if target.startswith("xl/") else f"xl/{target}",
                       "hidden": False if state == "visible" else state,
                       "active": index == active_tab})
    names = [{"name": node.get("name", ""), "formula": node.text or "",
              "local_sheet_id": int(node.get("localSheetId")) if node.get("localSheetId") else None,
              "hidden": node.get("hidden") in {"1", "true"}}
             for node in workbook.findall(f"{NS}definedNames/{NS}definedName")]
    return sheets, names


def _shared_strings(archive: zipfile.ZipFile, database: sqlite3.Connection) -> None:
    database.execute("CREATE TABLE strings (id INTEGER PRIMARY KEY, value TEXT NOT NULL)")
    if "xl/sharedStrings.xml" not in archive.namelist():
        return
    with archive.open("xl/sharedStrings.xml") as stream:
        stack, index = [], 0
        for event, element in SafeET.iterparse(stream, events=("start", "end")):
            if event == "start":
                stack.append(element)
                continue
            if element.tag == f"{NS}si":
                value = _visible_text(element)
                database.execute("INSERT INTO strings VALUES (?, ?)", (index, value))
                index += 1
                if len(stack) > 1:
                    stack[-2].remove(element)
            stack.pop()
    database.commit()


def _visible_text(element):
    if element is None:
        return ""
    parts = []
    for child in element:
        if child.tag == f"{NS}t":
            parts.append(child.text or "")
        elif child.tag == f"{NS}r":
            text = child.find(f"{NS}t")
            parts.append(text.text or "" if text is not None else "")
    return _decode_text("".join(unescape_ooxml(part) for part in parts))


def _decode_text(value):
    output, index = [], 0
    while index < len(value):
        high = ord(value[index])
        if 0xD800 <= high <= 0xDBFF and index + 1 < len(value):
            low = ord(value[index + 1])
            if 0xDC00 <= low <= 0xDFFF:
                output.append(chr(0x10000 + ((high - 0xD800) << 10) + low - 0xDC00))
                index += 2
                continue
        output.append(value[index])
        index += 1
    return "".join(output)


def _value(cell: ET.Element, database: sqlite3.Connection):
    kind, value = cell.get("t"), cell.findtext(f"{NS}v")
    if kind == "s" and value is not None:
        row = database.execute("SELECT value FROM strings WHERE id = ?", (int(value),)).fetchone()
        return row[0] if row else ""
    if kind == "b":
        return value == "1"
    if kind == "inlineStr":
        return _visible_text(cell.find(f"{NS}is"))
    if kind == "e":
        return {"error": value or "#VALUE!"}
    if value is None:
        return None
    try:
        return float(value) if any(marker in value for marker in ".eE") else int(value)
    except ValueError:
        return value


def _formula(cell: ET.Element, database: sqlite3.Connection, coordinate: tuple[int, int, int]) -> str | None:
    sheet, row, col = coordinate
    node = cell.find(f"{NS}f")
    if node is None:
        return None
    text, shared_id = node.text or "", node.get("si")
    if node.get("t") != "shared" or shared_id is None:
        return "=" + text.lstrip("=")
    if text:
        database.execute("INSERT OR REPLACE INTO formulas VALUES (?, ?, ?, ?, ?)",
                         (sheet, int(shared_id), row, col, "=" + text.lstrip("=")))
    return "@shared:" + shared_id


def _validation(node: ET.Element) -> dict:
    formulas = [node.findtext(f"{NS}formula1"), node.findtext(f"{NS}formula2")]
    ranges = [_bounds(ref) for ref in (node.get("sqref") or "").split()]
    result = {"type": node.get("type"), "ranges": ranges,
              "showErrorMessage": node.get("showErrorMessage") in {"1", "true"}}
    if node.get("errorStyle"):
        result["errorStyle"] = node.get("errorStyle")
    if node.get("operator"):
        result["operator"] = node.get("operator")
    if result["type"] == "list":
        if formulas[0] and formulas[0].startswith('"'):
            result["list"] = [_list_member(item) for item in formulas[0].strip('"').split(",")]
        elif formulas[0] is not None:
            result["formula1"] = formulas[0]
    elif formulas[0] is not None:
        result["min"] = _literal(formulas[0])
        if formulas[1] is not None:
            result["max"] = _literal(formulas[1])
    return result


def _list_member(item: str):
    # Excel enters a picked member as if typed: "1" becomes the number 1, "true" the boolean.
    if item.strip().upper() in {"TRUE", "FALSE"}:
        return item.strip().upper() == "TRUE"
    value = _literal(item)
    return value if not isinstance(value, float) or math.isfinite(value) else item


def _literal(value: str):
    try:
        return float(value) if any(marker in value for marker in ".eE") else int(value)
    except ValueError:
        return value


def _flush_block(directory: Path, block: int | None, rows: dict) -> None:
    if block is not None:
        _write_json(directory / f"{block}.json", rows)


def _relationships(archive: zipfile.ZipFile, part: str) -> dict[str, dict]:
    directory, name = posixpath.split(part)
    rels_part = posixpath.join(directory, "_rels", f"{name}.rels")
    if rels_part not in archive.namelist():
        return {}
    root = SafeET.fromstring(archive.read(rels_part))
    return {node.get("Id"): dict(node.attrib) for node in root.findall(f"{REL_NS}Relationship")
            if node.get("Id")}


def _internal_target(part: str, relationship: dict | None) -> str | None:
    if not relationship or relationship.get("TargetMode", "").lower() == "external":
        return None
    target = relationship.get("Target", "")
    if not target or "\\" in target or ":" in target:
        return None
    if target.startswith("/"):
        if not target.startswith("/xl/"):
            return None
        resolved = target.lstrip("/")
    else:
        resolved = posixpath.normpath(posixpath.join(posixpath.dirname(part), target))
    return resolved if resolved.startswith("xl/") and "../" not in resolved else None


def _table_item(context: SheetContext, sheet_part: str, relationship: dict):
    if not relationship.get("Type", "").endswith("/table"):
        return None
    part = _internal_target(sheet_part, relationship)
    if part is None or part not in context.archive.namelist():
        return None
    with context.archive.open(part) as stream:
        data = stream.read(TABLE_PART_BYTES + 1)
    if len(data) > TABLE_PART_BYTES:
        raise TemplateImportError("table part exceeds the import limit")
    root = SafeET.fromstring(data)
    bounds = _bounds(root.get("ref", ""))
    style_info = root.find(f"{NS}tableStyleInfo")
    style_name = "" if style_info is None else style_info.get("name", "")
    style = context.table_styles.get(style_name)
    if style is None:
        style = _preset_table_style(style_name, context.palette)
    header_row = int(root.get("headerRowCount", "1")) > 0
    totals_row = int(root.get("totalsRowCount", "0")) > 0
    columns = [item.get("name", "") for item in root.findall(f"{NS}tableColumns/{NS}tableColumn")]
    return {
        "id": root.get("displayName") or root.get("name") or posixpath.basename(part),
        "name": root.get("name", ""), "range": bounds, "columns": columns,
        "header_row": header_row, "totals_row": totals_row, "style_name": style_name,
        "show_first_column": _xml_bool(style_info, "showFirstColumn"),
        "show_last_column": _xml_bool(style_info, "showLastColumn"),
        "show_row_stripes": _xml_bool(style_info, "showRowStripes"),
        "show_column_stripes": _xml_bool(style_info, "showColumnStripes"),
        "filter": header_row and root.find(f"{NS}autoFilter") is not None,
        "style": style,
    }


def _xml_bool(node, name):
    return node is not None and node.get(name) in {"1", "true"}


def _table_items(context: SheetContext, sheet_part: str, relation_ids: list[str]):
    if len(relation_ids) > TABLE_COUNT:
        raise TemplateImportError("table count exceeds the import limit")
    relations = _relationships(context.archive, sheet_part)
    return [item for relation_id in relation_ids
            if (item := _table_item(context, sheet_part, relations.get(relation_id, {}))) is not None]


def _note_text(element: ET.Element) -> str:
    return _visible_text(element.find(f"{NS}text"))


def _record_note(context: SheetContext, sheet: int, element: ET.Element, authors: list[str]):
    if context.note_budget.count >= NOTE_COUNT:
        raise TemplateImportError("note count exceeds the import limit")
    row, col = _cell_ref(element.get("ref", ""))
    if row >= MAX_SHEET_ROWS or col >= MAX_SHEET_COLS:
        raise TemplateImportError("note coordinate is invalid")
    try:
        author_id = int(element.get("authorId", ""))
    except ValueError:
        raise TemplateImportError("note author is invalid") from None
    if not 0 <= author_id < len(authors):
        raise TemplateImportError("note author is invalid") from None
    author = authors[author_id]
    text = _note_text(element)
    if len(text.encode("utf-8")) > NOTE_TEXT_BYTES:
        raise TemplateImportError("note text exceeds the import limit")
    try:
        context.database.execute("INSERT INTO notes VALUES (?,?,?,?,?)", (sheet, row, col, author, text))
    except sqlite3.IntegrityError:
        raise TemplateImportError("duplicate note coordinate") from None
    context.note_budget.count += 1
    return row, col


def _stream_notes(context: SheetContext, part: str, sheet: int) -> tuple[int, int, int]:
    authors, count, max_row, max_col, stack = [], 0, -1, -1, []
    with context.archive.open(part) as stream:
        for event, element in SafeET.iterparse(stream, events=("start", "end")):
            if event == "start":
                stack.append(element)
                continue
            if element.tag == f"{NS}author":
                author = element.text or ""
                if len(authors) >= NOTE_COUNT or len(author.encode("utf-8")) > NOTE_TEXT_BYTES:
                    raise TemplateImportError("note authors exceed the import limit")
                authors.append(author)
            elif element.tag == f"{NS}comment":
                row, col = _record_note(context, sheet, element, authors)
                count, max_row, max_col = count + 1, max(max_row, row), max(max_col, col)
            if len(stack) > 1 and element.tag in {f"{NS}author", f"{NS}comment"}:
                stack[-2].remove(element)
            stack.pop()
    return count, max_row, max_col


def _write_note_blocks(destination: Path, database: sqlite3.Connection, sheet: int) -> None:
    directory, block, rows = destination / "notes" / str(sheet), None, {}
    query = database.execute(
        "SELECT row,col,author,text FROM notes WHERE sheet=? ORDER BY row,col", (sheet,)
    )
    for row, col, author, text in query:
        next_block = row // 256
        if block is not None and next_block != block:
            _flush_block(directory, block, rows)
            rows = {}
        block = next_block
        rows.setdefault(str(row % 256), {})[str(col)] = {"author": author, "text": text}
    _flush_block(directory, block, rows)


def _import_notes(context: SheetContext, sheet_part: str, metadata: dict) -> bool:
    total = 0
    for relationship in _relationships(context.archive, sheet_part).values():
        if not relationship.get("Type", "").endswith("/comments"):
            continue
        part = _internal_target(sheet_part, relationship)
        if part is None or part not in context.archive.namelist():
            continue
        count, max_row, max_col = _stream_notes(context, part, metadata["index"])
        total += count
        metadata["rows"] = max(metadata["rows"], max_row + 1)
        metadata["cols"] = max(metadata["cols"], max_col + 1)
    if total:
        _write_note_blocks(context.destination, context.database, metadata["index"])
    return total > 0


def _record_cell(element: ET.Element, context: CellContext) -> tuple[int, int, dict]:
    reference = element.get("r")
    if reference is None:
        raise TemplateImportError("cell coordinate is missing")
    row, col = _cell_ref(reference)
    style = int(element.get("s", "0"))
    formula = _formula(element, context.database, (context.sheet, row, col))
    record = {"v": _value(element, context.database), "s": style, "k": "f" if formula is not None else "c"}
    if formula is not None:
        record["formula"] = formula
    elif style in context.unlocked:
        context.database.execute("INSERT INTO candidates VALUES (?, ?, ?)", (context.sheet, row, col))
    return row, col, record


def _blank_sheet(sheet_info) -> dict:
    metadata = {**sheet_info, "rows": 0, "cols": 0, "freeze": {"rows": 0, "cols": 0},
                "col_widths": {}, "row_heights": {}, "hidden_rows": [], "hidden_cols": [],
                "col_styles": [], "row_styles": {}, "merges": [], "validations": [],
                "conditional_formats": [], "has_conditional_formats": False,
                "table_relationships": [], "sheet_filter": None,
                "tab_color": None, "show_gridlines": True, "zoom": 1.0,
                "scroll": {"row": 0, "col": 0}}
    metadata.pop("part")
    return metadata


def _sheet_pass(context: SheetContext, sheet_info) -> tuple[dict, bool]:
    index = sheet_info["index"]
    metadata = _blank_sheet(sheet_info)
    block_index, block_rows, protected = None, {}, False
    sheet_dir = context.destination / "sheets" / str(index)
    state = SheetState(block_rows, metadata, CellContext(context.database, index, context.unlocked))
    with context.archive.open(sheet_info["part"]) as stream:
        stack = []
        for event, element in SafeET.iterparse(stream, events=("start", "end")):
            if event == "start":
                stack.append(element)
                continue
            if element.tag == f"{NS}row":
                row_number = int(element.get("r", "1")) - 1
                next_block = row_number // 256
                if block_index is not None and next_block != block_index:
                    _flush_block(sheet_dir, block_index, block_rows)
                    block_rows = {}
                    state = SheetState(block_rows, metadata, state.cells)
                block_index = next_block
                _record_row(element, row_number, state)
            else:
                protected = _sheet_metadata(element, metadata, protected, context)
            if len(stack) > 1 and element.tag == f"{NS}row":
                stack[-2].remove(element)
            stack.pop()
    _flush_block(sheet_dir, block_index, block_rows)
    sheet_dir.mkdir(parents=True, exist_ok=True)
    metadata["protected"] = protected
    metadata["has_notes"] = _import_notes(context, sheet_info["part"], metadata)
    metadata["tables"] = _table_items(
        context, sheet_info["part"], metadata.pop("table_relationships"),
    )
    metadata["filters"] = _table_filters(metadata.pop("sheet_filter"), metadata["tables"])
    _extend_for_tables(metadata)
    return metadata, protected


def _record_row(element, row_number, state: SheetState) -> None:
    row_cells = {}
    for cell in element.findall(f"{NS}c"):
        row, col, record = _record_cell(cell, state.cells)
        row_cells[str(col)] = record
        state.metadata["rows"] = max(state.metadata["rows"], row + 1)
        state.metadata["cols"] = max(state.metadata["cols"], col + 1)
    if row_cells:
        state.rows[str(row_number % 256)] = row_cells
    if element.get("ht"):
        state.metadata["row_heights"][str(row_number)] = float(element.get("ht"))
    if element.get("hidden") in {"1", "true"}:
        state.metadata["hidden_rows"].append(row_number)
    if element.get("customFormat") in {"1", "true"} and element.get("s") is not None:
        state.metadata["row_styles"][str(row_number)] = int(element.get("s"))


def _sheet_metadata(element, metadata, protected, context):
    if element.tag == f"{NS}dimension" and element.get("ref"):
        _, _, row2, col2 = _bounds(element.get("ref"))
        metadata["rows"] = max(metadata["rows"], row2 + 1)
        metadata["cols"] = max(metadata["cols"], col2 + 1)
    elif element.tag == f"{NS}sheetProtection":
        protected = element.get("sheet") in {"1", "true"}
    elif element.tag == f"{NS}sheetFormatPr" and element.get("defaultRowHeight"):
        metadata["default_row_height"] = _positive_height(
            element.get("defaultRowHeight")
        )
    elif element.tag == f"{NS}tabColor":
        metadata["tab_color"] = context.palette.cell(element)
    elif element.tag == f"{NS}sheetView":
        metadata["show_gridlines"] = element.get("showGridLines", "1") not in {"0", "false"}
        metadata["zoom"] = _zoom(element.get("zoomScale"))
    elif element.tag == f"{NS}pane":
        _pane_metadata(element, metadata)
    elif element.tag == f"{NS}col":
        _column_metadata(element, metadata)
    elif element.tag == f"{NS}mergeCell" and element.get("ref"):
        metadata["merges"].append(_bounds(element.get("ref")))
    elif element.tag == f"{NS}dataValidation":
        metadata["validations"].append(_validation(element))
    elif element.tag == f"{NS}conditionalFormatting":
        metadata["conditional_formats"].extend(_conditional_formats(element, context.dxf_styles))
        metadata["has_conditional_formats"] = bool(metadata["conditional_formats"])
    elif element.tag == f"{NS}autoFilter" and element.get("ref"):
        metadata["sheet_filter"] = _bounds(element.get("ref"))
    elif element.tag == f"{NS}tablePart" and element.get(f"{RNS}id"):
        metadata["table_relationships"].append(element.get(f"{RNS}id"))
    return protected


def _positive_height(raw):
    try:
        value = float(raw)
    except (TypeError, ValueError) as error:
        raise TemplateImportError("default row height is invalid") from error
    if not math.isfinite(value) or value <= 0:
        raise TemplateImportError("default row height is invalid")
    return value


def _pane_metadata(element, metadata) -> None:
    origin = _cell_ref(element.get("topLeftCell")) if element.get("topLeftCell") else None
    if origin is not None:
        metadata["scroll"] = {"row": origin[0], "col": origin[1]}
    if element.get("state") != "frozen":
        return
    metadata["freeze"] = {"rows": _split(element.get("ySplit")),
                          "cols": _split(element.get("xSplit"))}
    if origin is not None:
        metadata["freeze"].update({"start_row": origin[0], "start_col": origin[1]})


def _split(value: str | None) -> int:
    try:
        return int(float(value or "0"))
    except ValueError:
        return 0


def _zoom(value: str | None) -> float:
    """OOXML stores zoomScale as a percentage; the client needs a ratio."""
    try:
        scale = float(value)
    except (TypeError, ValueError):
        return 1.0
    return round(min(max(scale / 100, 0.1), 4.0), 2)


def _table_filters(sheet_filter, tables):
    output = []
    if sheet_filter is not None:
        output.append({"id": "sheet-auto-filter", "source": "sheet",
                       "range": sheet_filter, "columns": []})
    for table in tables:
        if table.pop("filter"):
            output.append({"id": f"table:{table['id']}", "source": "table",
                           "range": table["range"], "columns": table["columns"]})
    return output


def _extend_for_tables(metadata):
    for item in [*metadata["tables"], *metadata["filters"]]:
        metadata["rows"] = max(metadata["rows"], item["range"][2] + 1)
        metadata["cols"] = max(metadata["cols"], item["range"][3] + 1)


def _conditional_formats(element, dxf_styles: list[dict]) -> list[dict]:
    ranges = [_cf_bounds(value) for value in (element.get("sqref") or "").split()]
    top_left = ranges[0][:2] if ranges else [0, 0]
    output = []
    for rule_index, rule in enumerate(element.findall(f"{NS}cfRule")):
        kind = rule.get("type")
        if kind not in CF_RULE_TYPES:
            continue
        try:
            dxf_id, priority = int(rule.get("dxfId", "-1")), int(rule.get("priority", str(2**31)))
        except ValueError:
            continue
        formulas = [item.text or "" for item in rule.findall(f"{NS}formula")]
        dxf = dxf_styles[dxf_id] if 0 <= dxf_id < len(dxf_styles) else {}
        for area_index, bounds in enumerate(ranges):
            output.append({"type": kind, "range": bounds, "top_left": top_left, "priority": priority,
                           "stop_if_true": rule.get("stopIfTrue") in {"1", "true"},
                           "operator": rule.get("operator"), "text": rule.get("text"),
                           "formulas": formulas, "dxf": dxf,
                           "rule_id": f"cf-{priority}-{rule_index}-{area_index}"})
    return output


def _cf_bounds(reference: str) -> list[int]:
    min_col, min_row, max_col, max_row = range_boundaries(reference)
    return [(min_row or 1) - 1, (min_col or 1) - 1,
            (max_row or 1_048_576) - 1, (max_col or 16_384) - 1]


def _column_metadata(element, metadata):
    first, last = int(element.get("min", "1")) - 1, int(element.get("max", "1")) - 1
    for col in range(first, last + 1):
        if element.get("width"):
            metadata["col_widths"][str(col)] = float(element.get("width"))
        if element.get("hidden") in {"1", "true"}:
            metadata["hidden_cols"].append(col)
    if element.get("style") is not None:
        metadata["col_styles"].append([first, last, int(element.get("style"))])
    metadata["cols"] = max(metadata["cols"], last + 1)


def _expand_shared_formulas(destination: Path, database: sqlite3.Connection, sheet: int) -> None:
    directory = destination / "sheets" / str(sheet)
    for path in sorted(directory.glob("*.json"), key=lambda item: int(item.stem)):
        rows = json.loads(path.read_text(encoding="utf-8"))
        changed = False
        for row_key, columns in rows.items():
            for col_key, record in columns.items():
                formula = record.get("formula", "")
                if not formula.startswith("@shared:"):
                    continue
                shared_id = int(formula.split(":", 1)[1])
                master = database.execute("SELECT row, col, formula FROM formulas WHERE sheet=? AND id=?",
                                          (sheet, shared_id)).fetchone()
                if master is None:
                    raise TemplateImportError("shared formula master is missing")
                origin = f"{get_column_letter(master[1] + 1)}{master[0] + 1}"
                target = f"{get_column_letter(int(col_key) + 1)}{int(path.stem) * 256 + int(row_key) + 1}"
                record["formula"] = Translator(master[2], origin=origin).translate_formula(target)
                changed = True
        if changed:
            _write_json(path, rows)


def _write_formula_block(sync: FormulaSync, key, cells) -> None:
    sheet, block = key
    path = sync.destination / "sheets" / str(sheet) / f"{block}.json"
    data = json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}
    for row, col, value, anchor in cells:
        row_data = data.setdefault(str(row % 256), {})
        record = row_data.setdefault(str(col), {"v": None, "s": 0, "k": "c"})
        record.update({"v": value, "k": "f"})
        if not anchor:
            record.pop("formula", None)
    _write_json(path, data)


def _synchronize_formula_results(source: Path, sync: FormulaSync) -> None:
    from ..engine.formualizer import EngineError, FormualizerAdapter

    try:
        engine = FormualizerAdapter.load(source)
    except EngineError as exc:
        raise TemplateImportError(f"template formula capability error: {exc}") from exc
    grouped = {}
    for sheet, row, col, value in engine.result_snapshot():
        anchor = (sheet, row, col) in engine.formula_anchors
        grouped.setdefault((sheet, row // 256), []).append((row, col, value, anchor))
        sync.database.execute("DELETE FROM candidates WHERE sheet=? AND row=? AND col=?",
                              (sheet, row, col))
        sync.metadata[sheet]["rows"] = max(sync.metadata[sheet]["rows"], row + 1)
        sync.metadata[sheet]["cols"] = max(sync.metadata[sheet]["cols"], col + 1)
    for key, cells in grouped.items():
        _write_formula_block(sync, key, cells)


def _apply_input_mask(destination: Path, database: sqlite3.Connection, protected: set[int]) -> None:
    for sheet in protected:
        rows = database.execute("SELECT row, col FROM candidates WHERE sheet=? ORDER BY row,col", (sheet,))
        by_block: dict[int, list[tuple[int, int]]] = {}
        for row, col in rows:
            by_block.setdefault(row // 256, []).append((row, col))
        for block, cells in by_block.items():
            path = destination / "sheets" / str(sheet) / f"{block}.json"
            data = json.loads(path.read_text(encoding="utf-8"))
            for row, col in cells:
                data[str(row % 256)][str(col)]["k"] = "i"
            _write_json(path, data)


def _write_input_mask(path: Path, database: sqlite3.Connection, protected: set[int]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as stream:
        stream.write("[")
        first = True
        for sheet in sorted(protected):
            for row, col in database.execute("SELECT row,col FROM candidates WHERE sheet=? ORDER BY row,col", (sheet,)):
                if not first:
                    stream.write(",")
                stream.write(json.dumps(f"{sheet}:{row}:{col}"))
                first = False
        stream.write("]")


VML = "{urn:schemas-microsoft-com:vml}"
EXCEL_VML = "{urn:schemas-microsoft-com:office:excel}"


def _macro_name(text: str | None) -> str:
    # "[0]!name" and "'Book.xlsm'!name" both name the same macro; "Sheet13.name" keeps its module.
    return (text or "").strip().rsplit("!", 1)[-1].strip()


def _drawing_buttons(root: ET.Element, part: str, archive: zipfile.ZipFile, palette: Palette) -> list[dict]:
    rels = related_parts(archive, part)
    buttons = []
    for anchor in anchors(root):
        for shape in anchor.iter():
            macro = _macro_name(shape.get("macro"))
            props = next(shape.iter(f"{DRAWING_XDR}cNvPr"), None)
            if not macro or (props is not None and props.get("hidden") in {"1", "true"}):
                continue
            label = "".join(text.text or "" for text in shape.iter(f"{DRAWING_A}t")).strip()
            button = {**anchor_box(anchor, shape), "label": label or (props.get("name", "") if props is not None else ""),
                      "macro": macro}
            if shape.tag == f"{DRAWING_XDR}sp":  # drawn with the shape's own fill, line and text
                try:
                    look = shape_look(shape, palette, rels)
                except (AttributeError, TypeError, ValueError):
                    look = None
                if look and (look["fill"] or look["line"]):  # a bare text shape keeps the plain button look
                    button["look"] = {key: value for key, value in look.items() if key not in {"kind", "link"}}
            buttons.append(button)
    return buttons


def _vml_buttons(root: ET.Element) -> list[dict]:
    buttons = []
    for shape in root.iter(f"{VML}shape"):
        client = shape.find(f"{EXCEL_VML}ClientData")
        if client is None or client.get("ObjectType") != "Button":
            continue
        macro = _macro_name(client.findtext(f"{EXCEL_VML}FmlaMacro"))
        style = "".join((shape.get("style") or "").split())
        sizes = dict(part.split(":", 1) for part in style.split(";") if ":" in part)
        zero = any(sizes.get(key, "1").rstrip("ptxin") in {"0", "0.0"} for key in ("width", "height"))
        if not macro or "visibility:hidden" in style or zero:  # zero-size leftovers are invisible in Excel
            continue
        try:
            left, left_px, top, top_px, right, right_px, bottom, bottom_px = (
                int(part) for part in client.findtext(f"{EXCEL_VML}Anchor", "").split(","))
        except ValueError:
            continue
        textbox = shape.find(f"{VML}textbox")
        label = "".join(textbox.itertext()).strip() if textbox is not None else ""
        buttons.append({"from": {"row": top, "col": left, "dx": left_px, "dy": top_px},
                        "to": {"row": bottom, "col": right, "dx": right_px, "dy": bottom_px},
                        "label": label, "macro": macro})
    return buttons


def _button_data(archive: zipfile.ZipFile, sheet_specs: list[dict]) -> list[dict]:
    palette = Palette(theme_colors(archive))
    found = []
    for spec in sheet_specs:
        for rel_type, part in sheet_drawing_parts(archive, spec["part"]):
            root = SafeET.fromstring(archive.read(part))
            extracted = _vml_buttons(root) if rel_type == VML_REL else _drawing_buttons(root, part, archive, palette)
            found.extend({"sheet": spec["index"], **button} for button in extracted)
    return [{"id": f"b{number}", **button} for number, button in enumerate(found)]


def _drawing_data(archive: zipfile.ZipFile, sheet_specs: list[dict]) -> list[dict]:
    palette = Palette(theme_colors(archive))
    found = []
    for spec in sheet_specs:
        for rel_type, part in sheet_drawing_parts(archive, spec["part"]):
            if rel_type == DRAWING_REL:
                root = SafeET.fromstring(archive.read(part))
                found.extend({"sheet": spec["index"], **item} for item in drawing_objects(root, part, archive, palette))
    return [{"id": f"d{number}", **item} for number, item in enumerate(found)]


def _extend_for_objects(metadata: list[dict], objects: list[dict]) -> None:
    """Grow each sheet's grid to reach its pictures and buttons, as Excel's endless grid does."""
    sheets = {info["index"]: info for info in metadata}
    for item in objects:
        info = sheets.get(item.get("sheet"))
        for corner in (item.get("from"), item.get("to")):
            if info is not None and corner:
                info["rows"] = max(info["rows"], corner["row"] + 1)
                info["cols"] = max(info["cols"], corner["col"] + 1)


def _has_data_connections(archive: zipfile.ZipFile) -> bool:
    """True when the workbook has at least one Power Query or other data connection Excel can refresh."""
    if "xl/connections.xml" not in archive.namelist():
        return False
    return SafeET.fromstring(archive.read("xl/connections.xml")).find(f"{NS}connection") is not None


def _dashboard_or_saved_values(archive: zipfile.ZipFile, sheet_specs: list[dict]) -> dict:
    """Pivot and slicer metadata, or none when the dashboard cannot be rebuilt from workbook cells.

    A pivot fed by an external connection (the real template's Power Query feed) keeps Excel's saved
    pivot values in the grid instead of failing the whole template import.
    """
    try:
        return read_dashboard_metadata(archive, sheet_specs)
    except DashboardUnsupported as exc:
        logging.getLogger(__name__).warning("dashboard left as saved values: %s", exc)
        return {spec["index"]: {"pivots": [], "slicers": []} for spec in sheet_specs}


def _import_archive(source: CapturedSource, staging: Path, results: str = "engine", only=None) -> None:
    database = sqlite3.connect(staging / "import.sqlite")
    database.execute("CREATE TABLE candidates (sheet INTEGER, row INTEGER, col INTEGER, PRIMARY KEY(sheet,row,col))")
    database.execute("CREATE TABLE formulas (sheet INTEGER, id INTEGER, row INTEGER, col INTEGER, formula TEXT, PRIMARY KEY(sheet,id))")
    database.execute("CREATE TABLE notes (sheet INTEGER, row INTEGER, col INTEGER, author TEXT, text TEXT, PRIMARY KEY(sheet,row,col))")
    try:
        with zipfile.ZipFile(source.path) as archive:
            _shared_strings(archive, database)
            styles, unlocked = _style_data(archive)
            dxf_styles = _dxf_data(archive)
            palette = Palette(theme_colors(archive), indexed_palette(
                SafeET.fromstring(archive.read("xl/styles.xml")) if "xl/styles.xml" in archive.namelist() else None))
            sheet_specs, names = _workbook_data(archive)
            dashboard = _dashboard_or_saved_values(archive, sheet_specs)
            metadata, protected = [], set()
            context = SheetContext(
                archive, database, unlocked, dxf_styles, staging, NoteBudget(),
                _custom_table_styles(archive, palette), palette,
            )
            for spec in sheet_specs:
                if only is not None and spec["name"].casefold() not in only:
                    info, is_protected = _blank_sheet(spec), False  # left empty: the caller never shows it
                else:
                    info, is_protected = _sheet_pass(context, spec)
                info.update(dashboard[spec["index"]])
                metadata.append(info)
                if is_protected:
                    protected.add(spec["index"])
                _write_json(staging / "progress.json", {"complete": False,
                            "sheets_completed": len(metadata), "sheets_total": len(sheet_specs)})
                LOGGER.info("Template import progress: sheet %d/%d", len(metadata), len(sheet_specs))
            database.commit()
            buttons = _button_data(archive, sheet_specs)
            drawings = _drawing_data(archive, sheet_specs)
            _extend_for_objects(metadata, [*buttons, *drawings])
            refreshable = _has_data_connections(archive)
        for info in metadata:
            _expand_shared_formulas(staging, database, info["index"])
        if results == "engine":  # "cached" and "saved" keep Excel's saved results
            _synchronize_formula_results(source.path, FormulaSync(staging, database, metadata))
        _apply_input_mask(staging, database, protected)
        _write_input_mask(staging / "input_mask.json", database, protected)
        _write_json(staging / "buttons.json", buttons)
        _write_json(staging / "drawings.json", drawings)
        _write_json(staging / "styles.json", styles)
        _write_json(staging / "meta.json", {"hash": source.digest, "name": source.name,
                    "import_version": IMPORT_VERSION, "results": results, "sheets": metadata,
                    "defined_names": names, "refreshable": refreshable})
        _write_json(staging / "progress.json", {"complete": True,
                    "sheets_completed": len(metadata), "sheets_total": len(metadata)})
    finally:
        database.close()
        (staging / "import.sqlite").unlink(missing_ok=True)


def _snapshot_current(destination: Path, results: str = "engine") -> bool:
    try:
        meta = json.loads((destination / "meta.json").read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return False
    return (meta.get("hash") == destination.name
            and meta.get("import_version") == IMPORT_VERSION
            and meta.get("results", "engine") == results
            and (destination / "source.xlsm").is_file())


def _build_snapshot(source: CapturedSource, destination: Path, results: str = "engine", only=None) -> Path:
    staging = staging_directory(destination.parent)
    try:
        _import_archive(source, staging, results, only)
        source.path.replace(staging / "source.xlsm")
        if destination.exists():
            return replace_snapshot(staging, destination)
        return publish(staging, destination)
    except Exception as exc:
        shutil.rmtree(staging, ignore_errors=True)
        if isinstance(exc, TemplateImportError):
            raise
        if isinstance(exc, (OSError, ValueError, IndexError, sqlite3.Error,
                            ET.ParseError, zipfile.BadZipFile, KeyError)):
            raise TemplateImportError("template import failed") from exc
        raise


def import_template(source: str | Path, data_dir: str | Path, results: str = "engine", sheets=None) -> Path:
    """Import a workbook into a snapshot. ``sheets`` (names) limits which tabs are read; the others stay empty."""
    source, data_dir = Path(source), Path(data_dir)
    captured = _capture_source(source, data_dir / "templates")
    destination = data_dir / "templates" / captured.digest
    try:
        only = None if sheets is None else {name.casefold() for name in sheets}
        if only is None and _snapshot_current(destination, results):
            return destination
        return _build_snapshot(captured, destination, results, only)
    finally:
        captured.path.unlink(missing_ok=True)
