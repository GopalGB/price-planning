"""Calculate copies in Microsoft Excel itself (office Windows server).

Excel recalculates natively and exactly. The engine writes a batch of input
values, asks Excel to calculate, then reads back only the formulas the static
dependency index says the batch can reach, in bulk row blocks, and diffs them
against the last value it knew (Excel's saved value until something changes).

The host is the only part that talks to Excel: ``ComExcelHost`` drives a
hidden, private Excel instance through pywin32; ``FormualizerHost`` stands in
for it in tests on machines without Excel.
"""

from __future__ import annotations

import hashlib
import logging
import os
import pickle
import shutil
import tempfile
import threading
import time
from dataclasses import dataclass, field
from pathlib import Path

from ops.typed import typed

from ..values import exactly_equal
from .depindex import DependencyIndex
from .formualizer import EngineError, MacroTimeout, _cached_formula_map, _json_value
from .native_load import prepared_template

BLOCK_ROWS = 256
LOG = logging.getLogger(__name__)
# COM reports a cell error as the HRESULT 0x800A0000 + Excel's CVErr number.
COM_ERROR_BASE = -2146828288
XL_CELL_TYPE_FORMULAS = -4123
XL_FORMULAS, XL_PART, XL_BY_COLUMNS, XL_PREVIOUS = -4123, 2, 2, 2  # Range.Find arguments
XL_MANUAL = -4135  # xlCalculationManual
NO_CELLS = -2146827284  # SpecialCells found no cells (0x800A03EC)
DIALOG_CLASS = "#32770"
REFRESH_ALL = ":refresh"  # not a valid VBA name, so no template button can collide with it
EXCEL_ERRORS = {2000: "#NULL!", 2007: "#DIV/0!", 2015: "#VALUE!", 2023: "#REF!", 2029: "#NAME?",
                2036: "#NUM!", 2042: "#N/A", 2043: "#GETTING_DATA", 2045: "#SPILL!",
                2047: "#BLOCKED!", 2048: "#UNKNOWN!", 2049: "#FIELD!", 2050: "#CALC!"}


def from_com(value):
    """Convert a COM Value2 into the app's JSON cell value."""
    if value is None or isinstance(value, (bool, str)):
        return value
    if isinstance(value, dict) and set(value) == {"error"}:
        return value
    if isinstance(value, int) and not isinstance(value, bool):
        code = value - COM_ERROR_BASE
        if code in EXCEL_ERRORS:
            return {"error": EXCEL_ERRORS[code]}
    if isinstance(value, float) and value.is_integer() and abs(value) < 2**53:
        return int(value)
    kind, normalized = typed(value)
    return {"error": normalized} if kind == "error" else normalized


def to_com(value):
    """Convert an input value for Value2; text is stored literally, never re-parsed."""
    if isinstance(value, dict) and set(value) == {"error"}:
        raise EngineError("errors cannot be typed into input cells")
    if isinstance(value, str) and value:
        return "'" + value
    return value


def _formula_text(value):
    """The formula of a ``{"formula": "=..."}`` value, or None for a plain value."""
    if isinstance(value, dict) and set(value) == {"formula"}:
        text = value["formula"]
        if not isinstance(text, str) or not text.startswith("="):
            raise EngineError("formulas must be text starting with =")
        return text
    return None


class ExcelHost:
    """What the engine needs from Excel. Coordinates are zero-based."""

    def write(self, sheet: int, cells: list[tuple[int, int, object]]) -> None: ...
    def write_formulas(self, sheet: int, cells: list[tuple[int, int, str]]) -> None: ...
    def calculate(self) -> None: ...
    def read(self, sheet: int, top: int, left: int, bottom: int, right: int) -> list[list]: ...
    def read_formulas(self, sheet: int, top: int, left: int, bottom: int, right: int) -> list[list]: ...
    def formula_areas(self, sheet: int) -> list[tuple[int, int, int, int]]: ...
    def used_range(self, sheet: int) -> tuple[int, int, int, int] | None: ...
    def run_macro(self, name: str, timeout: float) -> list[str]: ...
    def refresh_all(self, timeout: float) -> list[str]: ...
    def pivot_ranges(self) -> list[tuple[int, int, int, int, int]]: ...
    def refresh_pivots(self, timeout: float) -> list[str]: ...
    def set_slicer(self, cache: str, selected: list[str]) -> None: ...
    def memory_bytes(self) -> int: ...
    def close(self) -> None: ...

    def invalidate(self, cells) -> None:
        """Cells about to be read after an edit; Excel has already recalculated them."""


EXCEL_CLSID = "{00024500-0000-0000-C000-000000000046}"
CO_E_CLASSSTRING = -2147221005  # "Invalid class string": the Excel.Application name is not registered


def _new_excel(client):
    """By name first. A damaged Office install can lose the name but keep Excel's fixed class id."""
    try:
        return client.DispatchEx("Excel.Application")
    except Exception as exc:  # pywintypes.com_error carries the HRESULT first
        if not exc.args or exc.args[0] != CO_E_CLASSSTRING:
            raise
        return client.DispatchEx(EXCEL_CLSID)



def convert_to_xlsm(source: Path, target: Path) -> None:
    """Excel's own Save As .xlsm, for .xlsb and .xls uploads. Macros, events and links stay off, and an
    empty password makes a protected file fail at once instead of waiting for someone to type one."""
    import pythoncom
    import win32com.client
    import win32process

    pythoncom.CoInitialize()
    app, pid = None, None
    try:
        app = _new_excel(win32com.client)
        pid = win32process.GetWindowThreadProcessId(app.Hwnd)[1]
        for name, value in (("Visible", False), ("DisplayAlerts", False), ("EnableEvents", False),
                            ("AskToUpdateLinks", False),
                            ("AutomationSecurity", 3)):  # 3 = msoAutomationSecurityForceDisable
            setattr(app, name, value)
        book = app.Workbooks.Open(str(source), UpdateLinks=0, ReadOnly=True, Password="",
                                  AddToMru=False, Notify=False)
        try:
            book.SaveAs(str(target), FileFormat=52)  # 52 = xlOpenXMLWorkbookMacroEnabled
        finally:
            book.Close(SaveChanges=False)
    finally:
        try:
            if app is not None:
                app.Quit()
        except Exception:
            pass
        _kill(pid)
        pythoncom.CoUninitialize()

class ComExcelHost(ExcelHost):
    """A private hidden Excel with macros and events off, on its own file copy."""

    def __init__(self, workbook: Path, sheet_names: list[str]):
        import pythoncom
        import win32com.client
        import win32process

        pythoncom.CoInitialize()
        self.directory = tempfile.mkdtemp(prefix="pp-excel-")
        self.path = Path(self.directory) / workbook.name
        shutil.copy2(workbook, self.path)
        self.app = _new_excel(win32com.client)
        self.pid = win32process.GetWindowThreadProcessId(self.app.Hwnd)[1]
        try:
            self._quiet()
            # Calculation can only be set with a book open, and Excel recalculates a file saved by another
            # Excel build while opening it; a blank book first makes the template open already manual.
            blank = self.app.Workbooks.Add()
            self.app.Calculation = XL_MANUAL
            self.app.CalculateBeforeSave = False
            started = time.monotonic()
            self.book = self.app.Workbooks.Open(str(self.path), UpdateLinks=0, ReadOnly=False,
                                                IgnoreReadOnlyRecommended=True, AddToMru=False,
                                                Notify=False)
            blank.Close(SaveChanges=False)
            self.app.Calculation = XL_MANUAL
            LOG.info("Excel opened %s in %.1fs", workbook.name, time.monotonic() - started)
            self.sheets = [self.book.Worksheets(name) for name in sheet_names]
        except Exception:
            self.close()
            raise

    def _quiet(self):
        """Hidden, silent, manual Excel; template macros switch some of these back on, so it runs again after."""
        for name, value in (("Visible", False), ("DisplayAlerts", False), ("ScreenUpdating", False),
                            ("EnableEvents", False), ("AskToUpdateLinks", False),
                            ("AutomationSecurity", 1)):  # 1 = msoAutomationSecurityLow: macros allowed
            setattr(self.app, name, value)
        if getattr(self, "book", None) is not None:
            self.app.Calculation = XL_MANUAL

    def write(self, sheet, cells):
        self._put(sheet, [(row, col, to_com(value)) for row, col, value in cells], "Value2")

    def write_formulas(self, sheet, cells):
        self._put(sheet, cells, "Formula2")

    def _put(self, sheet, cells, kind):
        worksheet = self.sheets[sheet]
        for top, left, rows in _rectangles(cells):
            try:
                setattr(self._area(worksheet, top, left, rows), kind, tuple(tuple(row) for row in rows))
            except Exception:
                if len(rows) == 1 and len(rows[0]) == 1:
                    raise
                for r, line in enumerate(rows):  # one cell at a time, as before blocks were written whole
                    for c, item in enumerate(line):
                        setattr(worksheet.Cells(top + r + 1, left + c + 1), kind, item)

    @staticmethod
    def _area(worksheet, top, left, rows):
        return worksheet.Range(worksheet.Cells(top + 1, left + 1),
                               worksheet.Cells(top + len(rows), left + len(rows[0])))

    def calculate(self):
        self.app.Calculate()

    def read(self, sheet, top, left, bottom, right):
        worksheet = self.sheets[sheet]
        values = worksheet.Range(worksheet.Cells(top + 1, left + 1),
                                 worksheet.Cells(bottom + 1, right + 1)).Value2
        if not isinstance(values, tuple):
            return [[values]]
        return [list(row) for row in values]

    def read_formulas(self, sheet, top, left, bottom, right):
        worksheet = self.sheets[sheet]
        texts = worksheet.Range(worksheet.Cells(top + 1, left + 1),
                                worksheet.Cells(bottom + 1, right + 1)).Formula2
        if not isinstance(texts, tuple):
            texts = ((texts,),)
        return [[text if isinstance(text, str) and text.startswith("=") else None for text in row]
                for row in texts]

    def formula_areas(self, sheet):
        try:
            cells = self.sheets[sheet].UsedRange.SpecialCells(XL_CELL_TYPE_FORMULAS)
        except Exception as exc:
            if getattr(exc, "hresult", None) == NO_CELLS:
                return []
            raise
        areas = []
        for index in range(1, cells.Areas.Count + 1):
            area = cells.Areas.Item(index)
            top, left = area.Row - 1, area.Column - 1
            areas.append((top, left, top + area.Rows.Count - 1, left + area.Columns.Count - 1))
        return areas

    def used_range(self, sheet):
        """UsedRange, narrowed to the last column that holds content. Formatting alone stretches UsedRange
        across all 16384 columns of most template tabs: 230 million cells where content fills about 8 million.
        Rows stay UsedRange's, since Find can skip rows an AutoFilter hides."""
        worksheet = self.sheets[sheet]
        last = worksheet.Cells.Find(What="*", After=worksheet.Cells(1, 1), LookIn=XL_FORMULAS, LookAt=XL_PART,
                                    SearchOrder=XL_BY_COLUMNS, SearchDirection=XL_PREVIOUS)
        if last is None:
            return None
        used = worksheet.UsedRange
        top, left = used.Row - 1, used.Column - 1
        return (top, min(left, last.Column - 1), top + used.Rows.Count - 1, last.Column - 1)

    def save_copy(self, path):
        """Save this copy as Excel holds it now, to resume from after a restart."""
        started = time.monotonic()
        self.book.SaveCopyAs(str(path))
        LOG.info("Excel saved a resume point in %.1fs", time.monotonic() - started)

    def run_macro(self, name, timeout):
        """Run a workbook macro under the dialog watcher and the time limit."""
        try:
            return self._watched(lambda: self.app.Run(f"'{self.book.Name}'!{name}"), timeout)
        finally:
            if not getattr(self, "killed", False):
                try:
                    self._quiet()
                except Exception:
                    LOG.exception("Excel settings could not be restored after %s", name)

    def refresh_all(self, timeout):
        """Excel's own RefreshAll on this copy, then wait for the background queries to finish."""
        def refresh():
            self.book.RefreshAll()
            self.app.CalculateUntilAsyncQueriesDone()

        return self._watched(refresh, timeout)

    def pivot_ranges(self):
        """Every pivot table's output area on every sheet (TableRange2), as (sheet, top, left, bottom, right)."""
        ranges = []
        for index, worksheet in enumerate(self.sheets):
            tables = worksheet.PivotTables()
            for position in range(1, tables.Count + 1):
                area = tables.Item(position).TableRange2
                top, left = area.Row - 1, area.Column - 1
                ranges.append((index, top, left, top + area.Rows.Count - 1, left + area.Columns.Count - 1))
        return ranges

    def refresh_pivots(self, timeout):
        """Refresh every pivot table of this copy, then wait for the background queries they read."""
        def refresh():
            for worksheet in self.sheets:
                tables = worksheet.PivotTables()
                for position in range(1, tables.Count + 1):
                    tables.Item(position).RefreshTable()
            self.app.CalculateUntilAsyncQueriesDone()

        return self._watched(refresh, timeout)

    def set_slicer(self, cache, selected):
        """Select exactly the named items of one slicer, with the cache's pivots held until all are set."""
        wanted = set(selected)
        slicer = self.book.SlicerCaches(cache)
        tables = slicer.PivotTables
        pivots = [tables.Item(position) for position in range(1, tables.Count + 1)]
        for pivot in pivots:
            pivot.ManualUpdate = True
        try:
            items = slicer.SlicerItems
            found = set()
            for position in range(1, items.Count + 1):
                item = items.Item(position)
                item.Selected = item.Name in wanted
                if item.Name in wanted:
                    found.add(item.Name)
            if found != wanted:
                raise EngineError("slicer item is missing in Excel")
        finally:
            for pivot in pivots:
                pivot.ManualUpdate = False

    def _watched(self, action, timeout):
        """Run an Excel action; a watcher thread presses OK on its dialogs and kills Excel past the limit."""
        dialogs, done, timed_out = [], threading.Event(), threading.Event()
        deadline = time.monotonic() + timeout

        def watch():
            seen = set()
            while not done.is_set():
                if time.monotonic() >= deadline:
                    timed_out.set()
                    self.killed = True
                    _kill(self.pid)
                    return
                _press_dialogs(self.pid, seen, dialogs)
                done.wait(0.1)

        watcher = threading.Thread(target=watch, daemon=True)
        watcher.start()
        try:
            action()
        except Exception as exc:
            if timed_out.is_set():
                raise MacroTimeout("macro timed out") from exc
            raise EngineError("macro failed") from exc
        finally:
            done.set()
            watcher.join()
        if timed_out.is_set():
            raise MacroTimeout("macro timed out")
        return dialogs

    def memory_bytes(self):
        import win32api
        import win32con
        import win32process

        handle = win32api.OpenProcess(win32con.PROCESS_QUERY_INFORMATION | win32con.PROCESS_VM_READ,
                                      False, self.pid)
        try:
            return int(win32process.GetProcessMemoryInfo(handle)["PeakWorkingSetSize"])
        finally:
            win32api.CloseHandle(handle)

    def close(self):
        import pythoncom

        try:
            if getattr(self, "book", None) is not None:
                self.book.Close(SaveChanges=False)
            self.app.Quit()
        except Exception:
            pass
        finally:
            self.book = self.sheets = None
            self.app = None
            _kill(getattr(self, "pid", None))
            shutil.rmtree(self.directory, ignore_errors=True)
            pythoncom.CoUninitialize()


def _kill(pid):
    """Make sure a hung Excel never outlives its worker."""
    if not pid:
        return
    try:
        import win32api
        import win32con

        handle = win32api.OpenProcess(win32con.PROCESS_TERMINATE, False, pid)
        win32api.TerminateProcess(handle, 0)
        win32api.CloseHandle(handle)
    except Exception:
        pass


def _press_dialogs(pid, seen, dialogs):
    """Record each new dialog of the Excel process and press its OK; a poll that races a closing window is retried."""
    import win32con
    import win32gui
    import win32process

    try:
        handles = []

        def visit(hwnd, _):
            if win32gui.GetClassName(hwnd) == DIALOG_CLASS and win32process.GetWindowThreadProcessId(hwnd)[1] == pid:
                handles.append(hwnd)
            return True

        win32gui.EnumWindows(visit, None)
        for hwnd in handles:
            if hwnd in seen:
                continue
            texts = []

            def read(child, _):
                if win32gui.GetClassName(child) == "Static":
                    texts.append(win32gui.GetWindowText(child))
                return True

            win32gui.EnumChildWindows(hwnd, read, None)
            win32gui.PostMessage(hwnd, win32con.WM_COMMAND, win32con.IDOK, 0)
            seen.add(hwnd)
            dialogs.append("\n".join(text for text in texts if text))
    except Exception:
        pass


class FormualizerHost(ExcelHost):
    """Test stand-in: formualizer computes what Excel would for small workbooks."""

    def __init__(self, workbook: Path, sheet_names: list[str]):
        import formualizer
        from openpyxl import load_workbook

        self.book = formualizer.load_workbook(str(workbook))
        self.names = sheet_names
        self.formulas = {}  # (sheet, row, col) -> "=..."
        self.extent = {}  # sheet -> (rows, cols) that hold anything
        self.macros = {}  # name -> callable(host) returning the dialog texts
        self.pivots = []  # (sheet, top, left, bottom, right) output areas tests declare
        self.refresh_hook = None  # callable(host) that rewrites the pivot cells, for tests
        self.slicer_hook = None  # callable(host, cache, selected) that rewrites the pivot cells, for tests
        opened = load_workbook(workbook, read_only=True)
        try:
            for index, name in enumerate(sheet_names):
                for row, cells in enumerate(opened[name].iter_rows()):
                    for col, cell in enumerate(cells):
                        if isinstance(cell.value, str) and cell.value.startswith("="):
                            self.formulas[(index, row, col)] = cell.value
                        if cell.value is not None:
                            self._grow(index, row, col)
        finally:
            opened.close()

    def _grow(self, sheet, row, col):
        rows, cols = self.extent.get(sheet, (0, 0))
        self.extent[sheet] = (max(rows, row + 1), max(cols, col + 1))

    def write(self, sheet, cells):
        for row, col, value in cells:
            self.book.set_value(self.names[sheet], row + 1, col + 1, value)
            self.formulas.pop((sheet, row, col), None)
            if value is not None:
                self._grow(sheet, row, col)

    def write_formulas(self, sheet, cells):
        for row, col, formula in cells:
            self.book.set_formula(self.names[sheet], row + 1, col + 1, formula)
            self.formulas[(sheet, row, col)] = formula
            self._grow(sheet, row, col)

    def calculate(self):
        pass

    def read(self, sheet, top, left, bottom, right):
        name = self.names[sheet]
        cells = [(name, row + 1, col + 1) for row in range(top, bottom + 1)
                 for col in range(left, right + 1)]
        self.book.evaluate_cells(cells)
        width = right - left + 1
        values = [_json_value(self.book.get_value(*cell), False) for cell in cells]
        return [values[index:index + width] for index in range(0, len(values), width)]

    def read_formulas(self, sheet, top, left, bottom, right):
        return [[self.formulas.get((sheet, row, col)) for col in range(left, right + 1)]
                for row in range(top, bottom + 1)]

    def formula_areas(self, sheet):
        return [(row, col, row, col) for s, row, col in sorted(self.formulas) if s == sheet]

    def used_range(self, sheet):
        rows, cols = self.extent.get(sheet, (0, 0))
        return (0, 0, rows - 1, cols - 1) if rows else None

    def run_macro(self, name, timeout):
        return self.macros[name](self)

    def refresh_all(self, timeout):
        return self.macros[REFRESH_ALL](self)

    def pivot_ranges(self):
        return list(self.pivots)

    def refresh_pivots(self, timeout):
        return self.refresh_hook(self) if self.refresh_hook else []

    def set_slicer(self, cache, selected):
        if self.slicer_hook is None:
            raise EngineError("slicers need Excel")
        self.slicer_hook(self, cache, list(selected))

    def memory_bytes(self):
        return 0

    def close(self):
        self.book = None


def default_host():
    return ComExcelHost if os.name == "nt" else FormualizerHost


def _cached_index(prepared: Path, formulas: dict, names: list[str]) -> DependencyIndex:
    code = hashlib.sha256(Path(__file__).with_name("depindex.py").read_bytes()).hexdigest()[:12]
    path = prepared.with_name(f"{prepared.name}.index-{code}.pickle")
    try:
        with path.open("rb") as handle:
            return pickle.load(handle)  # nosec B301 # per-user engine cache written only by this app
    except FileNotFoundError:
        pass
    except Exception:
        path.unlink(missing_ok=True)
    index = DependencyIndex.build(prepared, formulas, names)
    pending = path.with_name(f"{path.name}.{os.getpid()}.tmp")
    try:
        with pending.open("wb") as handle:
            pickle.dump(index, handle, protocol=pickle.HIGHEST_PROTOCOL)
        os.replace(pending, path)
    except OSError:
        pending.unlink(missing_ok=True)
    return index


def warm_caches(template_path: str | Path, cache_dir: str | Path | None = None) -> None:
    """Build the prepared package, formula map and index once, before the first edit needs them."""
    prepared = prepared_template(Path(template_path), cache_dir)
    _, formulas, names = _cached_formula_map(prepared)[:3]
    _cached_index(prepared, formulas, names)


@dataclass
class ExcelEngine:
    host: ExcelHost
    index: DependencyIndex
    formula_anchors: frozenset
    baseline: dict  # Excel's saved values for formula cells
    current: dict = field(default_factory=dict)  # formula cells whose value now differs
    overlay_upserts: dict = field(default_factory=dict)
    overlay_removals: set = field(default_factory=set)
    prepared: Path | None = None
    names: list = field(default_factory=list)
    formula_overrides: dict = field(default_factory=dict)  # (sheet, row, col) -> "=..." set by macros or edits
    override_baseline: set = field(default_factory=set)  # override cells whose baseline we created
    override_index: DependencyIndex | None = None
    pivot_cells: frozenset = frozenset()  # pivot output cells: Excel's, published, never overrides (P6.4 decision 1)
    template_cells: frozenset = frozenset()  # the template's own formula cells
    handles_clears = True

    @classmethod
    def load(cls, template_path: str | Path, cache_dir: str | Path | None = None,
             host_factory=None, workbook: str | Path | None = None, state: dict | None = None) -> ExcelEngine:
        """Open the template, or with ``workbook`` and ``state`` a copy saved by ``checkpoint_state``."""
        path = Path(template_path)
        try:
            prepared = prepared_template(path, cache_dir)
            cells, formulas, names, cached = _cached_formula_map(prepared)[:4]
            index = _cached_index(prepared, formulas, names)
            host = (host_factory or default_host())(Path(workbook) if workbook else path, names)
            if state is None:
                pivot_cells, pivot_values = _pivot_outputs(host)
            else:
                pivot_cells, pivot_values = state["pivot_cells"], {}
        except EngineError:
            raise
        except Exception as exc:
            raise EngineError("Excel calculation engine failed to load") from exc
        baseline = {key: from_com(cached.get(key)) for key in cells}
        baseline.update(pivot_values)
        template = frozenset(cells)
        engine = cls(host, index, template | pivot_cells, baseline, prepared=prepared, names=names,
                     pivot_cells=frozenset(pivot_cells), template_cells=template)
        if state is not None:
            engine.baseline.update(state["baseline"])
            engine.formula_anchors = (template - state["anchors_removed"]) | state["anchors_added"]
            engine.current = dict(state["current"])
            engine.formula_overrides = dict(state["formula_overrides"])
            engine.override_baseline = set(state["override_baseline"])
        return engine

    def checkpoint_state(self) -> dict:
        """What ``load`` needs besides the saved workbook; the template supplies everything else."""
        return {"anchors_added": self.formula_anchors - self.template_cells,
                "anchors_removed": self.template_cells - self.formula_anchors,
                "current": dict(self.current), "formula_overrides": dict(self.formula_overrides),
                "override_baseline": set(self.override_baseline), "pivot_cells": self.pivot_cells,
                "baseline": {key: self.baseline.get(key) for key in self.override_baseline | self.pivot_cells}}

    @property
    def formula_cells(self):
        return self.formula_anchors

    def apply(self, changes: list[tuple[int, int, int, object]]) -> list[tuple[int, int, int, object]]:
        self.overlay_upserts, self.overlay_removals = {}, set()
        # Excel refuses writes inside a pivot table; its output comes from the pivot, never from a stored value.
        changes = [change for change in changes if change[:3] not in self.pivot_cells]
        if not changes:
            return []
        try:
            baselines = self._baselines(changes)
            plain, formulas = self._track(changes, baselines)
            started = time.monotonic()
            for sheet, cells in plain.items():
                self.host.write(sheet, cells)
            for sheet, cells in formulas.items():
                self.host.write_formulas(sheet, cells)
            wrote = time.monotonic()
            self.host.calculate()
            calculated = time.monotonic()
            affected = self._affected([(sheet, row, col) for sheet, row, col, _ in changes])
            traced = time.monotonic()
            published = self._publish(affected)
            LOG.info("Excel edit: write %.1fs (%d cells), calculate %.1fs, trace %.1fs, read %.1fs "
                     "(%d formulas read, %d changed)", wrote - started, len(changes), calculated - wrote,
                     traced - calculated, time.monotonic() - traced, len(affected), len(published))
            return published
        except EngineError:
            raise
        except Exception as exc:
            raise EngineError("Excel calculation engine failed to apply changes") from exc

    def run_macro(self, macro: str, timeout: float = 600) -> dict:
        """Run a template macro in Excel and record exactly what it changed (see P6.3 decision 5)."""
        return self._record(lambda: self.host.run_macro(macro, timeout))

    def refresh_data(self, timeout: float = 600) -> dict:
        """Excel's RefreshAll on this copy, recorded like a macro (see P6.5 refresh decision)."""
        return self._record(lambda: self.host.refresh_all(timeout))

    def set_slicer(self, cache: str, selected: list[str]) -> dict:
        """Select the named items of an Excel slicer and record what its pivots output (P6.4 decision 4)."""
        return self._record(lambda: self.host.set_slicer(cache, list(selected)))

    def restore_pivots(self, state: dict, timeout: float = 600) -> dict:
        """Re-apply stored pivot state to a freshly loaded copy: refresh, then slicer selections (decision 2)."""
        def replay():
            dialogs = self.host.refresh_pivots(timeout) if state.get("refreshed") else []
            for cache, selected in sorted(state.get("slicers", {}).items()):
                self.host.set_slicer(cache, list(selected))
            return dialogs

        return self._record(replay)

    def _record(self, action) -> dict:
        """Run an Excel action and record what it wrote, as a macro's delta.

        Overlay state is not reset here: the caller syncs the desired state first, and the
        delta it publishes must cover that sync plus the action.
        """
        try:
            started = time.monotonic()
            before = self._scan()
            scanned = time.monotonic()
            messages = action() or []
            ran = time.monotonic()
            self.host.calculate()
            now = set()
            areas = [(sheet, rect) for sheet in range(len(self.names))
                     for rect in self.host.formula_areas(sheet)]
            for sheet, (top, left, bottom, right) in areas:
                now.update((sheet, row, col) for row in range(top, bottom + 1) for col in range(left, right + 1))
            new = now - self.formula_anchors
            detached = self.formula_anchors - now - self.pivot_cells
            texts = {}
            for sheet, top, left, bottom, right, members in _blocks(new):
                grid = self.host.read_formulas(sheet, top, left, bottom, right)
                for row, col in members:
                    text = grid[row - top][col - left]
                    if text is None:
                        raise EngineError("macro formula could not be read")
                    texts[(sheet, row, col)] = text
            old, fresh = _diff(before, self._scan())
            differing = old.keys() | fresh.keys()
            # every other cell kept its value through the action; read the few that are needed
            same = self._read_cells(key for key in new | detached if key not in differing)

            def value(side, key):
                return side.get(key) if key in differing else same.get(key)

            changed = {key for key in differing if not exactly_equal(old.get(key), fresh.get(key))}
            moved = {key for key in changed if key not in now and key not in self.pivot_cells}
            # formula cells whose value moved, including ones whose text the action rewrote in place
            recalculated = {key for key in changed if key in now}
            contents = [(*key, {"formula": texts[key]}) for key in new]
            contents += [(*key, value(fresh, key)) for key in moved | detached]
            contents.sort(key=lambda item: item[:3])
            self._track(contents, {key: value(old, key) for key in new})
            # pivot outputs changed by the action are published, and so is whatever reads them
            reached = self._affected([item[:3] for item in contents] + sorted(self.pivot_cells))
            changes = self._publish(reached | (recalculated & self.formula_anchors))
            LOG.info("Excel action: scan %.1fs, run %.1fs, record %.1fs, %d cells written, %d values changed",
                     scanned - started, ran - scanned, time.monotonic() - ran, len(contents), len(changes))
            return {"messages": messages, "changes": changes, "contents": contents,
                    "before": {item[:3]: value(old, item[:3]) for item in contents}}
        except EngineError:
            raise
        except Exception as exc:
            raise EngineError("Excel calculation engine failed to run the macro") from exc

    def _scan(self) -> dict:
        """Every sheet's content box in row blocks, kept as bytes: (sheet, first row) -> (left, bytes).

        Bytes, not cells: a whole-workbook scan is millions of values, and only the blocks an action
        changed are ever unpacked.
        """
        blocks = {}
        for sheet in range(len(self.names)):
            box = self.host.used_range(sheet)
            if box is None:
                continue
            top, left, bottom, right = box
            for start in range(_block(top), bottom + 1, BLOCK_ROWS):
                rows = self.host.read(sheet, start, left, min(start + BLOCK_ROWS - 1, bottom), right)
                blocks[(sheet, start)] = (left, pickle.dumps(rows, protocol=pickle.HIGHEST_PROTOCOL))
        return blocks

    def _read_cells(self, keys) -> dict:
        found = {}
        for sheet, top, left, bottom, right, members in _blocks(keys):
            grid = self.host.read(sheet, top, left, bottom, right)
            for row, col in members:
                found[(sheet, row, col)] = from_com(grid[row - top][col - left])
        return found

    def _baselines(self, changes) -> dict:
        """The value a cell had before it becomes a formula override (read before the write)."""
        return {(sheet, row, col): self.get(sheet, row, col) for sheet, row, col, value in changes
                if _formula_text(value) is not None and (sheet, row, col) not in self.formula_anchors}

    def _track(self, changes, baselines) -> tuple[dict, dict]:
        """Update formula state for the changes; returns the plain and formula writes for the host."""
        plain, formulas, added, removed, overrides_changed = {}, {}, set(), set(), False
        for sheet, row, col, value in changes:
            key = (sheet, row, col)
            text = _formula_text(value)
            if text is None:
                plain.setdefault(sheet, []).append((row, col, value))
                if self.formula_overrides.pop(key, None) is not None:
                    overrides_changed = True
                    if key in self.override_baseline:
                        self.override_baseline.discard(key)
                        self.baseline.pop(key, None)
                if key in self.formula_anchors:
                    removed.add(key)
            else:
                formulas.setdefault(sheet, []).append((row, col, text))
                self.formula_overrides[key] = text
                overrides_changed = True
                if key not in self.formula_anchors:
                    added.add(key)
                    self.baseline[key] = baselines.get(key)
                    self.override_baseline.add(key)
        for key in removed:
            self.current.pop(key, None)
            self.overlay_upserts.pop(key, None)
            self.overlay_removals.add(key)
        if added or removed:
            self.formula_anchors = frozenset((self.formula_anchors | added) - removed)
        if overrides_changed:
            self.override_index = None
        return plain, formulas

    def _override_index(self) -> DependencyIndex | None:
        if not self.formula_overrides:
            return None
        if self.override_index is None:
            self.override_index = DependencyIndex.build(self.prepared, dict(self.formula_overrides), self.names)
        return self.override_index

    def _affected(self, edited) -> set:
        """Formulas an edit can reach, through the template index and the override index until nothing new."""
        found = {key for key in edited if key in self.formula_anchors}
        override = self._override_index()
        frontier = set(edited)
        while frontier:
            reached = set(self.index.cone(frontier))
            if override is not None:
                reached |= override.cone(frontier)
            frontier = {key for key in reached if key in self.formula_anchors} - found
            found |= frontier
            if override is None:
                break
        return found

    def _publish(self, affected) -> list[tuple[int, int, int, object]]:
        self.host.invalidate(affected)
        published = []
        for sheet, top, left, bottom, right, members in _blocks(affected):
            values = self.host.read(sheet, top, left, bottom, right)
            for row, col in members:
                value = from_com(values[row - top][col - left])
                key = (sheet, row, col)
                if exactly_equal(value, self.current.get(key, self.baseline.get(key))):
                    continue
                published.append((sheet, row, col, value))
                if exactly_equal(value, self.baseline.get(key)):
                    self.current.pop(key, None)
                    self.overlay_removals.add(key)
                    self.overlay_upserts.pop(key, None)
                else:
                    self.current[key] = value
                    self.overlay_upserts[key] = value
                    self.overlay_removals.discard(key)
        return sorted(published, key=lambda item: item[:3])

    def get(self, sheet: int, row: int, col: int):
        key = (sheet, row, col)
        if key in self.formula_anchors:
            return self.current.get(key, self.baseline.get(key))
        try:
            return from_com(self.host.read(sheet, row, col, row, col)[0][0])
        except Exception as exc:
            raise EngineError("Excel calculation engine failed to read a cell") from exc

    def overlay_snapshot(self):
        return [(*key, value) for key, value in sorted(self.current.items())]

    def overlay_delta(self):
        upserts = [(*key, value) for key, value in sorted(self.overlay_upserts.items())]
        return upserts, sorted(self.overlay_removals)

    def memory_bytes(self) -> int:
        return self.host.memory_bytes()

    def close(self) -> None:
        self.host.close()


def _pivot_outputs(host) -> tuple[set, dict]:
    """Every pivot output cell and its value as Excel saved it (the baseline for publishing changes)."""
    cells, values = set(), {}
    for sheet, top, left, bottom, right in host.pivot_ranges():
        for start in range(top, bottom + 1, BLOCK_ROWS):
            stop = min(start + BLOCK_ROWS - 1, bottom)
            for offset, row in enumerate(host.read(sheet, start, left, stop, right)):
                for position, raw in enumerate(row):
                    key = (sheet, start + offset, left + position)
                    cells.add(key)
                    values[key] = from_com(raw)
    return cells, values


def _block(row):
    return row - row % BLOCK_ROWS


def _diff(before, after):
    """Before and after values of the cells that may differ between two scans; any cell in neither
    dict holds the same value in both."""
    old, new = {}, {}
    for key in before.keys() | after.keys():
        first, second = before.get(key), after.get(key)
        if first == second:
            continue
        sheet, start = key
        if first is not None and second is not None and first[0] == second[0]:
            left = first[0]
            rows_before, rows_after = pickle.loads(first[1]), pickle.loads(second[1])
            for offset in range(max(len(rows_before), len(rows_after))):
                a = rows_before[offset] if offset < len(rows_before) else ()
                b = rows_after[offset] if offset < len(rows_after) else ()
                if a == b:
                    continue
                for position in range(max(len(a), len(b))):
                    x = a[position] if position < len(a) else None
                    y = b[position] if position < len(b) else None
                    if x != y or type(x) is not type(y):
                        cell = (sheet, start + offset, left + position)
                        if x is not None:
                            old[cell] = from_com(x)
                        if y is not None:
                            new[cell] = from_com(y)
            continue
        for side, out in ((first, old), (second, new)):
            if side is None:
                continue
            left = side[0]
            for offset, row in enumerate(pickle.loads(side[1])):
                for position, raw in enumerate(row):
                    if raw is not None:
                        out[(sheet, start + offset, left + position)] = from_com(raw)
    return old, new


def _rectangles(cells):
    """Cells as (top, left, rows) rectangles, so Excel takes a whole block in one call, not one per cell.

    Writing cell by cell costs one cross-process COM call each; an uploaded copy replays hundreds of
    thousands of cells. Later duplicates of a cell win, as they did when writes ran one by one.
    """
    latest = {(row, col): value for row, col, value in cells}
    runs = []
    for (row, col), value in sorted(latest.items()):
        if runs and runs[-1][0] == row and runs[-1][1] + len(runs[-1][2]) == col:
            runs[-1][2].append(value)
        else:
            runs.append((row, col, [value]))
    rectangles = []
    for row, col, values in runs:
        last = rectangles[-1] if rectangles else None
        if last and last[1] == col and len(last[2][0]) == len(values) and last[0] + len(last[2]) == row:
            last[2].append(values)
        else:
            rectangles.append((row, col, [values]))
    return rectangles


def _blocks(affected):
    """Group cells into row blocks with one bounding column span each."""
    grouped: dict[tuple[int, int], list] = {}
    for sheet, row, col in affected:
        grouped.setdefault((sheet, row // BLOCK_ROWS), []).append((row, col))
    for (sheet, _), members in sorted(grouped.items()):
        rows = [row for row, _ in members]
        cols = [col for _, col in members]
        yield sheet, min(rows), min(cols), max(rows), max(cols), members
