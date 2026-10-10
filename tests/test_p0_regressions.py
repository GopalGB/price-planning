"""Synthetic regressions for the ten P0 review findings, without repository writes."""
import contextlib
import hashlib
import io
import json
import os
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / 'ops'))
import bench_real
import content_guard as cg
import find_inputs
import install_env as ie
from typed import same, typed


def zip_bytes(members):
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, 'w', zipfile.ZIP_DEFLATED) as archive:
        for name, content in members.items():
            archive.writestr(name, content)
    return stream.getvalue()


def guard(files, deny=(), allow=None):
    return cg.check(list(files), files.__getitem__, deny, allow or {})


class Finding01(unittest.TestCase):
    def test_locked_tracked_vendor_wheels_keep_their_hash_exceptions(self):
        locks = (ROOT / 'requirements.lock').read_text() + (ROOT / 'requirements-dev.lock').read_text()
        allow = json.loads((ROOT / 'ops/large_file_allowlist.json').read_text())
        paths = ['vendor/waitress-3.0.2-py3-none-any.whl',
                 'vendor/wheels/et_xmlfile-2.0.0-py3-none-any.whl',
                 'vendor/wheels/openpyxl-3.1.5-py2.py3-none-any.whl',
                 'vendor/wheels/waitress-3.0.2-py3-none-any.whl']
        for name in paths:
            with self.subTest(wheel=name):
                data = (ROOT / name).read_bytes()
                digest = hashlib.sha256(data).hexdigest()
                self.assertIn('sha256:' + digest, locks)
                self.assertEqual(guard({name: data}, allow=allow), [])

    def test_renamed_binary_containers_are_blocked(self):
        samples = [zip_bytes({'xl/workbook.xml': '<workbook/>'}),
                   zip_bytes({'word/document.xml': '<document/>'}),
                   zip_bytes({'plain.txt': 'compressed confidential data'}),
                   b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1' + b'\0' * 32,
                   b'%PDF-1.7\n', b'SQLite format 3\0' + b'\0' * 32,
                   b'PK\x03\x04broken archive']
        for data in samples:
            with self.subTest(signature=data[:8]):
                self.assertTrue(guard({'renamed.txt': data}))

    def test_only_hash_verified_vendor_wheels_get_binary_exception(self):
        data = zip_bytes({'wheel.dist-info/METADATA': 'synthetic package'})
        name = 'vendor/wheels/package.whl'
        digest = hashlib.sha256(data).hexdigest()
        self.assertEqual(guard({name: data}, allow={name: digest}), [])
        self.assertTrue(guard({name: data}, allow={name: 'wrong'}))
        self.assertTrue(guard({'outside.whl': data}, allow={'outside.whl': digest}))


class Finding02(unittest.TestCase):
    def test_bom_text_and_invalid_encodings_cannot_bypass_denylist(self):
        for encoding in ('utf-16', 'utf-32', 'utf-8-sig'):
            with self.subTest(encoding=encoding):
                self.assertTrue(guard({'note.txt': 'Forbidden Phrase'.encode(encoding)},
                                      ['  forbidden phrase  ']))
                self.assertEqual(guard({'note.txt': 'permitted'.encode(encoding)}, ['forbidden']), [])
        self.assertTrue(guard({'note.txt': b'forbidden\xff phrase'}, ['forbidden phrase']))
        self.assertTrue(guard({'note.txt': b'\0forbidden'}, ['forbidden']))

    def test_utf8_bom_denylist_first_entry_is_active(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / 'ops').mkdir()
            (root / 'ops/large_file_allowlist.json').write_text('{}')
            (root / '.content-denylist').write_text('forbidden\n', encoding='utf-8-sig')
            (root / 'note.txt').write_text('forbidden')
            with patch.object(cg, 'ROOT', root), patch.object(cg, '_git', return_value=b'note.txt\0'):
                with contextlib.redirect_stdout(io.StringIO()):
                    self.assertEqual(cg.main(['--all']), 1)


class Finding03(unittest.TestCase):
    def test_staged_type_change_is_scanned_from_index(self):
        def git(*args):
            if args[0] == 'diff':
                # A symlink-to-regular-file change is T, not A/C/M/R.
                filters = [arg.split('=', 1)[1] for arg in args if arg.startswith('--diff-filter=')]
                return b'note.txt\0' if not filters or 'T' in filters[0] or filters[0] == 'd' else b''
            if args == ('show', ':note.txt'):
                return b'%PDF-1.7 confidential staged contents'
            raise AssertionError(args)
        with patch.object(cg, '_git', side_effect=git), contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(cg.main(['--staged']), 1)


class Finding04(unittest.TestCase):
    def test_force_added_private_paths_are_blocked_case_insensitively(self):
        paths = ['RESULTS-PRIVATE/report.json', 'nested/Data/value.json', 'BACKUPS/x.txt',
                 'Logs/x.txt', 'DEBUG.LOG', 'cache.DB-wal', 'cache.Db3', 'macro.BAS',
                 '.ENV', 'nested/.Env.local', 'results-private\\report.json']
        for name in paths:
            with self.subTest(path=name):
                self.assertTrue(guard({name: b'synthetic'}))
        self.assertEqual(guard({'metadata/x.py': b'pass', 'catalog.txt': b'clean'}), [])


class Finding05(unittest.TestCase):
    def test_reports_refuse_equal_nested_and_symlinked_destinations(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            public = root / 'public'
            public.mkdir()
            alias = root / 'alias'
            alias.symlink_to(public, target_is_directory=True)
            pairs = [(public, public), (public, public / 'private'),
                     (public / 'child', public), (public, alias)]
            for pub, priv in pairs:
                with self.subTest(public=pub, private=priv):
                    with self.assertRaises(ValueError):
                        bench_real.Report(SimpleNamespace(public_dir=pub, private_dir=priv))
            self.assertEqual(list(public.iterdir()), [])
            report = bench_real.Report(SimpleNamespace(public_dir=public, private_dir=root / 'private'))
            report.write()
            self.assertTrue((public / report.name).is_file())


ERROR_CASES = [('=1/0', 'Div', '#DIV/0!'), ('=NA()', 'Na', '#N/A'),
               ('=1+"x"', 'Value', '#VALUE!'), ('=#REF!', 'Ref', '#REF!'),
               ('=MissingName', 'Name', '#NAME?'), ('=SQRT(-1)', 'Num', '#NUM!')]


class Finding06(unittest.TestCase):
    def test_real_engine_error_dicts_compare_to_excel_cached_tags(self):
        import formualizer
        workbook = formualizer.Workbook()
        workbook.add_sheet('Probe')
        for row, (formula, kind, tag) in enumerate(ERROR_CASES, 1):
            with self.subTest(kind=kind):
                workbook.set_formula('Probe', row, 1, formula)
                actual = workbook.evaluate_cell('Probe', row, 1)
                self.assertEqual(actual, {'type': 'Error', 'kind': kind})
                self.assertEqual(typed(actual), ('error', tag))
                self.assertTrue(same(tag, workbook.get_value('Probe', row, 1)))
                self.assertFalse(same('#NULL!', actual))


class Finding07(unittest.TestCase):
    def test_protection_uses_style_inheritance_and_apply_flag(self):
        ns = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
        styles = f'''<styleSheet xmlns="{ns}"><cellStyleXfs count="2">
          <xf><protection locked="1"/></xf><xf><protection locked="0"/></xf>
          </cellStyleXfs><cellXfs count="6">
          <xf xfId="0" applyProtection="0"><protection locked="0"/></xf>
          <xf xfId="1"/><xf xfId="1" applyProtection="false"><protection locked="1"/></xf>
          <xf xfId="1" applyProtection="1"><protection locked="1"/></xf>
          <xf xfId="0" applyProtection="true"><protection locked="false"/></xf>
          <xf xfId="0"/></cellXfs></styleSheet>'''
        with zipfile.ZipFile(io.BytesIO(zip_bytes({'xl/styles.xml': styles}))) as archive:
            self.assertEqual(find_inputs._unlocked_style_ids(archive), {'1', '2', '4'})


def installer_process(python, state, install_fails):
    def run(args, **kwargs):
        if not python.exists() or not state['usable']:
            raise OSError('unusable interpreter')
        if args[1:] == ['--version']:
            return subprocess.CompletedProcess(args, 0)
        if args[1:4] == ['-m', 'ensurepip', '--upgrade']:
            state.update(pip=True, bootstrapped=True)
            return subprocess.CompletedProcess(args, 0)
        if args[1:4] == ['-m', 'pip', '--version']:
            if not state['pip'] and kwargs.get('check'):
                raise subprocess.CalledProcessError(1, args)
            return subprocess.CompletedProcess(args, 0 if state['pip'] else 1)
        if args[1:4] == ['-m', 'pip', 'install']:
            if install_fails:
                raise subprocess.CalledProcessError(1, args)
            state['installed'] = True
            return subprocess.CompletedProcess(args, 0)
        raise AssertionError(args)
    return run


class Finding08(unittest.TestCase):
    def run_installer(self, root, **scenario):
        env = root / '.venv'
        python = ie.venv_python(env)
        python.parent.mkdir(parents=True, exist_ok=True)
        if scenario.get('interpreter', True):
            python.write_text('synthetic interpreter')
        lock = root / 'requirements.lock'
        lock.write_text('synthetic lock')
        stamp = env / '.lock-sha256'
        stamp.write_text(ie.lock_sha256(lock))
        install_fails = scenario.get('install_fails', False)
        state = {'pip': scenario.get('pip', True), 'usable': not scenario.get('broken', False),
                 'installed': False, 'bootstrapped': False}

        def create(destination):
            python.write_text('recreated synthetic interpreter')
            state.update(usable=True, pip=True)

        run = installer_process(python, state, install_fails)
        with patch.object(ie, 'ROOT', root), patch.object(ie.subprocess, 'run', side_effect=run):
            with patch.object(ie.venv.EnvBuilder, 'create', side_effect=create):
                with contextlib.redirect_stdout(io.StringIO()):
                    if install_fails:
                        with self.assertRaises(subprocess.CalledProcessError):
                            ie.main()
                        self.assertFalse(stamp.exists(), 'stale stamp must not survive a failed repair')
                    else:
                        ie.main()
                        self.assertTrue(python.exists())
                        self.assertTrue(state['pip'])
                        self.assertTrue(state['usable'])
                        self.assertEqual(stamp.read_text(), ie.lock_sha256(lock))
        return state

    def test_matching_stamp_does_not_trust_missing_or_broken_interpreter(self):
        for kwargs in ({'interpreter': False}, {'broken': True}):
            with self.subTest(**kwargs), tempfile.TemporaryDirectory() as directory:
                self.assertTrue(self.run_installer(Path(directory), **kwargs)['installed'])

    def test_matching_stamp_does_not_trust_missing_pip(self):
        with tempfile.TemporaryDirectory() as directory:
            state = self.run_installer(Path(directory), pip=False)
            self.assertTrue(state['bootstrapped'])
            self.assertTrue(state['installed'])

    def test_failed_install_does_not_leave_stamp(self):
        with tempfile.TemporaryDirectory() as directory:
            self.run_installer(Path(directory), interpreter=False, install_fails=True)

    def test_healthy_matching_stamp_does_not_reinstall(self):
        with tempfile.TemporaryDirectory() as directory:
            self.assertFalse(self.run_installer(Path(directory))['installed'])


class Finding09(unittest.TestCase):
    def test_hook_supports_python_launcher_without_maxhq(self):
        for launcher in ('python3', 'python', 'py'):
            for code in (0, 1):
                with self.subTest(launcher=launcher, guard_exit=code), tempfile.TemporaryDirectory() as directory:
                    root = Path(directory)
                    shim = root / 'bin'
                    shim.mkdir()
                    trace = root / 'argv.txt'
                    script = shim / launcher
                    # Answer the hook's liveness probe; record and fail only the guard call.
                    script.write_text('#!/bin/sh\ncase " $* " in *" -c "*) exit 0;; esac\n'
                                      f'printf "%s\\n" "$@" > "{trace}"\nexit {code}\n')
                    script.chmod(0o755)
                    git = shim / 'git'
                    git.write_text(f'#!/bin/sh\nprintf "%s\\n" "{root}"\n')
                    git.chmod(0o755)
                    result = subprocess.run(['/bin/sh', str(ROOT / '.githooks/pre-commit')], cwd=root,
                                            env=dict(os.environ, PATH=str(shim)), capture_output=True, text=True)
                    self.assertEqual(result.returncode, code, result.stderr)
                    self.assertTrue(trace.exists(), 'hook must invoke the content guard')
                    expected = ['ops/content_guard.py', '--staged']
                    if launcher == 'py':
                        expected.insert(0, '-3')
                    self.assertEqual(trace.read_text().splitlines(), expected)


class Finding10(unittest.TestCase):
    def test_split_unescapes_quoted_sheet_names(self):
        self.assertEqual(bench_real.split("'O''Brien'!A1"), ("O'Brien", 1, 1))

    def test_quoted_apostrophe_input_roundtrips_and_has_engine_fanout(self):
        import formualizer
        from openpyxl import Workbook
        from openpyxl.styles import Protection
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'synthetic.xlsx'
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "O'Brien"
            sheet.protection.sheet = True
            sheet['A1'] = 2
            sheet['A1'].protection = Protection(locked=False)
            sheet['B1'] = '=A1*2'
            workbook.save(path)
            refs = find_inputs.find_inputs(path)
            self.assertEqual(refs, ["'O''Brien'!A1"])
            self.assertEqual(bench_real.split(refs[0]), ("O'Brien", 1, 1))
            engine = formualizer.load_workbook(str(path))
            engine.evaluate_all()
            picks = bench_real.pick_inputs(engine, refs)
            self.assertEqual(picks['worst'][0], 1)


if __name__ == '__main__':
    unittest.main()
