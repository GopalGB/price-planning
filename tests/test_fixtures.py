# tests/test_fixtures.py
import sys, tempfile, unittest, zipfile
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent / 'ops'))
import fixtures

class Fixtures(unittest.TestCase):
    def test_semantic_has_protected_inputs_and_modern_functions(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / 'semantic.xlsx'
            meta = fixtures.make_semantic(p)
            self.assertGreaterEqual(len(meta['inputs']), 5)
            joined = ' '.join(meta['formulas'].values())
            for fn in ('XLOOKUP', 'LET', 'FILTER', 'UNIQUE', 'SEQUENCE', 'SWITCH', 'IFS', 'NUMBERVALUE', 'IFERROR', 'SUMIFS'):
                self.assertIn(fn, joined)
            with zipfile.ZipFile(p) as archive:
                xml = archive.read('xl/worksheets/sheet1.xml').decode()
            self.assertIn('sheetProtection', xml)
    def test_scale_dimensions(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / 'scale.xlsx'
            fixtures.make_scale(p, rows=200, cols=30, lookup_rows=500)
            from openpyxl import load_workbook
            wb = load_workbook(p)
            self.assertEqual(wb['Grid'].max_row, 200)
            self.assertEqual(wb['Grid'].max_column, 30)
            self.assertEqual(wb['Lookup'].max_row, 500)
            wb.close()
