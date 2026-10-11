from pathlib import Path
import tempfile
import unittest
from unittest import mock
import xml.etree.ElementTree as ET

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

from ops.fixtures import make_semantic
from tests.v2.test_template_importer import rewrite_member


ROOT = Path(__file__).resolve().parents[2]
TMP = ROOT / "results-private" / "tmp"


def shared_isformula_fixture(path):
    workbook = Workbook()
    values = workbook.active
    values.title = "Values"
    values["A1"] = "=1"
    values["B1"] = "=2"
    formulas = workbook.create_sheet("Formulas")
    formulas["B2"] = "=ISFORMULA(Values!A1)"
    formulas["C2"] = "=ISFORMULA(Values!B1)"
    workbook.save(path)

    def make_shared(data):
        namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"

        def qualified(name):
            return f"{{{namespace}}}{name}"

        root = ET.fromstring(data)
        row = root.find(f".//{qualified('row')}[@r='2']")
        cells = {cell.get("r"): cell for cell in row.findall(qualified("c"))}
        master, follower = cells["B2"], cells["C2"]
        master.find(qualified("f")).attrib.update({"t": "shared", "ref": "B2:C2", "si": "9"})
        follower_formula = follower.find(qualified("f"))
        follower_formula.attrib.clear()
        follower_formula.attrib.update({"t": "shared", "si": "9"})
        follower_formula.text = None
        row.remove(follower)
        row.insert(list(row).index(master), follower)
        return ET.tostring(root, encoding="utf-8", xml_declaration=True)

    rewrite_member(path, "xl/worksheets/sheet2.xml", make_shared)


def array_formula_fixture(path):
    workbook = Workbook()
    sheet = workbook.active
    for row, value in enumerate((1, 2, 3), start=1):
        sheet.cell(row, 1, value)
    sheet["B1"] = "=A1*2"
    sheet["B2"] = 999
    sheet["B3"] = 999
    workbook.save(path)

    def make_array(data):
        namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        root = ET.fromstring(data)
        formula = root.find(f".//{{{namespace}}}c[@r='B1']/{{{namespace}}}f")
        formula.attrib.update({"t": "array", "ref": "B1:B3"})
        formula.text = "A1:A3*2"
        sheet_data = root.find(f"{{{namespace}}}sheetData")
        rows = {row.get("r"): row for row in sheet_data.findall(f"{{{namespace}}}row")}
        missing = rows["3"].find(f"{{{namespace}}}c[@r='B3']")
        rows["3"].remove(missing)
        for row in rows.values():
            sheet_data.remove(row)
        sheet_data.extend([rows["2"], rows["3"], rows["1"]])
        return ET.tostring(root, encoding="utf-8", xml_declaration=True)

    rewrite_member(path, "xl/worksheets/sheet1.xml", make_array)


class FormualizerAdapterTests(unittest.TestCase):
    def test_apply_recalculates_semantic_dependent_with_real_engine(self):
        """Replacing the engine call with cached fixture values must fail."""
        from pp.engine.formualizer import FormualizerAdapter

        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "semantic.xlsm"
            make_semantic(source)
            engine = FormualizerAdapter.load(source)

            changes = engine.apply([(0, 1, 1, 12)])

            self.assertEqual(engine.get(2, 1, 1), 24)
            self.assertIn((2, 1, 1, 24), changes)
            self.assertEqual(engine.get(2, 1, 1), 24)

    def test_template_formula_map_answers_isformula_without_engine_semantics(self):
        """Dropping formula metadata must make ISFORMULA compatibility wrong."""
        from pp.engine.formualizer import FormualizerAdapter

        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "semantic.xlsm"
            make_semantic(source)
            engine = FormualizerAdapter.load(source)

            self.assertTrue(engine.is_formula(2, 1, 1))
            self.assertFalse(engine.is_formula(0, 1, 1))

    def test_formula_map_and_dependency_interface_are_not_dimension_capped(self):
        """Scanning only a fixed 256 by 64 window must miss this formula."""
        from pp.engine.formualizer import FormualizerAdapter

        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "wide.xlsm"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Wide"
            sheet.cell(300, 70, "=A1+1")
            workbook.save(source)

            engine = FormualizerAdapter.load(source)

            self.assertTrue(engine.is_formula(0, 299, 69))
            self.assertIn("'Wide'!BR300", engine.dependents(0, 0, 0))

    def test_required_compatibility_functions_have_excel_semantics(self):
        """Trusting the engine's broken ISFORMULA result must fail."""
        from pp.engine.formualizer import FormualizerAdapter

        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "compat.xlsm"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Compat"
            sheet["A1"] = 10
            sheet["A2"] = 20
            sheet["B1"] = "=A1*2"
            sheet["C1"] = "=ISFORMULA(B1)"
            sheet["C2"] = '=CELL("address",A1)'
            sheet["C3"] = '=INDIRECT("A2")'
            sheet["C4"] = "=OFFSET(A1,1,0)"
            sheet["C5"] = '=SWITCH(A1,10,"ten","other")'
            workbook.save(source)

            engine = FormualizerAdapter.load(source)

            self.assertIs(engine.get(0, 0, 2), True)
            self.assertEqual(engine.get(0, 1, 2), "$A$1")
            self.assertEqual(engine.get(0, 2, 2), 20)
            self.assertEqual(engine.get(0, 3, 2), 20)
            self.assertEqual(engine.get(0, 4, 2), "ten")

    def test_shared_formula_follower_expands_before_isformula_rewrite(self):
        """Leaving a shared follower's formula text blank must not bypass ISFORMULA compatibility."""
        from pp.engine.formualizer import FormualizerAdapter

        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "shared-isformula.xlsm"
            shared_isformula_fixture(source)

            engine = FormualizerAdapter.load(source)

            self.assertIs(engine.get(1, 1, 1), True)
            self.assertIs(engine.get(1, 1, 2), True)

    def test_isformula_normalizes_lowercase_a1_and_sheet_name_case(self):
        """Formula lookup must follow Excel's case-insensitive reference rules."""
        from pp.engine.formualizer import FormualizerAdapter

        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "case-insensitive-isformula.xlsm"
            workbook = Workbook()
            inputs = workbook.active
            inputs.title = "Inputs"
            inputs["A1"] = "=1"
            formulas = workbook.create_sheet("Formulas")
            formulas["A1"] = "=ISFORMULA('inputs'!a1)"
            workbook.save(source)

            engine = FormualizerAdapter.load(source)

            self.assertIs(engine.get(1, 0, 0), True)

    def test_unsupported_isformula_forms_fail_closed_outside_literals(self):
        """Ranges, defined names, and expressions must not reach native FALSE semantics."""
        from pp.engine.formualizer import EngineError, FormualizerAdapter

        formulas = ["=ISFORMULA(A1:A2)", "=ISFORMULA(TargetCell)",
                    "=ISFORMULA(OFFSET(A1,0,0))"]
        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            for index, formula in enumerate(formulas):
                with self.subTest(formula=formula):
                    source = Path(temp) / f"unsupported-{index}.xlsm"
                    workbook = Workbook()
                    sheet = workbook.active
                    sheet["A1"] = "=1"
                    sheet["A2"] = formula
                    workbook.defined_names.add(
                        DefinedName("TargetCell", attr_text="'Sheet'!$A$1")
                    )
                    workbook.save(source)
                    with self.assertRaisesRegex(EngineError, "unsupported ISFORMULA"):
                        FormualizerAdapter.load(source)

    def test_native_sheet_names_must_match_xml_order_exactly(self):
        """A reordered or skipped native sheet must fail before indexed cell access."""
        from pp.engine.formualizer import EngineError, FormualizerAdapter

        class NativeWorkbook:
            sheet_names = ["NativeOnly"]

            def evaluate_all(self):
                pass

        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "sheet-order.xlsm"
            workbook = Workbook()
            workbook.active.title = "XML Name"
            workbook.save(source)
            with mock.patch("pp.engine.formualizer.formualizer.load_workbook",
                            return_value=NativeWorkbook()):
                with self.assertRaisesRegex(EngineError, "sheet names"):
                    FormualizerAdapter.load(source)

    def test_legacy_array_extent_fails_closed_instead_of_using_cached_followers(self):
        """Unsupported CSE followers must never be published as cached constants."""
        from pp.engine.formualizer import EngineError, FormualizerAdapter

        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "array.xlsm"
            array_formula_fixture(source)

            with self.assertRaisesRegex(EngineError, r"array.*B1:B3"):
                FormualizerAdapter.load(source)

    def test_legacy_array_extent_fails_template_import_with_capability_report(self):
        """Template import must reject an array type the native engine cannot safely calculate."""
        from pp.template.importer import TemplateImportError, import_template

        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            workspace = Path(temp)
            source = workspace / "array.xlsm"
            array_formula_fixture(source)

            with self.assertRaisesRegex(TemplateImportError, r"array.*B1:B3"):
                import_template(source, workspace / "data")


if __name__ == "__main__":
    unittest.main()
