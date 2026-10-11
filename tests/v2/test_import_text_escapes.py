import io
from pathlib import Path
import shutil
import unittest
import zipfile
import xml.etree.ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Protection
from openpyxl.utils.escape import unescape

from tests.v2.api_support import APIFixture


MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
CONTENT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
NS = f"{{{MAIN_NS}}}"
CR_TEXT = "a\rb"
SHARED_TEXT = "shared\rline emoji 😀 literal _x000D_ overlap _x005F_ cross _x000D_"
INLINE_TEXT = "inline\rline emoji 😀 literal _x000D_ overlap _x005F_ cross _x000D_"
FRAGMENT_TEXT = "a_x000D_b"


def rich_escape_fixture(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Escapes"
    sheet.protection.sheet = True
    for row in range(1, 6):
        sheet.cell(row, 1).value = "placeholder"
    sheet["A1"].protection = Protection(locked=False)
    workbook.save(path)
    _rewrite_rich_package(Path(path))


def _rewrite_rich_package(path):
    replacement = path.with_name("rich-escapes-rewritten.xlsm")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(replacement, "w") as target:
        names = set(source.namelist())
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                data = _rich_sheet(data)
            elif info.filename == "[Content_Types].xml":
                data = _shared_content_type(data)
            elif info.filename == "xl/_rels/workbook.xml.rels":
                data = _shared_relationship(data)
            target.writestr(info, data)
        if "xl/sharedStrings.xml" not in names:
            target.writestr("xl/sharedStrings.xml", _shared_strings_xml(), zipfile.ZIP_DEFLATED)
    shutil.move(replacement, path)


def _rich_sheet(data):
    root = ET.fromstring(data)
    cells = {cell.get("r"): cell for cell in root.findall(f".//{NS}c")}
    _set_inline(cells["A1"], ["a_x000D_", "b"])
    for child in list(cells["A2"]):
        cells["A2"].remove(child)
    cells["A2"].set("t", "s")
    ET.SubElement(cells["A2"], f"{NS}v").text = "0"
    _set_inline(cells["A3"], [
        "inline_x000D_line emoji _xD83D_",
        "_xDE00_ literal _x005F_x000D_ overlap _x005F_x005F_ cross _x005F_",
        "x000D_",
    ])
    _set_inline(cells["A4"], ["a_x00", "0D_b"])
    for child in list(cells["A5"]):
        cells["A5"].remove(child)
    cells["A5"].set("t", "s")
    ET.SubElement(cells["A5"], f"{NS}v").text = "1"
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _set_inline(cell, parts):
    for child in list(cell):
        cell.remove(child)
    cell.set("t", "inlineStr")
    inline = ET.SubElement(cell, f"{NS}is")
    for part in parts:
        run = ET.SubElement(inline, f"{NS}r")
        ET.SubElement(run, f"{NS}t").text = part


def _shared_strings_xml():
    root = ET.Element(f"{NS}sst", {"count": "2", "uniqueCount": "2"})
    item = ET.SubElement(root, f"{NS}si")
    for part in [
        "shared_x000D_line emoji _xD83D_",
        "_xDE00_ literal _x005F_x000D_ overlap _x005F_x005F_ cross _x005F_",
        "x000D_",
    ]:
        run = ET.SubElement(item, f"{NS}r")
        ET.SubElement(run, f"{NS}t").text = part
    fragmented = ET.SubElement(root, f"{NS}si")
    for part in ["a_x00", "0D_b"]:
        run = ET.SubElement(fragmented, f"{NS}r")
        ET.SubElement(run, f"{NS}t").text = part
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _shared_content_type(data):
    root = ET.fromstring(data)
    if not any(item.get("PartName") == "/xl/sharedStrings.xml" for item in root):
        ET.SubElement(root, f"{{{CONTENT_NS}}}Override", {
            "PartName": "/xl/sharedStrings.xml",
            "ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml",
        })
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _shared_relationship(data):
    root = ET.fromstring(data)
    shared_type = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings"
    if not any(item.get("Type") == shared_type for item in root):
        ET.SubElement(root, f"{{{REL_NS}}}Relationship", {
            "Id": "rIdSharedStrings", "Type": shared_type, "Target": "sharedStrings.xml",
        })
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


class ImportedTextEscapeTests(unittest.TestCase):
    def setUp(self):
        self.fixture = APIFixture(rich_escape_fixture)
        self.client, self.csrf, _ = self.fixture.signup(
            "text-escapes@example.test", "Text Escapes Owner"
        )
        _, self.copy = self.fixture.make_copy(self.client, self.csrf)

    def tearDown(self):
        self.fixture.close()

    def test_rich_text_import_undo_and_export_decode_ooxml_once(self):
        self.assertEqual(self.fixture.snapshot.cell(0, 0, 0)["v"], CR_TEXT)
        self.assertEqual(self.fixture.snapshot.cell(0, 1, 0)["v"], SHARED_TEXT)
        self.assertEqual(self.fixture.snapshot.cell(0, 2, 0)["v"], INLINE_TEXT)

        edited = self.client.post(f"/api/copies/{self.copy['id']}/ops", {
            "op_id": "replace-cr", "action": "edit",
            "cells": [{"sheet": 0, "r": 0, "c": 0, "v": "temporary", "expected_version": 0}],
        }, {"x-csrf-token": self.csrf})
        history = self.client.get(f"/api/copies/{self.copy['id']}/history").json()["entries"]
        original = next(entry["before"] for entry in history if entry["op_id"] == "replace-cr")
        self.assertEqual(original, CR_TEXT)
        undone = self.client.post(f"/api/copies/{self.copy['id']}/ops", {
            "op_id": "undo-replace-cr", "action": "undo",
            "cells": [{"sheet": 0, "r": 0, "c": 0, "v": original, "expected_version": 1}],
        }, {"x-csrf-token": self.csrf})
        exported = self.client.get(f"/api/copies/{self.copy['id']}/export.xlsm")

        self.assertEqual(edited.status_code, 200, edited.body)
        self.assertEqual(undone.status_code, 200, undone.body)
        self.assertEqual(exported.status_code, 200, exported.body)
        with zipfile.ZipFile(io.BytesIO(exported.body)) as archive:
            sheet_xml = archive.read("xl/worksheets/sheet1.xml")
        self.assertIn(b"a_x000D_b", sheet_xml)
        self.assertNotIn(b"a_x005F_x000D_b", sheet_xml)
        workbook = load_workbook(io.BytesIO(exported.body), keep_vba=True)
        self.addCleanup(workbook.close)
        self.assertEqual(unescape(workbook["Escapes"]["A1"].value), CR_TEXT)

    def test_escape_fragments_do_not_cross_text_element_boundaries(self):
        self.assertEqual(self.fixture.snapshot.cell(0, 3, 0)["v"], FRAGMENT_TEXT)
        self.assertEqual(self.fixture.snapshot.cell(0, 4, 0)["v"], FRAGMENT_TEXT)


if __name__ == "__main__":
    unittest.main()
