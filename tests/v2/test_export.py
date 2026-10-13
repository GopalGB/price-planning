import io
import re
import threading
import unittest
from unittest import mock
import zipfile
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
from openpyxl.styles import Protection

from ops.fixtures import make_semantic
from tests.v2.api_support import APIFixture
from tests.v2.test_template_importer import rewrite_member


NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


def _spill_fixture(path, child_value):
    make_semantic(path)
    workbook = load_workbook(path)
    inputs = workbook["Inputs"]
    inputs["D2"] = "=SEQUENCE(2)"
    inputs["D3"] = child_value
    inputs["D3"].protection = Protection(locked=False)
    workbook.save(path)


def spill_fixture(path):
    _spill_fixture(path, None)


def occupied_spill_fixture(path):
    _spill_fixture(path, 0)


def xml_cell(data, reference):
    match = re.search(fr'<c[^>]*\br="{reference}"[^>]*(?:/>|>.*?</c>)'.encode(), data, re.DOTALL)
    if match is None:
        raise AssertionError(f"Missing {reference} from worksheet XML")
    return match.group(0)


class ExportContractTests(unittest.TestCase):
    def setUp(self):
        self.fixture = APIFixture()
        self.client, self.csrf, _ = self.fixture.signup("export@example.test", "Export Owner")
        _, self.copy = self.fixture.make_copy(self.client, self.csrf)

    def tearDown(self):
        self.fixture.close()

    def test_public_export_preserves_package_and_patches_only_inputs_plus_calc_mode(self):
        """Rebuilding the workbook or dropping package members must fail."""
        from pp.export.xlsm import export_xlsm

        rewrite_member(
            self.fixture.source, "xl/workbook.xml",
            lambda data: re.sub(br'fullCalcOnLoad="[^"]*"', b'fullCalcOnLoad="0"', data),
        )
        output = self.fixture.root / "public-export.xlsm"
        export_xlsm(self.fixture.source, [(0, 1, 1, 12)], output)

        with zipfile.ZipFile(self.fixture.source) as original, zipfile.ZipFile(output) as exported:
            self.assertEqual(set(original.namelist()), set(exported.namelist()))
            changed = {name for name in original.namelist() if original.read(name) != exported.read(name)}
            self.assertEqual(changed, {"xl/workbook.xml", "xl/worksheets/sheet1.xml"})
            sheet = ET.fromstring(exported.read("xl/worksheets/sheet1.xml"))
            cell = sheet.find(f".//{NS}c[@r='B2']")
            self.assertEqual(cell.findtext(f"{NS}v"), "12")
            workbook = ET.fromstring(exported.read("xl/workbook.xml"))
            calc = workbook.find(f"{NS}calcPr")
            self.assertEqual(calc.get("fullCalcOnLoad"), "1")

    def test_calc_property_byte_patch_preserves_markup_compatibility_namespaces(self):
        """Changing fullCalcOnLoad must not reserialize workbook namespace prefixes."""
        from pp.export.xlsm import export_xlsm

        data = (
            b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            b'<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            b'xmlns:mc="urn:mc" xmlns:x15="urn:x15" xmlns:xr="urn:xr" '
            b'mc:Ignorable="x15 xr"><sheets/>'
            b'<calcPr calcId="191029" calcMode="manual" forceFullCalc="0" fullCalcOnLoad="0"/>'
            b'<mc:AlternateContent><mc:Choice Requires="x15"><x15:feature xr:uid="42"/>'
            b'</mc:Choice></mc:AlternateContent></workbook>'
        )
        rewrite_member(self.fixture.source, "xl/workbook.xml", lambda _: data)
        output = self.fixture.root / "prefixed-calc.xlsm"
        export_xlsm(self.fixture.source, [], output)

        with zipfile.ZipFile(output) as archive:
            self.assertEqual(
                archive.read("xl/workbook.xml"),
                data.replace(b'fullCalcOnLoad="0"', b'fullCalcOnLoad="1"'),
            )

    def test_missing_calc_property_is_inserted_after_defined_names(self):
        """A new calcPr must use schema order without reserializing the workbook."""
        from pp.export.xlsm import export_xlsm

        data = (
            b'<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            b'<sheets/><definedNames><definedName name="x">Sheet1!A1</definedName></definedNames>'
            b'<oleSize ref="A1"/></workbook>'
        )
        rewrite_member(self.fixture.source, "xl/workbook.xml", lambda _: data)
        output = self.fixture.root / "inserted-calc.xlsm"
        export_xlsm(self.fixture.source, [], output)

        with zipfile.ZipFile(output) as archive:
            self.assertEqual(
                archive.read("xl/workbook.xml"),
                data.replace(b'</definedNames>', b'</definedNames><calcPr fullCalcOnLoad="1"/>'),
            )

    def test_exported_inline_text_preserves_boundary_whitespace(self):
        """Excel must retain leading and trailing spaces in exported text inputs."""
        from pp.export.xlsm import export_xlsm

        output = self.fixture.root / "spaced-export.xlsm"
        export_xlsm(self.fixture.source, [(0, 1, 1, "  leading and trailing  ")], output)

        with zipfile.ZipFile(output) as archive:
            sheet = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        text = sheet.find(f".//{NS}c[@r='B2']/{NS}is/{NS}t")
        self.assertEqual(text.text, "  leading and trailing  ")
        self.assertEqual(text.get("{http://www.w3.org/XML/1998/namespace}space"), "preserve")

    def test_export_archive_work_runs_off_the_asgi_event_loop(self):
        """Large archive rewriting must not block edits and WebSocket traffic."""
        from pp.export.transfer import export_transfer

        threads = []

        def observed(*args):
            threads.append(threading.get_ident())
            return export_transfer(*args)

        with mock.patch("pp.server.export_transfer", side_effect=observed):
            response = self.client.get(f"/api/copies/{self.copy['id']}/export.xlsm")

        self.assertEqual(response.status_code, 200)
        self.assertNotEqual(threads, [threading.get_ident()])

    def test_server_uses_unique_outputs_and_removes_them_after_each_response(self):
        """Downloads must not collide or leave copy data behind on disk."""
        from pp.export.transfer import export_transfer

        outputs = []

        def observed(*args):
            outputs.append(args[0].out_path)
            return export_transfer(*args)

        with mock.patch("pp.server.export_transfer", side_effect=observed):
            first = self.client.get(f"/api/copies/{self.copy['id']}/export.xlsm")
            second = self.client.get(f"/api/copies/{self.copy['id']}/export.xlsm")

        self.assertEqual((first.status_code, second.status_code), (200, 200))
        self.assertEqual(len(set(outputs)), 2)
        self.assertTrue(all(not output.exists() for output in outputs))

    def test_server_removes_reserved_output_when_export_fails(self):
        """An exporter failure must not retain a partial response file."""
        outputs = []

        def broken_export(request):
            outputs.append(request.out_path)
            request.out_path.write_bytes(b"partial")
            raise RuntimeError("synthetic export failure")

        with mock.patch("pp.server.export_transfer", side_effect=broken_export):
            response = self.client.get(f"/api/copies/{self.copy['id']}/export.xlsm")

        self.assertEqual(response.status_code, 500)
        self.assertEqual(len(outputs), 1)
        self.assertFalse(outputs[0].exists())

    def test_export_rejects_forged_locked_override_without_output(self):
        """A legacy or forged non-input override must never reach the exporter."""
        with self.fixture.database.transaction() as connection:
            connection.execute(
                "INSERT INTO pp_overrides(copy_id,sheet,row,col,value_json,version,updated_by,updated_at) "
                "VALUES(?,?,?,?,?,?,?,?)",
                (self.copy["id"], 0, 1, 0, '"forged"', 1, "owner", "2026-10-11T00:00:00Z"),
            )

        response = self.client.get(f"/api/copies/{self.copy['id']}/export.xlsm")

        self.assertEqual(response.status_code, 422, response.body)
        self.assertEqual(response.json()["error"], "locked")
        self.assertFalse((self.fixture.root / "data" / "exports").exists())

    def test_export_keeps_template_input_override_even_with_current_result(self):
        """A result row must not invalidate an input defined by the immutable template."""
        with self.fixture.database.transaction() as connection:
            connection.execute(
                "INSERT INTO pp_overrides(copy_id,sheet,row,col,value_json,version,updated_by,updated_at) "
                "VALUES(?,?,?,?,?,?,?,?)",
                (self.copy["id"], 0, 1, 1, "12", 1, "owner", "2026-10-11T00:00:00Z"),
            )
            connection.execute(
                "INSERT INTO pp_results(copy_id,sheet,row,col,value_json,calculated_revision) "
                "VALUES(?,?,?,?,?,?)",
                (self.copy["id"], 0, 1, 1, "99", 1),
            )

        response = self.client.get(f"/api/copies/{self.copy['id']}/export.xlsm")

        self.assertEqual(response.status_code, 200, response.body)
        with zipfile.ZipFile(io.BytesIO(response.body)) as archive:
            cell = ET.fromstring(archive.read("xl/worksheets/sheet1.xml")).find(f".//{NS}c[@r='B2']")
        self.assertEqual(cell.findtext(f"{NS}v"), "12")


class SpillChildExportContractTests(unittest.TestCase):
    def setUp(self):
        self.fixture = APIFixture(spill_fixture)
        self.client, self.csrf, _ = self.fixture.signup("spill@example.test", "Spill Owner")
        _, self.copy = self.fixture.make_copy(self.client, self.csrf)

    def tearDown(self):
        self.fixture.close()

    def test_spill_child_without_formula_is_locked_and_never_exported_as_an_override(self):
        """An unlocked dynamic-array child without an f node must remain a calculated cell."""
        snapshot = self.fixture.snapshot
        self.assertEqual(snapshot.cell(0, 1, 3).get("formula"), "=SEQUENCE(2)")
        self.assertNotIn("0:2:3", snapshot.input_mask)
        self.assertFalse(snapshot.is_input(0, 2, 3))
        self.assertEqual(snapshot.cell(0, 2, 3)["k"], "f")
        with zipfile.ZipFile(self.fixture.source) as archive:
            original_child = xml_cell(archive.read("xl/worksheets/sheet1.xml"), "D3")
        self.assertNotIn(b"<f", original_child)
        self.assertNotIn(b"<v", original_child)
        self.assertRegex(original_child, br'\bs="[0-9]+"')

        with self.fixture.database.transaction() as connection:
            connection.execute(
                "DELETE FROM pp_results WHERE copy_id=? AND sheet=? AND row=? AND col=?",
                (self.copy["id"], 0, 2, 3),
            )
        rejected = self.client.post(f"/api/copies/{self.copy['id']}/ops", {
            "op_id": "spill-child-locked", "action": "edit",
            "cells": [
                {"sheet": 0, "r": 1, "c": 1, "v": 12, "expected_version": 0},
                {"sheet": 0, "r": 2, "c": 3, "v": 99, "expected_version": 0},
            ],
        }, {"x-csrf-token": self.csrf})

        self.assertEqual(rejected.status_code, 422, rejected.body)
        self.assertEqual(rejected.json()["error"], "locked")
        self.assertEqual(self.fixture.count("pp_overrides"), 0)
        with self.fixture.database.connect() as connection:
            revision = connection.execute(
                "SELECT input_revision FROM pp_copies WHERE id=?", (self.copy["id"],)
            ).fetchone()["input_revision"]
        self.assertEqual(revision, 0)

        accepted = self.client.post(f"/api/copies/{self.copy['id']}/ops", {
            "op_id": "ordinary-input", "action": "edit",
            "cells": [{"sheet": 0, "r": 1, "c": 1, "v": 12, "expected_version": 0}],
        }, {"x-csrf-token": self.csrf})
        self.assertEqual(accepted.status_code, 200)
        exported = self.client.get(f"/api/copies/{self.copy['id']}/export.xlsm")
        self.assertEqual(exported.status_code, 200)
        with zipfile.ZipFile(io.BytesIO(exported.body)) as archive:
            self.assertEqual(xml_cell(archive.read("xl/worksheets/sheet1.xml"), "D3"), original_child)


class OccupiedSpillTargetContractTests(unittest.TestCase):
    def setUp(self):
        self.fixture = APIFixture(occupied_spill_fixture)
        self.client, self.csrf, _ = self.fixture.signup("blocker@example.test", "Blocker Owner")
        _, self.copy = self.fixture.make_copy(self.client, self.csrf)

    def tearDown(self):
        self.fixture.close()

    def test_occupied_spill_target_remains_an_editable_input(self):
        """A real value blocking SEQUENCE must retain its template input classification."""
        snapshot = self.fixture.snapshot
        self.assertIn("0:2:3", snapshot.input_mask)
        self.assertTrue(snapshot.is_input(0, 2, 3))
        self.assertEqual(snapshot.cell(0, 2, 3)["k"], "i")

        response = self.client.post(f"/api/copies/{self.copy['id']}/ops", {
            "op_id": "edit-spill-blocker", "action": "edit",
            "cells": [{"sheet": 0, "r": 2, "c": 3, "v": 7, "expected_version": 0}],
        }, {"x-csrf-token": self.csrf})

        self.assertEqual(response.status_code, 200, response.body)


if __name__ == "__main__":
    unittest.main()
