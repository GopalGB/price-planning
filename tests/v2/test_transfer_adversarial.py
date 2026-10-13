"""Adversarial synthetic tests for workbook transfer package semantics."""

from __future__ import annotations

import io
import os
from pathlib import Path
import re
import resource
import shutil
import time
import unittest
from unittest import mock
import zipfile
import xml.etree.ElementTree as ET

import defusedxml.ElementTree as SafeET

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Protection
from openpyxl.worksheet.table import Table
from openpyxl.workbook.defined_name import DefinedName

from pp.export.transfer import TransferRequest, export_transfer, validate_package
from tests.v2.api_support import APIFixture


MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
CUSTOM = "http://schemas.openxmlformats.org/officeDocument/2006/custom-properties"
VT = "http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"


def _rewrite(path, replacements, additions=()):
    temporary = path.with_suffix(".rewrite")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(temporary, "w") as target:
        for info in source.infolist():
            if info.filename in replacements:
                target.writestr(info, replacements[info.filename])
            else:
                with source.open(info) as reader, target.open(info, "w") as writer:
                    shutil.copyfileobj(reader, writer, 1 << 20)
        for name, data in additions:
            target.writestr(name, data)
    temporary.replace(path)


def _formula_fixture(path):
    workbook = Workbook()
    keep_a = workbook.active
    keep_a.title = "KeepA"
    keep_b = workbook.create_sheet("KeepB")
    drop = workbook.create_sheet("Drop")
    drop["A1"] = 7
    unicode_drop = workbook.create_sheet("Résumé")
    unicode_drop["A1"] = 8
    for sheet in (keep_a, keep_b):
        sheet["A1"] = "=1+1"
        sheet["A2"] = "=1+1"
    keep_a["B1"] = "=Drop!A1"
    keep_a["B2"] = 0
    keep_a["B2"].number_format = "0.00%"
    keep_a["C1"] = "=Alias*2"
    keep_a["D1"] = "=Rate*2"
    keep_a["E1"] = '=\"drop!A1 Alias\"'
    keep_a["F1"] = "=SUM(keepa:drop!A1)"
    keep_a["G1"] = "=Résumé!A1*2"
    workbook.defined_names.add(DefinedName("Rate", attr_text="1"))
    workbook.defined_names.add(DefinedName("Rate", attr_text="'Drop'!$A$1", localSheetId=2))
    workbook.defined_names.add(DefinedName("Base", attr_text="'drop'!$A$1"))
    workbook.defined_names.add(DefinedName("Alias", attr_text="Base"))
    workbook.save(path)

    with zipfile.ZipFile(path) as archive:
        keep_a_xml = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        keep_b_xml = ET.fromstring(archive.read("xl/worksheets/sheet2.xml"))
        workbook_xml = archive.read("xl/workbook.xml")

    def formula(root, reference, text, **attributes):
        cell = root.find(f".//{{{MAIN}}}c[@r='{reference}']")
        node = cell.find(f"{{{MAIN}}}f")
        if node is None:
            node = ET.Element(f"{{{MAIN}}}f")
            cell.insert(0, node)
        node.attrib.clear()
        node.attrib.update(attributes)
        node.text = text

    formula(keep_a_xml, "A1", "Drop!A1", t="shared", ref="A1:A2", si="0")
    formula(keep_a_xml, "A2", None, t="shared", si="0")
    formula(keep_a_xml, "B1", "Drop!A1:A2", t="array", ref="B1:B2")
    follower = keep_a_xml.find(f".//{{{MAIN}}}c[@r='B2']")
    for node in list(follower):
        follower.remove(node)
    formula(keep_b_xml, "A1", "1+1", t="shared", ref="A1:A2", si="0")
    formula(keep_b_xml, "A2", None, t="shared", si="0")
    workbook_xml = workbook_xml.replace(
        b"<definedNames>", b'<definedNames><definedName name="Rate">1</definedName>', 1
    )
    _rewrite(path, {
        "xl/worksheets/sheet1.xml": ET.tostring(keep_a_xml, encoding="utf-8", xml_declaration=True),
        "xl/worksheets/sheet2.xml": ET.tostring(keep_b_xml, encoding="utf-8", xml_declaration=True),
        "xl/workbook.xml": workbook_xml,
    })


def _cell(root, reference):
    return root.find(f".//{{{MAIN}}}c[@r='{reference}']")


class FormulaPreservationTests(unittest.TestCase):
    def setUp(self):
        self.root = Path(self.id().replace(".", "-"))
        self.addCleanup(lambda: shutil.rmtree(self.root, ignore_errors=True))
        self.root.mkdir()
        self.source = self.root / "formula-source.xlsx"
        self.output = self.root / "formula-output.xlsx"
        _formula_fixture(self.source)

    def test_shared_array_and_scoped_transitive_names_freeze_only_affected_cells(self):
        results = {
            (0, 0, 0): 11,
            (0, 1, 0): 12,
            (0, 0, 1): 21,
            (0, 1, 1): 22,
            (0, 0, 2): 31,
            (0, 0, 5): 61,
            (0, 0, 6): 16,
        }
        with zipfile.ZipFile(self.source) as archive:
            original = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
            original_style = _cell(original, "B2").get("s")

        export_transfer(
            TransferRequest(
                template_path=self.source, cells=[], out_path=self.output,
                template_hash="template-hash", copy_id="copy-id", input_revision=3,
                fmt="xlsx", selected={"KeepA", "KeepB"}, results=results,
            )
        )

        with zipfile.ZipFile(self.output) as archive:
            keep_a = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
            keep_b = ET.fromstring(archive.read("xl/worksheets/sheet2.xml"))
            workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        for reference, value in {
            "A1": "11",
            "A2": "12",
            "B1": "21",
            "B2": "22",
            "C1": "31",
            "F1": "61",
            "G1": "16",
        }.items():
            cell = _cell(keep_a, reference)
            self.assertIsNone(cell.find(f"{{{MAIN}}}f"), reference)
            self.assertEqual(cell.findtext(f"{{{MAIN}}}v"), value, reference)
        self.assertEqual(_cell(keep_a, "B2").get("s"), original_style)
        self.assertIsNotNone(_cell(keep_a, "D1").find(f"{{{MAIN}}}f"))
        self.assertIsNotNone(_cell(keep_a, "E1").find(f"{{{MAIN}}}f"))
        self.assertIsNotNone(_cell(keep_b, "A1").find(f"{{{MAIN}}}f"))
        self.assertIsNotNone(_cell(keep_b, "A2").find(f"{{{MAIN}}}f"))
        names = [(item.get("name"), item.get("localSheetId"))
                 for item in workbook.iter() if item.tag.rsplit("}", 1)[-1] == "definedName"]
        self.assertIn(("Rate", None), names)
        self.assertNotIn(("Rate", "2"), names)
        self.assertNotIn(("Base", None), names)
        self.assertNotIn(("Alias", None), names)

    def test_array_range_refuses_partial_current_results(self):
        with self.assertRaisesRegex(ValueError, "KeepA!B2"):
            export_transfer(
                TransferRequest(
                    template_path=self.source, cells=[], out_path=self.output,
                    template_hash="template-hash", copy_id="copy-id", input_revision=3,
                    fmt="xlsx", selected={"KeepA", "KeepB"},
                    results={(0, 0, 0): 11, (0, 1, 0): 12, (0, 0, 1): 21,
                             (0, 0, 2): 31, (0, 0, 3): 41, (0, 0, 4): "literal",
                             (0, 0, 5): 61, (0, 0, 6): 16,
                             (1, 0, 0): 2, (1, 1, 0): 2},
                )
            )

    def test_single_cell_array_freezes_without_dynamic_array_metadata(self):
        with zipfile.ZipFile(self.source) as archive:
            root = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        cell = _cell(root, "B1")
        cell.set("cm", "1")
        cell.find(f"{{{MAIN}}}f").set("ref", "B1")
        _rewrite(self.source, {
            "xl/worksheets/sheet1.xml": ET.tostring(
                root, encoding="utf-8", xml_declaration=True,
            ),
        })
        results = {
            (0, 0, 0): 11, (0, 1, 0): 12, (0, 0, 1): 21,
            (0, 0, 2): 31, (0, 0, 5): 61, (0, 0, 6): 16,
        }

        export_transfer(TransferRequest(
            template_path=self.source, cells=[], out_path=self.output,
            template_hash="template-hash", copy_id="copy-id", input_revision=3,
            fmt="xlsx", selected={"KeepA", "KeepB"}, results=results,
        ))

        with zipfile.ZipFile(self.output) as archive:
            output = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        frozen = _cell(output, "B1")
        self.assertIsNone(frozen.find(f"{{{MAIN}}}f"))
        self.assertEqual(frozen.findtext(f"{{{MAIN}}}v"), "21")
        self.assertIsNone(frozen.get("cm"))


class IndirectReferenceTests(unittest.TestCase):
    def test_literal_removed_sheet_and_name_references_are_frozen(self):
        root = Path(self.id().replace(".", "-"))
        self.addCleanup(lambda: shutil.rmtree(root, ignore_errors=True))
        root.mkdir()
        source, output = root / "source.xlsx", root / "output.xlsx"
        workbook = Workbook()
        keep = workbook.active
        keep.title = "Keep"
        keep["A1"] = '=INDIRECT("Drop!A1")'
        keep["B1"] = '=INDIRECT("Gone")'
        keep["C1"] = '=INDIRECT("Keep!D1")'
        drop = workbook.create_sheet("Drop")
        drop["A1"] = 7
        workbook.defined_names.add(DefinedName("Gone", attr_text="'Drop'!$A$1"))
        workbook.save(source)

        export_transfer(TransferRequest(
            source, [], output, "hash", "copy", 0,
            selected={"Keep"}, results={(0, 0, 0): 7, (0, 0, 1): 7},
        ))

        loaded = load_workbook(output, data_only=False)
        self.addCleanup(loaded.close)
        self.assertEqual(loaded["Keep"]["A1"].value, 7)
        self.assertEqual(loaded["Keep"]["B1"].value, 7)
        self.assertEqual(loaded["Keep"]["C1"].value, '=INDIRECT("Keep!D1")')


class FormulaSerializationTests(unittest.TestCase):
    def test_literal_ooxml_escape_prefix_and_carriage_return_stay_distinct(self):
        root = Path(self.id().replace(".", "-"))
        self.addCleanup(lambda: shutil.rmtree(root, ignore_errors=True))
        root.mkdir()
        source, output = root / "source.xlsx", root / "output.xlsx"
        workbook = Workbook()
        keep = workbook.active
        keep.title = "Keep"
        keep["A1"] = "=Drop!A1"
        workbook.create_sheet("Drop")
        workbook.save(source)

        export_transfer(TransferRequest(
            source, [], output, "hash", "copy", 0, selected={"Keep"},
            results={(0, 0, 0): "literal_x000D_text\rnext"},
        ))

        with zipfile.ZipFile(output) as archive:
            sheet = archive.read("xl/worksheets/sheet1.xml")
        self.assertIn(b"literal_x005F_x000D_text_x000D_next", sheet)


class DefinedNameSerializationTests(unittest.TestCase):
    def test_self_closing_name_does_not_absorb_the_next_definition(self):
        root = Path(self.id().replace(".", "-"))
        self.addCleanup(lambda: shutil.rmtree(root, ignore_errors=True))
        root.mkdir()
        source, output = root / "source.xlsx", root / "output.xlsx"
        workbook = Workbook()
        workbook.active.title = "Keep"
        workbook.create_sheet("Drop")
        workbook.defined_names.add(DefinedName("Gone", attr_text="'Drop'!$A$1"))
        workbook.defined_names.add(DefinedName("KeepName", attr_text="'Keep'!$A$1"))
        workbook.save(source)
        with zipfile.ZipFile(source) as archive:
            xml = archive.read("xl/workbook.xml")
        xml = xml.replace(b"<definedNames>", b'<definedNames><definedName name="Empty"/>')
        _rewrite(source, {"xl/workbook.xml": xml})

        export_transfer(TransferRequest(
            source, [], output, "hash", "copy", 0, selected={"Keep"}, results={},
        ))

        with zipfile.ZipFile(output) as archive:
            names = archive.read("xl/workbook.xml")
        self.assertIn(b'<definedName name="Empty"/>', names)
        self.assertNotIn(b'name="Gone"', names)
        self.assertIn(b'name="KeepName"', names)


def _worksheet_reference_fixture(path):
    workbook = Workbook()
    workbook.active.title = "Keep"
    workbook.create_sheet("Drop")
    workbook.defined_names.add(DefinedName("Gone", attr_text="'Drop'!$A$1"))
    workbook.save(path)
    with zipfile.ZipFile(path) as archive:
        sheet = archive.read("xl/worksheets/sheet1.xml")
        styles = archive.read("xl/styles.xml")
    rules = (
        b'<conditionalFormatting sqref="A1"><cfRule type="expression" priority="1">'
        b'<formula>Drop!A1&gt;0</formula></cfRule></conditionalFormatting>'
        b'<dataValidations count="2"><dataValidation type="list" sqref="B1">'
        b'<formula1>Drop!$A$1</formula1></dataValidation>'
        b'<dataValidation type="list" sqref="B2"><formula1>Keep!$A$1</formula1>'
        b'</dataValidation></dataValidations>'
        b'<hyperlinks><hyperlink ref="C1" location="Gone"/>'
        b'<hyperlink ref="C2" location="Keep!A1" display="keep"/></hyperlinks>'
    )
    vendor = (
        b'<extLst><ext uri="urn:vendor"><v:dataValidation xmlns:v="urn:vendor">'
        b'<v:formula1>Drop!A1</v:formula1></v:dataValidation></ext></extLst>'
    )
    sheet = sheet.replace(b"<pageMargins", rules + b"<pageMargins", 1)
    sheet = sheet.replace(b"</worksheet>", vendor + b"</worksheet>", 1)
    _rewrite(path, {"xl/worksheets/sheet1.xml": sheet})
    return styles, vendor


def _chartsheet_fixture(path):
    workbook = Workbook()
    keep = workbook.active
    keep.title = "Keep"
    keep["A1"], keep["A2"] = 1, 2
    chart = BarChart()
    chart.add_data(Reference(keep, min_col=1, min_row=1, max_row=2))
    workbook.create_chartsheet("Chart").add_chart(chart)
    workbook.create_sheet("Drop")
    workbook.save(path)


class WorksheetReferenceCleanupTests(unittest.TestCase):
    def test_subset_drops_only_rules_and_links_to_removed_dependencies(self):
        root = Path(self.id().replace(".", "-"))
        self.addCleanup(lambda: shutil.rmtree(root, ignore_errors=True))
        root.mkdir()
        source, output = root / "source.xlsx", root / "output.xlsx"
        original_styles, vendor = _worksheet_reference_fixture(source)

        export_transfer(TransferRequest(
            source, [], output, "hash", "copy", 0, selected={"Keep"}, results={},
        ))

        with zipfile.ZipFile(output) as archive:
            cleaned = archive.read("xl/worksheets/sheet1.xml")
            self.assertEqual(archive.read("xl/styles.xml"), original_styles)
        worksheet = ET.fromstring(cleaned)
        validations = worksheet.findall(f".//{{{MAIN}}}dataValidation")
        self.assertEqual(len(validations), 1)
        self.assertEqual(validations[0].findtext(f"{{{MAIN}}}formula1"), "Keep!$A$1")
        self.assertIsNone(worksheet.find(f"{{{MAIN}}}dataValidations").get("count"))
        self.assertIsNone(worksheet.find(f"{{{MAIN}}}conditionalFormatting"))
        hyperlinks = worksheet.findall(f".//{{{MAIN}}}hyperlink")
        self.assertEqual([(item.get("ref"), item.get("location")) for item in hyperlinks], [
            ("C2", "Keep!A1"),
        ])
        self.assertIn(vendor, cleaned)

    def test_reference_cleanup_does_not_treat_a_chartsheet_as_a_worksheet(self):
        root = Path(self.id().replace(".", "-"))
        self.addCleanup(lambda: shutil.rmtree(root, ignore_errors=True))
        root.mkdir()
        source, output = root / "source.xlsx", root / "output.xlsx"
        _chartsheet_fixture(source)
        with zipfile.ZipFile(source) as archive:
            chart_sheet = archive.read("xl/chartsheets/sheet1.xml")

        export_transfer(TransferRequest(
            source, [], output, "hash", "copy", 0,
            selected={"Keep", "Chart"}, results={},
        ))

        with zipfile.ZipFile(output) as archive:
            self.assertEqual(archive.read("xl/chartsheets/sheet1.xml"), chart_sheet)
        loaded = load_workbook(output, read_only=True)
        self.addCleanup(loaded.close)
        self.assertEqual(loaded.sheetnames, ["Keep", "Chart"])


def _pivot_source_fixture(path):
    workbook = Workbook()
    keep = workbook.active
    keep.title = "Keep"
    keep["A1"] = 1
    keep["A2"] = 2
    chart = BarChart()
    chart.add_data(Reference(keep, min_col=1, min_row=1, max_row=2))
    keep.add_chart(chart, "H2")
    drop = workbook.create_sheet("Drop")
    drop["A1"] = 9
    drop["A2"] = 10
    drop["A1"] = "Cost"
    drop["A2"] = 9
    drop["A3"] = 10
    drop.add_table(Table(displayName="DropTable", ref="A1:A3"))
    workbook.defined_names.add(DefinedName("DropRange", attr_text="'Drop'!$A$1:$A$2"))
    workbook.defined_names.add(
        DefinedName("DropAlias", attr_text="DropRange", localSheetId=0)
    )
    removed_chart = BarChart()
    removed_chart.add_data(Reference(drop, min_col=1, min_row=1, max_row=2))
    keep.add_chart(removed_chart, "H18")
    workbook.save(path)
    with zipfile.ZipFile(path) as archive:
        workbook_xml = archive.read("xl/workbook.xml")
        workbook_rels = archive.read("xl/_rels/workbook.xml.rels")
        sheet_xml = archive.read("xl/worksheets/sheet1.xml")
        sheet_rels = archive.read("xl/worksheets/_rels/sheet1.xml.rels")
        drawing = archive.read("xl/drawings/drawing1.xml")
        drawing_rels = archive.read("xl/drawings/_rels/drawing1.xml.rels")
        chart2 = archive.read("xl/charts/chart2.xml")
        content_types = archive.read("[Content_Types].xml")
    workbook_xml = workbook_xml.replace(
        b"</workbook>",
        b'<pivotCaches><pivotCache cacheId="1" r:id="rIdBadCache"/></pivotCaches>'
        b'<extLst><ext uri="{removed-controls}"><x14:slicerCaches '
        b'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
        b'<x14:slicerCache r:id="rIdBadSlicerCache"/></x14:slicerCaches>'
        b'<x15:timelineCacheRefs xmlns:x15="http://schemas.microsoft.com/office/'
        b'spreadsheetml/2010/11/main"><x15:timelineCacheRef '
        b'r:id="rIdBadTimelineCache"/></x15:timelineCacheRefs></ext></extLst></workbook>',
    )
    workbook_rels = workbook_rels.replace(
        b"</Relationships>",
        b'<Relationship Id="rIdBadCache" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/pivotCacheDefinition" Target="pivotCache/pivotCacheDefinition1.xml"/></Relationships>',
    )
    workbook_rels = workbook_rels.replace(
        b"</Relationships>",
        b'<Relationship Id="rIdBadSlicerCache" Type="http://schemas.microsoft.com/'
        b'office/2007/relationships/slicerCache" Target="slicerCaches/slicerCache1.xml"/>'
        b'<Relationship Id="rIdBadTimelineCache" Type="http://schemas.microsoft.com/'
        b'office/2010/relationships/TimelineCache" '
        b'Target="timelineCaches/timelineCache1.xml"/></Relationships>',
    )
    sheet_xml = sheet_xml.replace(
        b"</worksheet>",
        b'<pivotTableParts count="1"><pivotTablePart r:id="rIdBadPivot"/></pivotTableParts>'
        b'<extLst><ext uri="{slicer-view}"><x14:slicerList r:id="rIdSlicerView"/>'
        b'</ext><ext uri="{timeline-view}"><x15:timelines xmlns:x15="http://schemas.'
        b'microsoft.com/office/spreadsheetml/2010/11/main" '
        b'r:id="rIdTimelineView"/></ext></extLst></worksheet>',
    )
    if b"xmlns:r=" not in sheet_xml.partition(b">")[0]:
        sheet_xml = sheet_xml.replace(
            b"<worksheet ",
            f'<worksheet xmlns:r="{REL}" '.encode(),
            1,
        )
    if b"xmlns:x14=" not in sheet_xml.partition(b">")[0]:
        sheet_xml = sheet_xml.replace(
            b"<worksheet ",
            b'<worksheet xmlns:x14="http://schemas.microsoft.com/office/'
            b'spreadsheetml/2009/9/main" ',
            1,
        )
    sheet_rels = sheet_rels.replace(
        b"</Relationships>",
        b'<Relationship Id="rIdBadPivot" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/pivotTable" Target="../pivotTables/pivotTable1.xml"/></Relationships>',
    )
    sheet_rels = sheet_rels.replace(
        b"</Relationships>",
        b'<Relationship Id="rIdSlicerView" Type="http://schemas.microsoft.com/office/'
        b'2007/relationships/slicer" Target="../slicers/slicer1.xml"/></Relationships>',
    )
    sheet_rels = sheet_rels.replace(
        b"</Relationships>",
        b'<Relationship Id="rIdTimelineView" Type="http://schemas.microsoft.com/office/'
        b'2010/relationships/Timeline" Target="../timelines/timeline1.xml"/>'
        b'</Relationships>',
    )
    bad_anchor = (
        b'<xdr:twoCellAnchor xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/'
        b'spreadsheetDrawing"><xdr:from><xdr:col>4</xdr:col><xdr:colOff>0</xdr:colOff>'
        b'<xdr:row>4</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from><xdr:to><xdr:col>5</xdr:col>'
        b'<xdr:colOff>0</xdr:colOff><xdr:row>5</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>'
        b'<xdr:graphicFrame><xdr:nvGraphicFramePr><xdr:cNvPr id="99" name="Bad slicer"/>'
        b'<xdr:cNvGraphicFramePr/></xdr:nvGraphicFramePr><xdr:xfrm/><a:graphic '
        b'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:graphicData '
        b'uri="http://schemas.microsoft.com/office/drawing/2010/slicer"><sle:slicer '
        b'xmlns:sle="http://schemas.microsoft.com/office/drawing/2010/slicer" '
        b'name="BadSlicer"/></a:graphicData></a:graphic>'
        b'</xdr:graphicFrame><xdr:clientData/></xdr:twoCellAnchor>'
    )
    timeline_anchor = bad_anchor.replace(
        b'Bad slicer', b'Bad timeline',
    ).replace(
        b'http://schemas.microsoft.com/office/drawing/2010/slicer',
        b'http://schemas.microsoft.com/office/drawing/2012/timeslicer',
    ).replace(
        b'sle:slicer', b'ts:timeslicer',
    ).replace(
        b'xmlns:sle=', b'xmlns:ts=',
    ).replace(b'BadSlicer', b'BadTimeline')
    vendor = (
        b'<extLst><ext uri="urn:vendor"><r:record xmlns:r="urn:vendor" '
        b'r:id="rIdBadSlicer">retain exactly</r:record></ext></extLst>'
    )
    drawing = drawing.replace(
        b"</wsDr>", bad_anchor + timeline_anchor + vendor + b"</wsDr>",
    )
    original_formula = b"'Drop'!$A$1:$A$2"
    if original_formula not in chart2:
        raise AssertionError("chart2 source formula was not found")
    chart2 = chart2.replace(original_formula, b"DropAlias")
    drawing_rels = drawing_rels.replace(
        b"</Relationships>",
        b'<Relationship Id="rIdBadSlicer" Type="http://schemas.microsoft.com/office/2007/relationships/slicer" Target="../slicers/slicer1.xml"/></Relationships>',
    )
    content_types = content_types.replace(
        b"</Types>",
        b'<Override PartName="/xl/pivotTables/pivotTable1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.pivotTable+xml"/>'
        b'<Override PartName="/xl/pivotCache/pivotCacheDefinition1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.pivotCacheDefinition+xml"/>'
        b'<Override PartName="/xl/pivotCache/pivotCacheRecords1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.pivotCacheRecords+xml"/>'
        b'<Override PartName="/xl/slicers/slicer1.xml" ContentType="application/vnd.ms-excel.slicer+xml"/>'
        b'<Override PartName="/xl/slicerCaches/slicerCache1.xml" ContentType="application/vnd.ms-excel.slicerCache+xml"/>'
        b'<Override PartName="/xl/timelines/timeline1.xml" ContentType="application/vnd.ms-excel.Timeline+xml"/>'
        b'<Override PartName="/xl/timelineCaches/timelineCache1.xml" ContentType="application/vnd.ms-excel.TimelineCache+xml"/></Types>',
    )
    additions = (
        ("xl/pivotTables/pivotTable1.xml", b'<pivotTableDefinition xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" name="BadPivot" cacheId="1" dataCaption="Values"><location ref="A1" firstHeaderRow="1" firstDataRow="1" firstDataCol="1"/><pivotFields count="0"/></pivotTableDefinition>'),
        ("xl/pivotTables/_rels/pivotTable1.xml.rels", b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rIdCache" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/pivotCacheDefinition" Target="../pivotCache/pivotCacheDefinition1.xml"/></Relationships>'),
        ("xl/pivotCache/pivotCacheDefinition1.xml", b'<pivotCacheDefinition xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" saveData="0" recordCount="0"><cacheSource type="worksheet"><worksheetSource name="DropTable"/></cacheSource><cacheFields count="0"/></pivotCacheDefinition>'),
        ("xl/slicers/slicer1.xml", b'<x14:slicers xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"><x14:slicer name="BadSlicer" cache="BadCache"/></x14:slicers>'),
        ("xl/slicerCaches/slicerCache1.xml", b'<x14:slicerCacheDefinition xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main" name="BadCache" r:id="rIdPivotCache" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"/>'),
        ("xl/slicerCaches/_rels/slicerCache1.xml.rels", b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rIdPivotCache" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/pivotCacheDefinition" Target="../pivotCache/pivotCacheDefinition1.xml"/></Relationships>'),
        ("xl/timelines/timeline1.xml", b'<x15:timelines xmlns:x15="http://schemas.microsoft.com/office/spreadsheetml/2010/11/main"><x15:timeline name="BadTimeline" cache="BadTimelineCache" level="0" selectionLevel="0"/></x15:timelines>'),
        ("xl/timelineCaches/timelineCache1.xml", b'<x15:timelineCacheDefinition xmlns:x15="http://schemas.microsoft.com/office/spreadsheetml/2010/11/main" name="BadTimelineCache"><x15:state pivotCacheId="1"/></x15:timelineCacheDefinition>'),
        ("xl/pivotCache/_rels/pivotCacheDefinition1.xml.rels", b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rIdRecords" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/pivotCacheRecords" Target="pivotCacheRecords1.xml"/></Relationships>'),
        ("xl/pivotCache/pivotCacheRecords1.xml", b'<pivotCacheRecords xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="1"><r><x v="synthetic"/></r></pivotCacheRecords>'),
    )
    _rewrite(path, {
        "xl/workbook.xml": workbook_xml,
        "xl/_rels/workbook.xml.rels": workbook_rels,
        "xl/worksheets/sheet1.xml": sheet_xml,
        "xl/worksheets/_rels/sheet1.xml.rels": sheet_rels,
        "xl/drawings/drawing1.xml": drawing,
        "xl/drawings/_rels/drawing1.xml.rels": drawing_rels,
        "xl/charts/chart2.xml": chart2,
        "[Content_Types].xml": content_types,
    }, additions)


class ObjectDependencyTests(unittest.TestCase):
    def test_retained_sheet_loses_objects_sourced_from_removed_sheet_but_keeps_chart(self):
        root = Path(self.id().replace(".", "-"))
        self.addCleanup(lambda: shutil.rmtree(root, ignore_errors=True))
        root.mkdir()
        source, output = root / "source.xlsx", root / "output.xlsx"
        _pivot_source_fixture(source)

        export_transfer(TransferRequest(
            template_path=source, cells=[], out_path=output,
            template_hash="hash", copy_id="copy", input_revision=0,
            fmt="xlsx", selected={"Keep"}, results={},
        ))

        with zipfile.ZipFile(output) as archive:
            names = set(archive.namelist())
            sheet = archive.read("xl/worksheets/sheet1.xml")
            drawing = archive.read("xl/drawings/drawing1.xml")
            drawing_rels = archive.read("xl/drawings/_rels/drawing1.xml.rels")
        self.assertNotIn("xl/pivotTables/pivotTable1.xml", names)
        self.assertNotIn("xl/pivotCache/pivotCacheDefinition1.xml", names)
        self.assertNotIn("xl/slicers/slicer1.xml", names)
        self.assertNotIn("xl/slicerCaches/slicerCache1.xml", names)
        self.assertNotIn("xl/timelines/timeline1.xml", names)
        self.assertNotIn("xl/timelineCaches/timelineCache1.xml", names)
        self.assertNotIn("xl/pivotCache/pivotCacheRecords1.xml", names)
        self.assertNotIn(b"rIdBadPivot", sheet)
        self.assertNotIn(b'name="BadSlicer"', drawing)
        self.assertNotIn(b'name="BadTimeline"', drawing)
        self.assertNotIn(b"rIdBadSlicer", drawing_rels)
        self.assertIn(
            b'<r:record xmlns:r="urn:vendor" r:id="rIdBadSlicer">retain exactly</r:record>',
            drawing,
        )
        self.assertIn("xl/charts/chart1.xml", names)
        self.assertNotIn("xl/charts/chart2.xml", names)
        self.assertIn(b"chart", drawing)

    def test_pivot_cache_records_are_not_tree_parsed(self):
        root = Path(self.id().replace(".", "-"))
        self.addCleanup(lambda: shutil.rmtree(root, ignore_errors=True))
        root.mkdir()
        source, output = root / "source.xlsx", root / "output.xlsx"
        _pivot_source_fixture(source)
        original = SafeET.parse

        def reject_records(stream, *args, **kwargs):
            if getattr(stream, "name", "").endswith("pivotCacheRecords1.xml"):
                raise AssertionError("pivot records reached the tree parser")
            return original(stream, *args, **kwargs)

        with mock.patch("pp.export.transfer.SafeET.parse", side_effect=reject_records):
            export_transfer(TransferRequest(
                source, [], output, "hash", "copy", 0,
                selected={"Keep"}, results={},
            ))

    def test_removed_cache_wrappers_are_not_left_empty(self):
        root = Path(self.id().replace(".", "-"))
        self.addCleanup(lambda: shutil.rmtree(root, ignore_errors=True))
        root.mkdir()
        source, output = root / "source.xlsx", root / "output.xlsx"
        _pivot_source_fixture(source)

        export_transfer(TransferRequest(
            source, [], output, "hash", "copy", 0,
            selected={"Keep"}, results={},
        ))

        with zipfile.ZipFile(output) as archive:
            workbook = archive.read("xl/workbook.xml")
        self.assertNotIn(b"pivotCaches", workbook)
        self.assertNotIn(b"slicerCaches", workbook)
        self.assertNotIn(b"timelineCacheRefs", workbook)


class WorkbookViewTests(unittest.TestCase):
    def test_subset_remaps_active_and_first_visible_sheet_indexes(self):
        root = Path(self.id().replace(".", "-"))
        self.addCleanup(lambda: shutil.rmtree(root, ignore_errors=True))
        root.mkdir()
        source, output = root / "source.xlsx", root / "output.xlsx"
        workbook = Workbook()
        workbook.active.title = "Drop"
        hidden = workbook.create_sheet("Hidden")
        hidden.sheet_state = "hidden"
        workbook.create_sheet("Keep")
        workbook.save(source)
        with zipfile.ZipFile(source) as archive:
            xml = archive.read("xl/workbook.xml")
        xml = re.sub(
            rb"<workbookView\b[^>]*/>",
            b'<workbookView activeTab="0" firstSheet="0"/>',
            xml,
        )
        _rewrite(source, {"xl/workbook.xml": xml})

        export_transfer(TransferRequest(
            source, [], output, "hash", "copy", 0,
            selected={"Hidden", "Keep"}, results={},
        ))

        with zipfile.ZipFile(output) as archive:
            root_xml = ET.fromstring(archive.read("xl/workbook.xml"))
        view = next(node for node in root_xml.iter() if node.tag.endswith("workbookView"))
        self.assertEqual(view.get("activeTab"), "1")
        self.assertEqual(view.get("firstSheet"), "1")


class FormulaReferenceCostTests(unittest.TestCase):
    def test_defined_name_detection_does_not_scan_once_per_known_name(self):
        from pp.export.transfer import _definition_references

        class MembershipOnlyNames(set):
            def __iter__(self):
                raise AssertionError("known names were traversed")

        known = MembershipOnlyNames(
            {f"unused_{index}" for index in range(1_000)} | {"target"}
        )
        references = _definition_references("=Target+1", known)
        self.assertEqual(references, {"target"})


class TableDependencyTests(unittest.TestCase):
    def test_removed_table_freezes_formulas_and_removes_named_chart_but_keeps_local_table(self):
        root = Path(self.id().replace(".", "-"))
        self.addCleanup(lambda: shutil.rmtree(root, ignore_errors=True))
        root.mkdir()
        source, output = root / "source.xlsx", root / "output.xlsx"
        workbook = Workbook()
        keep = workbook.active
        keep.title = "Keep"
        keep["A1"], keep["A2"] = "=SUM(DropTable[Cost])", "=TableAlias"
        keep["A1"].number_format = "0.00"
        keep["D1"], keep["D2"] = "Value", 1
        keep.add_table(Table(displayName="KeepTable", ref="D1:D2"))
        drop = workbook.create_sheet("Drop")
        drop["A1"], drop["A2"], drop["A3"] = "Cost", 9, 10
        drop.add_table(Table(displayName="DropTable", ref="A1:A3"))
        workbook.defined_names.add(DefinedName("TableAlias", attr_text="SUM(DropTable[Cost])"))
        chart = BarChart()
        chart.add_data(Reference(drop, min_col=1, min_row=2, max_row=3))
        keep.add_chart(chart, "H2")
        workbook.save(source)
        with zipfile.ZipFile(source) as archive:
            chart_xml = archive.read("xl/charts/chart1.xml")
        chart_xml, replacements = re.subn(
            rb"<f>[^<]*Drop[^<]*</f>", b"<f>DropTable[Cost]</f>", chart_xml,
        )
        self.assertGreater(replacements, 0)
        _rewrite(source, {"xl/charts/chart1.xml": chart_xml})

        export_transfer(TransferRequest(
            template_path=source, cells=[], out_path=output,
            template_hash="hash", copy_id="copy", input_revision=0,
            fmt="xlsx", selected={"Keep"}, results={(0, 0, 0): 19, (0, 1, 0): 19},
        ))

        with zipfile.ZipFile(output) as archive:
            names = set(archive.namelist())
            table_parts = [archive.read(name) for name in names if name.startswith("xl/tables/")]
        self.assertNotIn("xl/charts/chart1.xml", names)
        self.assertTrue(any(b'KeepTable' in item for item in table_parts))
        self.assertFalse(any(b'DropTable' in item for item in table_parts))
        loaded = load_workbook(output, data_only=False)
        self.addCleanup(loaded.close)
        self.assertEqual(loaded["Keep"]["A1"].value, 19)
        self.assertEqual(loaded["Keep"]["A1"].number_format, "0.00")
        self.assertEqual(loaded["Keep"]["A2"].value, 19)
        self.assertNotIn("TableAlias", loaded.defined_names)


def _simple_fixture(path):
    workbook = Workbook()
    workbook.active.title = "Keep"
    workbook.create_sheet("Drop")
    workbook.save(path)


def _inject_vba_project(path):
    payload = b"synthetic-vba-project\x00bytes"
    with zipfile.ZipFile(path) as archive:
        rels = archive.read("xl/_rels/workbook.xml.rels").replace(
            b"</Relationships>",
            f'<Relationship Id="rIdVba" Type="{REL}/vbaProject" '
            'Target="vbaProject.bin"/></Relationships>'.encode(),
        )
        content_types = archive.read("[Content_Types].xml")
    content_types = re.sub(
        rb'(<Override\b(?=[^>]*PartName="/xl/workbook\.xml")[^>]*ContentType=")[^"]*',
        rb'\1application/vnd.ms-excel.sheet.macroEnabled.main+xml',
        content_types,
    ).replace(
        b"</Types>",
        b'<Override PartName="/xl/vbaProject.bin" '
        b'ContentType="application/vnd.ms-office.vbaProject"/></Types>',
    )
    _rewrite(path, {
        "xl/_rels/workbook.xml.rels": rels,
        "[Content_Types].xml": content_types,
    }, (("xl/vbaProject.bin", payload),))
    return payload


def _prefix_tags(data, root, children, prefix):
    data = data.replace(f"<{root} ".encode(), f"<{prefix}:{root} ".encode(), 1)
    data = data.replace(f"</{root}>".encode(), f"</{prefix}:{root}>".encode())
    data = data.replace(
        f"<{prefix}:{root} ".encode(),
        f'<{prefix}:{root} xmlns:{prefix}="'.encode()
        + (PKG_REL if root == "Relationships" else "http://schemas.openxmlformats.org/package/2006/content-types").encode()
        + b'" ',
        1,
    )
    for child in children:
        data = data.replace(f"<{child} ".encode(), f"<{prefix}:{child} ".encode())
        data = data.replace(f"</{child}>".encode(), f"</{prefix}:{child}>".encode())
    data = re.sub(rb' xmlns="http://schemas.openxmlformats.org/package/2006/(?:relationships|content-types)"', b"", data, count=1)
    return data


def _namespaced_fixture(path):
    _simple_fixture(path)
    with zipfile.ZipFile(path) as archive:
        root_rels = archive.read("_rels/.rels")
        workbook_rels = archive.read("xl/_rels/workbook.xml.rels")
        content_types = archive.read("[Content_Types].xml")
    root_rels = root_rels.replace(
        b"</Relationships>",
        f'<Relationship Id="rIdCustom" Type="{REL}/custom-properties" Target="docProps/custom.xml"/></Relationships>'.encode(),
    )
    content_types = content_types.replace(
        b"</Types>",
        b'<Override PartName="/docProps/custom.xml" ContentType="application/vnd.openxmlformats-officedocument.custom-properties+xml"/></Types>',
    )
    custom = (
        f'<?xml version="1.0" encoding="UTF-8"?><x:Properties xmlns:x="{CUSTOM}" '
        f'xmlns:v="{VT}" xmlns:mc="urn:mc" xmlns:x15="urn:x15" mc:Ignorable="x15">'
        '<x:property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="2" name="Existing">'
        '<v:lpwstr>keep me</v:lpwstr></x:property></x:Properties>'
    ).encode()
    _rewrite(path, {
        "_rels/.rels": _prefix_tags(root_rels, "Relationships", ("Relationship",), "p"),
        "xl/_rels/workbook.xml.rels": _prefix_tags(
            workbook_rels, "Relationships", ("Relationship",), "p"
        ),
        "[Content_Types].xml": _prefix_tags(
            content_types, "Types", ("Default", "Override"), "ct"
        ),
    }, (("docProps/custom.xml", custom),))


class NamespaceAndPreservationTests(unittest.TestCase):
    def setUp(self):
        self.root = Path(self.id().replace(".", "-"))
        self.addCleanup(lambda: shutil.rmtree(self.root, ignore_errors=True))
        self.root.mkdir()

    def test_prefixed_package_metadata_and_custom_namespace_declarations_survive(self):
        source, output = self.root / "prefixed.xlsx", self.root / "output.xlsx"
        _namespaced_fixture(source)

        export_transfer(TransferRequest(
            template_path=source, cells=[], out_path=output,
            template_hash="hash", copy_id="copy", input_revision=7,
            fmt="xlsx", selected={"Keep"}, results={},
        ))

        with zipfile.ZipFile(output) as archive:
            custom = archive.read("docProps/custom.xml")
            workbook_rels = archive.read("xl/_rels/workbook.xml.rels")
            content_types = archive.read("[Content_Types].xml")
        self.assertIn(b'xmlns:x15="urn:x15"', custom)
        self.assertIn(b'mc:Ignorable="x15"', custom)
        self.assertIn(b"keep me", custom)
        self.assertNotIn(b"sheet2.xml", workbook_rels)
        self.assertIn(b"ct:Override", content_types)
        validate_package(output, require_provenance=True)

    def test_all_sheet_transform_keeps_original_unreferenced_member_bytes(self):
        source, output = self.root / "orphan.xlsx", self.root / "output.xlsx"
        _simple_fixture(source)
        payload = b'<audit xmlns="urn:test">preserve exactly</audit>'
        _rewrite(source, {}, (("customXml/unreferenced.xml", payload),))

        export_transfer(TransferRequest(
            template_path=source, cells=[], out_path=output,
            template_hash="hash", copy_id="copy", input_revision=0,
            fmt="xlsm", results={},
        ))

        with zipfile.ZipFile(output) as archive:
            self.assertEqual(archive.read("customXml/unreferenced.xml"), payload)


class MacroTransferTests(unittest.TestCase):
    def setUp(self):
        self.root = Path(self.id().replace(".", "-"))
        self.addCleanup(lambda: shutil.rmtree(self.root, ignore_errors=True))
        self.root.mkdir()
        self.source = self.root / "macro-source.xlsx"
        _simple_fixture(self.source)
        self.payload = _inject_vba_project(self.source)

    def transfer(self, fmt, selected=None):
        output = self.root / f"output.{fmt}"
        export_transfer(TransferRequest(
            template_path=self.source, cells=[], out_path=output,
            template_hash="hash", copy_id="copy", input_revision=0,
            fmt=fmt, selected=selected, results={},
        ))
        return output

    def test_xlsx_removes_vba_part_relationship_and_content_type(self):
        output = self.transfer("xlsx")
        with zipfile.ZipFile(output) as archive:
            names = set(archive.namelist())
            rels = archive.read("xl/_rels/workbook.xml.rels")
            content_types = archive.read("[Content_Types].xml")
        self.assertNotIn("xl/vbaProject.bin", names)
        self.assertNotIn(b"vbaProject", rels)
        self.assertNotIn(b"vbaProject", content_types)
        self.assertIn(b"spreadsheetml.sheet.main+xml", content_types)

    def test_xlsm_retains_vba_bytes_and_subset_remains_valid(self):
        for selected in (None, {"Keep"}):
            with self.subTest(selected=selected):
                output = self.transfer("xlsm", selected)
                with zipfile.ZipFile(output) as archive:
                    self.assertEqual(archive.read("xl/vbaProject.bin"), self.payload)
                    self.assertIn(b"vbaProject", archive.read("xl/_rels/workbook.xml.rels"))
                    self.assertIn(b"vbaProject", archive.read("[Content_Types].xml"))
                validate_package(output, require_provenance=True)
                package = io.BytesIO(output.read_bytes())
                workbook = load_workbook(package, read_only=True, keep_vba=True)
                self.assertEqual(workbook.sheetnames, ["Keep", "Drop"] if selected is None else ["Keep"])
                workbook.vba_archive.close()
                workbook.close()
                package.close()

    def test_xlsm_from_xlsx_uses_macro_enabled_workbook_content_type(self):
        source, output = self.root / "plain.xlsx", self.root / "plain.xlsm"
        _simple_fixture(source)

        export_transfer(TransferRequest(
            source, [], output, "hash", "copy", 0, fmt="xlsm", results={},
        ))

        with zipfile.ZipFile(output) as archive:
            content_types = archive.read("[Content_Types].xml")
        self.assertIn(b"application/vnd.ms-excel.sheet.macroEnabled.main+xml", content_types)


class PackageValidationTests(unittest.TestCase):
    def setUp(self):
        self.root = Path(self.id().replace(".", "-"))
        self.addCleanup(lambda: shutil.rmtree(self.root, ignore_errors=True))
        self.root.mkdir()

    def package(self, name):
        path = self.root / name
        _simple_fixture(path)
        return path

    def add_root_relationship(self, path, target, additions=()):
        with zipfile.ZipFile(path) as archive:
            relationships = archive.read("_rels/.rels")
        addition = (
            b'<Relationship Id="rIdSynthetic" Type="urn:synthetic" Target="'
            + target.encode() + b'"/>'
        )
        _rewrite(path, {"_rels/.rels": relationships.replace(
            b"</Relationships>", addition + b"</Relationships>",
        )}, additions)

    def test_in_part_relationship_reference_requires_a_matching_relationship(self):
        path = self.package("missing-ref.xlsx")
        dangling = f'<item xmlns:r="{REL}" r:id="rIdMissing"/>'.encode()
        _rewrite(path, {}, (("customXml/item1.xml", dangling),))
        with self.assertRaisesRegex(ValueError, "relationship reference"):
            validate_package(path)

    def test_vendor_relationship_namespace_is_not_an_ooxml_reference(self):
        path = self.package("vendor-namespace.xlsx")
        metadata = b'<item xmlns:r="urn:vendor:relationships:metadata" r:id="vendor-record-1"/>'
        _rewrite(path, {}, (("customXml/item1.xml", metadata),))

        validate_package(path)

    def test_content_type_override_cannot_point_to_a_missing_part(self):
        path = self.package("dangling-content-type.xlsx")
        with zipfile.ZipFile(path) as archive:
            content_types = archive.read("[Content_Types].xml")
        content_types = content_types.replace(
            b"</Types>",
            b'<Override PartName="/missing.xml" ContentType="application/xml"/></Types>',
        )
        _rewrite(path, {"[Content_Types].xml": content_types})
        with self.assertRaisesRegex(ValueError, "content type override"):
            validate_package(path)

    def test_relationship_ids_are_unique_per_source(self):
        path = self.package("duplicate-id.xlsx")
        with zipfile.ZipFile(path) as archive:
            root_rels = archive.read("_rels/.rels")
        first_id = re.search(rb'Id="([^"]+)"', root_rels).group(1)
        addition = (
            b'<Relationship Id="' + first_id
            + b'" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" '
              b'Target="https://example.test" TargetMode="External"/>'
        )
        _rewrite(path, {"_rels/.rels": root_rels.replace(
            b"</Relationships>", addition + b"</Relationships>"
        )})
        with self.assertRaisesRegex(ValueError, "duplicate relationship id"):
            validate_package(path)

    def test_percent_encoded_internal_relationship_target_is_resolved(self):
        path = self.package("percent-target.xlsx")
        self.add_root_relationship(
            path, "customXml/My%20Item.xml",
            (("customXml/My Item.xml", b'<item xmlns="urn:synthetic"/>'),),
        )

        validate_package(path)

    def test_percent_encoded_traversal_is_rejected_after_decoding(self):
        path = self.package("encoded-traversal.xlsx")
        self.add_root_relationship(path, "%2e%2e/%2e%2e/escape.xml")

        with self.assertRaisesRegex(ValueError, "unsafe OOXML relationship target"):
            validate_package(path)

    def test_zip_directory_entries_do_not_require_content_types(self):
        path = self.package("directory-entry.xlsx")
        _rewrite(path, {}, (("xl/worksheets/", b""),))

        validate_package(path)


def _api_formula_fixture(path):
    workbook = Workbook()
    keep = workbook.active
    keep.title = "Keep"
    keep["B1"] = "=Drop!A1*2"
    drop = workbook.create_sheet("Drop")
    drop["A1"] = 21
    drop["A1"].protection = Protection(locked=False)
    workbook.save(path)


class ExportCalculationRouteTests(unittest.TestCase):
    def setUp(self):
        self.fixture = APIFixture(_api_formula_fixture)
        self.addCleanup(self.fixture.close)
        self.client, self.csrf, _ = self.fixture.signup("formula@example.test", "Formula Owner")
        _, self.copy = self.fixture.make_copy(self.client, self.csrf)

    def test_subset_route_freezes_an_actual_cross_sheet_calculation_result(self):
        response = self.client.get(
            f"/api/copies/{self.copy['id']}/export.xlsx?sheets=Keep"
        )
        self.assertEqual(response.status_code, 200, response.body)
        workbook = load_workbook(io.BytesIO(response.body), data_only=False, read_only=True)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook["Keep"]["B1"].value, 42)

    def test_pending_calculation_is_rejected_before_package_export(self):
        with self.fixture.database.transaction() as connection:
            connection.execute(
                "UPDATE pp_copies SET input_revision=1,calculated_revision=0 WHERE id=?",
                (self.copy["id"],),
            )
        with mock.patch("pp.server.export_transfer") as exporter:
            response = self.client.get(
                f"/api/copies/{self.copy['id']}/export.xlsx?sheets=Keep"
            )
        self.assertEqual(response.status_code, 409, response.body)
        self.assertEqual(response.json()["error"], "calculation_required")
        self.assertIn("pending", response.json()["message"].lower())
        exporter.assert_not_called()


def _large_xml_fixture(path, sheet_path):
    _simple_fixture(path)
    with sheet_path.open("wb") as stream:
        stream.write(
            f'<worksheet xmlns="{MAIN}"><sheetData>'.encode()
        )
        for row in range(1, 1_000_001):
            stream.write(
                f'<row r="{row}"><c r="A{row}" s="0" t="n"><v>{row % 10}</v></c></row>'.encode()
            )
        stream.write(b"</sheetData></worksheet>")
    temporary = path.with_suffix(".large")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(temporary, "w") as target:
        for info in source.infolist():
            if info.filename == "xl/worksheets/sheet1.xml":
                target.write(sheet_path, info.filename, compress_type=zipfile.ZIP_STORED)
            else:
                with source.open(info) as reader, target.open(info, "w") as writer:
                    shutil.copyfileobj(reader, writer, 1 << 20)
    temporary.replace(path)


class LargeXmlSubsetBenchmarkTests(unittest.TestCase):
    def test_55mb_xml_subset_path_stays_within_budget(self):
        root = Path(self.id().replace(".", "-"))
        self.addCleanup(lambda: shutil.rmtree(root, ignore_errors=True))
        root.mkdir()
        source, output, sheet = root / "source.xlsx", root / "output.xlsx", root / "sheet.xml"
        _large_xml_fixture(source, sheet)
        self.assertGreaterEqual(source.stat().st_size, 55 * 1024 * 1024)
        started = time.monotonic()
        export_transfer(TransferRequest(
            template_path=source, cells=[], out_path=output,
            template_hash="hash", copy_id="copy", input_revision=0,
            fmt="xlsx", selected={"Keep"}, results={},
        ))
        elapsed = time.monotonic() - started
        rss = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
        rss_mb = rss / (1024 * 1024) if os.uname().sysname == "Darwin" else rss / 1024
        print(f"TRANSFER_XML_BENCH input_mb={source.stat().st_size / 1024 / 1024:.2f} "
              f"elapsed_s={elapsed:.3f} rss_mb={rss_mb:.2f}")
        self.assertLess(elapsed, 30)
        self.assertLess(rss_mb, 500)


if __name__ == "__main__":
    unittest.main()


class HardenedXmlParseTests(unittest.TestCase):
    def test_relationship_parts_with_entities_are_refused(self):
        """The shared transfer parser must refuse entity tricks even where no upload screen ran first."""
        from defusedxml import EntitiesForbidden
        from pp.export.transfer import _relationship_graph

        bomb = ('<?xml version="1.0"?><!DOCTYPE r [<!ENTITY a "aaaaaaaaaa"><!ENTITY b "&a;&a;&a;&a;">]>'
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Id="&b;" Type="x" Target="sheet1.xml"/></Relationships>')
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr("xl/_rels/workbook.xml.rels", bomb)
        with zipfile.ZipFile(buffer) as archive, self.assertRaises(EntitiesForbidden):
            _relationship_graph(archive)


class InDocumentLinkTests(unittest.TestCase):
    def test_link_to_a_cell_in_the_same_workbook_is_not_a_missing_part(self):
        """Real templates link drawings to cells with Target="#'Sheet'!A1"; uploads failed on it."""
        from tempfile import TemporaryDirectory

        with TemporaryDirectory() as directory:
            source, patched = Path(directory) / "in.xlsx", Path(directory) / "out.xlsx"
            book = Workbook()
            book.active["A1"] = 1
            book.save(source)
            with zipfile.ZipFile(source) as original, zipfile.ZipFile(patched, "w") as out:
                for info in original.infolist():
                    out.writestr(info, original.read(info))
                out.writestr("xl/worksheets/_rels/sheet1.xml.rels", (
                    '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.'
                    'openxmlformats.org/package/2006/relationships"><Relationship Id="rIdJump" '
                    'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
                    'hyperlink" Target="#\'Sheet A\'!B2"/></Relationships>'))
            validate_package(patched)
