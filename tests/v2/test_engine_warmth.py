"""Warm copy workers must reuse their engine and reconstruct exact snapshots."""

import json

from pathlib import Path
import tempfile
import unittest
from unittest import mock

from openpyxl import Workbook
from openpyxl.styles import Protection

from pp.engine.pool import CalculationRequest, EnginePool
from tests.v2.support import BackendFixture, TMP


class WarmEngineTests(unittest.TestCase):
    def test_zero_revision_open_has_an_empty_sparse_overlay(self):
        from pp.engine.worker import WarmCopy

        fixture = BackendFixture()
        self.addCleanup(fixture.close)

        opened = WarmCopy("empty-overlay").apply(0, str(fixture.source), [])

        self.assertEqual(opened["results"], [])

    def test_sparse_overlay_excludes_unrelated_formula_results(self):
        from pp.engine.worker import WarmCopy

        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "sparse-overlay.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = 1
            sheet["B1"] = "=A1*2"
            for column in range(3, 103):
                sheet.cell(1, column, f"={column}")
            workbook.save(source)
            copy = WarmCopy("sparse-overlay")

            changed = copy.apply(1, str(source), [(0, 0, 0, 4)])

            self.assertEqual(changed["results"], [(0, 0, 1, 8)])

    def test_warm_delta_after_broad_edit_contains_only_changed_overlay_rows(self):
        from pp.engine.worker import WarmCopy, WorkerBatch

        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "delta-overlay.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = 1
            sheet["A2"] = 1
            for row in range(1, 101):
                sheet.cell(row, 2, f"=$A$1*{row}")
            sheet["C1"] = "=A2*3"
            workbook.save(source)
            copy = WarmCopy("delta-overlay")

            broad = copy.apply(1, str(source), WorkerBatch([(0, 0, 0, 2)], 0))
            ordinary = copy.apply(
                2, str(source), WorkerBatch([(0, 0, 0, 2), (0, 1, 0, 3)], 1)
            )

            self.assertFalse(broad["delta"])
            self.assertEqual(len(broad["results"]), 100)
            self.assertTrue(ordinary["delta"])
            self.assertEqual(ordinary["base_revision"], 1)
            self.assertEqual(ordinary["results"], [(0, 0, 2, 9)])
            self.assertEqual(ordinary["removals"], [])

    def test_worker_base_mismatch_falls_back_to_one_full_overlay(self):
        from pp.engine.worker import WarmCopy, WorkerBatch

        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "delta-mismatch.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = 1
            sheet["B1"] = "=A1*2"
            workbook.save(source)
            copy = WarmCopy("delta-mismatch")
            copy.apply(1, str(source), WorkerBatch([(0, 0, 0, 2)], 0))

            mismatched = copy.apply(2, str(source), WorkerBatch([(0, 0, 0, 3)], 0))

            self.assertFalse(mismatched["delta"])
            self.assertEqual(mismatched["base_revision"], 0)
            self.assertEqual(mismatched["results"], [(0, 0, 1, 6)])

    def test_restoring_baseline_removes_sparse_overlay_row(self):
        from pp.engine.worker import WarmCopy

        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "restore-overlay.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = 1
            sheet["B1"] = "=A1*2"
            workbook.save(source)
            copy = WarmCopy("restore-overlay")

            self.assertEqual(copy.apply(1, str(source), [(0, 0, 0, 4)])["results"],
                             [(0, 0, 1, 8)])
            restored = copy.apply(2, str(source), [])

            self.assertEqual(restored["results"], [])

    def test_shrunk_baseline_spill_uses_null_tombstone_until_restored(self):
        from pp.engine.worker import WarmCopy

        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "baseline-spill.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = 2
            sheet["B1"] = "=SEQUENCE(A1)"
            workbook.save(source)
            copy = WarmCopy("baseline-spill")

            self.assertEqual(copy.apply(0, str(source), [])["results"], [])
            shrunk = copy.apply(1, str(source), [(0, 0, 0, 1)])
            restored = copy.apply(2, str(source), [])

            self.assertIn((0, 1, 1, None), shrunk["results"])
            self.assertEqual(restored["results"], [])

    def test_removed_input_spill_member_reveals_immutable_input(self):
        from pp.engine.worker import WarmCopy

        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "input-spill.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.protection.sheet = True
            sheet["A1"] = 2
            sheet["A1"].protection = Protection(locked=False)
            sheet["B1"] = "=SEQUENCE(A1)"
            sheet["B2"] = 7
            sheet["B2"].protection = Protection(locked=False)
            workbook.save(source)
            copy = WarmCopy("input-spill")

            expanded = copy.apply(1, str(source), [(0, 1, 1, None)])
            restored = copy.apply(2, str(source), [])

            self.assertIn((0, 1, 1, 2), expanded["results"])
            self.assertNotIn((0, 1, 1, None), restored["results"])
            self.assertEqual(restored["results"], [])

    def test_warm_revision_does_not_reload_source_and_restores_removed_overrides(self):
        fixture = BackendFixture()
        self.addCleanup(fixture.close)
        pool = EnginePool(max_workers=1)
        self.addCleanup(pool.close)
        # the template's own value: an empty first request would need no engine at all
        original = pool.apply(CalculationRequest("warm-copy", 0, fixture.source, [(0, 1, 1, 10)]))
        fixture.source.rename(fixture.source.with_suffix(".parked"))
        edited = pool.apply(CalculationRequest("warm-copy", 1, fixture.source, [(0, 1, 1, 15)]))
        restored = pool.apply(CalculationRequest("warm-copy", 2, fixture.source, []))
        self.assertEqual((original["pid"], edited["pid"]), (restored["pid"], restored["pid"]))
        self.assertEqual(self.cell(edited), 30)
        self.assertEqual(self.cell(restored), 20)

    def test_untouched_copy_is_published_without_starting_an_engine(self):
        pool = EnginePool(max_workers=1)
        self.addCleanup(pool.close)

        result = pool.apply(CalculationRequest("new-copy", 0, Path("missing.xlsm"), [], 0))

        self.assertEqual((result["results"], result["delta"], result["base_revision"]), ([], False, 0))
        self.assertEqual(pool.order(), [])

    def cell(self, response):
        values = {(sheet, row, col): value for sheet, row, col, value in response["changes"]}
        return values[(2, 1, 1)]

    def test_warm_overrides_distinguish_boolean_true_from_number_one(self):
        """Typed override comparison must send a boolean after a numeric one."""
        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "typed-warm.xlsm"
            workbook = Workbook()
            inputs = workbook.active
            inputs.title = "Inputs"
            inputs["A1"] = 0
            formulas = workbook.create_sheet("Formulas")
            formulas["A1"] = "=Inputs!A1"
            workbook.save(source)
            pool = EnginePool(max_workers=1)
            self.addCleanup(pool.close)

            pool.apply(CalculationRequest("typed-copy", 1, source, [(0, 0, 0, 1)]))
            boolean = pool.apply(CalculationRequest("typed-copy", 2, source, [(0, 0, 0, True)]))

            self.assertIn("results", boolean, "workers must expose an exact complete overlay")
            results = {(sheet, row, col): value for sheet, row, col, value in boolean["results"]}
            self.assertIs(results[(1, 0, 0)], True)

    def test_warm_overrides_preserve_small_numeric_change(self):
        """Parity tolerance must not suppress a real input override."""
        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "exact-warm.xlsm"
            workbook = Workbook()
            inputs = workbook.active
            inputs.title = "Inputs"
            inputs["A1"] = 0
            formulas = workbook.create_sheet("Formulas")
            formulas["A1"] = "=Inputs!A1"
            workbook.save(source)
            pool = EnginePool(max_workers=1)
            self.addCleanup(pool.close)

            pool.apply(CalculationRequest("exact-copy", 1, source, [(0, 0, 0, 1)]))
            changed = pool.apply(CalculationRequest(
                "exact-copy", 2, source, [(0, 0, 0, 1.0000000001)]
            ))

            results = {(sheet, row, col): value for sheet, row, col, value in changed["results"]}
            self.assertEqual(results[(1, 0, 0)], 1.0000000001)

    def test_error_to_identical_looking_text_updates_engine_and_publication(self):
        """An API error object must differ from text that renders like that object."""
        from pp.calc_service import CalculationPublication, CalculationService

        fixture = BackendFixture()
        self.addCleanup(fixture.close)
        pool = EnginePool(max_workers=1)
        self.addCleanup(pool.close)
        with fixture.database.transaction() as connection:
            connection.execute(
                "UPDATE pp_copies SET input_revision=2,calculated_revision=0 WHERE id='copy-a'"
            )
            connection.executemany(
                "INSERT INTO pp_calc_jobs(copy_id,revision,status,created_at,updated_at) "
                "VALUES('copy-a',?,'pending','2026-10-10T00:00:00Z','2026-10-10T00:00:00Z')",
                [(1,), (2,)],
            )
        service = CalculationService(fixture.database)
        error = pool.apply(CalculationRequest(
            "copy-a", 1, fixture.source, [(0, 1, 1, {"error": "#DIV/0!"})]
        ))
        self.assertTrue(service.complete(CalculationPublication("copy-a", 1, error["results"])))

        literal = pool.apply(CalculationRequest(
            "copy-a", 2, fixture.source, [(0, 1, 1, "{'error': '#DIV/0!'}")]
        ))
        self.assertTrue(service.complete(CalculationPublication("copy-a", 2, literal["results"])))

        results = {(sheet, row, col): value for sheet, row, col, value in literal["results"]}
        with fixture.database.connect() as connection:
            row = connection.execute(
                "SELECT payload_json FROM pp_events WHERE copy_id='copy-a' "
                "AND event_type='calc' AND input_revision=2"
            ).fetchone()
        changes = {(item["sheet"], item["r"], item["c"]): item["v"]
                   for item in json.loads(row["payload_json"])["changes"]}
        self.assertEqual(results[(2, 1, 1)], {"error": "#VALUE!"})
        self.assertEqual(changes[(2, 1, 1)], {"error": "#VALUE!"})

    def test_failed_apply_discards_partially_mutated_warm_engine(self):
        """A failed native write must not poison the next exact snapshot replay."""
        from pp.engine.formualizer import EngineError
        from pp.engine.worker import worker_loop

        class MutatingEngine:
            def __init__(self):
                self.value = 10

            def get(self, sheet, row, col):
                return self.value

            def apply(self, updates):
                for _, _, _, value in updates:
                    self.value = value
                if self.value == 13:
                    raise EngineError("failure after native mutation")
                return [(1, 0, 0, self.value * 2)]

            def overlay_snapshot(self):
                return [(1, 0, 0, self.value * 2)]

        class ScriptedConnection:
            def __init__(self):
                self.requests = iter([
                    (1, "synthetic.xlsm", [(0, 0, 0, 12)]),
                    (2, "synthetic.xlsm", [(0, 0, 0, 13)]),
                    (3, "synthetic.xlsm", [(0, 0, 0, 12)]),
                    None,
                ])
                self.responses = []

            def recv(self):
                return next(self.requests)

            def send(self, response):
                self.responses.append(response)

            def close(self):
                pass

        connection = ScriptedConnection()
        with mock.patch("pp.engine.worker.FormualizerAdapter.load",
                        side_effect=lambda path: MutatingEngine()):
            worker_loop(connection, "exact-copy")

        self.assertEqual([success for success, _ in connection.responses], [True, False, True])
        self.assertEqual(connection.responses[0][1]["results"], [(1, 0, 0, 24)])
        self.assertEqual(connection.responses[2][1]["results"], [(1, 0, 0, 24)])
