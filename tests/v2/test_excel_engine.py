"""Excel engine: Excel calculates, the engine publishes exactly what changed."""

import sys
import tempfile
import types
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

from tests.v2.support import TMP
from tests.v2.test_engine_cached import _replace_formula_values


def workbook(directory):
    path = Path(directory) / "plan.xlsx"
    book = Workbook()
    calc = book.active
    calc.title = "Calc"
    data = book.create_sheet("Data Src")
    calc["A1"] = 2
    calc["B1"] = "=A1*2"
    calc["C1"] = "=SUM(B:B)"
    calc["D1"] = "=Rate*10"
    calc["E1"] = "=7"
    data["B2"] = 3
    book.defined_names["Rate"] = DefinedName("Rate", attr_text="'Data Src'!$B$2")
    book.save(path)
    _replace_formula_values(path, {"xl/worksheets/sheet1.xml": {
        "B1": ("4", None), "C1": ("4", None), "D1": ("30", None), "E1": ("7", None)}})
    return path


class ExcelEngineTests(unittest.TestCase):
    def load(self):
        from pp.engine.excel import ExcelEngine, FormualizerHost

        TMP.mkdir(parents=True, exist_ok=True)
        directory = tempfile.TemporaryDirectory(dir=TMP)
        self.addCleanup(directory.cleanup)
        source = workbook(directory.name)
        engine = ExcelEngine.load(source, cache_dir=Path(directory.name) / "cache",
                                  host_factory=FormualizerHost)
        self.addCleanup(engine.close)
        return engine

    def test_edit_publishes_reachable_changes_and_reverting_removes_them(self):
        engine = self.load()
        self.assertEqual(engine.get(0, 0, 1), 4)

        changes = engine.apply([(0, 0, 0, 5)])

        self.assertEqual(changes, [(0, 0, 1, 10), (0, 0, 2, 10)])
        self.assertEqual(engine.overlay_snapshot(), [(0, 0, 1, 10), (0, 0, 2, 10)])
        engine.apply([(0, 0, 0, 2)])
        self.assertEqual(engine.overlay_snapshot(), [])
        self.assertEqual(engine.overlay_delta(), ([], [(0, 0, 1), (0, 0, 2)]))

    def test_cross_sheet_name_and_clear(self):
        engine = self.load()
        self.assertEqual(engine.apply([(1, 1, 1, 4)]), [(0, 0, 3, 40)])
        self.assertEqual(engine.apply([(0, 0, 0, None)]), [(0, 0, 1, 0), (0, 0, 2, 0)])
        self.assertTrue(engine.handles_clears)

    def test_unrelated_formulas_are_never_read(self):
        engine = self.load()
        with mock.patch.object(engine.host, "read", wraps=engine.host.read) as read:
            engine.apply([(1, 1, 1, 9)])
        self.assertEqual([call.args[0] for call in read.call_args_list], [0])
        self.assertEqual(engine.get(0, 0, 4), 7)


class ExcelWorkerTests(unittest.TestCase):
    def test_worker_uses_excel_engine_and_clears_without_reloading(self):
        from pp.engine.worker import WarmCopy

        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as directory, \
                mock.patch.dict("os.environ", {"PP_ENGINE_CACHE_DIR": str(Path(directory) / "cache")}):
            source = workbook(directory)
            copy = WarmCopy("copy", "excel")
            first = copy.apply(1, str(source), [(0, 0, 0, 5)])
            with mock.patch.object(WarmCopy, "_reload_with_clears", side_effect=AssertionError("reload")):
                cleared = copy.apply(2, str(source), [(0, 0, 0, None)])
            engine = copy.engine
            with mock.patch.object(engine, "close", wraps=engine.close) as close:
                copy.close()

        self.assertEqual(type(engine).__name__, "ExcelEngine")
        self.assertEqual(first["changes"], [(0, 0, 1, 10), (0, 0, 2, 10)])
        self.assertEqual(cleared["changes"], [(0, 0, 1, 0), (0, 0, 2, 0)])
        close.assert_called_once()


def saving_host():
    """FormualizerHost that can save itself, the way Excel's SaveCopyAs keeps the copy's current cells."""
    from openpyxl import load_workbook

    from pp.engine.excel import FormualizerHost

    class SavingHost(FormualizerHost):
        opened = []

        def __init__(self, workbook, sheet_names):
            super().__init__(workbook, sheet_names)
            self.source, self.written = Path(workbook), {}
            SavingHost.opened.append(self.source)

        def write(self, sheet, cells):
            super().write(sheet, cells)
            self.written.update(((sheet, row, col), value) for row, col, value in cells)

        def write_formulas(self, sheet, cells):
            super().write_formulas(sheet, cells)
            self.written.update(((sheet, row, col), text) for row, col, text in cells)

        def save_copy(self, path):
            book = load_workbook(self.source)
            for (sheet, row, col), value in self.written.items():
                book[self.names[sheet]].cell(row + 1, col + 1).value = value
            book.save(path)

    return SavingHost


class ResumeTests(unittest.TestCase):
    def test_restarted_copy_opens_its_saved_point_and_replays_only_newer_edits(self):
        from pp.engine.worker import WarmCopy

        host = saving_host()
        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as directory, \
                mock.patch.dict("os.environ", {"PP_ENGINE_CACHE_DIR": str(Path(directory) / "cache")}), \
                mock.patch("pp.engine.excel.default_host", return_value=host):
            source = workbook(directory)
            resume = Path(directory) / "resume"
            first = WarmCopy("copy", "excel", resume)
            first.apply(1, str(source), [(0, 0, 0, 5)])
            self.assertIsNotNone(first.resume_due())
            first.save_resume_point()
            self.assertIsNone(first.resume_due())
            first.close()

            again = WarmCopy("copy", "excel", resume)
            with mock.patch("pp.engine.excel.ExcelEngine.apply", autospec=True,
                            side_effect=lambda engine, updates: (applied.append(updates), [])[1]) as _:
                applied = []
                again.apply(2, str(source), [(0, 0, 0, 5), (1, 1, 1, 4)])
            again.close()

            fresh = WarmCopy("other", "excel", resume)
            restarted = WarmCopy("copy", "excel", resume)
            expected = fresh.apply(3, str(source), [(0, 0, 0, 5), (1, 1, 1, 4)])
            resumed = restarted.apply(3, str(source), [(0, 0, 0, 5), (1, 1, 1, 4)])
            fresh.close()
            restarted.close()

        self.assertEqual(applied, [[(1, 1, 1, 4)]])  # the edit saved at the resume point is not replayed
        self.assertEqual([path.name for path in host.opened], ["plan.xlsx"] * 4)
        self.assertEqual(host.opened[1].parent.parent, resume / "copy")  # restarts open the saved point
        self.assertEqual(host.opened[2], source)  # a copy with no saved point starts from the template
        self.assertEqual(host.opened[3].parent.parent, resume / "copy")
        self.assertEqual(resumed["results"], expected["results"])

    def test_resume_point_from_another_template_file_is_ignored(self):
        from pp.engine.worker import WarmCopy

        host = saving_host()
        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as directory, \
                mock.patch.dict("os.environ", {"PP_ENGINE_CACHE_DIR": str(Path(directory) / "cache")}), \
                mock.patch("pp.engine.excel.default_host", return_value=host):
            source = workbook(directory)
            resume = Path(directory) / "resume"
            copy = WarmCopy("copy", "excel", resume)
            copy.apply(1, str(source), [(0, 0, 0, 5)])
            copy.save_resume_point()
            copy.close()
            workbook(directory)  # the template file is replaced
            again = WarmCopy("copy", "excel", resume)
            result = again.apply(2, str(source), [(0, 0, 0, 5)])
            again.close()

        self.assertEqual(result["changes"], [(0, 0, 1, 10), (0, 0, 2, 10)])  # recalculated from the template


class IdleSaveTests(unittest.TestCase):
    def test_worker_saves_a_resume_point_only_while_idle(self):
        from pp.engine import worker

        class Connection:
            def __init__(self):
                self.events, self.sent = ["request", "idle", None], []

            def poll(self, timeout):
                if self.events[0] == "idle":
                    self.events.pop(0)
                    return False
                return True

            def recv(self):
                return self.events.pop(0)

            def send(self, message):
                self.sent.append(message)

            def close(self):
                pass

        saves = []
        state = {"unsaved": False}

        def serve(copy, request):
            state["unsaved"] = True
            return {}

        def save(copy):
            saves.append(len(connection.events))
            state["unsaved"] = False

        connection = Connection()
        with mock.patch.object(worker, "_serve", serve), \
                mock.patch.object(worker.WarmCopy, "resume_due", lambda copy: 5.0 if state["unsaved"] else None), \
                mock.patch.object(worker.WarmCopy, "save_resume_point", save), \
                mock.patch.object(worker.WarmCopy, "engine_bytes", lambda copy: 0):
            worker.worker_loop(connection, "copy", "excel")

        self.assertEqual(saves, [1])  # once, after the request, while no request was waiting
        self.assertEqual(connection.sent, [(True, {"engine_bytes": 0})])


class ExcelRuntimeTests(unittest.TestCase):
    def test_excel_runtime_warms_template_caches_and_builds_an_excel_pool(self):
        import json

        from pp.runtime import create_runtime
        from tests.v2.test_end_to_end import runtime_root

        temporary, root = runtime_root()
        self.addCleanup(temporary.cleanup)
        settings = json.loads((root / "settings.json").read_text())
        (root / "settings.json").write_text(json.dumps({**settings, "engine": "excel"}))
        with mock.patch("pp.engine.excel.warm_caches") as warm, \
                mock.patch("threading.Thread.start", lambda self: self.run()):
            runtime = create_runtime(root)
        self.addCleanup(runtime.close)
        warm.assert_called_once_with(runtime.snapshot.source_path)
        self.assertEqual(runtime.snapshot.results, "saved")  # no import engine pass, edits stay on
        self.assertEqual(runtime.pool._engine, "excel")


class ValueConversionTests(unittest.TestCase):
    def test_com_values(self):
        from pp.engine.excel import COM_ERROR_BASE, from_com, to_com

        self.assertEqual(from_com(COM_ERROR_BASE + 2007), {"error": "#DIV/0!"})
        self.assertEqual(from_com(COM_ERROR_BASE + 2042), {"error": "#N/A"})
        self.assertEqual(from_com(5.0), 5)
        self.assertEqual(from_com(2.5), 2.5)
        self.assertIs(from_com(True), True)
        self.assertEqual(from_com("x"), "x")
        self.assertIsNone(from_com(None))
        self.assertEqual(from_com({"error": "#REF!"}), {"error": "#REF!"})
        self.assertEqual(to_com("1-2"), "'1-2")
        self.assertEqual(to_com(""), "")
        self.assertEqual(to_com(3), 3)


class FakeRange:
    def __init__(self, sheet, cells):
        self.sheet, self.cells = sheet, cells
        sheet.calls += 1

    @property
    def Value2(self):
        (top, left), (bottom, right) = self.cells
        rows = tuple(tuple(self.sheet.values.get((r, c)) for c in range(left, right + 1))
                     for r in range(top, bottom + 1))
        return rows[0][0] if (top, left) == (bottom, right) else rows

    @Value2.setter
    def Value2(self, rows):
        (top, left), _ = self.cells
        for r, line in enumerate(rows):
            for c, item in enumerate(line):
                self.sheet.values[(top + r, left + c)] = item


class FakeCell:
    def __init__(self, sheet, row, col):
        self.sheet, self.at = sheet, (row, col)

    @property
    def Value2(self):
        return self.sheet.values.get(self.at)

    @Value2.setter
    def Value2(self, value):
        self.sheet.values[self.at] = value


class FakeSheet:
    def __init__(self):
        self.values = {}
        self.calls = 0

    def Cells(self, row, col):
        return FakeCell(self, row, col)

    def Range(self, first, last):
        return FakeRange(self, (first.at, last.at))


class FakeApp:
    def __init__(self):
        self.Hwnd, self.calculated, self.quit, self.sheets = 77, 0, False, {}
        self.Workbooks = types.SimpleNamespace(Open=self.open, Add=self.add)

    def add(self):
        self.added = types.SimpleNamespace(closed=False)
        self.added.Close = lambda SaveChanges: setattr(self.added, "closed", True)
        self.calculation_when_opened = None
        return self.added

    def open(self, path, **options):
        self.opened = (path, options)
        self.calculation_when_opened = getattr(self, "Calculation", None)
        return types.SimpleNamespace(Worksheets=lambda name: self.sheets.setdefault(name, FakeSheet()),
                                     Close=lambda SaveChanges: setattr(self, "closed", SaveChanges))

    def Calculate(self):
        self.calculated += 1

    def Quit(self):
        self.quit = True


class ConvertTests(unittest.TestCase):
    def test_save_as_xlsm_with_macros_off_and_excel_always_closed(self):
        calls, killed = [], []
        book = types.SimpleNamespace(SaveAs=lambda path, FileFormat: calls.append(("save", FileFormat)),
                                     Close=lambda SaveChanges: calls.append(("close", SaveChanges)))
        app = types.SimpleNamespace(Hwnd=5, Quit=lambda: calls.append(("quit",)))
        app.Workbooks = types.SimpleNamespace(
            Open=lambda path, **options: calls.append(("open", options["ReadOnly"], options["Password"])) or book)
        modules = {
            "pythoncom": types.SimpleNamespace(CoInitialize=lambda: None, CoUninitialize=lambda: None),
            "win32com.client": types.SimpleNamespace(DispatchEx=lambda name: app),
            "win32process": types.SimpleNamespace(GetWindowThreadProcessId=lambda hwnd: (1, 99)),
        }
        modules["win32com"] = types.SimpleNamespace(client=modules["win32com.client"])
        with mock.patch.dict(sys.modules, modules), \
                mock.patch("pp.engine.excel._kill", side_effect=killed.append):
            from pp.engine.excel import convert_to_xlsm

            convert_to_xlsm(Path("in.xlsb"), Path("out.xlsm"))

        self.assertEqual(calls, [("open", True, ""), ("save", 52), ("close", False), ("quit",)])
        self.assertEqual((app.AutomationSecurity, app.EnableEvents, app.DisplayAlerts), (3, False, False))
        self.assertEqual(killed, [99])


class ComHostTests(unittest.TestCase):
    def test_unregistered_excel_name_falls_back_to_the_class_id(self):
        from pp.engine.excel import CO_E_CLASSSTRING, EXCEL_CLSID, _new_excel

        app, asked = object(), []

        def dispatch(name):
            asked.append(name)
            if name == "Excel.Application":
                raise RuntimeError(CO_E_CLASSSTRING, "Invalid class string", None, None)
            return app

        self.assertIs(_new_excel(types.SimpleNamespace(DispatchEx=dispatch)), app)
        self.assertEqual(asked, ["Excel.Application", EXCEL_CLSID])

        def other_failure(name):
            raise RuntimeError(-2147024891, "Access is denied", None, None)

        with self.assertRaises(RuntimeError):
            _new_excel(types.SimpleNamespace(DispatchEx=other_failure))

    def test_hidden_macro_free_manual_excel_on_a_private_copy(self):
        app = FakeApp()
        killed = []
        modules = {
            "pythoncom": types.SimpleNamespace(CoInitialize=lambda: None, CoUninitialize=lambda: None),
            "win32com.client": types.SimpleNamespace(DispatchEx=lambda name: app),
            "win32process": types.SimpleNamespace(GetWindowThreadProcessId=lambda hwnd: (1, 4242)),
        }
        modules["win32com"] = types.SimpleNamespace(client=modules["win32com.client"])
        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as directory, \
                mock.patch.dict(sys.modules, modules), \
                mock.patch("pp.engine.excel._kill", side_effect=killed.append):
            from pp.engine.excel import ComExcelHost

            source = Path(directory) / "plan.xlsm"
            source.write_bytes(b"template")
            host = ComExcelHost(source, ["Calc", "Data"])
            sheet = app.sheets["Calc"]
            sheet.calls = 0
            host.write(0, [(0, 0, 5), (1, 2, "1-2")])
            separate_writes = sheet.calls
            sheet.calls = 0
            host.write(0, [(r, c, r * 10 + c) for r in range(10, 20) for c in range(3)])
            block_writes = sheet.calls
            host.calculate()
            app.sheets["Calc"].values[(2, 2)] = 7.0  # Excel row 2, column B
            block = host.read(0, 0, 0, 1, 2)
            single = host.read(0, 0, 0, 0, 0)
            opened = Path(app.opened[0])
            self.assertNotEqual(opened, source)
            self.assertEqual(opened.read_bytes(), b"template")
            host.close()

        self.assertEqual((app.Visible, app.DisplayAlerts, app.EnableEvents, app.AutomationSecurity,
                          app.Calculation), (False, False, False, 1, -4135))
        self.assertEqual(app.opened[1]["UpdateLinks"], 0)
        self.assertEqual(app.calculation_when_opened, -4135)  # never recalculates while opening
        self.assertTrue(app.added.closed)
        self.assertEqual((separate_writes, block_writes), (2, 1))  # 30 cells, one Excel call
        self.assertEqual(app.sheets["Calc"].values[(20, 3)], 192)
        self.assertEqual(app.calculated, 1)
        self.assertEqual(block, [[5, None, None], [None, 7.0, "'1-2"]])
        self.assertEqual(single, [[5]])
        self.assertTrue(app.quit)
        self.assertEqual(killed, [4242])
        self.assertFalse(opened.parent.exists())


if __name__ == "__main__":
    unittest.main()
