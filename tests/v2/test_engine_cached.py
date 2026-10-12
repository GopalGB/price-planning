import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest import mock
import xml.etree.ElementTree as ET
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Protection
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table

from ops.engine_scale_fixture import generate
from tests.v2.support import TMP


class CachedStartupTests(unittest.TestCase):
    def test_contiguous_formula_registrations_use_row_batches(self):
        from pp.engine.formualizer import _set_native_formulas

        class WorkbookSpy:
            def __init__(self):
                self.batches = []
                self.singles = []

            def set_formulas_batch(self, sheet, row, col, formulas):
                self.batches.append((sheet, row, col, formulas))

            def set_formula(self, sheet, row, col, formula):
                self.singles.append((sheet, row, col, formula))

        workbook = WorkbookSpy()
        registrations = {
            (0, 0, 0): "=A2+$D$1",
            (0, 0, 1): "=B2+$D$1",
            (0, 0, 3): "=D2+$D$1",
            (0, 1, 0): "=A3+$D$1",
            (0, 1, 1): "=B3+$D$1",
            (1, 0, 0): "=A2+$D$1",
        }

        _set_native_formulas(workbook, ["One", "Two"], registrations)

        self.assertEqual(workbook.batches, [
            ("One", 1, 1, [["=A2+$D$1", "=B2+$D$1"]]),
            ("One", 2, 1, [["=A3+$D$1", "=B3+$D$1"]]),
        ])
        self.assertEqual(workbook.singles, [
            ("One", 1, 4, "=D2+$D$1"),
            ("Two", 1, 1, "=A2+$D$1"),
        ])

    def test_relative_shared_follower_is_traced_from_its_own_precedent(self):
        from pp.engine.formualizer import FormualizerAdapter

        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "relative-shared.xlsx"
            follower = _relative_shared_fixture(source)
            engine = FormualizerAdapter.load(source)

            changes = engine.apply([(0, 0, 10, 10)])

            self.assertIn((*follower, 17), changes)
            self.assertEqual(engine.get(*follower), 17)

    def test_absolute_shared_family_is_traced_from_common_precedent(self):
        from pp.engine.formualizer import FormualizerAdapter

        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "absolute-shared.xlsx"
            generate(source, 0.0001, 20261011)
            engine = FormualizerAdapter.load(source)

            changes = engine.apply([(0, 0, 0, 8)])

            self.assertEqual(len(changes), 150)
            self.assertTrue(all(isinstance(item[3], (int, float)) for item in changes))

    def test_cached_open_never_evaluates_all_and_preserves_typed_cache(self):
        from pp.engine.formualizer import FormualizerAdapter

        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "cached.xlsm"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = 3
            sheet["B1"] = "=A1*2"
            sheet["C1"] = "=TRUE"
            sheet["D1"] = '=""'
            sheet["E1"] = "=NA()"
            workbook.save(source)
            _replace_formula_values(source, {"xl/worksheets/sheet1.xml": {
                "B1": ("6", None), "C1": ("1", "b"), "D1": ("", "str"),
                "E1": ("#N/A", "e"),
            }})
            with mock.patch("pp.engine.formualizer.formualizer.Workbook.evaluate_all",
                            side_effect=AssertionError("evaluate_all boundary")):
                with mock.patch("pp.engine.formualizer.formualizer.Workbook.evaluate_cells",
                                side_effect=AssertionError("unexpected cached startup evaluation")):
                    engine = FormualizerAdapter.load(source)
                engine.apply([(0, 0, 0, 4)])
            self.assertEqual(engine.get(0, 0, 1), 8)
            self.assertIs(engine.get(0, 0, 2), True)
            self.assertEqual(engine.get(0, 0, 3), "")
            self.assertEqual(engine.get(0, 0, 4), {"error": "#N/A"})

    def test_cached_chain_invalidates_transitive_native_dependents_only(self):
        from pp.engine.formualizer import FormualizerAdapter

        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "cached-chain.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Calc"
            sheet["A1"] = 1
            sheet["B1"] = "=A1+1"
            sheet["C1"] = "=B1+1"
            sheet["D1"] = "=40+2"
            workbook.save(source)
            _replace_formula_values(source, {"xl/worksheets/sheet1.xml": {
                "B1": ("2", None), "C1": ("3", None), "D1": ("999", None),
            }})
            engine = FormualizerAdapter.load(source)
            engine.workbook.get_eval_plan([("Calc", 1, 1)])

            changes = engine.apply([(0, 0, 0, 5)])

            self.assertEqual(engine.get(0, 0, 2), 7)
            self.assertEqual(engine.get(0, 0, 3), 999)
            self.assertIn((0, 0, 2, 7), changes)
            self.assertNotIn((0, 0, 3, 42), changes)

    def test_cached_cross_sheet_range_edit_recalculates_range_and_downstream(self):
        from pp.engine.formualizer import FormualizerAdapter

        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "cached-range.xlsx"
            workbook = Workbook()
            inputs = workbook.active
            inputs.title = "Inputs"
            inputs.append([1])
            inputs.append([2])
            inputs.append([3])
            calc = workbook.create_sheet("Calc")
            calc["A1"] = "=SUM(Inputs!A1:A3)"
            audit = workbook.create_sheet("Audit")
            audit["A1"] = "=Calc!A1*2"
            workbook.save(source)
            _replace_formula_values(source, {
                "xl/worksheets/sheet2.xml": {"A1": ("6", None)},
                "xl/worksheets/sheet3.xml": {"A1": ("12", None)},
            })
            engine = FormualizerAdapter.load(source)
            engine.workbook.get_eval_plan([("Calc", 2, 1)])

            changes = engine.apply([(0, 1, 0, 9)])

            self.assertEqual(engine.get(1, 0, 0), 13)
            self.assertEqual(engine.get(2, 0, 0), 26)
            self.assertIn((1, 0, 0, 13), changes)
            self.assertIn((2, 0, 0, 26), changes)

    def test_dependency_trace_budget_covers_all_range_members(self):
        from pp.engine.formualizer import FormualizerAdapter

        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "range-budget.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            for row in range(1, 11):
                sheet.cell(row, 1, row)
            sheet["B1"] = "=SUM(A1:A10)"
            workbook.save(source)
            _replace_formula_values(
                source, {"xl/worksheets/sheet1.xml": {"B1": ("55", None)}}
            )
            engine = FormualizerAdapter.load(source)

            changes = engine.apply([(0, 9, 0, 20)])

            self.assertIn((0, 0, 1, 65), changes)

    def test_whole_column_dependency_recalculates_downstream_after_edit(self):
        from pp.engine.formualizer import FormualizerAdapter

        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "whole-column.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"], sheet["A2"] = 1, 2
            sheet["B1"], sheet["C1"] = "=SUM(A:A)", "=B1*2"
            workbook.save(source)
            _replace_formula_values(source, {"xl/worksheets/sheet1.xml": {
                "B1": ("3", None), "C1": ("6", None),
            }})

            changes = FormualizerAdapter.load(source).apply([(0, 1, 0, 4)])

            self.assertIn((0, 0, 1, 5), changes)
            self.assertIn((0, 0, 2, 10), changes)

    def test_whole_row_dependency_recalculates_downstream_after_edit(self):
        from pp.engine.formualizer import FormualizerAdapter

        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "whole-row.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"], sheet["B1"] = 1, 2
            sheet["A2"], sheet["B2"] = "=SUM(1:1)", "=A2*2"
            workbook.save(source)
            _replace_formula_values(source, {"xl/worksheets/sheet1.xml": {
                "A2": ("3", None), "B2": ("6", None),
            }})

            changes = FormualizerAdapter.load(source).apply([(0, 0, 1, 4)])

            self.assertIn((0, 1, 0, 5), changes)
            self.assertIn((0, 1, 1, 10), changes)

    def test_quoted_sheet_whole_column_recalculates_downstream_after_edit(self):
        from pp.engine.formualizer import FormualizerAdapter

        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "quoted-sheet-column.xlsx"
            workbook = Workbook()
            inputs = workbook.active
            inputs.title = "Sheet 2"
            inputs["A1"], inputs["A2"] = 1, 2
            calc = workbook.create_sheet("Calc")
            calc["A1"], calc["B1"] = "=SUM('Sheet 2'!$A:$A)", "=A1*2"
            workbook.save(source)
            _replace_formula_values(source, {"xl/worksheets/sheet2.xml": {
                "A1": ("3", None), "B1": ("6", None),
            }})

            changes = FormualizerAdapter.load(source).apply([(0, 1, 0, 4)])

            self.assertIn((1, 0, 0, 5), changes)
            self.assertIn((1, 0, 1, 10), changes)

    def test_escaped_sheet_name_whole_column_recalculates_after_edit(self):
        from pp.engine.formualizer import FormualizerAdapter

        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "escaped-sheet-column.xlsx"
            workbook = Workbook()
            inputs = workbook.active
            inputs.title = "Owner's Data"
            inputs["A1"], inputs["A2"] = 1, 2
            calc = workbook.create_sheet("Calc")
            calc["A1"], calc["B1"] = "=SUM('Owner''s Data'!$A:$A)", "=A1*2"
            workbook.save(source)
            _replace_formula_values(source, {"xl/worksheets/sheet2.xml": {
                "A1": ("3", None), "B1": ("6", None),
            }})

            changes = FormualizerAdapter.load(source).apply([(0, 1, 0, 4)])

            self.assertIn((1, 0, 0, 5), changes)
            self.assertIn((1, 0, 1, 10), changes)

    def test_dependency_trace_budget_covers_defined_name_members(self):
        from pp.engine.formualizer import FormualizerAdapter

        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "named-range-budget.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            for row in range(1, 11):
                sheet.cell(row, 1, row)
            workbook.defined_names.add(
                DefinedName("PriceValues", attr_text="'Sheet'!$A$1:$A$10")
            )
            sheet["B1"] = "=SUM(PriceValues)"
            sheet["C1"] = "=B1*2"
            workbook.save(source)
            _replace_formula_values(
                source, {"xl/worksheets/sheet1.xml": {
                    "B1": ("55", None), "C1": ("110", None),
                }}
            )
            engine = FormualizerAdapter.load(source)

            changes = engine.apply([(0, 9, 0, 20)])

            self.assertIn((0, 0, 1, 65), changes)
            self.assertIn((0, 0, 2, 130), changes)

    def test_defined_name_does_not_capture_let_local(self):
        from pp.engine.formualizer import FormualizerAdapter

        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "let-name-shadow.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = 2
            sheet["A10"] = 100
            workbook.defined_names.add(
                DefinedName("v", attr_text="'Sheet'!$A$10")
            )
            sheet["B1"] = "=LET(v,A1+1,v*2)"
            workbook.save(source)
            _replace_formula_values(
                source, {"xl/worksheets/sheet1.xml": {"B1": ("6", None)}}
            )

            engine = FormualizerAdapter.load(source)
            changes = engine.apply([(0, 0, 0, 3)])

            self.assertIn((0, 0, 1, 8), changes)

    def test_short_defined_name_expands_without_capturing_let_local(self):
        from pp.engine.formualizer import FormualizerAdapter

        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "short-name-shadow.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = 2
            workbook.defined_names.add(
                DefinedName("Tax", attr_text="'Sheet'!$A$1")
            )
            sheet["B1"] = "=Tax*2"
            sheet["C1"] = "=B1+1"
            sheet["D1"] = "=LET(Tax,A1+1,Tax*3)"
            workbook.save(source)
            _replace_formula_values(
                source, {"xl/worksheets/sheet1.xml": {
                    "B1": ("4", None), "C1": ("5", None), "D1": ("9", None),
                }}
            )

            engine = FormualizerAdapter.load(source)
            changes = engine.apply([(0, 0, 0, 4)])

            self.assertIn((0, 0, 1, 8), changes)
            self.assertIn((0, 0, 2, 9), changes)
            self.assertIn((0, 0, 3, 15), changes)

    def test_out_of_bounds_cell_like_defined_name_expands(self):
        from pp.engine.formualizer import FormualizerAdapter

        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "out-of-bounds-name.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = 2
            workbook.defined_names.add(
                DefinedName("ZZZ1", attr_text="'Sheet'!$A$1")
            )
            sheet["B1"] = "=ZZZ1*2"
            sheet["C1"] = "=B1+1"
            workbook.save(source)
            _replace_formula_values(
                source, {"xl/worksheets/sheet1.xml": {
                    "B1": ("4", None), "C1": ("5", None),
                }}
            )

            changes = FormualizerAdapter.load(source).apply([(0, 0, 0, 4)])

            self.assertIn((0, 0, 1, 8), changes)
            self.assertIn((0, 0, 2, 9), changes)

    def test_complex_defined_name_recalculates_on_consecutive_edits(self):
        from pp.engine.formualizer import FormualizerAdapter

        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "complex-name.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"], sheet["A2"] = 1, 2
            workbook.defined_names.add(
                DefinedName("Total", attr_text="SUM(Sheet!$A$1:$A$2)")
            )
            sheet["B1"] = "=Total*2"
            sheet["C1"] = "=B1+1"
            workbook.save(source)
            _replace_formula_values(
                source, {"xl/worksheets/sheet1.xml": {
                    "B1": ("6", None), "C1": ("7", None),
                }}
            )
            engine = FormualizerAdapter.load(source)

            first = engine.apply([(0, 0, 0, 3)])
            second = engine.apply([(0, 1, 0, 4)])

            self.assertIn((0, 0, 1, 10), first)
            self.assertIn((0, 0, 2, 11), first)
            self.assertIn((0, 0, 1, 14), second)
            self.assertIn((0, 0, 2, 15), second)

    def test_structured_table_reference_fails_before_cached_value_can_go_stale(self):
        from pp.engine.formualizer import EngineError, FormualizerAdapter

        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "structured-reference.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Price", "Total"])
            sheet.append([1])
            sheet.append([2])
            sheet.add_table(Table(displayName="Prices", ref="A1:A3"))
            sheet["B1"] = "=SUM(Prices[Price])"
            workbook.save(source)
            _replace_formula_values(
                source, {"xl/worksheets/sheet1.xml": {"B1": ("3", None)}}
            )

            with self.assertRaisesRegex(
                EngineError, "structured or external formula reference"
            ):
                FormualizerAdapter.load(source)

    def test_named_range_operator_fails_before_cached_value_can_go_stale(self):
        from pp.engine.formualizer import EngineError, FormualizerAdapter

        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "named-range-operator.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"], sheet["B1"] = 1, 2
            workbook.defined_names.add(
                DefinedName("PriceStart", attr_text="'Sheet'!$A$1")
            )
            sheet["C1"] = "=SUM(PriceStart:B1)"
            workbook.save(source)
            _replace_formula_values(
                source, {"xl/worksheets/sheet1.xml": {"C1": ("3", None)}}
            )

            with self.assertRaisesRegex(EngineError, "named range operator"):
                FormualizerAdapter.load(source)

    def test_ordinary_let_local_is_not_a_dynamic_target(self):
        from pp.engine.formualizer import FormualizerAdapter

        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "ordinary-let.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = 2
            sheet["B1"] = "=LET(v,A1+1,v*2)"
            workbook.save(source)
            _replace_formula_values(
                source, {"xl/worksheets/sheet1.xml": {"B1": ("6", None)}}
            )

            engine = FormualizerAdapter.load(source)

            self.assertNotIn((0, 0, 1), engine.dynamic_targets)

    def test_spill_bootstrap_failure_is_exposed_as_engine_error(self):
        from pp.engine.formualizer import EngineError, FormualizerAdapter

        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "spill-bootstrap.xlsx"
            workbook = Workbook()
            workbook.active["A1"] = "=SEQUENCE(2)"
            workbook.save(source)
            with mock.patch(
                "pp.engine.formualizer._spill_map",
                side_effect=RuntimeError("native spill inspection failed"),
            ):
                with self.assertRaisesRegex(EngineError, "failed to load"):
                    FormualizerAdapter.load(source)

    def test_dependency_trace_budget_covers_dense_reference_fanout(self):
        from pp.engine.formualizer import FormualizerAdapter

        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "link-budget.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = 1
            for row in range(1, 7):
                sheet.cell(row, 24, "=$A$1")
            for row in range(1, 11):
                sheet.cell(row, 26, "=SUM($X$1:$X$6)")
            workbook.save(source)
            cached = {f"X{row}": ("1", None) for row in range(1, 7)}
            cached.update({f"Z{row}": ("6", None) for row in range(1, 11)})
            _replace_formula_values(source, {"xl/worksheets/sheet1.xml": cached})
            engine = FormualizerAdapter.load(source)

            changes = engine.apply([(0, 0, 0, 2)])

            self.assertEqual(engine.get(0, 9, 25), 12)
            self.assertIn((0, 9, 25, 12), changes)

    def test_trace_budget_beyond_native_u32_is_capped(self):
        """Whole-column references on a large template overflow the native budget type."""
        from pp.engine.formualizer import FormualizerAdapter, TraceBudget

        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "huge-budget.xlsx"
            workbook = Workbook()
            workbook.active["A1"] = 1
            workbook.active["B1"] = "=A1*2"
            workbook.save(source)
            _replace_formula_values(source, {"xl/worksheets/sheet1.xml": {"B1": ("2", None)}})
            engine = FormualizerAdapter.load(source)
            engine.trace_budget = TraceBudget(2**40, 2**40, 2**40)

            changes = engine.apply([(0, 0, 0, 5)])

            self.assertIn((0, 0, 1, 10), changes)

    def test_cached_dynamic_formula_recalculates_its_downstream_closure(self):
        from pp.engine.formualizer import FormualizerAdapter

        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "cached-dynamic.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Calc"
            sheet["A2"] = 1
            sheet["B2"] = '=INDIRECT("A"&ROW())'
            sheet["C2"] = "=B2+1"
            workbook.save(source)
            _replace_formula_values(source, {"xl/worksheets/sheet1.xml": {
                "B2": ("1", None), "C2": ("2", None),
            }})
            engine = FormualizerAdapter.load(source)
            engine.workbook.get_eval_plan([("Calc", 2, 2)])

            changes = engine.apply([(0, 1, 0, 5)])

            self.assertEqual(engine.get(0, 1, 1), 5)
            self.assertEqual(engine.get(0, 1, 2), 6)
            self.assertIn((0, 1, 2, 6), changes)

    def test_grown_spill_recalculates_cached_downstream_formula(self):
        from pp.engine.formualizer import FormualizerAdapter

        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "cached-spill-growth.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = 1
            sheet["B1"] = "=SEQUENCE(A1)"
            sheet["C1"] = "=B2"
            workbook.save(source)
            _replace_formula_values(source, {"xl/worksheets/sheet1.xml": {
                "B1": ("1", None), "C1": ("0", None),
            }})
            engine = FormualizerAdapter.load(source)
            self.assertEqual(engine.get(0, 0, 2), 0)

            changes = engine.apply([(0, 0, 0, 2)])

            self.assertEqual(engine.get(0, 1, 1), 2)
            self.assertEqual(engine.get(0, 0, 2), 2)
            self.assertIn((0, 0, 2, 2), changes)

    def test_excel_settled_array_formula_skips_bootstrap_and_still_grows(self):
        """A large template must not evaluate every potential spill at load."""
        from openpyxl.worksheet.formula import ArrayFormula

        from pp.engine.formualizer import FormualizerAdapter, FormualizerEngine

        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "settled-array.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = 1
            sheet["B1"] = ArrayFormula("B1", "=SEQUENCE(A1)")
            sheet["C1"] = "=B2"
            workbook.save(source)
            _replace_formula_values(source, {"xl/worksheets/sheet1.xml": {
                "B1": ("1", None), "C1": ("0", None),
            }})
            with mock.patch.object(FormualizerEngine, "evaluate_dirty",
                                   side_effect=AssertionError("bootstrap evaluated")):
                engine = FormualizerAdapter.load(source)
            self.assertEqual(engine.get(0, 0, 2), 0)

            changes = engine.apply([(0, 0, 0, 2)])

            self.assertEqual(engine.get(0, 1, 1), 2)
            self.assertIn((0, 0, 2, 2), changes)

    def test_grown_spill_rechecks_same_batch_direct_dependent(self):
        from pp.engine.formualizer import FormualizerAdapter

        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "same-batch-spill.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["D1"] = 1
            sheet["B1"] = "=SEQUENCE(D1)"
            sheet["A1"] = "=B2+D1"
            workbook.save(source)
            _replace_formula_values(source, {"xl/worksheets/sheet1.xml": {
                "A1": ("1", None), "B1": ("1", None),
            }})
            engine = FormualizerAdapter.load(source)

            changes = engine.apply([(0, 0, 3, 2)])

            self.assertEqual(engine.get(0, 0, 0), 4)
            self.assertIn((0, 0, 0, 4), changes)

    def test_load_closes_dependents_after_cached_spill_growth(self):
        from pp.engine.formualizer import FormualizerAdapter

        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "load-spill-closure.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["D1"] = 2
            sheet["B1"] = "=SEQUENCE(D1)"
            sheet["A1"] = "=B2+D1"
            workbook.save(source)
            _replace_formula_values(source, {"xl/worksheets/sheet1.xml": {
                "A1": ("1", None), "B1": ("1", None),
            }})

            engine = FormualizerAdapter.load(source)

            self.assertEqual(engine.get(0, 0, 0), 4)

    def test_shrunk_spill_updates_cached_downstream_overlay_and_restores(self):
        from pp.engine.worker import WarmCopy

        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "cached-spill-shrink.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = 2
            sheet["B1"] = "=SEQUENCE(A1)"
            sheet["C1"] = "=B2"
            workbook.save(source)
            _replace_formula_values(source, {"xl/worksheets/sheet1.xml": {
                "B1": ("1", None), "C1": ("2", None),
            }})
            copy = WarmCopy("cached-spill-shrink")
            self.assertEqual(copy.apply(0, str(source), [])["results"], [])

            shrunk = copy.apply(1, str(source), [(0, 0, 0, 1)])
            restored = copy.apply(2, str(source), [])

            self.assertEqual(shrunk["results"], [(0, 0, 2, 0), (0, 1, 1, None)])
            self.assertEqual(restored["results"], [])
            self.assertEqual(copy.engine.get(0, 1, 1), 2)
            self.assertEqual(copy.engine.get(0, 0, 2), 2)

    def test_chained_spill_growth_reaches_cached_downstream_formula(self):
        from pp.engine.formualizer import FormualizerAdapter

        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "cached-chained-spill.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = 1
            sheet["B1"] = "=SEQUENCE(A1)"
            sheet["C1"] = "=SEQUENCE(IF(B2=0,1,B2))"
            sheet["D1"] = "=C2"
            workbook.save(source)
            _replace_formula_values(source, {"xl/worksheets/sheet1.xml": {
                "B1": ("1", None), "C1": ("1", None), "D1": ("0", None),
            }})
            engine = FormualizerAdapter.load(source)

            changes = engine.apply([(0, 0, 0, 2)])

            self.assertEqual(engine.get(0, 1, 1), 2)
            self.assertEqual(engine.get(0, 1, 2), 2)
            self.assertEqual(engine.get(0, 0, 3), 2)
            self.assertIn((0, 0, 3, 2), changes)

    def test_spill_closure_does_not_repeat_unrelated_dynamic_target(self):
        from pp.engine.formualizer import FormualizerAdapter

        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "spill-with-dynamic.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = 1
            sheet["B1"] = "=SEQUENCE(A1)"
            sheet["C1"] = "=SEQUENCE(IF(B2=0,1,B2))"
            sheet["D1"] = "=C2"
            sheet["E1"] = "=NOW()"
            workbook.save(source)
            _replace_formula_values(source, {"xl/worksheets/sheet1.xml": {
                "B1": ("1", None), "C1": ("1", None), "D1": ("0", None),
                "E1": ("46000", None),
            }})
            engine = FormualizerAdapter.load(source)
            recording = _RecordingEvaluationWorkbook(engine.workbook)
            engine.workbook = recording

            engine.apply([(0, 0, 0, 2)])

            evaluated = [target for batch in recording.calls for target in batch]
            self.assertEqual(evaluated.count(("Sheet", 1, 5)), 1)

    def test_spill_closure_trace_truncation_fails_closed(self):
        from pp.engine.formualizer import EngineError, FormualizerAdapter

        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "truncated-spill-closure.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = 1
            sheet["B1"] = "=SEQUENCE(A1)"
            sheet["C1"] = "=B2"
            workbook.save(source)
            _replace_formula_values(source, {"xl/worksheets/sheet1.xml": {
                "B1": ("1", None), "C1": ("0", None),
            }})
            engine = FormualizerAdapter.load(source)
            engine.workbook = _TruncatedAfterTrace(engine.workbook, complete_calls=1)

            with self.assertRaisesRegex(EngineError, "dependency trace truncated"):
                engine.apply([(0, 0, 0, 2)])

    def test_nonconverging_spill_closure_fails_closed(self):
        from pp.engine.formualizer import EngineError, FormualizerAdapter

        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "nonconverging-spill-closure.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = 1
            sheet["B1"] = "=SEQUENCE(A1)"
            sheet["C1"] = "=SEQUENCE(A1)"
            workbook.save(source)
            engine = FormualizerAdapter.load(source)
            engine.workbook = _CyclingSpillWorkbook(engine.workbook)

            with self.assertRaisesRegex(EngineError, "spill dependency closure did not converge"):
                engine.apply([(0, 0, 0, 2)])

    def test_empty_apply_never_evaluates_or_inspects_spills(self):
        from pp.engine.formualizer import FormualizerAdapter

        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "empty-apply.xlsx"
            workbook = Workbook()
            workbook.active["A1"] = "=1"
            workbook.save(source)
            _replace_formula_values(source, {"xl/worksheets/sheet1.xml": {"A1": ("1", None)}})
            native_load = __import__("pp.engine.formualizer", fromlist=["formualizer"]).formualizer.load_workbook
            counts = {"evaluations": 0, "inspections": 0}

            class CountingWorkbook:
                def __init__(self, native):
                    self.native = native

                def evaluate_cells(self, targets):
                    counts["evaluations"] += 1
                    return self.native.evaluate_cells(targets)

                def inspect_cell(self, address):
                    counts["inspections"] += 1
                    return self.native.inspect_cell(address)

                def __getattr__(self, name):
                    return getattr(self.native, name)

            with mock.patch("pp.engine.formualizer.formualizer.load_workbook",
                            side_effect=lambda path: CountingWorkbook(native_load(path))):
                engine = FormualizerAdapter.load(source)
            counts.update(evaluations=0, inspections=0)

            self.assertEqual(engine.apply([]), [])
            self.assertEqual(counts, {"evaluations": 0, "inspections": 0})

    def test_cached_1904_scalar_cache_keeps_excel_serial(self):
        from openpyxl.utils.datetime import MAC_EPOCH
        from pp.engine.formualizer import FormualizerAdapter

        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "cached-1904.xlsx"
            workbook = Workbook()
            workbook.epoch = MAC_EPOCH
            workbook.active["A1"] = "=DATE(2026,1,1)"
            workbook.save(source)
            _replace_formula_values(source, {"xl/worksheets/sheet1.xml": {"A1": ("44561", None)}})

            engine = FormualizerAdapter.load(source)

            self.assertEqual(engine.get(0, 0, 0), 44561)

    def test_warm_clear_recalculates_formula_after_rebuild_with_old_xml_cache(self):
        from pp.engine.worker import WarmCopy

        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "warm-clear.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = 3
            sheet["B1"] = "=A1*2"
            workbook.save(source)
            _replace_formula_values(source, {"xl/worksheets/sheet1.xml": {"B1": ("6", None)}})
            copy = WarmCopy("cached-clear")
            copy.apply(0, str(source), [])

            cleared = copy.apply(1, str(source), [(0, 0, 0, None)])

            results = {(sheet, row, col): value for sheet, row, col, value in cleared["results"]}
            self.assertEqual(results[(0, 0, 1)], 0)

    def test_warm_clear_preserves_baseline_spill_topology_and_downstream_restore(self):
        from pp.engine.worker import WarmCopy

        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "warm-clear-spill.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = 1
            sheet["B1"] = "=SEQUENCE(A1+1)"
            sheet["D1"] = "=B2"
            workbook.save(source)
            _replace_formula_values(source, {"xl/worksheets/sheet1.xml": {
                "B1": ("1", None), "D1": ("2", None),
            }})
            copy = WarmCopy("warm-clear-spill")
            self.assertEqual(copy.apply(0, str(source), [])["results"], [])

            cleared = copy.apply(1, str(source), [(0, 0, 0, None)])
            restored = copy.apply(2, str(source), [])

            self.assertEqual(cleared["results"], [(0, 0, 3, 0), (0, 1, 1, None)])
            self.assertEqual(restored["results"], [])
            self.assertEqual(copy.engine.get(0, 0, 0), 1)
            self.assertEqual(copy.engine.get(0, 1, 1), 2)
            self.assertEqual(copy.engine.get(0, 0, 3), 2)

    def test_restoring_blank_input_unblocks_grown_template_spill(self):
        from pp.engine.worker import WarmCopy

        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "restore-blank-spill-blocker.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.protection.sheet = True
            sheet["A1"] = 1
            sheet["A1"].protection = Protection(locked=False)
            sheet["B1"] = "=SEQUENCE(A1)"
            sheet["B2"].protection = Protection(locked=False)
            workbook.save(source)
            copy = WarmCopy("restore-blank-spill-blocker")
            copy.apply(0, str(source), [])
            copy.apply(1, str(source), [(0, 0, 0, 2)])

            copy.apply(2, str(source), [(0, 0, 0, 2), (0, 1, 1, 5)])
            restored = copy.apply(3, str(source), [(0, 0, 0, 2)])

            self.assertIsNone(copy.baseline[(0, 1, 1)])
            self.assertEqual(restored["results"], [(0, 1, 1, 2)])
            self.assertEqual(copy.engine.get(0, 0, 0), 2)
            self.assertEqual(copy.engine.get(0, 0, 1), 1)
            self.assertEqual(copy.engine.get(0, 1, 1), 2)

    def test_first_touch_clear_records_blank_for_a_current_spill_ghost(self):
        from pp.engine.worker import WarmCopy

        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "clear-first-spill-ghost.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = 1
            sheet["B1"] = "=SEQUENCE(A1)"
            workbook.save(source)
            copy = WarmCopy("clear-first-spill-ghost")
            copy.apply(0, str(source), [])
            copy.apply(1, str(source), [(0, 0, 0, 2)])

            copy.apply(2, str(source), [(0, 0, 0, 2), (0, 1, 1, None)])

            self.assertIsNone(copy.baseline[(0, 1, 1)])

    def test_cached_potential_spill_child_is_not_imported_as_an_input(self):
        import json
        from openpyxl.styles import Protection
        from pp.template.importer import import_template

        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "cached-spill.xlsx"
            snapshot = Path(temp) / "snapshot"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Spill"
            sheet.protection.sheet = True
            sheet["A1"] = 2
            sheet["B1"] = "=SEQUENCE(A1)"
            sheet["B2"].protection = Protection(locked=False)
            workbook.save(source)
            _replace_formula_values(source, {"xl/worksheets/sheet1.xml": {"B1": ("1", None)}})

            imported = import_template(source, snapshot)

            mask = json.loads((imported / "input_mask.json").read_text())
            self.assertNotIn("0:1:1", mask)

    def test_incomplete_native_trace_fails_with_named_dependency_error(self):
        from pp.engine.formualizer import EngineError, FormualizerAdapter

        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "truncated.xlsx"
            workbook = Workbook()
            workbook.active["A1"] = 1
            workbook.active["B1"] = "=A1+1"
            workbook.save(source)
            engine = FormualizerAdapter.load(source)
            truncated = SimpleNamespace(truncation=SimpleNamespace(incomplete=True), nodes={})
            native = engine.workbook

            class TruncatedTraceWorkbook:
                def trace(self, *args, **kwargs):
                    return truncated

                def __getattr__(self, name):
                    return getattr(native, name)

            engine.workbook = TruncatedTraceWorkbook()
            with self.assertRaisesRegex(EngineError, "dependency trace truncated"):
                engine.apply([(0, 0, 0, 2)])


class _TruncatedAfterTrace:
    def __init__(self, native, complete_calls):
        self.native = native
        self.complete_calls = complete_calls
        self.calls = 0

    def trace(self, *args, **kwargs):
        self.calls += 1
        if self.calls <= self.complete_calls:
            return self.native.trace(*args, **kwargs)
        return SimpleNamespace(truncation=SimpleNamespace(incomplete=True), nodes={})

    def __getattr__(self, name):
        return getattr(self.native, name)


class _RecordingEvaluationWorkbook:
    def __init__(self, native):
        self.native = native
        self.calls = []

    def evaluate_cells(self, targets):
        self.calls.append(list(targets))
        return self.native.evaluate_cells(targets)

    def __getattr__(self, name):
        return getattr(self.native, name)


class _CyclingSpillWorkbook:
    def __init__(self, native):
        self.native = native
        self.calls = 0

    def evaluate_cells(self, targets):
        self.calls += 1
        self.native.set_value("Sheet", 1, 1, 2 if self.calls % 2 else 1)
        return self.native.evaluate_cells(targets)

    def __getattr__(self, name):
        return getattr(self.native, name)


def _replace_formula_values(path, values):
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    with zipfile.ZipFile(path) as source:
        members = {name: source.read(name) for name in source.namelist()}
    from xml.etree import ElementTree as ET
    for member, replacements in values.items():
        root = ET.fromstring(members[member])
        for cell in root.findall(f".//{{{namespace}}}c"):
            address = cell.get("r")
            if address not in replacements:
                continue
            value, kind = replacements[address]
            if kind is None:
                cell.attrib.pop("t", None)
            else:
                cell.set("t", kind)
            node = cell.find(f"{{{namespace}}}v")
            if node is None:
                node = ET.SubElement(cell, f"{{{namespace}}}v")
            node.text = value
        members[member] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    with zipfile.ZipFile(path, "w") as target:
        for name, data in members.items():
            target.writestr(name, data)


def _relative_shared_fixture(path):
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"

    def tag(name):
        return f"{{{namespace}}}{name}"

    generate(path, 0.0001, 20261011)
    with zipfile.ZipFile(path) as source:
        members = {name: source.read(name) for name in source.namelist()}
    inputs = ET.fromstring(members["xl/worksheets/sheet1.xml"])
    input_cells = {cell.get("r"): cell for cell in inputs.findall(f".//{tag('c')}")}
    for reference, value in (("J1", "1"), ("K1", "2")):
        input_cells[reference].find(tag("v")).text = value
    members["xl/worksheets/sheet1.xml"] = ET.tostring(
        inputs, encoding="utf-8", xml_declaration=True
    )
    root = ET.fromstring(members["xl/worksheets/sheet2.xml"])
    shared = [cell for cell in root.findall(f".//{tag('c')}")
              if (formula := cell.find(tag("f"))) is not None and formula.get("t") == "shared"]
    master = next(cell for cell in shared if cell.find(tag("f")).text)
    shared_id = master.find(tag("f")).get("si")
    followers = [cell for cell in shared
                 if cell.find(tag("f")).get("si") == shared_id and not cell.find(tag("f")).text]
    follower = followers[0]
    master.find(tag("f")).text = "S01!J1+S01!$A$1"
    master.find(tag("v")).text = "8"
    follower.find(tag("v")).text = "9"
    members["xl/worksheets/sheet2.xml"] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    with zipfile.ZipFile(path, "w") as target:
        for name, data in members.items():
            target.writestr(name, data)
    letters = "".join(character for character in follower.get("r") if character.isalpha())
    row = int("".join(character for character in follower.get("r") if character.isdigit()))
    column = 0
    for letter in letters:
        column = column * 26 + ord(letter) - 64
    return 1, row - 1, column - 1


if __name__ == "__main__":
    unittest.main()
