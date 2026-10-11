"""Synthetic contract tests for safe workbook download transforms."""

import io
import os
from pathlib import Path
import shutil
import time
import unittest
from unittest import mock
import zipfile
import xml.etree.ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName

from pp.export.transfer import TransferRequest, export_transfer
from tests.v2.api_support import APIFixture
from tests.v2.support import SYNTHETIC_LOGIN


MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
OFFICE_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _replace_members(path, replacements, additions=()):
    """Rewrite a synthetic ZIP package while keeping its original member order."""
    temporary = path.with_suffix(".rewrite")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(temporary, "w") as target:
        for info in source.infolist():
            target.writestr(info, replacements.get(info.filename, source.read(info.filename)))
        for name, data in additions:
            target.writestr(name, data)
    shutil.move(temporary, path)


def _package_fixture(path):
    """Create a workbook plus parts that must disappear with the Drop sheet."""
    workbook = Workbook()
    keep = workbook.active
    keep.title = "Keep"
    keep["A1"] = "kept"
    keep["B2"] = "=Drop!A1*2"
    keep["B2"].number_format = "0.00%"
    dropped = workbook.create_sheet("Drop")
    dropped["A1"] = 21
    hidden = workbook.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    workbook.defined_names.add(DefinedName("DropGlobal", attr_text="'Drop'!$A$1"))
    workbook.defined_names.add(DefinedName("DropLocal", attr_text="'Drop'!$A$1", localSheetId=1))
    workbook.defined_names.add(DefinedName("_xlnm.Print_Area", attr_text="'Drop'!$A$1", localSheetId=1))
    workbook.save(path)

    with zipfile.ZipFile(path) as archive:
        workbook_xml = archive.read("xl/workbook.xml")
        workbook_rels = archive.read("xl/_rels/workbook.xml.rels")
        sheet_xml = archive.read("xl/worksheets/sheet2.xml")
        content_types = archive.read("[Content_Types].xml")
    pivot_caches = (
        b'<pivotCaches><pivotCache cacheId="1" r:id="rIdPivotCache"/>'
        b'</pivotCaches>'
    )
    workbook_xml = workbook_xml.replace(b"</workbook>", pivot_caches + b"</workbook>")
    workbook_rels = workbook_rels.replace(
        b"</Relationships>",
        b'<Relationship Id="rIdPivotCache" '
        b'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/pivotCacheDefinition" '
        b'Target="pivotCache/pivotCacheDefinition1.xml"/></Relationships>',
    )
    sheet_xml = sheet_xml.replace(
        b"</worksheet>",
        b'<drawing r:id="rIdDrawing"/><pivotTableParts count="1">'
        b'<pivotTablePart r:id="rIdPivot"/></pivotTableParts>'
        b'<extLst><ext uri="{slicer}"><x14:slicerList r:id="rIdSlicer"/></ext>'
        b'</extLst></worksheet>',
    )
    sheet_xml = sheet_xml.replace(
        b"<worksheet ",
        b'<worksheet xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        b'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main" ',
        1,
    )
    content_types = content_types.replace(
        b"</Types>",
        b'<Override PartName="/xl/drawings/drawing1.xml" '
        b'ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/>'
        b'<Override PartName="/xl/pivotTables/pivotTable1.xml" '
        b'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.pivotTable+xml"/>'
        b'<Override PartName="/xl/pivotCache/pivotCacheDefinition1.xml" '
        b'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.pivotCacheDefinition+xml"/>'
        b'<Override PartName="/xl/pivotCache/pivotCacheRecords1.xml" '
        b'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.pivotCacheRecords+xml"/>'
        b'<Override PartName="/xl/slicers/slicer1.xml" '
        b'ContentType="application/vnd.ms-excel.slicer+xml"/>'
        b'<Override PartName="/xl/slicerCaches/slicerCache1.xml" '
        b'ContentType="application/vnd.ms-excel.slicerCache+xml"/>'
        b'<Override PartName="/xl/calcChain.xml" '
        b'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml"/>'
        b"</Types>",
    )
    sheet_rels = (
        b'<?xml version="1.0" encoding="UTF-8"?>'
        b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        b'<Relationship Id="rIdDrawing" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing" Target="../drawings/drawing1.xml"/>'
        b'<Relationship Id="rIdPivot" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/pivotTable" Target="../pivotTables/pivotTable1.xml"/>'
        b'<Relationship Id="rIdSlicer" Type="http://schemas.microsoft.com/office/2007/relationships/slicer" Target="../slicers/slicer1.xml"/>'
        b'</Relationships>'
    )
    additions = (
        ("xl/worksheets/_rels/sheet2.xml.rels", sheet_rels),
        ("xl/drawings/drawing1.xml", b'<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"/>'),
        ("xl/pivotTables/pivotTable1.xml", b'<pivotTableDefinition xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"/>'),
        ("xl/pivotTables/_rels/pivotTable1.xml.rels", b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rIdCache" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/pivotCacheDefinition" Target="../pivotCache/pivotCacheDefinition1.xml"/></Relationships>'),
        ("xl/pivotCache/pivotCacheDefinition1.xml", b'<pivotCacheDefinition xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><cacheSource type="worksheet"><worksheetSource sheet="Drop" ref="A1"/></cacheSource></pivotCacheDefinition>'),
        ("xl/pivotCache/_rels/pivotCacheDefinition1.xml.rels", b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rIdRecords" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/pivotCacheRecords" Target="pivotCacheRecords1.xml"/></Relationships>'),
        ("xl/pivotCache/pivotCacheRecords1.xml", b'<pivotCacheRecords xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="0"/>'),
        ("xl/slicers/slicer1.xml", b'<x14:slicers xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"><x14:slicer name="DropSlicer"/></x14:slicers>'),
        ("xl/slicers/_rels/slicer1.xml.rels", b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rIdSlicerCache" Type="http://schemas.microsoft.com/office/2007/relationships/slicerCache" Target="../slicerCaches/slicerCache1.xml"/></Relationships>'),
        ("xl/slicerCaches/slicerCache1.xml", b'<x14:slicerCacheDefinition xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"/>'),
        ("xl/calcChain.xml", b'<calcChain xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><c r="B2" i="0"/></calcChain>'),
    )
    _replace_members(path, {
        "xl/workbook.xml": workbook_xml,
        "xl/_rels/workbook.xml.rels": workbook_rels,
        "xl/worksheets/sheet2.xml": sheet_xml,
        "[Content_Types].xml": content_types,
    }, additions)


def _selection_fixture(path):
    workbook = Workbook()
    workbook.active.title = "Visible"
    workbook.create_sheet("Cost, Fees")
    hidden = workbook.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    workbook.save(path)


class PackageDownloadTests(unittest.TestCase):
    def setUp(self):
        self.temp = Path(self.id().replace(".", "-"))
        self.addCleanup(lambda: shutil.rmtree(self.temp, ignore_errors=True))
        self.temp.mkdir()
        self.source = self.temp / "source.xlsx"
        self.output = self.temp / "download.xlsx"
        _package_fixture(self.source)

    def download(self, selected={"Keep"}, results={(0, 1, 1): 42}):
        return export_transfer(
            TransferRequest(
                template_path=self.source, cells=[], out_path=self.output,
                template_hash="template-hash", copy_id="copy-id", input_revision=3,
                fmt="xlsx", selected=selected, results=results,
            )
        )

    def members(self):
        with zipfile.ZipFile(self.output) as archive:
            return set(archive.namelist())

    def xml(self, name):
        with zipfile.ZipFile(self.output) as archive:
            return archive.read(name)

    def test_removed_sheet_part_and_relationship_are_absent(self):
        self.download()
        self.assertNotIn("xl/worksheets/sheet2.xml", self.members())
        self.assertNotIn(b'name="Drop"', self.xml("xl/workbook.xml"))
        self.assertNotIn(b'Target="worksheets/sheet2.xml"', self.xml("xl/_rels/workbook.xml.rels"))

    def test_removed_global_defined_name_is_absent(self):
        self.download()
        self.assertNotIn(b'DropGlobal', self.xml("xl/workbook.xml"))

    def test_removed_local_defined_name_is_absent(self):
        self.download()
        self.assertNotIn(b'DropLocal', self.xml("xl/workbook.xml"))

    def test_removed_print_area_is_absent(self):
        self.download()
        self.assertNotIn(b'Print_Area', self.xml("xl/workbook.xml"))

    def test_removed_pivot_cache_is_absent(self):
        self.download()
        self.assertNotIn("xl/pivotCache/pivotCacheDefinition1.xml", self.members())
        self.assertNotIn(b'rIdPivotCache', self.xml("xl/workbook.xml"))

    def test_removed_pivot_table_is_absent(self):
        self.download()
        self.assertNotIn("xl/pivotTables/pivotTable1.xml", self.members())

    def test_removed_slicer_is_absent(self):
        self.download()
        self.assertNotIn("xl/slicers/slicer1.xml", self.members())
        self.assertNotIn("xl/slicerCaches/slicerCache1.xml", self.members())

    def test_removed_drawing_is_absent(self):
        self.download()
        self.assertNotIn("xl/drawings/drawing1.xml", self.members())

    def test_calc_chain_is_absent(self):
        self.download()
        self.assertNotIn("xl/calcChain.xml", self.members())

    def test_cross_sheet_formula_is_frozen_to_current_result_with_its_style(self):
        self.download()
        with zipfile.ZipFile(self.output) as archive:
            root = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        cell = root.find(f".//{{{MAIN_NS}}}c[@r='B2']")
        self.assertIsNone(cell.find(f"{{{MAIN_NS}}}f"))
        self.assertEqual(cell.findtext(f"{{{MAIN_NS}}}v"), "42")
        self.assertEqual(cell.get("s"), "1")

    def test_missing_current_result_refuses_to_substitute_a_template_cache(self):
        with self.assertRaisesRegex(ValueError, "current calculation result"):
            self.download(results={})

    def test_output_has_custom_property_root_relationship_and_content_type(self):
        self.download()
        root_rels = ET.fromstring(self.xml("_rels/.rels"))
        rels = root_rels.findall(f"{{{PKG_REL_NS}}}Relationship")
        self.assertTrue(any(
            rel.get("Type") == f"{OFFICE_REL_NS}/custom-properties"
            and rel.get("Target") == "docProps/custom.xml"
            for rel in rels
        ))
        properties = ET.fromstring(self.xml("docProps/custom.xml"))
        names = {item.get("name") for item in properties}
        self.assertEqual(names & {"TemplateHash", "CopyId", "InputRevision"},
                         {"TemplateHash", "CopyId", "InputRevision"})
        self.assertIn(b'PartName="/docProps/custom.xml"', self.xml("[Content_Types].xml"))

    def test_output_is_structurally_valid_and_openpyxl_openable(self):
        from pp.export.transfer import validate_package

        self.download()
        validate_package(self.output, require_provenance=True)
        workbook = load_workbook(self.output, read_only=True)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.sheetnames, ["Keep"])

    def test_all_sheet_transform_keeps_workbook_namespace_bytes(self):
        with zipfile.ZipFile(self.source) as archive:
            original = archive.read("xl/workbook.xml")
        marked = original.replace(
            b"<workbook ",
            b'<workbook xmlns:mc="urn:mc" xmlns:x15="urn:x15" mc:Ignorable="x15" ',
            1,
        )
        _replace_members(self.source, {"xl/workbook.xml": marked})
        export_transfer(
            TransferRequest(
                template_path=self.source, cells=[], out_path=self.output,
                template_hash="template-hash", copy_id="copy-id", input_revision=3,
                fmt="xlsm", results={},
            )
        )
        self.assertEqual(self.xml("xl/workbook.xml"), marked)

    def test_writer_calls_structural_validation_before_returning(self):
        with mock.patch("pp.export.transfer.validate_package", side_effect=ValueError("broken package")):
            with self.assertRaisesRegex(ValueError, "broken package"):
                self.download()


class DownloadRouteSelectionTests(unittest.TestCase):
    def setUp(self):
        self.fixture = APIFixture(_selection_fixture)
        self.client, self.csrf, _ = self.fixture.signup("transfer@example.test", "Transfer Owner")
        _, self.copy = self.fixture.make_copy(self.client, self.csrf)

    def tearDown(self):
        self.fixture.close()

    def prepare(self, client=None, csrf=None, sheets=None):
        return (client or self.client).post(
            f"/api/copies/{self.copy['id']}/downloads",
            {"format": "xlsx", "sheets": sheets or ["Visible"]},
            {"x-csrf-token": csrf or self.csrf},
        )

    def prepared_path(self, response):
        token = response.json()["download_url"].rsplit("/", 1)[-1]
        return self.fixture.app.state.services.downloads._items[token].path

    def test_repeated_parameters_select_comma_named_sheet(self):
        response = self.client.get(
            f"/api/copies/{self.copy['id']}/export.xlsx?sheets=Visible&sheets=Cost%2C%20Fees"
        )
        self.assertEqual(response.status_code, 200, response.body)
        workbook = load_workbook(io.BytesIO(response.body), read_only=True)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.sheetnames, ["Visible", "Cost, Fees"])

    def test_json_selection_preserves_comma_named_sheet(self):
        response = self.client.get(
            f"/api/copies/{self.copy['id']}/export.xlsx?sheets=%5B%22Visible%22%2C%22Cost%2C%20Fees%22%5D"
        )
        self.assertEqual(response.status_code, 200, response.body)
        workbook = load_workbook(io.BytesIO(response.body), read_only=True)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.sheetnames, ["Visible", "Cost, Fees"])

    def test_empty_unknown_and_all_hidden_selection_are_user_validation_errors(self):
        base = f"/api/copies/{self.copy['id']}/export.xlsx?sheets="
        for query in (base, base + "Unknown", base + "Hidden"):
            with self.subTest(query=query):
                response = self.client.get(query)
                self.assertEqual(response.status_code, 422, response.body)
                self.assertEqual(response.json()["error"], "invalid_selection")

    def test_explicit_hidden_sheet_is_preserved_when_a_visible_sheet_is_selected(self):
        response = self.client.get(
            f"/api/copies/{self.copy['id']}/export.xlsx?sheets=Visible&sheets=Hidden"
        )
        self.assertEqual(response.status_code, 200, response.body)
        workbook = load_workbook(io.BytesIO(response.body), read_only=True)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.sheetnames, ["Visible", "Hidden"])
        self.assertEqual(workbook["Hidden"].sheet_state, "hidden")

    def test_preflight_validates_selection_without_generating_a_file(self):
        with mock.patch("pp.server.export_transfer") as exporter:
            accepted = self.client.get(
                f"/api/copies/{self.copy['id']}/export.xlsx?check=1&sheets=Visible"
            )
            rejected = self.client.get(
                f"/api/copies/{self.copy['id']}/export.xlsx?check=1&sheets=Hidden"
            )

        self.assertEqual(accepted.status_code, 204, accepted.body)
        self.assertEqual(rejected.status_code, 422, rejected.body)
        exporter.assert_not_called()

    def test_prepared_download_survives_a_later_revision_and_generates_once(self):
        with mock.patch("pp.server.export_transfer", wraps=export_transfer) as exporter:
            prepared = self.prepare()
            self.assertEqual(prepared.status_code, 201, prepared.body)
            path = self.prepared_path(prepared)
            expected = path.read_bytes()
            with self.fixture.database.transaction() as connection:
                connection.execute(
                    "UPDATE pp_copies SET input_revision=1,calculated_revision=0 WHERE id=?",
                    (self.copy["id"],),
                )
            current = self.client.get(
                f"/api/copies/{self.copy['id']}/export.xlsx?sheets=Visible"
            )
            downloaded = self.client.get(prepared.json()["download_url"])

        self.assertEqual(current.status_code, 409, current.body)
        self.assertEqual(downloaded.status_code, 200, downloaded.body)
        self.assertEqual(downloaded.body, expected)
        self.assertIn("price-plan.xlsx", downloaded.headers["content-disposition"])
        self.assertEqual(exporter.call_count, 1)
        self.assertFalse(path.exists())
        replay = self.client.get(prepared.json()["download_url"])
        self.assertEqual(replay.status_code, 404, replay.body)

    def test_prepared_download_rejects_other_user_session_expiry_and_revoked_access(self):
        prepared = self.prepare()
        path = self.prepared_path(prepared)
        anonymous = self.fixture.client()
        denied = anonymous.get(prepared.json()["download_url"])
        self.assertEqual(denied.status_code, 401, denied.body)
        self.assertFalse(path.exists())

        other, _, _ = self.fixture.signup("download-other@example.test", "Other")
        prepared = self.prepare()
        path = self.prepared_path(prepared)
        denied = other.get(prepared.json()["download_url"])
        self.assertEqual(denied.status_code, 404, denied.body)
        self.assertFalse(path.exists())

        prepared = self.prepare()
        path = self.prepared_path(prepared)
        second_session = self.fixture.client()
        login = second_session.post("/api/auth/login", {
            "email": "transfer@example.test", "password": SYNTHETIC_LOGIN,
        })
        self.assertEqual(login.status_code, 200, login.body)
        denied = second_session.get(prepared.json()["download_url"])
        self.assertEqual(denied.status_code, 404, denied.body)
        self.assertFalse(path.exists())

        with mock.patch("pp.server.PREPARED_DOWNLOAD_SECONDS", -1):
            prepared = self.prepare()
        path = self.prepared_path(prepared)
        expired = self.client.get(prepared.json()["download_url"])
        self.assertEqual(expired.status_code, 404, expired.body)
        self.assertFalse(path.exists())

        viewer, viewer_csrf, _ = self.fixture.signup("download-viewer@example.test", "Viewer")
        shared = self.client.put(
            f"/api/copies/{self.copy['id']}/shares",
            {"email": "download-viewer@example.test", "role": "viewer"},
            {"x-csrf-token": self.csrf},
        )
        self.assertEqual(shared.status_code, 200, shared.body)
        prepared = self.prepare(viewer, viewer_csrf)
        path = self.prepared_path(prepared)
        self.client.put(
            f"/api/copies/{self.copy['id']}/shares",
            {"email": "download-viewer@example.test", "role": None},
            {"x-csrf-token": self.csrf},
        )
        denied = viewer.get(prepared.json()["download_url"])
        self.assertEqual(denied.status_code, 404, denied.body)
        self.assertFalse(path.exists())

    def test_prepared_download_store_is_bounded_and_cleans_evicted_file(self):
        with mock.patch("pp.server.MAX_PREPARED_DOWNLOADS", 1):
            first = self.prepare()
            first_path = self.prepared_path(first)
            second = self.prepare(sheets=["Cost, Fees"])
        self.assertFalse(first_path.exists())
        self.assertEqual(self.client.get(first.json()["download_url"]).status_code, 404)
        accepted = self.client.get(second.json()["download_url"])
        self.assertEqual(accepted.status_code, 200, accepted.body)

        pending = self.prepare()
        pending_path = self.prepared_path(pending)
        self.fixture.app.state.services.downloads.close()
        self.assertFalse(pending_path.exists())
        self.assertEqual(self.client.get(pending.json()["download_url"]).status_code, 404)

    def test_preparation_error_does_not_create_artifact_or_generate(self):
        exports = self.fixture.config.data_dir / "exports"
        before = {item for item in exports.glob("**/*") if item.is_file()}
        with mock.patch("pp.server.export_transfer") as exporter:
            response = self.prepare(sheets=["Hidden"])
        self.assertEqual(response.status_code, 422, response.body)
        self.assertEqual(response.json()["error"], "invalid_selection")
        self.assertEqual({item for item in exports.glob("**/*") if item.is_file()}, before)
        exporter.assert_not_called()

        with mock.patch("pp.server.export_transfer", side_effect=ValueError("broken package")):
            failed = self.prepare()
        self.assertEqual(failed.status_code, 500, failed.body)
        self.assertEqual({item for item in exports.glob("**/*") if item.is_file()}, before)


class DownloadStreamingBenchmarkTests(unittest.TestCase):
    def test_55mb_template_stays_within_download_budget(self):
        root = Path(self.id().replace(".", "-"))
        self.addCleanup(lambda: shutil.rmtree(root, ignore_errors=True))
        root.mkdir()
        source = root / "55mb.xlsx"
        output = root / "download.xlsx"
        _package_fixture(source)
        payload = root / "payload.bin"
        with payload.open("wb") as stream:
            remaining = 55 * 1024 * 1024
            block = b"\0" * (1 << 20)
            while remaining:
                count = min(len(block), remaining)
                stream.write(block[:count])
                remaining -= count
        _replace_members(
            source,
            {
                "xl/_rels/workbook.xml.rels": _large_payload_relationship(source),
                "[Content_Types].xml": _large_payload_content_type(source),
            },
            (("xl/media/large.bin", payload.read_bytes()),),
        )
        self.assertGreaterEqual(source.stat().st_size, 55 * 1024 * 1024)
        started = time.monotonic()
        export_transfer(TransferRequest(
            template_path=source, cells=[], out_path=output,
            template_hash="template-hash", copy_id="copy-id", input_revision=3,
            fmt="xlsx", results={},
        ))
        self.assertLess(time.monotonic() - started, 30)
        self.assertLess(_rss_mb(), 500)


def _large_payload_relationship(path):
    with zipfile.ZipFile(path) as archive:
        data = archive.read("xl/_rels/workbook.xml.rels")
    return data.replace(
        b"</Relationships>",
        b'<Relationship Id="rIdLarge" Type="urn:price-planning:test-large" '
        b'Target="media/large.bin"/></Relationships>',
    )


def _large_payload_content_type(path):
    with zipfile.ZipFile(path) as archive:
        data = archive.read("[Content_Types].xml")
    start = data.index(b"<Types")
    end = data.index(b">", start) + 1
    return data[:end] + b'<Default Extension="bin" ContentType="application/octet-stream"/>' + data[end:]


def _rss_mb():
    import resource

    value = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
    return value / (1024 * 1024) if os.uname().sysname == "Darwin" else value / 1024


if __name__ == "__main__":
    unittest.main()
