import io
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

from ops.fixtures import make_semantic
from pp.engine.formualizer import EngineError, MacroTimeout
from pp.workbook_service import MACRO_MESSAGE, SLICER_MESSAGE
from pp.export.xlsm import export_xlsm
from tests.v2.api_support import APIFixture
from tests.v2.support import TMP
from tests.v2.test_macro_buttons import add_buttons, write_button_fixture


def semantic_with_buttons(path):
    make_semantic(path)
    add_buttons(path)


class MacroEndpointContractTests(unittest.TestCase):
    def setUp(self):
        self.fixture = APIFixture(builder=write_button_fixture)
        self.client, self.csrf, _ = self.fixture.signup("macros@example.test", "Macro Owner")
        _, self.copy = self.fixture.make_copy(self.client, self.csrf)
        self.workbook = self.client.get(f"/api/copies/{self.copy['id']}/workbook").json()
        self.copy_id = self.copy["id"]

    def tearDown(self):
        self.fixture.close()

    def test_workbook_payload_lists_buttons_with_position_and_macro_only(self):
        buttons = self.workbook["buttons"]
        self.assertEqual([button["id"] for button in buttons], ["b0", "b1", "b2"])
        self.assertEqual(set(buttons[0]), {"id", "sheet", "from", "to", "label", "macro"})
        self.assertNotIn("buttons", self.workbook["sheets"][0])

    def test_owner_pressing_known_button_gets_macros_unavailable_until_office_run_exists(self):
        response = self.client.post(f"/api/copies/{self.copy_id}/macros", {"button": "b2", "op_id": "press-1"},
                                    {"x-csrf-token": self.csrf})
        self.assertEqual(response.status_code, 409)
        self.assertEqual(response.json()["error"], "macros_unavailable")
        self.assertEqual(response.json()["message"],
                         MACRO_MESSAGE)

    def test_unknown_button_is_404(self):
        response = self.client.post(f"/api/copies/{self.copy_id}/macros", {"button": "b9", "op_id": "press-2"},
                                    {"x-csrf-token": self.csrf})
        self.assertEqual((response.status_code, response.json()["error"]), (404, "unknown_button"))

    def test_viewer_is_forbidden_even_for_known_button(self):
        viewer, viewer_csrf, _ = self.fixture.signup("macro-viewer@example.test", "Viewer")
        self.client.put(f"/api/copies/{self.copy_id}/shares", {"email": "macro-viewer@example.test", "role": "viewer"},
                        {"x-csrf-token": self.csrf})
        response = viewer.post(f"/api/copies/{self.copy_id}/macros", {"button": "b2", "op_id": "press-3"},
                               {"x-csrf-token": viewer_csrf})
        self.assertEqual((response.status_code, response.json()["error"]), (403, "forbidden"))

    def test_missing_csrf_and_anonymous_requests_are_rejected(self):
        anonymous = self.fixture.client()
        self.assertEqual(anonymous.post(f"/api/copies/{self.copy_id}/macros", {"button": "b2", "op_id": "x"},
                                        {"x-csrf-token": self.csrf}).status_code, 401)
        self.assertEqual(self.client.post(f"/api/copies/{self.copy_id}/macros", {"button": "b2", "op_id": "x"}
                                          ).json()["error"], "csrf")


class FakeMacroPool:
    """Canned Excel outcomes; records the macro requests the server sends."""

    kind = "excel"

    def __init__(self, outcome=None, error=None):
        self.outcome, self.error, self.requests = outcome, error, []

    def apply(self, request):
        return {"copy_id": request.copy_id, "revision": request.revision, "pid": 0, "changes": [],
                "results": [], "removals": [], "delta": False, "base_revision": request.base_revision}

    def run_macro(self, request):
        self.requests.append(request)
        if self.error is not None:
            raise self.error
        return {"copy_id": request.copy_id, "revision": request.revision, "pid": 0,
                "changes": [], "results": self.outcome["results"], "removals": [], "delta": False,
                "base_revision": request.base_revision, "messages": self.outcome["messages"],
                "contents": self.outcome["contents"]}


MACRO_OUTCOME = {
    "messages": ["Rebuilt the plan"],
    "results": [(0, 10, 0, 66)],
    "contents": [
        (0, 1, 1, 33),                              # Inputs!B2, an input cell
        (0, 9, 0, "Hi"),                            # Inputs!A10, a plain template blank
        (0, 10, 0, {"formula": "=Inputs!B2*2"}),    # Inputs!A11, a formula the macro wrote
    ],
}


class MacroRunTests(unittest.TestCase):
    def setUp(self):
        self.fixture = APIFixture(builder=semantic_with_buttons)
        self.client, self.csrf, self.user = self.fixture.signup("macro-run@example.test", "Macro Runner")
        _, self.copy = self.fixture.make_copy(self.client, self.csrf)
        self.copy_id = self.copy["id"]
        self.pool = FakeMacroPool(MACRO_OUTCOME)
        self.fixture.app.state.services.engine_pool = self.pool

    def tearDown(self):
        self.fixture.close()

    def press(self, button="b1", op_id="press-1"):
        return self.client.post(f"/api/copies/{self.copy_id}/macros", {"button": button, "op_id": op_id},
                                {"x-csrf-token": self.csrf})

    def workbook(self):
        return self.client.get(f"/api/copies/{self.copy_id}/workbook").json()

    def history(self):
        with self.fixture.database.connect() as connection:
            return [tuple(row) for row in connection.execute(
                "SELECT sheet,row,col,action,revision FROM pp_history WHERE copy_id=? ORDER BY sheet,row,col",
                (self.copy_id,)).fetchall()]

    def cell(self, row, col):
        return self.client.get(f"/api/copies/{self.copy_id}/sheets/0/cell?r={row}&c={col}").json()

    def test_press_commits_inputs_and_non_input_contents_as_one_macro_revision(self):
        response = self.press()

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), {"messages": ["Rebuilt the plan"], "changed": 3, "input_revision": 1})
        request = self.pool.requests[0]
        self.assertEqual((request.macro, request.revision, request.base_revision), ("Plain", 1, 0))
        self.assertEqual(self.history(), [(0, 1, 1, "macro", 1), (0, 9, 0, "macro", 1), (0, 10, 0, "macro", 1)])
        self.assertEqual(self.workbook()["calculated_revision"], 1)

    def test_outputs_render_by_kind_and_formulas_are_never_sent_as_values(self):
        self.press()

        self.assertEqual(self.cell(1, 1)["v"], 33)
        self.assertEqual(self.cell(1, 1)["k"], "i")
        self.assertEqual(self.cell(1, 1)["ver"], 1)
        plain = self.cell(9, 0)
        self.assertEqual((plain["v"], plain["k"]), ("Hi", "c"))
        self.assertNotIn("ver", plain)
        formula = self.cell(10, 0)
        self.assertEqual(formula["v"], 66)
        self.assertNotIn("ver", formula)
        rows = self.client.get(f"/api/copies/{self.copy_id}/sheets/0/rows?start=9&count=2").json()["rows"]
        self.assertEqual(rows["10"]["0"]["v"], 66)
        self.assertNotIn("ver", rows["10"]["0"])

    def test_formula_override_on_an_input_shows_the_formula_as_its_input_value(self):
        with self.fixture.database.transaction() as connection:
            connection.execute(
                "INSERT INTO pp_overrides(copy_id,sheet,row,col,value_json,version,updated_by,updated_at) "
                "VALUES(?,?,?,?,?,?,?,?)",
                (self.copy_id, 0, 1, 1, '{"formula":"=Inputs!B5*3"}', 1, self.user["id"], "2026-10-12T00:00:00Z"))

        record = self.cell(1, 1)

        self.assertEqual(record["v"], 10)
        self.assertEqual((record["input_value"], record["input_formula"], record["k"]),
                         ("=Inputs!B5*3", True, "i"))

    def test_repeating_an_op_id_with_the_same_button_replays_the_ack_without_running_again(self):
        first = self.press(op_id="same")
        second = self.press(op_id="same")

        self.assertEqual(second.json(), first.json())
        self.assertEqual(len(self.pool.requests), 1)

    def test_reusing_an_op_id_for_another_button_is_refused(self):
        self.press(button="b1", op_id="shared")

        response = self.press(button="b0", op_id="shared")

        self.assertEqual((response.status_code, response.json()["error"]), (409, "op_id_reused"))
        self.assertEqual(len(self.pool.requests), 1)

    def test_macro_that_wrote_nothing_commits_nothing(self):
        self.pool.outcome = {"messages": ["Nothing to do"], "results": [], "contents": []}

        response = self.press()

        self.assertEqual(response.json(), {"messages": ["Nothing to do"], "changed": 0, "input_revision": 0})
        self.assertEqual(self.history(), [])
        self.assertEqual(self.workbook()["input_revision"], 0)

    def test_engine_failure_is_502_with_a_plain_message_and_commits_nothing(self):
        self.pool.error = EngineError("calculation worker failed: EngineError")

        response = self.press()

        self.assertEqual((response.status_code, response.json()["error"]), (502, "macro_failed"))
        self.assertEqual(response.json()["message"], "The macro did not finish in Excel.")
        self.assertEqual(self.workbook()["input_revision"], 0)

    def test_timeout_is_502_with_its_own_message(self):
        self.pool.error = MacroTimeout("macro timed out")

        response = self.press()

        self.assertEqual((response.status_code, response.json()["error"]), (502, "macro_failed"))
        self.assertEqual(response.json()["message"], "The macro ran too long and was stopped.")

    def test_without_an_office_engine_or_with_saved_values_only_it_is_409(self):
        self.fixture.app.state.services.engine_pool = None
        self.assertEqual(self.press().json()["error"], "macros_unavailable")

        self.fixture.app.state.services.engine_pool = self.pool
        with mock.patch.object(self.fixture.snapshot, "results", "cached"):
            self.assertEqual(self.press(op_id="cached").json()["error"], "macros_unavailable")
        self.assertEqual(self.pool.requests, [])

    def test_macro_override_passes_the_export_guard_but_a_hand_edited_locked_cell_does_not(self):
        self.press()
        self.assertEqual(self.client.get(f"/api/copies/{self.copy_id}/export.xlsm?check=1").status_code, 204)

        with self.fixture.database.transaction() as connection:
            connection.execute(
                "INSERT INTO pp_overrides(copy_id,sheet,row,col,value_json,version,updated_by,updated_at) "
                "VALUES(?,?,?,?,?,?,?,?)",
                (self.copy_id, 0, 9, 1, '"typed"', 1, self.user["id"], "2026-10-12T00:00:00Z"))
        response = self.client.get(f"/api/copies/{self.copy_id}/export.xlsm?check=1")
        self.assertEqual((response.status_code, response.json()["error"]), (422, "locked"))

    def test_export_after_a_press_writes_the_formula_and_the_input_value(self):
        self.press()

        exported = self.client.get(f"/api/copies/{self.copy_id}/export.xlsm")

        with zipfile.ZipFile(io.BytesIO(exported.body)) as archive:
            sheet = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        self.assertEqual(sheet.findtext(f".//{ns}c[@r='A11']/{ns}f"), "Inputs!B2*2")
        self.assertEqual(sheet.findtext(f".//{ns}c[@r='B2']/{ns}v"), "33")

    def test_fork_keeps_formula_overrides(self):
        self.press()

        response = self.client.post(f"/api/copies/{self.copy_id}/fork", {"name": "Fork of macro plan"},
                                    {"x-csrf-token": self.csrf})
        fork_id = response.json()["copy"]["id"]

        with self.fixture.database.connect() as connection:
            value = connection.execute(
                "SELECT value_json FROM pp_overrides WHERE copy_id=? AND sheet=0 AND row=10 AND col=0",
                (fork_id,)).fetchone()["value_json"]
        self.assertEqual(value, '{"formula":"=Inputs!B2*2"}')


class MacroExportTests(unittest.TestCase):
    def test_export_writes_a_formula_override_as_a_formula(self):
        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "semantic.xlsx"
            make_semantic(source)
            out = Path(temp) / "out.xlsx"
            export_xlsm(source, [(0, 10, 0, {"formula": "=Inputs!B2*2"}), (0, 11, 0, "text")], out)

            sheet = load_workbook(out)["Inputs"]

        self.assertEqual(sheet["A11"].value, "=Inputs!B2*2")
        self.assertEqual(sheet["A12"].value, "text")


if __name__ == "__main__":
    unittest.main()
