import hashlib
import json
import io
import shutil
import sqlite3
from pathlib import Path
import tempfile
import unittest
from unittest import mock
import zipfile
import xml.etree.ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

from ops.fixtures import make_semantic


ROOT = Path(__file__).resolve().parents[2]
TMP = ROOT / "results-private" / "tmp"
NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


def rewrite_member(source, member, transform):
    replacement = source.with_name("rewritten.xlsm")
    with zipfile.ZipFile(source) as source_zip, zipfile.ZipFile(replacement, "w") as target_zip:
        for info in source_zip.infolist():
            data = source_zip.read(info.filename)
            if info.filename == member:
                data = transform(data)
            target_zip.writestr(info, data)
    shutil.move(replacement, source)


def shared_formula_fixture(path):
    workbook = Workbook()
    values = workbook.active
    values.title = "Values"
    values["A1"] = 3
    values["B1"] = 4
    formulas = workbook.create_sheet("Formulas")
    formulas["B2"] = "=Values!A1*2"
    formulas["C2"] = "=Values!B1*2"
    workbook.save(path)

    def shared_formula_xml(data):
        namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        def qualified(name):
            return f"{{{namespace}}}{name}"
        root = ET.fromstring(data)
        row = root.find(f".//{qualified('row')}[@r='2']")
        cells = {cell.get("r"): cell for cell in row.findall(qualified("c"))}
        master = cells["B2"]
        follower = cells["C2"]
        master_formula = master.find(qualified("f"))
        follower_formula = follower.find(qualified("f"))
        master_formula.attrib.update({"t": "shared", "ref": "B2:C2", "si": "7"})
        follower_formula.attrib.clear()
        follower_formula.attrib.update({"t": "shared", "si": "7"})
        follower_formula.text = None
        row.remove(follower)
        row.insert(list(row).index(master), follower)
        return ET.tostring(root, encoding="utf-8", xml_declaration=True)

    rewrite_member(path, "xl/worksheets/sheet2.xml", shared_formula_xml)


def metadata_fixture(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Visible"
    sheet.protection.sheet = True
    sheet.freeze_panes = "C3"
    sheet.column_dimensions["B"].width = 18
    sheet.column_dimensions["D"].hidden = True
    sheet.row_dimensions[4].height = 24
    sheet.row_dimensions[6].hidden = True
    sheet.merge_cells("A8:B9")
    sheet["B2"] = 10
    sheet["B2"].protection = Protection(locked=False)
    sheet["C2"] = "=B2*2"
    sheet["C2"].protection = Protection(locked=False)
    sheet["B2"].number_format = "0.00%"
    sheet["B2"].font = Font(bold=True, italic=True, color="FF112233", size=14)
    sheet["B2"].fill = PatternFill("solid", fgColor="FF445566")
    edge = Side(style="thin", color="FF000000")
    sheet["B2"].border = Border(top=edge, right=edge, bottom=edge, left=edge)
    sheet["B2"].alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
    validation = DataValidation(type="list", formula1='"A,B,C"',
                                showErrorMessage=True, errorStyle="stop")
    validation.add("B2:B4")
    sheet.add_data_validation(validation)
    hidden = workbook.create_sheet("Hidden")
    hidden.sheet_state = "veryHidden"
    hidden["A1"] = "A"
    reference = DataValidation(type="list", formula1="'Hidden'!$A$1:$A$1",
                               showErrorMessage=True, errorStyle="warning")
    reference.add("B5")
    sheet.add_data_validation(reference)
    workbook.defined_names.add(DefinedName("PlanInput", attr_text="'Visible'!$B$2"))
    sheet.conditional_formatting.add("B2", CellIsRule(operator="greaterThan", formula=["0"]))
    workbook.save(path)


def default_style_fixture(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Styled Inputs"
    sheet.protection.sheet = True
    sheet["A1"] = "edge"
    sheet["C4"] = "edge"
    sheet.column_dimensions["B"].protection = Protection(locked=False)
    sheet.row_dimensions[3].protection = Protection(locked=False)
    workbook.save(path)


class TemplateImporterTests(unittest.TestCase):
    def test_import_writes_atomic_snapshot_with_input_mask_and_formula_map(self):
        """Removing protection/formula parsing must make this contract fail."""
        from pp.template.importer import import_template

        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            workspace = Path(temp)
            source = workspace / "semantic.xlsm"
            make_semantic(source)

            snapshot = import_template(source, workspace / "data")

            meta = json.loads((snapshot / "meta.json").read_text(encoding="utf-8"))
            mask = json.loads((snapshot / "input_mask.json").read_text(encoding="utf-8"))
            blocks = json.loads((snapshot / "sheets" / "0" / "0.json").read_text(encoding="utf-8"))
            self.assertEqual([sheet["name"] for sheet in meta["sheets"]], ["Inputs", "Lookup", "Calc"])
            self.assertIn("0:1:1", mask)
            self.assertNotIn("2:1:1", mask)
            self.assertEqual(blocks["1"]["1"]["k"], "i")
            self.assertEqual(blocks["1"]["1"]["v"], 10)

    def test_failed_import_does_not_publish_partial_snapshot(self):
        """Publishing staging data after a malformed archive would be corruption."""
        from pp.template.importer import TemplateImportError, import_template

        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            workspace = Path(temp)
            source = workspace / "broken.xlsm"
            source.write_bytes(b"not a zip")

            with self.assertRaises(TemplateImportError):
                import_template(source, workspace / "data")

            templates = workspace / "data" / "templates"
            self.assertEqual(list(templates.iterdir()), [])

    def test_outdated_snapshot_format_is_reimported_at_the_same_content_hash(self):
        """A same-source cache must not preserve old masks after importer semantics change."""
        from pp.template.importer import IMPORT_VERSION, import_template

        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            workspace = Path(temp)
            source = workspace / "semantic.xlsm"
            make_semantic(source)
            snapshot = import_template(source, workspace / "data")
            meta_path = snapshot / "meta.json"
            meta = json.loads(meta_path.read_text(encoding="utf-8"))
            meta["import_version"] = 6
            meta_path.write_text(json.dumps(meta), encoding="utf-8")
            (snapshot / "input_mask.json").write_text('["stale"]', encoding="utf-8")

            refreshed = import_template(source, workspace / "data")

            current = json.loads((refreshed / "meta.json").read_text(encoding="utf-8"))
            mask = json.loads((refreshed / "input_mask.json").read_text(encoding="utf-8"))
            self.assertEqual(refreshed, snapshot)
            self.assertEqual(current["import_version"], IMPORT_VERSION)
            self.assertNotIn("stale", mask)

    def test_outdated_copy_snapshot_requires_owner_reimport(self):
        from pp.errors import AppError
        from pp.template.importer import import_template
        from pp.workbook_service import TemplateSnapshot

        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            workspace = Path(temp)
            source = workspace / "old-copy.xlsm"
            make_semantic(source)
            snapshot = import_template(source, workspace / "data")
            meta_path = snapshot / "meta.json"
            meta = json.loads(meta_path.read_text(encoding="utf-8"))
            meta["import_version"] = 6
            meta_path.write_text(json.dumps(meta), encoding="utf-8")

            with self.assertRaises(AppError) as raised:
                TemplateSnapshot(snapshot)

        self.assertEqual(raised.exception.status, 409)
        self.assertEqual(raised.exception.code, "template_reimport_required")
        self.assertIn("reimport", raised.exception.message.lower())

    def test_import_uses_one_immutable_source_capture_across_path_replacement(self):
        """Replacing the configured path mid-import must not mix hash, cells, and source bytes."""
        from pp.template import importer

        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            workspace = Path(temp)
            source = workspace / "semantic.xlsm"
            replacement = workspace / "replacement.xlsm"
            make_semantic(source)
            make_semantic(replacement)
            workbook = load_workbook(replacement)
            workbook["Inputs"]["B2"] = 99
            workbook.save(replacement)
            original_hash, swapped = importer._source_hash, False

            def swap_after_hash(path):
                nonlocal swapped
                digest = original_hash(path)
                if not swapped:
                    replacement.replace(source)
                    swapped = True
                return digest

            with mock.patch.object(importer, "_source_hash", side_effect=swap_after_hash):
                snapshot = importer.import_template(source, workspace / "data")

            meta = json.loads((snapshot / "meta.json").read_text(encoding="utf-8"))
            block = json.loads((snapshot / "sheets" / "0" / "0.json").read_text(encoding="utf-8"))
            actual = hashlib.sha256((snapshot / "source.xlsm").read_bytes()).hexdigest()
            self.assertEqual((meta["hash"], block["1"]["1"]["v"]), (actual, 10))
            self.assertEqual(meta["name"], "semantic.xlsm")

    def test_shared_formula_follower_before_master_is_expanded(self):
        """Resolving followers only when their XML is visited must fail."""
        from pp.template.importer import import_template

        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            workspace = Path(temp)
            source = workspace / "shared.xlsm"
            shared_formula_fixture(source)

            snapshot = import_template(source, workspace / "data")
            block = json.loads((snapshot / "sheets" / "1" / "0.json").read_text())

            self.assertEqual(block["1"]["1"]["formula"], "=Values!A1*2")
            self.assertEqual(block["1"]["2"]["formula"], "=Values!B1*2")

    def test_import_reports_numeric_progress_for_each_sheet(self):
        """A long import must expose sheet progress without workbook content."""
        from pp.template.importer import import_template

        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            workspace = Path(temp)
            source = workspace / "progress.xlsm"
            make_semantic(source)

            with self.assertLogs("pp.template.importer", level="INFO") as captured:
                import_template(source, workspace / "data")

            messages = [record.getMessage() for record in captured.records]
            self.assertEqual(messages, [
                "Template import progress: sheet 1/3",
                "Template import progress: sheet 2/3",
                "Template import progress: sheet 3/3",
            ])
            self.assertNotIn("Inputs", " ".join(messages))

    def test_template_import_error_removes_staging_directory(self):
        """A named shared-formula failure must not leave a partial staging tree."""
        from pp.template.importer import TemplateImportError, import_template

        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            workspace = Path(temp)
            source = workspace / "missing-master.xlsm"
            shared_formula_fixture(source)

            def remove_master(data):
                namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
                root = ET.fromstring(data)
                formula = root.find(f".//{{{namespace}}}c[@r='B2']/{{{namespace}}}f")
                formula.text = None
                return ET.tostring(root, encoding="utf-8", xml_declaration=True)

            rewrite_member(source, "xl/worksheets/sheet2.xml", remove_master)
            with self.assertRaises(TemplateImportError):
                import_template(source, workspace / "data")

            templates = workspace / "data" / "templates"
            self.assertEqual(list(templates.glob(".staging-*")), [])

    def test_engine_formula_scan_detaches_processed_cells(self):
        """Formula discovery must not retain cleared cell nodes under their row."""
        from pp.engine import formualizer

        helper = getattr(formualizer, "_formula_cells", None)
        self.assertIsNotNone(helper, "formula scan must expose its bounded iterator")
        namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        row = ET.Element(f"{{{namespace}}}row")
        retained = []

        def parser(stream, events):
            yield "start", row
            for index in range(100):
                cell = ET.Element(f"{{{namespace}}}c", {"r": f"A{index + 1}"})
                ET.SubElement(cell, f"{{{namespace}}}f").text = "1+1"
                row.append(cell)
                yield "start", cell
                yield "end", cell
                retained.append(len(row))
            yield "end", row

        cells = list(helper(mock.Mock(), parser))
        self.assertEqual(len(cells), 100)
        self.assertEqual(retained, [0] * 100)

    def test_metadata_styles_protection_and_progress_match_api_contract(self):
        """Dropping render metadata or treating unlocked formulas as inputs must fail."""
        from pp.template.importer import import_template

        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            workspace = Path(temp)
            source = workspace / "metadata.xlsm"
            metadata_fixture(source)

            snapshot = import_template(source, workspace / "data")
            meta = json.loads((snapshot / "meta.json").read_text())
            styles = json.loads((snapshot / "styles.json").read_text())
            mask = json.loads((snapshot / "input_mask.json").read_text())
            progress = json.loads((snapshot / "progress.json").read_text())
            visible = meta["sheets"][0]
            hidden = meta["sheets"][1]
            block = json.loads((snapshot / "sheets" / "0" / "0.json").read_text())
            style = styles[block["1"]["1"]["s"]]

            self.assertEqual((visible["rows"], visible["cols"]), (9, 4))
            self.assertEqual(visible["freeze"], {"rows": 2, "cols": 2, "start_row": 2, "start_col": 2})
            self.assertEqual(visible["scroll"], {"row": 2, "col": 2})
            self.assertGreater(visible["col_widths"]["1"], 0)
            self.assertGreater(visible["row_heights"]["3"], 0)
            self.assertEqual(visible["hidden_rows"], [5])
            self.assertEqual(visible["hidden_cols"], [3])
            self.assertEqual(visible["merges"], [[7, 0, 8, 1]])
            self.assertEqual(visible["validations"][0]["list"], ["A", "B", "C"])
            self.assertEqual(visible["validations"][0]["ranges"], [[1, 1, 3, 1]])
            self.assertIs(visible["validations"][0]["showErrorMessage"], True)
            self.assertEqual(visible["validations"][0]["errorStyle"], "stop")
            self.assertEqual(visible["validations"][1]["formula1"], "'Hidden'!$A$1:$A$1")
            self.assertIs(visible["validations"][1]["showErrorMessage"], True)
            self.assertEqual(visible["validations"][1]["errorStyle"], "warning")
            self.assertEqual(hidden["hidden"], "veryHidden")
            self.assertEqual(meta["defined_names"][0]["name"], "PlanInput")
            self.assertEqual(progress, {"complete": True, "sheets_completed": 2, "sheets_total": 2})
            self.assertIn("0:1:1", mask)
            self.assertNotIn("0:1:2", mask)
            self.assertEqual(block["1"]["2"]["k"], "f")
            self.assertEqual(style["nf"], "0.00%")
            self.assertTrue(style["bold"])
            self.assertTrue(style["italic"])
            self.assertEqual(style["align"], "center")
            self.assertEqual(style["valign"], "top")
            self.assertTrue(style["wrap"])
            edge = {"style": "thin", "color": "#000000"}
            self.assertEqual(style["border"], {"t": edge, "r": edge, "b": edge, "l": edge})
            self.assertEqual((style["fill"], style["color"]), ("#445566", "#112233"))

    def test_blank_row_and_column_default_styles_resolve_as_lazy_inputs(self):
        """Unlocked blank styled cells must stay editable without materializing a full-sheet mask."""
        from pp.template.importer import import_template
        from pp.workbook_service import TemplateSnapshot

        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            workspace = Path(temp)
            source = workspace / "default-styles.xlsm"
            default_style_fixture(source)

            snapshot = TemplateSnapshot(import_template(source, workspace / "data"))

            self.assertEqual(snapshot.input_mask, frozenset())
            self.assertTrue(snapshot.is_input(0, 1, 1))
            self.assertTrue(snapshot.is_input(0, 2, 2))
            self.assertFalse(snapshot.is_input(0, 1, 0))
            self.assertEqual(snapshot.cell(0, 1, 1)["k"], "i")
            self.assertEqual(snapshot.cell(0, 2, 2)["k"], "i")
            self.assertEqual(snapshot.cell(0, 1, 1)["s"], 1)

    def test_phonetic_annotations_are_excluded_from_shared_and_inline_text(self):
        """Ruby pronunciation runs must not be appended to the visible cell value."""
        from pp.template.importer import _shared_strings, import_template

        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            workspace = Path(temp)
            source = workspace / "phonetic.xlsm"
            workbook = Workbook()
            workbook.active["A1"] = "Visible"
            workbook.save(source)

            def add_phonetic(data):
                root = ET.fromstring(data)
                inline = root.find(f".//{NS}c[@r='A1']/{NS}is")
                phonetic = ET.SubElement(inline, f"{NS}rPh")
                ET.SubElement(phonetic, f"{NS}t").text = "PHONETIC"
                return ET.tostring(root, encoding="utf-8", xml_declaration=True)

            rewrite_member(source, "xl/worksheets/sheet1.xml", add_phonetic)
            snapshot = import_template(source, workspace / "data")
            block = json.loads((snapshot / "sheets" / "0" / "0.json").read_text())

            archive = io.BytesIO()
            with zipfile.ZipFile(archive, "w") as target:
                target.writestr("xl/sharedStrings.xml", (
                    f'<sst xmlns="{NS[1:-1]}"><si><t>Shared</t>'
                    f'<rPh><t>PHONETIC</t></rPh></si></sst>'
                ))
            database = sqlite3.connect(":memory:")
            with zipfile.ZipFile(archive) as synthetic:
                _shared_strings(synthetic, database)
            shared = database.execute("SELECT value FROM strings WHERE id=0").fetchone()[0]
            database.close()

            self.assertEqual(block["0"]["0"]["v"], "Visible")
            self.assertEqual(shared, "Shared")

    def test_sheet_blocks_are_split_at_256_rows(self):
        """Keeping a single sheet dictionary instead of publishing row blocks must fail."""
        from pp.template.importer import import_template

        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            workspace = Path(temp)
            source = workspace / "blocks.xlsm"
            workbook = Workbook(write_only=True)
            sheet = workbook.create_sheet("Rows")
            for row in range(520):
                sheet.append([row])
            workbook.save(source)

            snapshot = import_template(source, workspace / "data")

            files = sorted(path.name for path in (snapshot / "sheets" / "0").glob("*.json"))
            self.assertEqual(files, ["0.json", "1.json", "2.json"])

    def test_style_apply_flags_use_base_number_format_and_alignment(self):
        """Ignoring applyNumberFormat or applyAlignment inheritance must fail."""
        from pp.template.importer import _style_data

        namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        styles = f'''<styleSheet xmlns="{namespace}">
          <fonts count="1"><font/></fonts><fills count="1"><fill/></fills>
          <borders count="1"><border/></borders>
          <cellStyleXfs count="1"><xf numFmtId="10"><alignment horizontal="left"/></xf></cellStyleXfs>
          <cellXfs count="1"><xf xfId="0" numFmtId="2" applyNumberFormat="0" applyAlignment="0">
            <alignment horizontal="right"/></xf></cellXfs></styleSheet>'''
        archive_bytes = io.BytesIO()
        with zipfile.ZipFile(archive_bytes, "w") as archive:
            archive.writestr("xl/styles.xml", styles)
        archive_bytes.seek(0)
        with zipfile.ZipFile(archive_bytes) as archive:
            parsed, _ = _style_data(archive)
        self.assertEqual(parsed[0]["nf"], "0.00%")
        self.assertEqual(parsed[0]["align"], "left")


if __name__ == "__main__":
    unittest.main()


class LiteralListValidationTests(unittest.TestCase):
    def test_literal_list_members_take_the_type_excel_enters_on_pick(self):
        """Picking "1" from an Excel dropdown enters the number 1, so SUM and lookups see a number."""
        from pp.template.importer import NS, _validation

        node = ET.fromstring(
            f'<dataValidation xmlns="{NS[1:-1]}" type="list" sqref="B2">'
            '<formula1>"1,2.5,A,TRUE,007,1E400"</formula1></dataValidation>'
        )
        self.assertEqual(_validation(node)["list"], [1, 2.5, "A", True, 7, "1E400"])

    def test_ops_accept_the_numeric_member_and_reject_its_text_twin(self):
        from pp.ops import _validate_rule

        rule = {"type": "list", "list": [1, 2], "showErrorMessage": True}
        _validate_rule(rule, 1, {"sheet": 0, "r": 1, "c": 1})
        with self.assertRaises(Exception):
            _validate_rule(rule, "1", {"sheet": 0, "r": 1, "c": 1})
