"""Structured table references are resolved to plain ranges before native load."""

import shutil
import tempfile
import unittest
import zipfile
from pathlib import Path

TABLES = {"salestbl": {"sheet": "Data Src", "ref": "B2:D5", "headers": 1, "totals": 0,
                       "columns": ["Item", "Qty Sold", "Price"]}}


class RewriteTests(unittest.TestCase):
    def rewrite(self, text, tables=None):
        from pp.engine.native_load import Table, rewrite_structured
        tables = tables or {key: Table(**value) for key, value in TABLES.items()}
        return rewrite_structured(text, tables)

    def test_specifiers_and_columns(self):
        cases = {
            "SUM(SalesTbl[[#All],[Qty Sold]])": "SUM('Data Src'!$C$2:$C$5)",
            "SUM(salestbl[Qty Sold])": "SUM('Data Src'!$C$3:$C$5)",
            "SUM(SalesTbl[[#Data],[Price]])": "SUM('Data Src'!$D$3:$D$5)",
            "SalesTbl[[#Headers],[Price]]": "'Data Src'!$D$2",
            "COUNTA(SalesTbl[#All])": "COUNTA('Data Src'!$B$2:$D$5)",
            "ROWS(SalesTbl)": "ROWS('Data Src'!$B$3:$D$5)",
            "SUM(SalesTbl[[Qty Sold]:[Price]])": "SUM('Data Src'!$C$3:$D$5)",
            "SUM(SalesTbl[[#All],[Item]:[Qty Sold]])": "SUM('Data Src'!$B$2:$C$5)",
        }
        for text, expected in cases.items():
            with self.subTest(text=text):
                self.assertEqual(self.rewrite(text), expected)

    def test_strings_sheet_names_and_other_names_are_untouched(self):
        for text in ('"SalesTbl[Item]"&A1', "'SalesTbl [x]'!A1+1", "Other[Item]", "SalesTbl!A1"):
            with self.subTest(text=text):
                self.assertEqual(self.rewrite(text), text)

    def test_unsupported_forms_raise(self):
        from pp.engine.native_load import SharedFormulaError
        for text in ("SalesTbl[@Price]", "SalesTbl[[#This Row],[Price]]", "SalesTbl[Missing]",
                     "SalesTbl[#Totals]"):
            with self.subTest(text=text), self.assertRaises(SharedFormulaError):
                self.rewrite(text)

    def test_escaped_column_names(self):
        from pp.engine.native_load import Table
        tables = {"t": Table("S", "A1:B3", 1, 0, ["Price [USD]", "#Units"])}
        self.assertEqual(self.rewrite("SUM(T[Price '[USD']])", tables), "SUM('S'!$A$2:$A$3)")
        self.assertEqual(self.rewrite("SUM(T['#Units])", tables), "SUM('S'!$B$2:$B$3)")


class EngineTests(unittest.TestCase):
    def test_engine_evaluates_table_formulas(self):
        from openpyxl import Workbook
        from openpyxl.worksheet.table import Table as XlTable

        from pp.engine.formualizer import FormualizerEngine
        directory = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, directory, True)
        book = Workbook()
        calc = book.active
        calc.title = "Calc"
        data = book.create_sheet("Data Src")
        for row in (["Item", "Qty Sold", "Price"], ["a", 2, 10], ["b", 3, 20], ["c", 5, 30]):
            data.append(row)
        data.add_table(XlTable(displayName="SalesTbl", ref="A1:C4"))
        calc["A1"] = "=SUM(SalesTbl[[#All],[Qty Sold]])"
        calc["A2"] = "=COUNTA(SalesTbl[[#All],[Item]])"
        calc["A3"] = "=SUMPRODUCT(SalesTbl[Qty Sold],SalesTbl[Price])"
        calc["A4"] = "=SalesTbl[[#Headers],[Price]]"
        source = directory / "tables.xlsx"
        book.save(source)
        with zipfile.ZipFile(source) as archive:
            self.assertTrue(any(name.startswith("xl/tables/") for name in archive.namelist()))
        engine = FormualizerEngine.load(source, cache_dir=directory / "cache")
        self.assertEqual([engine.get(0, row, 0) for row in range(4)], [10, 4, 230, "Price"])


if __name__ == "__main__":
    unittest.main()
