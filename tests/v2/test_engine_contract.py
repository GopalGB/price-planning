"""The documented engine names must load and calculate a real fixture."""

from datetime import date, datetime, time
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook
from openpyxl.utils.datetime import MAC_EPOCH

import pp.engine as engines
from ops.typed import ENGINE_ERRORS, ERRORS
from tests.v2.support import BackendFixture, TMP


def _temporal_workbook(source, epoch=None):
    workbook = Workbook()
    if epoch is not None:
        workbook.epoch = epoch
    sheet = workbook.active
    sheet["A1"] = "=DATE(2026,1,1)"
    sheet["A2"] = "=TIME(12,30,0)"
    sheet["A3"] = "=DATE(2026,1,1)+TIME(6,0,0)"
    if epoch is None:
        sheet["A4"] = "=DATE(1900,2,28)"
        sheet["A5"] = "=DATE(1900,3,1)"
    workbook.save(source)


class PublicEngineContractTests(unittest.TestCase):
    def test_excel_prefixed_isformula_uses_template_formula_map(self):
        from openpyxl import load_workbook

        fixture = BackendFixture()
        self.addCleanup(fixture.close)
        workbook = load_workbook(fixture.source)
        workbook["Calc"]["C1"] = "=_xlfn.ISFORMULA(B2)"
        workbook.save(fixture.source)
        engine = engines.FormualizerEngine.load(fixture.source)
        self.assertIs(engine.get(2, 0, 2), True)

    def test_isformula_shim_preserves_quoted_literal_text(self):
        from openpyxl import load_workbook

        fixture = BackendFixture()
        self.addCleanup(fixture.close)
        workbook = load_workbook(fixture.source)
        workbook["Calc"]["C1"] = '=IF(ISFORMULA(B2),"ISFORMULA(B2)","no")'
        workbook.save(fixture.source)
        engine = engines.FormualizerEngine.load(fixture.source)
        self.assertEqual(engine.get(2, 0, 2), "ISFORMULA(B2)")

    def test_named_engine_load_apply_and_get_contract(self):
        engine_class = getattr(engines, "FormualizerEngine", None)
        self.assertIsNotNone(engine_class, "FormualizerEngine must be publicly available")
        fixture = BackendFixture()
        self.addCleanup(fixture.close)
        engine = engine_class.load(fixture.source)
        engine.apply([(0, 1, 1, 13)])
        self.assertEqual(engine.get(2, 1, 1), 26)
        self.assertIsInstance(engine, engines.Engine)

    def test_apply_reports_changed_formulas_only(self):
        fixture = BackendFixture()
        self.addCleanup(fixture.close)
        engine = engines.FormualizerEngine.load(fixture.source)
        changes = engine.apply([(0, 1, 1, 13)])
        addresses = {(sheet, row, col) for sheet, row, col, value in changes}
        self.assertIn((2, 1, 1), addresses)
        self.assertNotIn((2, 16, 1), addresses, "constant TRIM formula did not change")

    def test_apply_delta_distinguishes_boolean_true_from_number_one(self):
        """A cached numeric one must not suppress a later boolean true result."""
        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "typed.xlsm"
            workbook = Workbook()
            inputs = workbook.active
            inputs.title = "Inputs"
            inputs["A1"] = 1
            formulas = workbook.create_sheet("Formulas")
            formulas["A1"] = "=Inputs!A1"
            workbook.save(source)
            engine = engines.FormualizerEngine.load(source)

            changes = engine.apply([(0, 0, 0, True)])

            values = {(sheet, row, col): value for sheet, row, col, value in changes}
            self.assertIn((1, 0, 0), values, "a type-only change must be published")
            self.assertIs(values[(1, 0, 0)], True)

    def test_error_input_recalculates_to_native_formula_error(self):
        """Passing an API error dictionary directly to native set_value must fail."""
        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "error-input.xlsm"
            workbook = Workbook()
            inputs = workbook.active
            inputs.title = "Inputs"
            inputs["A1"] = 10
            formulas = workbook.create_sheet("Formulas")
            formulas["A1"] = "=Inputs!A1*2"
            workbook.save(source)
            engine = engines.FormualizerEngine.load(source)

            changes = engine.apply([(0, 0, 0, {"error": "#DIV/0!"})])

            values = {(sheet, row, col): value for sheet, row, col, value in changes}
            self.assertEqual(values[(1, 0, 0)], {"error": "#DIV/0!"})

    def test_temporal_formulas_use_1900_excel_serials(self):
        """Stringifying native temporal values must not destroy Excel numeric semantics."""
        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "dates.xlsm"
            _temporal_workbook(source)
            engine = engines.FormualizerEngine.load(source)

            native = [engine.workbook.get_value("Sheet", row, 1) for row in range(1, 4)]
            self.assertIs(type(native[0]), date)
            self.assertIs(type(native[1]), time)
            self.assertIs(type(native[2]), datetime)
            self.assertEqual(engine.get(0, 0, 0), 46023)
            self.assertAlmostEqual(engine.get(0, 1, 0), 0.5208333333333333, places=15)
            self.assertEqual(engine.get(0, 2, 0), 46023.25)
            self.assertEqual((engine.get(0, 3, 0), engine.get(0, 4, 0)), (59, 61))

    def test_temporal_formulas_use_the_source_workbook_1904_date_system(self):
        """Ignoring date1904 must not shift native date and datetime results by 1,462 days."""
        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "dates-1904.xlsm"
            _temporal_workbook(source, MAC_EPOCH)

            engine = engines.FormualizerEngine.load(source)

            self.assertEqual(engine.get(0, 0, 0), 44561)
            self.assertAlmostEqual(engine.get(0, 1, 0), 0.5208333333333333, places=15)
            self.assertEqual(engine.get(0, 2, 0), 44561.25)

    def test_api_error_tokens_are_all_supported_by_the_native_engine(self):
        self.assertLessEqual(ERRORS, set(ENGINE_ERRORS.values()))

    def test_apply_delta_preserves_small_numeric_change(self):
        """Parity tolerance must not suppress a real formula value change."""
        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "exact-number.xlsm"
            workbook = Workbook()
            inputs = workbook.active
            inputs.title = "Inputs"
            inputs["A1"] = 1
            formulas = workbook.create_sheet("Formulas")
            formulas["A1"] = "=Inputs!A1"
            workbook.save(source)
            engine = engines.FormualizerEngine.load(source)

            changes = engine.apply([(0, 0, 0, 1.0000000001)])

            values = {(sheet, row, col): value for sheet, row, col, value in changes}
            self.assertEqual(values[(1, 0, 0)], 1.0000000001)

    def test_error_looking_text_stays_text_while_native_error_stays_typed(self):
        """Only native error dictionaries may become API error objects."""
        TMP.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=TMP) as temp:
            source = Path(temp) / "error-looking-text.xlsm"
            workbook = Workbook()
            inputs = workbook.active
            inputs.title = "Inputs"
            inputs["A1"] = "seed"
            formulas = workbook.create_sheet("Formulas")
            formulas["A1"] = "=Inputs!A1"
            workbook.save(source)
            engine = engines.FormualizerEngine.load(source)

            text_change = engine.apply([(0, 0, 0, "#DIV/0!")])
            error_change = engine.apply([(0, 0, 0, {"error": "#DIV/0!"})])

            text_values = {(sheet, row, col): value for sheet, row, col, value in text_change}
            error_values = {(sheet, row, col): value for sheet, row, col, value in error_change}
            self.assertEqual(text_values[(1, 0, 0)], "#DIV/0!")
            self.assertEqual(error_values[(1, 0, 0)], {"error": "#DIV/0!"})
