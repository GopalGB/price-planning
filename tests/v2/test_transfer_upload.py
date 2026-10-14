"""Synthetic upload preview and confirmation contracts."""

from __future__ import annotations

import io
import json
import os
import re
from base64 import b64encode
from pathlib import Path
from types import SimpleNamespace
import unittest
from unittest import mock
import warnings
import zipfile
import xml.etree.ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Protection

import pp.upload as upload_module
from pp.errors import AppError
from pp.upload import PreviewStore
from tests.v2.api_support import APIFixture
from tests.v2.test_template_importer import default_style_fixture


def _workbook_bytes(data, root, edits=(), locked_edit=False, new_sheet=False):
    source, output = root / "upload-source.xlsx", root / "upload-edited.xlsx"
    source.write_bytes(data)
    workbook = load_workbook(source, keep_vba=False, data_only=False)
    for sheet, cell, value in edits:
        workbook[sheet][cell] = value
    if locked_edit:
        workbook["Calc"]["B2"] = 999
    if new_sheet:
        workbook.create_sheet("Unexpected")["A1"] = "ignored"
    workbook.save(output)
    workbook.close()
    return output.read_bytes()


def _shifted_rows(data, root, operation):
    source, output = root / "shift-source.xlsx", root / "shift-upload.xlsx"
    source.write_bytes(data)
    workbook = load_workbook(source, keep_vba=False, data_only=False)
    if operation == "insert":
        workbook["Inputs"].insert_rows(2)
        workbook["Inputs"]["B3"] = 12
    elif operation == "same_extent":
        workbook["Inputs"].insert_rows(2)
        workbook["Inputs"].delete_rows(7)
        workbook["Inputs"]["B3"] = 12
    else:
        workbook["Inputs"].delete_rows(2)
    workbook.save(output)
    workbook.close()
    return output.read_bytes()


def _shifted_columns(data, root, operation):
    source, output = root / "column-source.xlsx", root / "column-upload.xlsx"
    source.write_bytes(data)
    workbook = load_workbook(source, keep_vba=False, data_only=False)
    if operation == "insert":
        workbook["Inputs"].insert_cols(2)
        workbook["Inputs"]["C2"] = 12
    else:
        workbook["Inputs"].delete_cols(2)
    workbook.save(output)
    workbook.close()
    return output.read_bytes()


def _input_only_fixture(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Inputs"
    sheet.protection.sheet = True
    for reference, value in (("B2", 10), ("B5", 20)):
        sheet[reference] = value
        sheet[reference].protection = Protection(locked=False)
    workbook.save(path)


def _open_table_fixture(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Plan"
    sheet.append(["Item", "Price", "Double"])
    for row, (item, price) in enumerate((("S-1", 10), ("S-2", 20), ("S-3", 30)), start=2):
        sheet.append([item, price, f"=B{row}*2"])
    workbook.save(path)


def _wide_layout_fixture(path):
    # A widened column past the last cell: the template's sheet size counts it, its cells do not
    _input_only_fixture(path)
    workbook = load_workbook(path)
    workbook["Inputs"].column_dimensions["H"].width = 30
    workbook.save(path)


def _formula_input_fixture(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Inputs"
    sheet.protection.sheet = True
    sheet["B2"] = 10
    sheet["B2"].protection = Protection(locked=False)
    sheet["C2"], sheet["C3"] = "=B2*2", "=B2*3"
    workbook.save(path)


def _forge_declared_compressed_sizes(data):
    """Rewrite every central-directory compress_size to match its file_size.

    zipfile trusts these declared values when it reports member sizes, so a
    forged archive looks perfectly compressed to any check that reads them.
    """
    forged = bytearray(data)
    start = forged.find(b"PK\x05\x06")
    if start < 0:
        raise AssertionError("no end of central directory record")
    count = int.from_bytes(forged[start + 10:start + 12], "little")
    offset = int.from_bytes(forged[start + 16:start + 20], "little")
    for _ in range(count):
        if forged[offset:offset + 4] != b"PK\x01\x02":
            raise AssertionError("no central directory header")
        file_size = int.from_bytes(forged[offset + 24:offset + 28], "little")
        forged[offset + 20:offset + 24] = file_size.to_bytes(4, "little")
        name_length = int.from_bytes(forged[offset + 28:offset + 30], "little")
        extra_length = int.from_bytes(forged[offset + 30:offset + 32], "little")
        comment_length = int.from_bytes(forged[offset + 32:offset + 34], "little")
        offset += 46 + name_length + extra_length + comment_length
    return bytes(forged)


def _real_stored_ratio(path):
    """Highest expanded/stored ratio measured from real member footprints."""
    with zipfile.ZipFile(path) as archive:
        members = sorted(archive.infolist(), key=lambda item: item.header_offset)
        total = path.stat().st_size
        worst = 0.0
        for index, info in enumerate(members):
            end = members[index + 1].header_offset if index + 1 < len(members) else total
            stored = max(end - info.header_offset - 30 - len(info.filename), 1)
            worst = max(worst, info.file_size / stored)
    return worst


def _rewrite_zip(data, root, transform=None, additions=()):
    source, output = root / "zip-source.xlsx", root / "zip-output.xlsx"
    source.write_bytes(data)
    with zipfile.ZipFile(source) as archive, zipfile.ZipFile(output, "w") as target:
        for info in archive.infolist():
            if transform:
                replacement = transform(info.filename, archive.read(info))
                if replacement is None:
                    continue
            else:
                replacement = archive.read(info)
            target.writestr(info, replacement)
        for name, content in additions:
            target.writestr(name, content)
    return output.read_bytes()


def _without_provenance(data, root):
    def transform(name, content):
        if name == "docProps/custom.xml":
            return None
        if name == "_rels/.rels":
            return re.sub(
                rb'<Relationship\b[^>]*Type="[^"]*/custom-properties"[^>]*/>', b"", content
            )
        if name == "[Content_Types].xml":
            return re.sub(
                rb'<Override\b[^>]*PartName="/docProps/custom\.xml"[^>]*/>', b"", content
            )
        return content

    return _rewrite_zip(data, root, transform)


def _with_shared_strings(data, root, xml):
    relationship = (
        b'<Relationship Id="rIdSharedStrings" '
        b'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" '
        b'Target="sharedStrings.xml"/>'
    )
    content_type = (
        b'<Override PartName="/xl/sharedStrings.xml" '
        b'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>'
    )

    def transform(name, content):
        if name == "xl/_rels/workbook.xml.rels":
            return content.replace(b"</Relationships>", relationship + b"</Relationships>")
        if name == "[Content_Types].xml":
            return content.replace(b"</Types>", content_type + b"</Types>")
        return content

    return _rewrite_zip(data, root, transform, (("xl/sharedStrings.xml", xml),))


class UploadAPITests(unittest.TestCase):
    def setUp(self):
        self.fixture = APIFixture()
        self.addCleanup(self.fixture.close)
        self.owner, self.owner_csrf, self.owner_user = self.fixture.signup(
            "upload-owner@example.test", "Upload Owner"
        )
        self.category, self.copy = self.fixture.make_copy(
            self.owner, self.owner_csrf, visibility="category", default_role="viewer"
        )
        exported = self.owner.get(f"/api/copies/{self.copy['id']}/export.xlsx")
        self.assertEqual(exported.status_code, 200, exported.body)
        self.exported = exported.body

    def preview(self, client, csrf, data, filename="changed.xlsx", copy_id=None):
        target = copy_id or self.copy["id"]
        return client.raw(
            "POST",
            f"/api/copies/{target}/uploads/preview",
            data,
            {
                "x-csrf-token": csrf,
                "x-upload-filename": filename,
                "content-type": "application/octet-stream",
            },
        )

    def preview_home(self, client, csrf, data, category_id, copy_id=None):
        query = f"?category_id={category_id}"
        if copy_id:
            query += f"&copy_id={copy_id}"
        return client.raw(
            "POST",
            f"/api/uploads/preview{query}",
            data,
            {
                "x-csrf-token": csrf,
                "x-upload-filename": "home-upload.xlsx",
                "content-type": "application/octet-stream",
            },
        )

    def confirm(self, client, csrf, token, body):
        return client.post(
            f"/api/uploads/{token}/confirm", body, {"x-csrf-token": csrf}
        )

    def test_non_numeric_content_length_is_a_client_error_not_a_crash(self):
        for target in (f"/api/copies/{self.copy['id']}/uploads/preview",
                       f"/api/uploads/preview?category_id={self.category['id']}"):
            with self.subTest(target=target):
                response = self.owner.raw(
                    "POST", target, self.exported,
                    {"x-csrf-token": self.owner_csrf, "x-upload-filename": "changed.xlsx",
                     "content-type": "application/octet-stream", "content-length": "not-a-number"},
                )

                self.assertEqual(response.status_code, 400, response.body)
                self.assertEqual(response.json()["error"], "invalid_length")

    def test_unparsable_confirm_body_length_is_a_client_error_not_a_crash(self):
        response = self.owner.raw(
            "POST", "/api/uploads/not-a-token/confirm", b"{}",
            {"x-csrf-token": self.owner_csrf, "content-type": "application/json",
             "content-length": "12 megabytes"},
        )

        self.assertEqual(response.status_code, 400, response.body)
        self.assertEqual(response.json()["error"], "invalid_length")

    def test_preview_uses_content_not_extension_and_reports_changes_and_ignored_edits(self):
        upload = _workbook_bytes(
            self.exported,
            self.fixture.root,
            (("Inputs", "B2", 12), ("Inputs", "B3", 25)),
            locked_edit=True,
            new_sheet=True,
        )

        before = (self.fixture.count("pp_overrides"), self.fixture.count("pp_history"))
        response = self.preview(self.owner, self.owner_csrf, upload, "renamed.bin")

        self.assertEqual(response.status_code, 200, response.body)
        body = response.json()
        self.assertEqual(
            [(item["sheet"], item["cell"], item["old"], item["new"])
             for item in body["changes"]],
            [("Inputs", "B2", 10, 12), ("Inputs", "B3", 20, 25)],
        )
        self.assertTrue(any(item["cell"] is None and item.get("count") == 1
                            and "locked" in item["reason"]
                            for item in body["ignored"] if item["sheet"] == "Calc"))
        self.assertTrue(any(item["sheet"] == "Unexpected" and "new sheet" in item["reason"]
                            for item in body["ignored"]))
        self.assertEqual(body["filename"], "renamed.bin")
        self.assertEqual(
            (self.fixture.count("pp_overrides"), self.fixture.count("pp_history")), before,
        )

    def test_missing_provenance_uses_structure_and_formula_fingerprint(self):
        upload = _workbook_bytes(
            _without_provenance(self.exported, self.fixture.root),
            self.fixture.root,
            (("Inputs", "B2", 13),),
        )
        accepted = self.preview(self.owner, self.owner_csrf, upload)
        self.assertEqual(accepted.status_code, 200, accepted.body)
        self.assertEqual(accepted.json()["identification"], "structure")
        self.assertTrue(any(
            "later edits cannot be checked for conflicts" in item["reason"]
            for item in accepted.json()["ignored"]
        ))

        changed = _workbook_bytes(upload, self.fixture.root, (("Calc", "B3", "=1+1"),))
        noted = self.preview(self.owner, self.owner_csrf, changed)
        self.assertEqual(noted.status_code, 200, noted.body)
        self.assertEqual([(item["cell"], item["new"]) for item in noted.json()["changes"]], [("B2", 13)])
        self.assertTrue(any("formulas differ from the template on: Calc" in item["reason"]
                            for item in noted.json()["ignored"]))

        source, output = self.fixture.root / "renamed-source.xlsx", self.fixture.root / "renamed.xlsx"
        source.write_bytes(upload)
        workbook = load_workbook(source)
        workbook["Lookup"].title = "Prices"
        workbook.save(output)
        refused = self.preview(self.owner, self.owner_csrf, output.read_bytes())
        self.assertEqual(refused.status_code, 422, refused.body)
        self.assertEqual(refused.json()["error"], "different_template")
        self.assertIn("no sheet named Lookup", refused.json()["message"])

    def test_tampered_or_incomplete_provenance_is_rejected(self):
        def rewrite_property(data, name, value):
            root = ET.fromstring(data)
            item = next(node for node in root if node.get("name") == name)
            if value is None:
                root.remove(item)
            else:
                next(iter(item)).text = value
            return ET.tostring(root, encoding="utf-8", xml_declaration=True)

        cases = {
            "different_template": lambda data: rewrite_property(data, "TemplateHash", "other"),
            "invalid_provenance": lambda data: rewrite_property(data, "CopyId", None),
        }
        for expected, change in cases.items():
            with self.subTest(expected=expected):
                data = _rewrite_zip(
                    self.exported, self.fixture.root,
                    lambda name, content: change(content) if name == "docProps/custom.xml" else content,
                )
                response = self.preview(self.owner, self.owner_csrf, data)
                self.assertEqual(response.status_code, 422, response.body)
                self.assertEqual(response.json()["error"], expected)

    def test_unknown_template_hash_on_a_new_copy_is_a_client_error(self):
        def rewrite_property(data, name, value):
            root = ET.fromstring(data)
            item = next(node for node in root if node.get("name") == name)
            next(iter(item)).text = value
            return ET.tostring(root, encoding="utf-8", xml_declaration=True)

        def tamper(content):
            return rewrite_property(
                rewrite_property(content, "CopyId", "00000000-0000-4000-8000-000000000000"),
                "TemplateHash", "0" * 64,
            )

        data = _rewrite_zip(
            self.exported, self.fixture.root,
            lambda name, content: tamper(content) if name == "docProps/custom.xml" else content,
        )

        response = self.preview_home(self.owner, self.owner_csrf, data, self.category["id"])

        self.assertEqual(response.status_code, 422, response.body)
        self.assertEqual(response.json()["error"], "different_template")

    def test_subset_export_can_be_uploaded_back_to_its_copy(self):
        subset = self.owner.get(
            f"/api/copies/{self.copy['id']}/export.xlsx?sheets=Inputs"
        )
        self.assertEqual(subset.status_code, 200, subset.body)
        upload = _workbook_bytes(
            subset.body, self.fixture.root, (("Inputs", "B2", 16),)
        )

        response = self.preview(self.owner, self.owner_csrf, upload)

        self.assertEqual(response.status_code, 200, response.body)
        self.assertEqual([(item["sheet"], item["cell"], item["new"])
                          for item in response.json()["changes"]], [("Inputs", "B2", 16)])

    def test_unsafe_archives_are_rejected_before_workbook_parsing(self):
        cases = {
            "not-a-zip": b"not a workbook",
            "path traversal": _rewrite_zip(
                self.exported, self.fixture.root, additions=(("../outside.xml", b"<x/>"),)
            ),
            "unsafe XML": _rewrite_zip(
                self.exported,
                self.fixture.root,
                additions=(("customXml/evil.xml", b'<!DOCTYPE x [<!ENTITY e SYSTEM "file:///etc/passwd">]><x>&e;</x>'),),
            ),
        }
        for name, data in cases.items():
            with self.subTest(name=name):
                response = self.preview(self.owner, self.owner_csrf, data)
                self.assertIn(response.status_code, {400, 413, 422}, response.body)
                self.assertIn(response.json()["error"],
                              {"invalid_workbook", "unsafe_xml", "archive_bomb", "unsupported_workbook"})

    def test_duplicate_encrypted_and_excessive_members_are_rejected(self):
        from pp.upload import _safe_member

        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            duplicate = _rewrite_zip(
                self.exported, self.fixture.root,
                additions=(("xl/workbook.xml", b"<duplicate/>"),),
            )
        response = self.preview(self.owner, self.owner_csrf, duplicate)
        self.assertEqual(response.status_code, 422, response.body)
        self.assertEqual(response.json()["error"], "invalid_workbook")
        encrypted = SimpleNamespace(
            filename="xl/encrypted.bin", flag_bits=1, file_size=1, compress_size=1,
        )
        with self.assertRaises(AppError) as raised:
            _safe_member(encrypted)
        self.assertEqual(raised.exception.code, "invalid_workbook")
        with mock.patch("pp.upload.MAX_MEMBERS", 1):
            excessive = self.preview(self.owner, self.owner_csrf, self.exported)
        self.assertEqual(excessive.status_code, 422, excessive.body)
        self.assertEqual(excessive.json()["error"], "invalid_workbook")

    def test_missing_content_types_is_a_controlled_workbook_error(self):
        data = _rewrite_zip(
            self.exported, self.fixture.root,
            lambda name, content: None if name == "[Content_Types].xml" else content,
        )

        response = self.preview(self.owner, self.owner_csrf, data)

        self.assertEqual(response.status_code, 422, response.body)
        self.assertEqual(response.json()["error"], "invalid_workbook")

    def test_archive_and_tree_budgets_run_before_elementtree_materialization(self):
        from pp.upload import _validate_archive

        candidate = self.fixture.root / "budget.xlsx"
        candidate.write_bytes(self.exported)
        with zipfile.ZipFile(candidate) as archive:
            sizes = [item.file_size for item in archive.infolist()]
            ratios = [item.file_size / max(item.compress_size, 1) for item in archive.infolist()]
        cases = (
            {"MAX_TREE_XML": 16},
            {"MAX_UNCOMPRESSED": max(sizes) - 1},
            {"MAX_UNCOMPRESSED": max(sizes) + 1},
            {"RATIO_MIN_BYTES": 0, "MAX_RATIO": max(ratios) / 2},
        )
        self.assertGreater(sum(sizes), max(sizes) + 1)
        for settings in cases:
            with self.subTest(settings=settings), mock.patch.multiple("pp.upload", **settings), \
                    mock.patch("pp.upload.SafeET.parse", side_effect=AssertionError("tree parser reached")):
                with self.assertRaises(AppError) as raised:
                    _validate_archive(candidate)
                self.assertEqual((raised.exception.status, raised.exception.code), (413, "archive_bomb"))

    def test_forged_compressed_sizes_cannot_hide_an_expansion_ratio(self):
        from pp.upload import _validate_archive

        candidate = self.fixture.root / "forged-ratio.xlsx"
        candidate.write_bytes(_forge_declared_compressed_sizes(self.exported))
        true_ratio = _real_stored_ratio(candidate)
        self.assertGreater(true_ratio, 1.5, "fixture must actually be compressed")
        with zipfile.ZipFile(candidate) as archive:
            declared = max(
                item.file_size / max(item.compress_size, 1) for item in archive.infolist()
            )
        self.assertLessEqual(declared, 1.5, "declared sizes must look uncompressed")

        for threshold in (2.0, (true_ratio + 1.5) / 2):
            with self.subTest(threshold=threshold), \
                    mock.patch.multiple(
                        "pp.upload", RATIO_MIN_BYTES=0, MAX_RATIO=threshold,
                    ):
                with self.assertRaises(AppError) as raised:
                    _validate_archive(candidate)
                self.assertEqual(
                    (raised.exception.status, raised.exception.code), (413, "archive_bomb"),
                )

    def test_chart_and_drawing_parts_are_screened_without_an_xml_suffix(self):
        chart = (
            '<?xml version="1.0" encoding="UTF-16"?>'
            '<!DOCTYPE chartSpace [<!ENTITY unsafe "1">]>'
            '<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart">'
            '<c:chart><c:plotArea><c:val><c:numRef><c:f>&unsafe;</c:f></c:numRef>'
            '</c:val></c:plotArea></c:chart></c:chartSpace>'
        ).encode("utf-16")
        drawing = (
            '<?xml version="1.0" encoding="UTF-16"?>'
            '<!DOCTYPE wsDr [<!ENTITY unsafe "1">]>'
            '<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing">'
            '<xdr:twoCellAnchor/></xdr:wsDr>'
        ).encode("utf-16")

        cases = (
            ("chart", "xl/_rels/workbook.xml.rels", "charts/chart.bin",
             "xl/charts/chart.bin", chart),
            ("chartsheet", "xl/_rels/workbook.xml.rels", "chartsheets/sheet.bin",
             "xl/chartsheets/sheet.bin", chart),
            ("drawing", None, "../drawings/drawing.bin",
             "xl/drawings/drawing.bin", drawing),
        )
        for role, rels, target, member, payload in cases:
            with self.subTest(role=role):
                relationship = (
                    '<Relationship Id="rIdRole" Type="http://schemas.openxmlformats.org/'
                    f'officeDocument/2006/relationships/{role}" Target="{target}"/>'
                )
                source = rels or "xl/worksheets/_rels/sheet1.xml.rels"
                additions = [(member, payload)]
                if rels is None:
                    additions.insert(0, (source, (
                        '<Relationships xmlns="http://schemas.openxmlformats.org/package/'
                        f'2006/relationships">{relationship}</Relationships>'
                    ).encode()))

                def relocate(name, content, rels=rels, relationship=relationship):
                    if name == (rels or "xl/worksheets/_rels/sheet1.xml.rels"):
                        return content.replace(
                            b"</Relationships>", relationship.encode() + b"</Relationships>")
                    return content

                data = _rewrite_zip(self.exported, self.fixture.root, relocate, additions)

                response = self.preview(self.owner, self.owner_csrf, data)

                self.assertEqual(response.status_code, 422, response.body)
                self.assertEqual(response.json()["error"], "unsafe_xml")

    def test_a_dangling_chart_relationship_is_refused(self):
        chart = (
            b'<?xml version="1.0" encoding="UTF-8"?>'
            b'<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart">'
            b'<c:chart><c:plotArea><c:val><c:numRef><c:f>Calc!$B$2</c:f></c:numRef></c:val>'
            b"</c:plotArea></c:chart></c:chartSpace>"
        )

        def relocate(name, content):
            if name == "xl/_rels/workbook.xml.rels":
                return content.replace(
                    b"</Relationships>",
                    b'<Relationship Id="rIdGone" Type="http://schemas.openxmlformats.org/'
                    b'officeDocument/2006/relationships/chart" Target="charts/missing.bin"/>'
                    b"</Relationships>",
                )
            return content

        data = _rewrite_zip(self.exported, self.fixture.root, relocate,
                            (("xl/charts/chart1.xml", chart),))

        response = self.preview(self.owner, self.owner_csrf, data)

        self.assertEqual(response.status_code, 422, response.body)
        self.assertEqual(response.json()["error"], "invalid_workbook")

    def test_chart_and_drawing_roles_are_capped_by_the_tree_limit(self):
        from pp.upload import _validate_archive

        payload = b64encode(os.urandom(700_000)).rstrip(b"=")
        points = b"".join(b"<c:pt>" + payload + b"</c:pt>" for _ in range(10))
        chart = (
            b'<?xml version="1.0" encoding="UTF-8"?>'
            b'<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart">'
            b'<c:chart><c:plotArea><c:val>' + points + b"</c:val></c:plotArea>"
            b"</c:chart></c:chartSpace>"
        )
        self.assertGreater(len(chart), upload_module.MAX_TREE_XML)
        self.assertLess(len(payload), upload_module.MAX_XML_TOKEN_BYTES)

        def relocate(name, content):
            if name == "xl/_rels/workbook.xml.rels":
                return content.replace(
                    b"</Relationships>",
                    b'<Relationship Id="rIdChart" Type="http://schemas.openxmlformats.org/'
                    b'officeDocument/2006/relationships/chart" Target="charts/chart1.xml"/>'
                    b"</Relationships>",
                )
            return content

        candidate = self.fixture.root / "oversized-chart.xlsx"
        candidate.write_bytes(
            _rewrite_zip(self.exported, self.fixture.root, relocate,
                         (("xl/charts/chart1.xml", chart),))
        )
        with mock.patch("pp.upload.validate_package",
                        side_effect=AssertionError("tree parser reached")):
            with self.assertRaises(AppError) as raised:
                _validate_archive(candidate)
        self.assertEqual((raised.exception.status, raised.exception.code), (413, "archive_bomb"))
        self.assertEqual(raised.exception.message, "Workbook XML metadata is too large")

    def test_a_real_chart_part_still_previews(self):
        chart = (
            b'<?xml version="1.0" encoding="UTF-8"?>'
            b'<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart">'
            b'<c:chart><c:plotArea><c:val><c:numRef><c:f>Calc!$B$2</c:f></c:numRef></c:val>'
            b"</c:plotArea></c:chart></c:chartSpace>"
        )

        def relocate(name, content):
            if name == "xl/_rels/workbook.xml.rels":
                return content.replace(
                    b"</Relationships>",
                    b'<Relationship Id="rIdChart" Type="http://schemas.openxmlformats.org/'
                    b'officeDocument/2006/relationships/chart" Target="charts/chart1.xml"/>'
                    b"</Relationships>",
                )
            if name == "[Content_Types].xml":
                return content.replace(
                    b"</Types>",
                    b'<Override PartName="/xl/charts/chart1.xml" ContentType="application/'
                    b'vnd.openxmlformats-officedocument.drawingml.chart+xml"/></Types>',
                )
            return content

        data = _rewrite_zip(self.exported, self.fixture.root, relocate,
                            (("xl/charts/chart1.xml", chart),))

        response = self.preview(self.owner, self.owner_csrf, data)

        self.assertEqual(response.status_code, 200, response.body)

    def test_styles_theme_and_renamed_shared_strings_have_tree_limits(self):
        from pp.upload import _validate_archive

        shared_xml = (
            b'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            b'<si><t>shared</t></si></sst>'
        )
        shared = _with_shared_strings(self.exported, self.fixture.root, shared_xml)

        def rename_shared(name, content):
            if name == "xl/sharedStrings.xml":
                return None
            return content.replace(b"sharedStrings.xml", b"metadata/strings.bin")

        renamed = _rewrite_zip(
            shared, self.fixture.root, rename_shared,
            (("xl/metadata/strings.bin", shared_xml),),
        )
        cases = (
            ("styles", self.exported, "MAX_STYLES_XML"),
            ("theme", self.exported, "MAX_THEME_XML"),
            ("renamed shared strings", renamed, "MAX_SHARED_STRINGS_XML"),
        )
        for name, data, limit in cases:
            with self.subTest(name=name):
                candidate = self.fixture.root / f"{name.replace(' ', '-')}.xlsx"
                candidate.write_bytes(data)
                with mock.patch(f"pp.upload.{limit}", 16, create=True), \
                        mock.patch("pp.upload.validate_package",
                                   side_effect=AssertionError("tree parser reached")):
                    with self.assertRaises(AppError) as raised:
                        _validate_archive(candidate)
                self.assertEqual((raised.exception.status, raised.exception.code),
                                 (413, "archive_bomb"))

    def test_shared_string_count_and_decoded_bytes_are_bounded(self):
        from pp.upload import _validate_archive

        xml = (
            b'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            b'<si><t>abcd</t></si><si><t>efgh</t></si></sst>'
        )
        data = _with_shared_strings(self.exported, self.fixture.root, xml)
        candidate = self.fixture.root / "shared-string-limits.xlsx"
        candidate.write_bytes(data)
        cases = (("MAX_SHARED_STRING_COUNT", 1), ("MAX_SHARED_STRING_BYTES", 3))
        for limit, value in cases:
            with self.subTest(limit=limit), mock.patch(f"pp.upload.{limit}", value), \
                    mock.patch("pp.upload.validate_package",
                               side_effect=AssertionError("package parser reached")):
                with self.assertRaises(AppError) as raised:
                    _validate_archive(candidate)
                self.assertEqual((raised.exception.status, raised.exception.code),
                                 (413, "archive_bomb"))

    def test_shared_string_bytes_are_bounded_while_each_item_streams(self):
        from pp.upload import _iter_shared_strings

        root = ET.Element(f"{{{upload_module.MAIN_NS}}}sst")
        item = ET.SubElement(root, f"{{{upload_module.MAIN_NS}}}si")
        text = ET.SubElement(item, f"{{{upload_module.MAIN_NS}}}t")
        text.text = "abcd"

        def events():
            yield "start", root
            yield "start", item
            yield "start", text
            yield "end", text
            raise AssertionError("parser advanced beyond the decoded-byte limit")

        archive = mock.Mock()
        archive.namelist.return_value = ["xl/sharedStrings.xml"]
        archive.open.return_value = io.BytesIO()
        with mock.patch("pp.upload.MAX_SHARED_STRING_BYTES", 3), \
                mock.patch("pp.upload.SafeET.iterparse", return_value=events()):
            with self.assertRaises(AppError) as raised:
                list(_iter_shared_strings(archive, "xl/sharedStrings.xml"))
        self.assertEqual((raised.exception.status, raised.exception.code),
                         (413, "archive_bomb"))

    def test_shared_strings_join_rich_text_without_phonetic_runs(self):
        from pp.upload import _shared_strings

        xml = (
            b'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            b'<si><r><t>ab</t></r><rPh><t>phonetic</t></rPh><r><t>cd</t></r></si>'
            b'</sst>'
        )
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, "w") as archive:
            archive.writestr("xl/sharedStrings.xml", xml)
        with zipfile.ZipFile(io.BytesIO(stream.getvalue())) as archive:
            self.assertEqual(_shared_strings(archive, "xl/sharedStrings.xml"), ["abcd"])

    def test_worksheet_structure_records_and_bytes_are_bounded(self):
        from pp.upload import _empty_structure, _iter_sheet

        xml = (
            b'<worksheet xmlns="http://schemas.openxmlformats.org/'
            b'spreadsheetml/2006/main"><sheetData><row r="1" ht="22" '
            b'customHeight="1"/></sheetData></worksheet>'
        )
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, "w") as archive:
            archive.writestr("xl/worksheets/sheet1.xml", xml)
        cases = (("MAX_STRUCTURE_RECORDS", 0), ("MAX_STRUCTURE_BYTES", 1))
        for limit, value in cases:
            with self.subTest(limit=limit), \
                    zipfile.ZipFile(io.BytesIO(stream.getvalue())) as archive, \
                    mock.patch(f"pp.upload.{limit}", value, create=True):
                with self.assertRaises(AppError) as raised:
                    with archive.open("xl/worksheets/sheet1.xml") as part:
                        list(_iter_sheet(part, _empty_structure()))
                self.assertEqual((raised.exception.status, raised.exception.code),
                                 (413, "archive_bomb"))

    def test_worksheet_structure_budget_is_shared_across_sheets(self):
        from pp.upload import _read_workbook

        candidate = self.fixture.root / "aggregate-structure.xlsx"
        workbook = Workbook()
        workbook.active.row_dimensions[1].height = 20
        workbook.create_sheet("Second").row_dimensions[1].height = 20
        workbook.save(candidate)
        workbook.close()

        with mock.patch("pp.upload.MAX_STRUCTURE_RECORDS", 1):
            with self.assertRaises(AppError) as raised:
                _read_workbook(candidate, provenance=False)
        self.assertEqual((raised.exception.status, raised.exception.code),
                         (413, "archive_bomb"))

    def test_relationship_budget_is_shared_before_graph_materialization(self):
        from pp.upload import _validate_archive

        candidate = self.fixture.root / "aggregate-relationships.xlsx"
        candidate.write_bytes(self.exported)
        with zipfile.ZipFile(candidate) as archive:
            sizes = [item.file_size for item in archive.infolist()
                     if item.filename.endswith(".rels")]
        self.assertGreater(len(sizes), 1)
        self.assertGreater(sum(sizes), max(sizes))

        with mock.patch("pp.upload.MAX_RELATIONSHIP_XML_TOTAL", max(sizes), create=True), \
                mock.patch("pp.upload._relationship_graph",
                           side_effect=AssertionError("graph parser reached")):
            with self.assertRaises(AppError) as raised:
                _validate_archive(candidate)
        self.assertEqual((raised.exception.status, raised.exception.code),
                         (413, "archive_bomb"))

    def test_row_spans_are_not_treated_as_structural_changes(self):
        def add_spans(name, content):
            if name == "xl/worksheets/sheet1.xml":
                return content.replace(b"<row ", b'<row spans="1:8" ')
            return content

        upload = _rewrite_zip(self.exported, self.fixture.root, add_spans)

        response = self.preview(self.owner, self.owner_csrf, upload)

        self.assertEqual(response.status_code, 200, response.body)
        self.assertFalse(any(
            item["sheet"] == "Inputs" and "row, column" in item["reason"]
            for item in response.json()["ignored"]
        ))

    def test_utf16_and_chunk_split_xml_declarations_are_rejected(self):
        utf16 = ('<?xml version="1.0" encoding="UTF-16"?>'
                 '<!DOCTYPE x [<!ENTITY e "unsafe">]><x>&e;</x>').encode("utf-16")
        prefix = b"<!--" + b"x" * ((1 << 20) - 10) + b"-->"
        split = prefix + b'<!DOCTYPE x [<!ENTITY e "unsafe">]><x>&e;</x>'
        for name, xml in (("utf16", utf16), ("split", split)):
            with self.subTest(name=name):
                data = _rewrite_zip(
                    self.exported, self.fixture.root,
                    additions=((f"customXml/{name}.xml", xml),),
                )
                response = self.preview(self.owner, self.owner_csrf, data)
                self.assertEqual(response.status_code, 422, response.body)
                self.assertEqual(response.json()["error"], "unsafe_xml")

    def test_xml_depth_is_bounded_before_package_parsing(self):
        from pp.upload import _validate_archive

        data = _rewrite_zip(
            self.exported, self.fixture.root,
            additions=(("customXml/deep.xml", b"<one><two><three/></two></one>"),),
        )
        candidate = self.fixture.root / "deep-xml.xlsx"
        candidate.write_bytes(data)
        with mock.patch("pp.upload.MAX_XML_DEPTH", 2, create=True), \
                mock.patch("pp.upload.validate_package",
                           side_effect=AssertionError("package parser reached")):
            with self.assertRaises(AppError) as raised:
                _validate_archive(candidate)
        self.assertEqual((raised.exception.status, raised.exception.code),
                         (413, "archive_bomb"))

    def test_xml_attributes_and_text_nodes_have_token_limits(self):
        from pp.upload import _check_xml_token, _screen_xml, _xml_parser

        cases = (
            ("attribute", b'<one value="abcd"/>'),
            ("text", b"<one>abcd</one>"),
        )
        for name, xml in cases:
            stream = io.BytesIO()
            with zipfile.ZipFile(stream, "w") as archive:
                archive.writestr(f"customXml/{name}.xml", xml)
            with self.subTest(name=name), \
                    zipfile.ZipFile(io.BytesIO(stream.getvalue())) as archive, \
                    mock.patch("pp.upload.MAX_XML_TOKEN_BYTES", 3, create=True):
                with self.assertRaises(AppError) as raised:
                    _screen_xml(archive, archive.getinfo(f"customXml/{name}.xml"), 1024)
                self.assertEqual((raised.exception.status, raised.exception.code),
                                 (413, "archive_bomb"))

        parser = _xml_parser()
        unfinished = b'<one value="'
        with mock.patch("pp.upload.MAX_XML_TOKEN_BYTES", 3), \
                self.assertRaises(AppError) as raised:
            parser.Parse(unfinished, False)
            _check_xml_token(parser, len(unfinished))
        self.assertEqual((raised.exception.status, raised.exception.code),
                         (413, "archive_bomb"))

    def test_relationship_targeted_xml_is_screened_even_without_xml_suffix(self):
        main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        worksheet = (
            '<?xml version="1.0" encoding="UTF-16"?>'
            '<!DOCTYPE worksheet [<!ENTITY unsafe "1">]>'
            f'<worksheet xmlns="{main}"><sheetData><row r="1">'
            '<c r="A1" t="inlineStr"><is><t>&unsafe;</t></is></c>'
            '</row></sheetData></worksheet>'
        ).encode("utf-16")

        def relocate(name, content):
            if name == "xl/worksheets/sheet1.xml":
                return None
            if name == "xl/_rels/workbook.xml.rels":
                return content.replace(b"worksheets/sheet1.xml", b"worksheets/sheet.bin")
            if name == "[Content_Types].xml":
                return content.replace(b"/xl/worksheets/sheet1.xml", b"/xl/worksheets/sheet.bin")
            return content

        data = _rewrite_zip(
            self.exported,
            self.fixture.root,
            relocate,
            (("xl/worksheets/sheet.bin", worksheet),),
        )

        response = self.preview(self.owner, self.owner_csrf, data)

        self.assertEqual(response.status_code, 422, response.body)
        self.assertEqual(response.json()["error"], "unsafe_xml")

    def test_xml_suffix_is_screened_when_content_type_is_mislabeled_binary(self):
        def mislabel(name, content):
            if name == "[Content_Types].xml":
                return re.sub(
                    rb'(<Default\b(?=[^>]*\bExtension="xml")[^>]*\bContentType=")[^"]*',
                    rb'\1application/octet-stream', content,
                )
            return content

        data = _rewrite_zip(
            self.exported,
            self.fixture.root,
            mislabel,
            (("customXml/payload.xml", b'<!DOCTYPE x [<!ENTITY e "unsafe">]><x>&e;</x>'),),
        )

        response = self.preview(self.owner, self.owner_csrf, data)

        self.assertEqual(response.status_code, 422, response.body)
        self.assertEqual(response.json()["error"], "unsafe_xml")

    def test_formula_in_an_input_is_ignored_without_clearing_the_input(self):
        upload = _workbook_bytes(
            self.exported, self.fixture.root, (("Inputs", "B2", "=1+1"),)
        )

        response = self.preview(self.owner, self.owner_csrf, upload)

        self.assertEqual(response.status_code, 200, response.body)
        self.assertFalse(any(item["cell"] == "B2" for item in response.json()["changes"]))
        self.assertTrue(any(item["cell"] == "B2" and "formula" in item["reason"]
                            for item in response.json()["ignored"]))

    def test_inserted_or_deleted_rows_are_matched_back_to_template_rows(self):
        # The user typed 12 next to "Cost" after inserting a row above it: it lands on the Cost row
        expected = {"insert": ([("B2", 10, 12)], None), "delete": ([], "Row 2 of the template"),
                    "same_extent": ([("B2", 10, 12)], "Row 6 of the template")}
        for operation, (changes, missing) in expected.items():
            with self.subTest(operation=operation):
                upload = _shifted_rows(self.exported, self.fixture.root, operation)

                response = self.preview(self.owner, self.owner_csrf, upload)

                self.assertEqual(response.status_code, 200, response.body)
                body = response.json()
                self.assertEqual([(item["cell"], item["old"], item["new"]) for item in body["changes"]
                                  if item["sheet"] == "Inputs"], changes)
                reasons = [item["reason"] for item in body["ignored"] if item["sheet"] == "Inputs"]
                self.assertTrue(any("matched back to their template rows" in reason for reason in reasons))
                self.assertEqual(any(missing in reason for reason in reasons) if missing else False,
                                 bool(missing))

    def test_binary_old_and_foreign_files_say_what_they_are(self):
        xlsb = self.fixture.root / "book.xlsb"
        with zipfile.ZipFile(xlsb, "w") as archive:
            archive.writestr("[Content_Types].xml", "<Types/>")
            archive.writestr("xl/workbook.bin", b"\x83\x01\x00")
        cases = {
            "xlsb": (xlsb.read_bytes(), ".xlsb"),
            "xls": (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 512, "old .xls"),
            "other": (b"not a workbook at all", "not an Excel workbook"),
        }
        for kind, (data, phrase) in cases.items():
            with self.subTest(kind=kind):
                response = self.preview(self.owner, self.owner_csrf, data)

                self.assertEqual(response.status_code, 422, response.body)
                self.assertEqual(response.json()["error"], "unsupported_workbook")
                self.assertIn(phrase, response.json()["message"])

    def test_xlsb_is_saved_as_xlsm_by_excel_when_the_app_runs_on_excel(self):
        edited = _workbook_bytes(self.exported, self.fixture.root, (("Inputs", "B2", 12),))
        seen = []

        def excel_save_as(source, target):
            seen.append((source.suffix, source.read_bytes()[:2], target.suffix))
            target.write_bytes(edited)

        xlsb = self.fixture.root / "plan.xlsb"
        with zipfile.ZipFile(xlsb, "w") as archive:
            archive.writestr("[Content_Types].xml", "<Types/>")
            archive.writestr("xl/workbook.bin", b"\x83\x01\x00")
        uploads = self.fixture.app.state.services.uploads
        with mock.patch.object(uploads, "converter", excel_save_as):
            response = self.preview(self.owner, self.owner_csrf, xlsb.read_bytes())

        self.assertEqual(response.status_code, 200, response.body)
        self.assertEqual([(item["cell"], item["new"]) for item in response.json()["changes"]], [("B2", 12)])
        self.assertEqual(seen, [(".xlsb", b"PK", ".xlsm")])
        previews = Path(self.fixture.config.data_dir) / "uploads"
        self.assertFalse([path for path in previews.rglob("*") if path.suffix in {".xlsb", ".xls"}
                          or path.name.endswith("-converted.xlsm")])

    def test_rows_filled_below_the_table_are_new_values_not_moved_rows(self):
        fixture = APIFixture(_open_table_fixture)
        self.addCleanup(fixture.close)
        client, csrf, _ = fixture.signup("grow@example.test", "Grown Table")
        _, copy = fixture.make_copy(client, csrf)
        exported = client.get(f"/api/copies/{copy['id']}/export.xlsx").body
        # Sir's case: prices filled in, more rows added at the end, the formula copied down to them
        upload = _workbook_bytes(exported, fixture.root, (
            ("Plan", "B2", 11), ("Plan", "A5", "S-4"), ("Plan", "B5", 40), ("Plan", "C5", "=B5*2"),
            ("Plan", "A6", "S-5"), ("Plan", "B6", 50),
        ))

        response = client.raw(
            "POST", f"/api/copies/{copy['id']}/uploads/preview", upload,
            {"x-csrf-token": csrf, "x-upload-filename": "grown.xlsx",
             "content-type": "application/octet-stream"},
        )

        self.assertEqual(response.status_code, 200, response.body)
        body = response.json()
        self.assertEqual(sorted((item["cell"], item["new"]) for item in body["changes"]),
                         [("A5", "S-4"), ("A6", "S-5"), ("B2", 11), ("B5", 40), ("B6", 50)])
        self.assertFalse(any("moved" in item["reason"] for item in body["ignored"]))

    def test_column_inserted_in_an_open_table_is_reported_not_applied(self):
        fixture = APIFixture(_open_table_fixture)
        self.addCleanup(fixture.close)
        client, csrf, _ = fixture.signup("column-table@example.test", "Inserted Column")
        _, copy = fixture.make_copy(client, csrf)
        exported = client.get(f"/api/copies/{copy['id']}/export.xlsx").body
        source, output = fixture.root / "column-table-source.xlsx", fixture.root / "column-table.xlsx"
        source.write_bytes(exported)
        workbook = load_workbook(source)
        workbook["Plan"].insert_cols(2)
        workbook["Plan"]["C2"] = 15
        workbook.save(output)

        response = client.raw(
            "POST", f"/api/copies/{copy['id']}/uploads/preview", output.read_bytes(),
            {"x-csrf-token": csrf, "x-upload-filename": "column.xlsx",
             "content-type": "application/octet-stream"},
        )

        self.assertEqual(response.status_code, 200, response.body)
        body = response.json()
        self.assertEqual(body["changes"], [])
        self.assertTrue(any("columns look inserted" in item["reason"] for item in body["ignored"]))

    def test_row_inserted_inside_a_formula_table_keeps_values_on_their_items(self):
        fixture = APIFixture(_open_table_fixture)
        self.addCleanup(fixture.close)
        client, csrf, _ = fixture.signup("insert-table@example.test", "Inserted Row")
        _, copy = fixture.make_copy(client, csrf)
        exported = client.get(f"/api/copies/{copy['id']}/export.xlsx").body
        source, output = fixture.root / "table-source.xlsx", fixture.root / "table-upload.xlsx"
        source.write_bytes(exported)
        workbook = load_workbook(source)
        sheet = workbook["Plan"]
        sheet.insert_rows(3)  # a new item between S-1 and S-2, formula copied in by hand
        sheet["A3"], sheet["B3"], sheet["C3"] = "NEW", 99, "=B3*2"
        sheet["B4"] = 21  # S-2's price, now on row 4
        workbook.save(output)

        response = client.raw(
            "POST", f"/api/copies/{copy['id']}/uploads/preview", output.read_bytes(),
            {"x-csrf-token": csrf, "x-upload-filename": "inserted.xlsx",
             "content-type": "application/octet-stream"},
        )

        self.assertEqual(response.status_code, 200, response.body)
        body = response.json()
        self.assertEqual([(item["cell"], item["old"], item["new"]) for item in body["changes"]],
                         [("B3", 20, 21)])
        self.assertTrue(any("Row 3 of your file is a new row" in item["reason"] for item in body["ignored"]))

    def test_inserted_or_deleted_columns_skip_affected_sheet_inputs(self):
        for operation in ("insert", "delete"):
            with self.subTest(operation=operation):
                upload = _shifted_columns(self.exported, self.fixture.root, operation)

                response = self.preview(self.owner, self.owner_csrf, upload)

                self.assertEqual(response.status_code, 200, response.body)
                self.assertFalse(any(
                    item["sheet"] == "Inputs" for item in response.json()["changes"]
                ))
                self.assertTrue(any(
                    item["sheet"] == "Inputs" and "may have moved"
                    in item["reason"] for item in response.json()["ignored"]
                ))

    def test_input_only_sheet_shift_is_skipped_without_locked_anchors(self):
        fixture = APIFixture(_input_only_fixture)
        self.addCleanup(fixture.close)
        client, csrf, _ = fixture.signup("input-only@example.test", "Input Only")
        _, copy = fixture.make_copy(client, csrf)
        exported = client.get(f"/api/copies/{copy['id']}/export.xlsx").body
        upload = _shifted_columns(exported, fixture.root, "insert")

        response = client.raw(
            "POST", f"/api/copies/{copy['id']}/uploads/preview", upload,
            {"x-csrf-token": csrf, "x-upload-filename": "shifted.xlsx",
             "content-type": "application/octet-stream"},
        )

        self.assertEqual(response.status_code, 200, response.body)
        self.assertEqual(response.json()["changes"], [])
        self.assertTrue(any(
            item["sheet"] == "Inputs" and "may have moved"
            in item["reason"] for item in response.json()["ignored"]
        ))

    def test_formatted_columns_past_the_last_cell_do_not_count_as_moved(self):
        fixture = APIFixture(_wide_layout_fixture)
        self.addCleanup(fixture.close)
        client, csrf, _ = fixture.signup("wide@example.test", "Wide Layout")
        _, copy = fixture.make_copy(client, csrf)
        exported = client.get(f"/api/copies/{copy['id']}/export.xlsx").body
        upload = _workbook_bytes(exported, fixture.root, (("Inputs", "B2", 77),))

        response = client.raw(
            "POST", f"/api/copies/{copy['id']}/uploads/preview", upload,
            {"x-csrf-token": csrf, "x-upload-filename": "wide.xlsx",
             "content-type": "application/octet-stream"},
        )

        self.assertEqual(response.status_code, 200, response.body)
        self.assertEqual([(item["cell"], item["new"]) for item in response.json()["changes"]], [("B2", 77)])
        self.assertFalse(any("may have moved" in item["reason"] for item in response.json()["ignored"]))

    def test_changed_locked_layout_skips_sheet_safely(self):
        upload = _workbook_bytes(
            self.exported, self.fixture.root,
            (("Inputs", "A2", "changed one"), ("Inputs", "A3", "changed two"),
             ("Inputs", "B2", 12)),
        )

        response = self.preview(self.owner, self.owner_csrf, upload)

        self.assertEqual(response.status_code, 200, response.body)
        self.assertFalse(any(item["sheet"] == "Inputs" for item in response.json()["changes"]))
        self.assertTrue(any(
            item["sheet"] == "Inputs" and "may have moved"
            in item["reason"] for item in response.json()["ignored"]
        ))

    def test_single_changed_locked_label_marks_layout_ambiguous(self):
        upload = _workbook_bytes(
            self.exported, self.fixture.root,
            (("Inputs", "A2", "changed locked label"), ("Inputs", "B2", 12)),
        )

        response = self.preview(self.owner, self.owner_csrf, upload)

        self.assertEqual(response.status_code, 200, response.body)
        self.assertFalse(any(item["sheet"] == "Inputs" for item in response.json()["changes"]))
        self.assertTrue(any(
            item["sheet"] == "Inputs" and "may have moved"
            in item["reason"] for item in response.json()["ignored"]
        ))

    def test_frozen_formulas_do_not_exhaust_displacement_candidates(self):
        fixture = APIFixture(_formula_input_fixture)
        self.addCleanup(fixture.close)
        client, csrf, _ = fixture.signup("frozen-formulas@example.test", "Frozen Formulas")
        _, copy = fixture.make_copy(client, csrf)
        exported = client.get(f"/api/copies/{copy['id']}/export.xlsx")
        self.assertEqual(exported.status_code, 200, exported.body)
        upload = _workbook_bytes(
            exported.body, fixture.root,
            (("Inputs", "B2", 12), ("Inputs", "C2", 20), ("Inputs", "C3", 30)),
        )

        response = client.raw(
            "POST", f"/api/copies/{copy['id']}/uploads/preview", upload,
            {"x-csrf-token": csrf, "x-upload-filename": "frozen.xlsx",
             "content-type": "application/octet-stream"},
        )

        self.assertEqual(response.status_code, 200, response.body)
        self.assertEqual(
            [(item["cell"], item["new"]) for item in response.json()["changes"]],
            [("B2", 12)],
        )
        self.assertTrue(any("formula" in item["reason"] for item in response.json()["ignored"]))

    def test_same_address_formula_edit_does_not_hide_valid_input_change(self):
        upload = _workbook_bytes(
            self.exported, self.fixture.root,
            (("Inputs", "B2", 12), ("Calc", "B2", "=Inputs!B2*3")),
        )

        response = self.preview(self.owner, self.owner_csrf, upload)

        self.assertEqual(response.status_code, 200, response.body)
        self.assertEqual(
            [(item["sheet"], item["cell"], item["new"]) for item in response.json()["changes"]],
            [("Inputs", "B2", 12)],
        )
        self.assertTrue(any(
            item["sheet"] == "Calc" and "formula" in item["reason"]
            for item in response.json()["ignored"]
        ))
        self.assertFalse(any(
            "may have moved" in item["reason"]
            for item in response.json()["ignored"]
        ))

    def test_worksheet_structure_and_sheet_order_changes_are_reported_as_ignored(self):
        source = self.fixture.root / "structure-source.xlsx"
        output = self.fixture.root / "structure-upload.xlsx"
        source.write_bytes(self.exported)
        workbook = load_workbook(source)
        workbook["Inputs"].row_dimensions[2].hidden = True
        workbook["Inputs"].merge_cells("D2:E2")
        workbook._sheets = [workbook["Calc"], workbook["Inputs"], workbook["Lookup"]]
        workbook.save(output)
        workbook.close()

        response = self.preview(self.owner, self.owner_csrf, output.read_bytes())

        self.assertEqual(response.status_code, 200, response.body)
        reasons = [item["reason"] for item in response.json()["ignored"]]
        self.assertTrue(any("sheet order" in reason for reason in reasons))
        self.assertTrue(any("row heights, column widths, merged cells" in reason for reason in reasons))

    def test_missing_implicit_input_cell_clears_its_existing_override(self):
        fixture = APIFixture(default_style_fixture)
        self.addCleanup(fixture.close)
        client, csrf, _ = fixture.signup("implicit-upload@example.test", "Implicit Upload")
        _, copy = fixture.make_copy(client, csrf)
        edited = client.post(f"/api/copies/{copy['id']}/ops", {
            "op_id": "implicit-value", "action": "edit",
            "cells": [{"sheet": 0, "r": 1, "c": 1, "v": "set", "expected_version": 0}],
        }, {"x-csrf-token": csrf})
        self.assertEqual(edited.status_code, 200, edited.body)
        exported = client.get(f"/api/copies/{copy['id']}/export.xlsx").body

        def remove_input(name, content):
            if name == "xl/worksheets/sheet1.xml":
                return re.sub(rb'<c r="B2"[^>]*>.*?</c>', b"", content)
            return content

        upload = _rewrite_zip(exported, fixture.root, remove_input)
        response = client.raw("POST", f"/api/copies/{copy['id']}/uploads/preview", upload, {
            "x-csrf-token": csrf, "x-upload-filename": "clear.xlsx",
            "content-type": "application/octet-stream",
        })

        self.assertEqual(response.status_code, 200, response.body)
        change = next(item for item in response.json()["changes"] if item["cell"] == "B2")
        self.assertEqual((change["old"], change["new"]), ("set", None))

    def test_preview_size_limit_is_streamed_and_configurable(self):
        self.fixture.config.upload_max_mb = 1
        response = self.preview(
            self.owner, self.owner_csrf, b"x" * (1024 * 1024 + 1)
        )
        self.assertEqual(response.status_code, 413, response.body)
        self.assertEqual(response.json()["error"], "too_large")

    def test_busy_transfer_does_not_stage_upload_body(self):
        services = self.fixture.app.state.services
        admission = mock.Mock()
        admission.acquire.return_value = False
        with mock.patch.object(services, "transfer_admission", admission, create=True), \
                mock.patch("pp.server.tempfile.NamedTemporaryFile") as staged, \
                mock.patch.object(services.uploads, "preview") as preview:
            response = self.preview(self.owner, self.owner_csrf, self.exported)
        self.assertEqual(response.status_code, 503, response.body)
        self.assertEqual(response.json()["error"], "transfer_busy")
        staged.assert_not_called()
        preview.assert_not_called()

    def test_preview_temporary_file_is_removed_after_parser_failure(self):
        paths = []

        def fail(request):
            paths.append(request.path)
            raise AppError(422, "invalid_workbook", "Synthetic parser failure")

        with mock.patch.object(self.fixture.app.state.services.uploads, "preview", side_effect=fail):
            response = self.preview(self.owner, self.owner_csrf, self.exported)
        self.assertEqual(response.status_code, 422, response.body)
        self.assertEqual(len(paths), 1)
        self.assertFalse(paths[0].exists())

    def test_conflict_includes_latest_other_actor_even_after_change_and_revert(self):
        other, other_csrf, _ = self.fixture.signup("other@example.test", "Other Editor")
        shared = self.owner.put(
            f"/api/copies/{self.copy['id']}/shares",
            {"email": "other@example.test", "role": "editor"},
            {"x-csrf-token": self.owner_csrf},
        )
        self.assertEqual(shared.status_code, 200, shared.body)
        for op_id, value, version in (("other-change", 11, 0), ("other-revert", 10, 1)):
            changed = other.post(
                f"/api/copies/{self.copy['id']}/ops",
                {"op_id": op_id, "action": "edit", "cells": [
                    {"sheet": 0, "r": 1, "c": 1, "v": value, "expected_version": version}
                ]},
                {"x-csrf-token": other_csrf},
            )
            self.assertEqual(changed.status_code, 200, changed.body)
        upload = _workbook_bytes(self.exported, self.fixture.root, (("Inputs", "B2", 12),))

        response = self.preview(self.owner, self.owner_csrf, upload)

        self.assertEqual(response.status_code, 200, response.body)
        conflict = response.json()["conflicts"][0]
        self.assertEqual((conflict["sheet"], conflict["cell"], conflict["by"]["name"]),
                         ("Inputs", "B2", "Other Editor"))
        self.assertRegex(conflict["at"], r"^\d{4}-\d{2}-\d{2}T")

    def test_unticked_conflict_is_excluded_while_other_changes_apply(self):
        other, other_csrf, _ = self.fixture.signup("exclude@example.test", "Conflict Editor")
        self.owner.put(
            f"/api/copies/{self.copy['id']}/shares",
            {"email": "exclude@example.test", "role": "editor"},
            {"x-csrf-token": self.owner_csrf},
        )
        other.post(
            f"/api/copies/{self.copy['id']}/ops",
            {"op_id": "conflict", "action": "edit", "cells": [
                {"sheet": 0, "r": 1, "c": 1, "v": 99, "expected_version": 0}
            ]},
            {"x-csrf-token": other_csrf},
        )
        upload = _workbook_bytes(
            self.exported, self.fixture.root,
            (("Inputs", "B2", 12), ("Inputs", "B3", 25)),
        )
        preview = self.preview(self.owner, self.owner_csrf, upload)
        conflict_key = preview.json()["conflicts"][0]["key"]

        confirmed = self.confirm(self.owner, self.owner_csrf, preview.json()["token"], {
            "mode": "update", "excluded_conflicts": [conflict_key],
        })

        self.assertEqual(confirmed.status_code, 200, confirmed.body)
        self.assertEqual(confirmed.json()["accepted"], 1)
        with self.fixture.database.connect() as connection:
            rows = connection.execute(
                "SELECT row,value_json FROM pp_overrides WHERE copy_id=? ORDER BY row",
                (self.copy["id"],),
            ).fetchall()
        self.assertEqual([(row["row"], json.loads(row["value_json"])) for row in rows],
                         [(1, 99), (2, 25)])

    def test_untouched_export_value_skips_update_but_remains_in_new_copy(self):
        other, other_csrf, _ = self.fixture.signup("stale@example.test", "Stale Editor")
        self.owner.put(
            f"/api/copies/{self.copy['id']}/shares",
            {"email": "stale@example.test", "role": "editor"},
            {"x-csrf-token": self.owner_csrf},
        )
        first = self.owner.post(f"/api/copies/{self.copy['id']}/ops", {
            "op_id": "exported-value", "action": "edit", "cells": [
                {"sheet": 0, "r": 1, "c": 1, "v": 14, "expected_version": 0},
            ],
        }, {"x-csrf-token": self.owner_csrf})
        self.assertEqual(first.status_code, 200, first.body)
        exported = self.owner.get(f"/api/copies/{self.copy['id']}/export.xlsx").body
        later = other.post(f"/api/copies/{self.copy['id']}/ops", {
            "op_id": "later-value", "action": "edit", "cells": [
                {"sheet": 0, "r": 1, "c": 1, "v": 99, "expected_version": 1},
            ],
        }, {"x-csrf-token": other_csrf})
        self.assertEqual(later.status_code, 200, later.body)

        preview = self.preview(self.owner, self.owner_csrf, exported)
        self.assertEqual((preview.json()["changes"], preview.json()["conflicts"]), ([], []))
        updated = self.confirm(
            self.owner, self.owner_csrf, preview.json()["token"], {"mode": "update"},
        )
        self.assertEqual(updated.json()["accepted"], 0)
        second = self.preview(self.owner, self.owner_csrf, exported)
        created = self.confirm(self.owner, self.owner_csrf, second.json()["token"], {
            "mode": "new_copy", "category_id": self.category["id"], "name": "Export baseline",
        })
        with self.fixture.database.connect() as connection:
            values = connection.execute(
                "SELECT copy_id,value_json FROM pp_overrides WHERE sheet=0 AND row=1 AND col=1 "
                "AND copy_id IN (?,?) ORDER BY copy_id",
                (self.copy["id"], created.json()["copy"]["id"]),
            ).fetchall()
        self.assertEqual(sorted(json.loads(row["value_json"]) for row in values), [14, 99])

    def test_own_post_export_edit_is_a_true_conflict_and_upload_stays_default(self):
        changed = self.owner.post(f"/api/copies/{self.copy['id']}/ops", {
            "op_id": "own-later-value", "action": "edit", "cells": [
                {"sheet": 0, "r": 1, "c": 1, "v": 77, "expected_version": 0},
            ],
        }, {"x-csrf-token": self.owner_csrf})
        self.assertEqual(changed.status_code, 200, changed.body)
        upload = _workbook_bytes(self.exported, self.fixture.root, (("Inputs", "B2", 12),))

        preview = self.preview(self.owner, self.owner_csrf, upload)

        self.assertEqual(len(preview.json()["conflicts"]), 1)
        self.assertEqual(preview.json()["conflicts"][0]["by"]["name"], "Upload Owner")
        confirmed = self.confirm(
            self.owner, self.owner_csrf, preview.json()["token"], {"mode": "update"},
        )
        self.assertEqual(confirmed.json()["accepted"], 1)
        with self.fixture.database.connect() as connection:
            value = connection.execute(
                "SELECT value_json FROM pp_overrides WHERE copy_id=? AND sheet=0 AND row=1 AND col=1",
                (self.copy["id"],),
            ).fetchone()
        self.assertEqual(json.loads(value["value_json"]), 12)

    def test_update_is_one_atomic_operation_with_upload_source_and_uploader(self):
        upload = _workbook_bytes(
            self.exported,
            self.fixture.root,
            (("Inputs", "B2", 12), ("Inputs", "B3", 25)),
        )
        preview = self.preview(self.owner, self.owner_csrf, upload, "quarter plan.xlsx")
        self.assertEqual(preview.status_code, 200, preview.body)

        confirmed = self.confirm(
            self.owner, self.owner_csrf, preview.json()["token"], {"mode": "update"}
        )

        self.assertEqual(confirmed.status_code, 200, confirmed.body)
        self.assertEqual(confirmed.json()["accepted"], 2)
        self.assertEqual(len(confirmed.json()["operation"]["cells"]), 2)
        with self.fixture.database.connect() as connection:
            history = connection.execute(
                "SELECT DISTINCT op_id,action,actor_id FROM pp_history WHERE copy_id=?",
                (self.copy["id"],),
            ).fetchall()
            values = connection.execute(
                "SELECT row,value_json FROM pp_overrides WHERE copy_id=? ORDER BY row",
                (self.copy["id"],),
            ).fetchall()
        self.assertEqual(len(history), 1)
        self.assertEqual((history[0]["action"], history[0]["actor_id"]),
                         ("upload: quarter plan.xlsx", self.owner_user["id"]))
        self.assertEqual([(row["row"], json.loads(row["value_json"])) for row in values],
                         [(1, 12), (2, 25)])

    def test_undo_then_redo_then_undo_keeps_the_upload_lineage(self):
        def value():
            response = self.owner.get(f"/api/copies/{self.copy['id']}/sheets/0/cell?r=1&c=1")
            self.assertEqual(response.status_code, 200, response.body)
            return response.json()["v"]

        upload = _workbook_bytes(self.exported, self.fixture.root, (("Inputs", "B2", 12),))
        preview = self.preview(self.owner, self.owner_csrf, upload, "lineage.xlsx")
        self.assertEqual(preview.status_code, 200, preview.body)
        confirmed = self.confirm(
            self.owner, self.owner_csrf, preview.json()["token"], {"mode": "update"}
        )
        self.assertEqual(confirmed.status_code, 200, confirmed.body)
        source_op_id = confirmed.json()["op_id"]
        self.assertEqual(value(), 12)

        undo = self.owner.post(
            f"/api/copies/{self.copy['id']}/ops",
            {"op_id": "lineage-undo", "action": "undo", "source_op_id": source_op_id},
            {"x-csrf-token": self.owner_csrf},
        )
        self.assertEqual(undo.status_code, 200, undo.body)
        self.assertEqual(value(), 10)

        redo = self.owner.post(
            f"/api/copies/{self.copy['id']}/ops",
            {"op_id": "lineage-redo", "action": "undo", "source_op_id": "lineage-undo"},
            {"x-csrf-token": self.owner_csrf},
        )
        self.assertEqual(redo.status_code, 200, redo.body)
        self.assertEqual(value(), 12)

        again = self.owner.post(
            f"/api/copies/{self.copy['id']}/ops",
            {"op_id": "lineage-undo-2", "action": "undo", "source_op_id": source_op_id},
            {"x-csrf-token": self.owner_csrf},
        )
        self.assertEqual(again.status_code, 409, again.body)
        self.assertEqual(value(), 12)

    def test_unrelated_concurrent_edit_does_not_block_upload(self):
        upload = _workbook_bytes(self.exported, self.fixture.root, (("Inputs", "B2", 12),))
        preview = self.preview(self.owner, self.owner_csrf, upload)
        self.assertEqual(preview.status_code, 200, preview.body)
        concurrent = self.owner.post(f"/api/copies/{self.copy['id']}/ops", {
            "op_id": "unrelated-race", "action": "edit", "cells": [
                {"sheet": 0, "r": 2, "c": 1, "v": 99, "expected_version": 0},
            ],
        }, {"x-csrf-token": self.owner_csrf})
        self.assertEqual(concurrent.status_code, 200, concurrent.body)

        confirmed = self.confirm(
            self.owner, self.owner_csrf, preview.json()["token"], {"mode": "update"},
        )

        self.assertEqual(confirmed.status_code, 200, confirmed.body)
        self.assertEqual(confirmed.json()["accepted"], 1)
        with self.fixture.database.connect() as connection:
            values = connection.execute(
                "SELECT row,value_json FROM pp_overrides WHERE copy_id=? ORDER BY row",
                (self.copy["id"],),
            ).fetchall()
        self.assertEqual([(row["row"], json.loads(row["value_json"])) for row in values],
                         [(1, 12), (2, 99)])

    def test_invalid_confirm_mode_does_not_consume_preview(self):
        upload = _workbook_bytes(self.exported, self.fixture.root, (("Inputs", "B2", 12),))
        preview = self.preview(self.owner, self.owner_csrf, upload)
        token = preview.json()["token"]

        invalid = self.confirm(self.owner, self.owner_csrf, token, {"mode": "invalid"})
        confirmed = self.confirm(self.owner, self.owner_csrf, token, {"mode": "update"})

        self.assertEqual(invalid.status_code, 422, invalid.body)
        self.assertEqual(confirmed.status_code, 200, confirmed.body)
        self.assertEqual(confirmed.json()["accepted"], 1)

    def test_confirm_rechecks_versions_access_and_csrf(self):
        other, other_csrf, _ = self.fixture.signup("race@example.test", "Race Editor")
        self.owner.put(
            f"/api/copies/{self.copy['id']}/shares",
            {"email": "race@example.test", "role": "editor"},
            {"x-csrf-token": self.owner_csrf},
        )
        upload = _workbook_bytes(self.exported, self.fixture.root, (("Inputs", "B2", 12),))
        preview = self.preview(other, other_csrf, upload)
        self.assertEqual(preview.status_code, 200, preview.body)
        raced = self.owner.post(
            f"/api/copies/{self.copy['id']}/ops",
            {"op_id": "race", "action": "edit", "cells": [
                {"sheet": 0, "r": 1, "c": 1, "v": 15, "expected_version": 0}
            ]},
            {"x-csrf-token": self.owner_csrf},
        )
        self.assertEqual(raced.status_code, 200, raced.body)

        stale = self.confirm(other, other_csrf, preview.json()["token"], {"mode": "update"})
        self.assertEqual(stale.status_code, 409, stale.body)
        self.assertEqual(stale.json()["error"], "conflict")

        fresh = self.preview(other, other_csrf, upload)
        self.owner.put(
            f"/api/copies/{self.copy['id']}/shares",
            {"email": "race@example.test", "role": None},
            {"x-csrf-token": self.owner_csrf},
        )
        denied = self.confirm(other, other_csrf, fresh.json()["token"], {"mode": "update"})
        self.assertEqual(denied.status_code, 403, denied.body)
        no_csrf = other.post(
            f"/api/uploads/{fresh.json()['token']}/confirm", {"mode": "update"}
        )
        self.assertEqual(no_csrf.status_code, 403, no_csrf.body)

    def test_viewer_can_preview_but_cannot_update_existing_copy(self):
        viewer, viewer_csrf, _ = self.fixture.signup(
            "upload-viewer@example.test", "Upload Viewer",
        )
        self.owner.put(
            f"/api/categories/{self.category['id']}/members",
            {"email": "upload-viewer@example.test", "role": "member"},
            {"x-csrf-token": self.owner_csrf},
        )
        upload = _workbook_bytes(self.exported, self.fixture.root, (("Inputs", "B2", 18),))
        preview = self.preview(viewer, viewer_csrf, upload)
        self.assertEqual(preview.status_code, 200, preview.body)
        self.assertFalse(preview.json()["can_update"])

        denied = self.confirm(
            viewer, viewer_csrf, preview.json()["token"], {"mode": "update"},
        )

        self.assertEqual(denied.status_code, 403, denied.body)

    def test_latest_preview_replaces_older_preview_for_the_same_session(self):
        first = self.preview(self.owner, self.owner_csrf, self.exported)
        second = self.preview(self.owner, self.owner_csrf, self.exported)

        expired = self.confirm(
            self.owner, self.owner_csrf, first.json()["token"], {"mode": "update"},
        )

        self.assertEqual(expired.status_code, 404, expired.body)
        self.assertEqual(expired.json()["error"], "preview_expired")
        self.assertNotEqual(first.json()["token"], second.json()["token"])

    def test_preview_store_evicts_oldest_entry_at_capacity(self):
        store = PreviewStore(self.fixture.root / "preview-store")
        with mock.patch("pp.upload.MAX_ACTIVE_PREVIEWS", 2):
            first = store.put({"user_id": "one", "session_hash": "one"})
            second = store.put({"user_id": "two", "session_hash": "two"})
            third = store.put({"user_id": "three", "session_hash": "three"})

        with self.assertRaisesRegex(AppError, "expired"):
            store.take(first, "one", "one")
        self.assertEqual(store.take(second, "two", "two")["user_id"], "two")
        self.assertEqual(store.take(third, "three", "three")["user_id"], "three")

    def test_irrelevant_cells_are_aggregated_and_preview_cache_is_disk_bounded(self):
        source = self.fixture.root / "many-source.xlsx"
        output = self.fixture.root / "many-upload.xlsx"
        source.write_bytes(self.exported)
        workbook = load_workbook(source)
        for row in range(100, 3_100):
            workbook["Calc"].cell(row, 1, f"=ROW()+{row}")
        extra = workbook.create_sheet("Bulk ignored")
        for row in range(1, 3_001):
            extra.cell(row, 1, f"=ROW()+{row}")
        workbook.save(output)
        workbook.close()

        retained = []
        original_fingerprint = upload_module._fingerprint

        def record_retained(sheet_names, cells):
            retained.append({name: len(cells[name]) for name in sheet_names})
            return original_fingerprint(sheet_names, cells)

        with mock.patch("pp.upload._fingerprint", side_effect=record_retained):
            response = self.preview(self.owner, self.owner_csrf, output.read_bytes())

        self.assertEqual(response.status_code, 200, response.body)
        # A file the app wrote is identified by provenance, so it may skip the formula fingerprint entirely
        self.assertTrue(all(item.get("Bulk ignored", 0) == 0 for item in retained))
        self.assertLess(max((item.get("Calc", 0) for item in retained), default=0), 100)
        ignored = response.json()["ignored"]
        self.assertLess(len(ignored), 10)
        self.assertTrue(any(item["sheet"] == "Calc" and item.get("count") == 3_000
                            and "locked" in item["reason"] for item in ignored))
        self.assertTrue(any(item["sheet"] == "Bulk ignored" and "new sheet" in item["reason"]
                            for item in ignored))
        store = self.fixture.app.state.services.uploads.previews
        cached = store._items[response.json()["token"]]
        self.assertFalse(any(
            key in cached.__dict__ for key in ("changes", "all_inputs", "ignored")
        ))
        self.assertTrue(cached.path.is_file())

    def test_preview_cache_weight_limit_is_a_controlled_413_and_leaves_no_file(self):
        directory = self.fixture.root / "small-preview-store"
        store = PreviewStore(directory)
        self.addCleanup(store.close)
        with mock.patch("pp.upload.MAX_PREVIEW_CACHE_BYTES", 64), self.assertRaises(AppError) as error:
            store.put({"user_id": "one", "session_hash": "one", "value": "x" * 256})
        self.assertEqual(error.exception.status, 413)
        self.assertEqual(list(directory.glob("*")), [])

    def test_uploaded_vba_is_discarded_from_the_copy_pipeline(self):
        from tests.v2.test_transfer_adversarial import _inject_vba_project

        source = self.fixture.root / "macro-upload-source.xlsx"
        upload = self.fixture.root / "macro-upload.xlsm"
        source.write_bytes(self.exported)
        _inject_vba_project(source)
        workbook = load_workbook(source, keep_vba=True)
        workbook["Inputs"]["B2"] = 19
        workbook.save(upload)
        workbook.vba_archive.close()
        workbook.close()
        preview = self.preview(self.owner, self.owner_csrf, upload.read_bytes(), "macro.xlsm")
        self.assertEqual(preview.status_code, 200, preview.body)
        confirmed = self.confirm(
            self.owner, self.owner_csrf, preview.json()["token"], {"mode": "update"},
        )
        self.assertEqual(confirmed.status_code, 200, confirmed.body)

        downloaded = self.owner.get(f"/api/copies/{self.copy['id']}/export.xlsx")
        with zipfile.ZipFile(io.BytesIO(downloaded.body)) as archive:
            self.assertNotIn("xl/vbaProject.bin", archive.namelist())

    def test_category_member_can_create_atomic_new_copy_with_template_differences(self):
        member, member_csrf, member_user = self.fixture.signup("member@example.test", "Category Member")
        added = self.owner.put(
            f"/api/categories/{self.category['id']}/members",
            {"email": "member@example.test", "role": "member"},
            {"x-csrf-token": self.owner_csrf},
        )
        self.assertEqual(added.status_code, 200, added.body)
        current = self.owner.post(
            f"/api/copies/{self.copy['id']}/ops",
            {"op_id": "same-as-upload", "action": "edit", "cells": [
                {"sheet": 0, "r": 1, "c": 1, "v": 14, "expected_version": 0}
            ]},
            {"x-csrf-token": self.owner_csrf},
        )
        self.assertEqual(current.status_code, 200, current.body)
        upload = _workbook_bytes(self.exported, self.fixture.root, (("Inputs", "B2", 14),))
        preview = self.preview(member, member_csrf, upload)
        self.assertEqual(preview.status_code, 200, preview.body)
        before_count = self.fixture.count("pp_copies")

        created = self.confirm(member, member_csrf, preview.json()["token"], {
            "mode": "new_copy", "category_id": self.category["id"], "name": "Uploaded scenario"
        })

        self.assertEqual(created.status_code, 201, created.body)
        self.assertEqual(self.fixture.count("pp_copies"), before_count + 1)
        new_id = created.json()["copy"]["id"]
        with self.fixture.database.connect() as connection:
            inputs = connection.execute(
                "SELECT row,value_json,updated_by FROM pp_overrides WHERE copy_id=? ORDER BY row",
                (new_id,),
            ).fetchall()
            original = connection.execute(
                "SELECT value_json FROM pp_overrides WHERE copy_id=? AND sheet=0 AND row=1 AND col=1",
                (self.copy["id"],),
            ).fetchone()
        self.assertEqual(json.loads(original["value_json"]), 14)
        self.assertEqual([(row["row"], json.loads(row["value_json"])) for row in inputs],
                         [(1, 14)])
        self.assertTrue(all(row["updated_by"] == member_user["id"] for row in inputs))

        second = self.preview(member, member_csrf, upload)
        with mock.patch.object(
            self.fixture.app.state.services.operations,
            "apply_upload_in_transaction",
            side_effect=AppError(422, "invalid", "Synthetic validation failure"),
        ):
            failed = self.confirm(member, member_csrf, second.json()["token"], {
                "mode": "new_copy", "category_id": self.category["id"], "name": "Must roll back"
            })
        self.assertEqual(failed.status_code, 422, failed.body)
        self.assertEqual(self.fixture.count("pp_copies"), before_count + 1)

    def test_category_member_without_copy_access_can_preview_for_new_copy(self):
        private_category, private_copy = self.fixture.make_copy(
            self.owner, self.owner_csrf, name="Private source", visibility="private"
        )
        exported = self.owner.get(f"/api/copies/{private_copy['id']}/export.xlsx").body
        no_provenance = _without_provenance(exported, self.fixture.root)
        upload = _workbook_bytes(no_provenance, self.fixture.root, (("Inputs", "B2", 17),))
        member, member_csrf, _ = self.fixture.signup("new-only@example.test", "New Copy Only")
        added = self.owner.put(
            f"/api/categories/{private_category['id']}/members",
            {"email": "new-only@example.test", "role": "member"},
            {"x-csrf-token": self.owner_csrf},
        )
        self.assertEqual(added.status_code, 200, added.body)
        self.assertEqual(member.get(
            f"/api/copies?category={private_category['id']}"
        ).json()["copies"], [])

        preview = self.preview_home(member, member_csrf, upload, private_category["id"])

        self.assertEqual(preview.status_code, 200, preview.body)
        self.assertIsNone(preview.json()["target_copy_id"])
        self.assertFalse(preview.json()["can_update"])
        self.assertEqual(preview.json()["identification"], "structure")
        created = self.confirm(member, member_csrf, preview.json()["token"], {
            "mode": "new_copy", "category_id": private_category["id"], "name": "Imported private plan",
        })
        self.assertEqual(created.status_code, 201, created.body)
        self.assertEqual(created.json()["accepted"], 1)

    def test_unchanged_upload_creates_copy_without_an_empty_operation(self):
        member, member_csrf, _ = self.fixture.signup("unchanged@example.test", "Unchanged Member")
        self.owner.put(
            f"/api/categories/{self.category['id']}/members",
            {"email": "unchanged@example.test", "role": "member"},
            {"x-csrf-token": self.owner_csrf},
        )
        preview = self.preview(member, member_csrf, self.exported)
        before = self.fixture.count("pp_copies")

        created = self.confirm(member, member_csrf, preview.json()["token"], {
            "mode": "new_copy", "category_id": self.category["id"], "name": "Clean template",
        })

        self.assertEqual(created.status_code, 201, created.body)
        self.assertEqual((created.json()["accepted"], created.json()["operation"]), (0, None))
        self.assertEqual(created.json()["copy"]["calculated_revision"], 0)
        self.assertEqual(self.fixture.count("pp_copies"), before + 1)

    def test_new_copy_refuses_an_active_template_change_after_preview(self):
        preview = self.preview(self.owner, self.owner_csrf, self.exported)
        before = self.fixture.count("pp_copies")
        self.fixture.app.state.services.copies.template_id = "different-active-template"

        response = self.confirm(self.owner, self.owner_csrf, preview.json()["token"], {
            "mode": "new_copy", "category_id": self.category["id"], "name": "Stale template",
        })

        self.assertEqual(response.status_code, 409, response.body)
        self.assertEqual(response.json()["error"], "refresh_preview")
        self.assertEqual(self.fixture.count("pp_copies"), before)

    def test_large_upload_requires_trusted_pipeline_and_has_server_derived_undo(self):
        from pp.ops import UploadOperationRequest, _shape
        from pp.upload import _operation_payload

        items = [
            {"sheet_index": 0, "r": row, "c": 0, "new": row, "expected_version": 0}
            for row in range(10_001)
        ]
        payload = _operation_payload("large.xlsx", items)

        with self.assertRaises(AppError):
            _shape(payload)
        cells, source_op_id, action = _shape(payload, "upload")
        self.assertEqual(len(cells), 10_001)
        self.assertEqual(action, "upload: large.xlsx")
        forged = self.owner.post(
            f"/api/copies/{self.copy['id']}/ops", payload,
            {"x-csrf-token": self.owner_csrf},
        )
        self.assertEqual(forged.status_code, 422, forged.body)
        payload["action"] = "undo"
        oversized = self.owner.post(
            f"/api/copies/{self.copy['id']}/ops", payload,
            {"x-csrf-token": self.owner_csrf},
        )
        self.assertEqual(oversized.status_code, 422, oversized.body)

        snapshot = SimpleNamespace(
            sheets=[{"rows": 10_001, "cols": 1}],
            is_input=lambda *_: True,
            cell=lambda *_: {"v": None},
            validation=lambda *_: None,
        )
        services = self.fixture.app.state.services
        old_pool, services.engine_pool = services.engine_pool, None
        self.addCleanup(setattr, services, "engine_pool", old_pool)
        with mock.patch.object(services.operations.templates, "get", return_value=snapshot):
            with self.fixture.database.transaction() as connection:
                ack, created, _ = services.operations.apply_upload_in_transaction(
                    UploadOperationRequest(
                        connection, self.copy["id"], self.owner_user["id"],
                        {**payload, "action": "upload: large.xlsx"}, None,
                    )
                )
            self.assertTrue(created)
            self.assertEqual(ack["op_id"], source_op_id)

            other, other_csrf, _ = self.fixture.signup("lineage@example.test", "Lineage Editor")
            shared = self.owner.put(
                f"/api/copies/{self.copy['id']}/shares",
                {"email": "lineage@example.test", "role": "editor"},
                {"x-csrf-token": self.owner_csrf},
            )
            self.assertEqual(shared.status_code, 200, shared.body)
            altered = other.post(
                f"/api/copies/{self.copy['id']}/ops",
                {"op_id": "wrong-author", "action": "undo", "source_op_id": source_op_id},
                {"x-csrf-token": other_csrf},
            )
            self.assertEqual(altered.status_code, 409, altered.body)

            undo = self.owner.post(
                f"/api/copies/{self.copy['id']}/ops",
                {"op_id": "undo-large", "action": "undo", "source_op_id": source_op_id},
                {"x-csrf-token": self.owner_csrf},
            )
            self.assertEqual(undo.status_code, 200, undo.body)
            self.assertEqual(undo.json()["accepted"], 10_001)
            replay = self.owner.post(
                f"/api/copies/{self.copy['id']}/ops",
                {"op_id": "undo-large-again", "action": "undo", "source_op_id": source_op_id},
                {"x-csrf-token": self.owner_csrf},
            )
            self.assertEqual(replay.status_code, 409, replay.body)
        self.assertEqual(self.fixture.count("pp_history"), 20_002)


if __name__ == "__main__":
    unittest.main()
