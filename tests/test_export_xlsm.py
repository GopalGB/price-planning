import re, sys, tempfile, unittest, zipfile
from pathlib import Path
from unittest import mock
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT)); sys.path.insert(0, str(ROOT / 'ops'))
import fixtures
from openpyxl import load_workbook
from pp.export import xlsm

VBA = b'\xd0\xcf\x11\xe0fake-vba\x00\x01\x02' * 50
CHAIN_CT = '<Override PartName="/xl/calcChain.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml"/>'
CHAIN_REL = '<Relationship Id="rId99" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/calcChain" Target="calcChain.xml"/>'


def make_template(path):
    """Semantic fixture plus a fake vbaProject.bin and calcChain, and no calcPr."""
    raw = path.with_suffix('.raw')
    meta = fixtures.make_semantic(raw)
    with zipfile.ZipFile(raw) as zin, zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info)
            if info.filename == 'xl/workbook.xml':
                data = re.sub(rb'<calcPr[^>]*/>', b'', data)
            elif info.filename == '[Content_Types].xml':
                data = data.replace(b'</Types>', CHAIN_CT.encode() + b'</Types>')
            elif info.filename == 'xl/_rels/workbook.xml.rels':
                data = data.replace(b'</Relationships>', CHAIN_REL.encode() + b'</Relationships>')
            zout.writestr(info, data)
        zout.writestr('xl/vbaProject.bin', VBA)
        zout.writestr('xl/calcChain.xml', '<calcChain xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><c r="B2" i="3"/></calcChain>')
    raw.unlink()
    return meta


def members(path):
    with zipfile.ZipFile(path) as z:
        return {i.filename: z.read(i) for i in z.infolist()}


def order(path):
    with zipfile.ZipFile(path) as z:
        return z.namelist()


class ExportXlsm(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        d = Path(self.tmp.name)
        self.src, self.out = d / 'template.xlsm', d / 'out.xlsm'
        self.meta = make_template(self.src)

    def tearDown(self):
        self.tmp.cleanup()

    def export(self, cells):
        return xlsm.export_xlsm(self.src, cells, self.out)

    def test_round_trip_values_and_escaping(self):
        text = 'Tom & Jerry <b> "quoted" \'single\''
        res = self.export([(0, 1, 1, 99.5), (0, 2, 1, text), (0, 3, 1, True), (0, 5, 1, None), (0, 1, 0, 7)])
        self.assertEqual(res, {'cells_written': 5, 'sheets_touched': [0]})
        ws = load_workbook(self.out)['Inputs']
        self.assertEqual(ws['B2'].value, 99.5)
        self.assertEqual(ws['B3'].value, text)
        self.assertIs(ws['B4'].value, True)
        self.assertIsNone(ws['B6'].value)
        self.assertEqual(ws['A2'].value, 7)
        raw = members(self.out)['xl/worksheets/sheet1.xml'].decode()
        self.assertIn('&amp;', raw); self.assertIn('&lt;b&gt;', raw)

    def test_formulas_untouched(self):
        self.export([(0, 1, 1, 1)])
        ws = load_workbook(self.out)['Calc']
        for ref, f in self.meta['formulas'].items():
            self.assertEqual(ws[ref.split('!')[1]].value, f)

    def test_styles_kept_and_clear_keeps_style(self):
        before = members(self.src)['xl/worksheets/sheet1.xml'].decode()
        s = re.search(r'<c r="B2" s="(\d+)"', before).group(1)
        self.export([(0, 1, 1, 5), (0, 2, 1, None)])
        after = members(self.out)['xl/worksheets/sheet1.xml'].decode()
        self.assertRegex(after, rf'<c r="B2" s="{s}"><v>5</v></c>')
        self.assertRegex(after, rf'<c r="B3" s="{s}"\s*/>')
        self.assertFalse(load_workbook(self.out)['Inputs']['B2'].protection.locked)

    def test_untouched_members_identical_and_order_kept(self):
        self.export([(0, 1, 1, 3)])
        a, b = members(self.src), members(self.out)
        changed = {'xl/worksheets/sheet1.xml', 'xl/workbook.xml', '[Content_Types].xml', 'xl/_rels/workbook.xml.rels'}
        for name, data in a.items():
            if name not in changed | {'xl/calcChain.xml'}:
                self.assertEqual(b[name], data, name)
        self.assertEqual(b['xl/vbaProject.bin'], VBA)
        self.assertEqual(order(self.out), [n for n in order(self.src) if n != 'xl/calcChain.xml'])

    def test_insert_missing_rows_and_cells_in_order(self):
        cells = [(0, 20, 3, 5), (0, 0, 0, 'top'), (0, 1, 3, 7), (0, 1, 2, 'mid'), (0, 3, 25, 1.25)]
        for chunk in (1 << 20, 7):  # tiny chunks exercise buffer boundaries
            with mock.patch.object(xlsm, 'CHUNK', chunk):
                self.export(cells)
            raw = members(self.out)['xl/worksheets/sheet1.xml'].decode()
            rows = [int(r) for r in re.findall(r'<row r="(\d+)"', raw)]
            self.assertEqual(rows, sorted(rows))
            self.assertEqual(rows, [1, 2, 3, 4, 5, 6, 21])
            row2 = re.search(r'<row r="2".*?</row>', raw).group(0)
            self.assertEqual(re.findall(r'<c r="([A-Z]+)2"', row2), ['A', 'B', 'C', 'D'])
            self.assertIn('<dimension ref="A1:Z21"', raw)
            ws = load_workbook(self.out)['Inputs']
            self.assertEqual((ws['A1'].value, ws['C2'].value, ws['D2'].value, ws['D21'].value, ws['Z4'].value),
                             ('top', 'mid', 7, 5, 1.25))
            self.assertEqual(ws['B2'].value, 10)

    def test_calcchain_removed_and_full_calc_on_load(self):
        self.export([(0, 1, 1, 3)])
        m = members(self.out)
        self.assertNotIn('xl/calcChain.xml', m)
        self.assertNotIn(b'calcChain', m['[Content_Types].xml'])
        self.assertNotIn(b'calcChain', m['xl/_rels/workbook.xml.rels'])
        self.assertEqual(m['xl/workbook.xml'].count(b'<calcPr fullCalcOnLoad="1"/>'), 1)

    def test_existing_calcpr_gets_attribute_once(self):
        m = members(self.src)
        wb = m['xl/workbook.xml'].replace(b'</workbook>', b'<calcPr calcId="191029" fullCalcOnLoad="0"/></workbook>')
        with zipfile.ZipFile(self.src) as zin, zipfile.ZipFile(self.src.with_name('b.xlsm'), 'w') as zout:
            for info in zin.infolist():
                zout.writestr(info, wb if info.filename == 'xl/workbook.xml' else zin.read(info))
        xlsm.export_xlsm(self.src.with_name('b.xlsm'), [(0, 1, 1, 3)], self.out)
        out = members(self.out)['xl/workbook.xml']
        self.assertIn(b'<calcPr calcId="191029" fullCalcOnLoad="1"/>', out)
        self.assertEqual(out.count(b'fullCalcOnLoad'), 1)

    def test_formula_cell_refused(self):
        with self.assertRaises(ValueError):
            self.export([(0, 1, 1, 1), (2, 1, 1, 5)])
        self.assertFalse(self.out.exists())

    def test_bad_input_refused(self):
        for cell in [(9, 0, 0, 1), (0, -1, 0, 1), (0, 0, 16384, 1), (0, 0, 0, float('nan')), (0, 0, 0, object()),
                     (0, 0, 0, 'bad\x01char')]:
            with self.assertRaises(ValueError, msg=repr(cell)):
                self.export([cell])
        with self.assertRaises(ValueError):
            xlsm.export_xlsm(self.src, [(0, 1, 1, 1)], self.src)

    def with_sheet2(self, xml):
        alt = self.src.with_name('alt.xlsm')
        with zipfile.ZipFile(self.src) as zin, zipfile.ZipFile(alt, 'w', zipfile.ZIP_DEFLATED) as zout:
            for info in zin.infolist():
                data = zin.read(info)
                if info.filename == 'xl/worksheets/sheet2.xml':
                    data = ('<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                            + xml + '</worksheet>').encode()
                zout.writestr(info, data)
        return alt

    def test_implicit_refs_and_empty_sheetdata(self):
        alt = self.with_sheet2('<sheetData><row><c><v>1</v></c><c><v>2</v></c></row><row><c><v>3</v></c></row></sheetData>')
        xlsm.export_xlsm(alt, [(1, 0, 1, 9), (1, 1, 2, 'x'), (1, 2, 0, 5)], self.out)
        ws = load_workbook(self.out)['Lookup']
        self.assertEqual([[c.value for c in row] for row in ws.iter_rows(max_col=3)],
                         [[1, 9, None], [3, None, 'x'], [5, None, None]])
        alt = self.with_sheet2('<sheetData/>')
        xlsm.export_xlsm(alt, [(1, 3, 1, 'only')], self.out)
        self.assertEqual(load_workbook(self.out)['Lookup']['B4'].value, 'only')

    def test_text_over_excel_limit_refused(self):
        with self.assertRaises(ValueError):
            self.export([(0, 1, 1, 'x' * 32768)])

    def test_big_sheet_is_streamed(self):
        import tracemalloc
        big = self.src.with_name('big.xlsm')
        with zipfile.ZipFile(self.src) as zin, zipfile.ZipFile(big, 'w', zipfile.ZIP_DEFLATED) as zout:
            for info in zin.infolist():
                if info.filename != 'xl/worksheets/sheet2.xml':
                    zout.writestr(info, zin.read(info))
                    continue
                with zout.open(info.filename, 'w') as out:
                    out.write(b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>')
                    for r in range(1, 200001):
                        out.write(f'<row r="{r}"><c r="A{r}" t="inlineStr"><is><t>K{r}</t></is></c><c r="B{r}"><v>{r}</v></c><c r="C{r}"><v>{r}.5</v></c></row>'.encode())
                    out.write(b'</sheetData></worksheet>')
        with zipfile.ZipFile(big) as z:
            self.assertGreater(z.getinfo('xl/worksheets/sheet2.xml').file_size, 20_000_000)
        tracemalloc.start()
        with mock.patch.object(xlsm, 'CHUNK', 1 << 16):
            xlsm.export_xlsm(big, [(1, 99999, 1, -1), (1, 199999, 3, 'end')], self.out)
        peak = tracemalloc.get_traced_memory()[1]
        tracemalloc.stop()
        self.assertLess(peak, 3_000_000)  # the sheet part is > 20 MB
        with zipfile.ZipFile(self.out) as z:
            raw = z.read('xl/worksheets/sheet2.xml').decode()
        self.assertIn('<row r="100000"><c r="A100000" t="inlineStr"><is><t>K100000</t></is></c><c r="B100000"><v>-1</v></c>', raw)
        self.assertIn('<c r="D200000" t="inlineStr"><is><t xml:space="preserve">end</t></is></c></row></sheetData>', raw)


if __name__ == '__main__':
    unittest.main()
