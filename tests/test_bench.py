# tests/test_bench.py
import json, os, subprocess, sys, tempfile, unittest
from pathlib import Path
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / 'ops'))
import fixtures, find_inputs

PY = sys.executable

class Bench(unittest.TestCase):
    def setUp(self):
        self.d = tempfile.TemporaryDirectory(); self.dir = Path(self.d.name)
        self.wb = self.dir / 'semantic.xlsx'; self.meta = fixtures.make_semantic(self.wb)
    def tearDown(self):
        self.d.cleanup()
    def test_find_inputs_lists_unlocked_cells_on_protected_sheet(self):
        self.assertEqual(find_inputs.find_inputs(self.wb),
                         ["'Inputs'!B2", "'Inputs'!B3", "'Inputs'!B4", "'Inputs'!B5", "'Inputs'!B6"])
    def run_bench(self, *extra):
        manifest = self.dir / 'm.json'
        manifest.write_text(json.dumps({'parity_sheets': ['Calc'], 'parity_sample': 50,
                                        'candidate_inputs': self.meta['inputs'], 'edit_values': [12, 30], 'repeats': 3}))
        r = subprocess.run([PY, str(ROOT / 'ops' / 'bench_real.py'), '--xlsm', str(self.wb), '--manifest', str(manifest),
                            '--public-dir', str(self.dir / 'pub'), '--private-dir', str(self.dir / 'priv'), *extra],
                           capture_output=True, text=True, timeout=300)
        reports = list((self.dir / 'pub').glob('*.json'))
        self.assertEqual(len(reports), 1, r.stdout + r.stderr)
        pub = json.loads(reports[0].read_text())
        return r, pub
    def test_public_report_has_aggregates_only(self):
        r, pub = self.run_bench()
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertEqual(pub['status'], 'ok', r.stderr)
        self.assertIn('narrow', pub['edits'])
        text = json.dumps(pub)
        for leak in ('XLOOKUP', 'twenty', 'K3', 'Cost', 'Calc'):
            self.assertNotIn(leak, text)
        self.assertTrue(list((self.dir / 'priv').glob('*.json')))
    def test_fanout_reports_truncation(self):
        from openpyxl import load_workbook
        wb = load_workbook(self.wb)
        for row in range(20, 10_320):
            wb['Calc'].cell(row, 2, '=Inputs!B2*2')
        wb.save(self.wb)
        wb.close()
        r, pub = self.run_bench()
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertIn('fanout_truncated', pub['edits']['worst'])
        self.assertGreater(pub['edits']['worst']['fanout_direct'], 256)
        self.assertTrue(pub['edits']['worst']['fanout_truncated'])
    def test_time_watchdog_writes_partial_report(self):
        r, pub = self.run_bench('--max-seconds', '0.001')
        self.assertEqual(r.returncode, 3, r.stderr)
        self.assertEqual(pub['status'], 'aborted')
        self.assertIn('time', pub['aborted_reason'])
    def test_watchdog_writes_partial_report(self):
        r, pub = self.run_bench('--max-rss-mb', '1')
        self.assertEqual(pub['status'], 'aborted')
        self.assertIn('rss', pub['aborted_reason'])
