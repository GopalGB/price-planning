# ops/bench_real.py
"""Benchmark Formualizer on a workbook. Public report = aggregates only; details stay private."""
import argparse, hashlib, json, os, platform, socket, statistics, subprocess, sys, threading, time
from datetime import datetime, timezone
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
from typed import same

DEPENDENTS_LIMIT = 10_000

def rss_mb():
    if os.name == 'nt':
        import ctypes, ctypes.wintypes as w
        class PMC(ctypes.Structure):
            _fields_ = [('cb', w.DWORD), ('PageFaultCount', w.DWORD), ('PeakWorkingSetSize', ctypes.c_size_t),
                        ('WorkingSetSize', ctypes.c_size_t), ('QuotaPeakPagedPoolUsage', ctypes.c_size_t),
                        ('QuotaPagedPoolUsage', ctypes.c_size_t), ('QuotaPeakNonPagedPoolUsage', ctypes.c_size_t),
                        ('QuotaNonPagedPoolUsage', ctypes.c_size_t), ('PagefileUsage', ctypes.c_size_t),
                        ('PeakPagefileUsage', ctypes.c_size_t)]
        c = PMC(); c.cb = ctypes.sizeof(PMC)
        ctypes.windll.psapi.GetProcessMemoryInfo(ctypes.windll.kernel32.GetCurrentProcess(), ctypes.byref(c), c.cb)
        return c.WorkingSetSize / 1e6
    if Path('/proc/self/statm').exists():
        return int(Path('/proc/self/statm').read_text().split()[1]) * os.sysconf('SC_PAGE_SIZE') / 1e6
    import resource
    # macOS reports bytes; Linux reports KiB. Peak RSS is conservative for the guard.
    peak = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
    return peak / (1e6 if sys.platform == 'darwin' else 1e3)

class Report:
    def __init__(self, a):
        public, private = Path(a.public_dir).resolve(), Path(a.private_dir).resolve()
        if public == private or public in private.parents or private in public.parents:
            raise ValueError('public and private report directories must be separate, non-nested trees')
        a.public_dir, a.private_dir = public, private
        self.a, self.pub, self.priv = a, {'status': 'running'}, {'mismatches': []}
        stamp = datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%SZ')
        host = hashlib.sha256(socket.gethostname().encode()).hexdigest()[:8]
        self.name = f'bench-{host}-{stamp}.json'
        self.pub.update(host=host, python=platform.python_version(), os=platform.platform(terse=True))
    def write(self):
        for d, data in ((self.a.public_dir, self.pub), (self.a.private_dir, self.priv)):
            Path(d).mkdir(parents=True, exist_ok=True)
            (Path(d) / self.name).write_text(json.dumps(data, indent=2, default=str))

def watchdog(rep, limits, t0):
    limit_mb, limit_s = limits
    while rep.pub['status'] == 'running':
        m = rss_mb()
        rep.pub['rss_mb_peak'] = max(rep.pub.get('rss_mb_peak', 0), round(m, 1))
        reason = f'rss {m:.0f} MB > {limit_mb} MB' if m > limit_mb else (
            f'time > {limit_s} s' if time.time() - t0 > limit_s else None)
        if reason:
            rep.pub.update(status='aborted', aborted_reason=reason); rep.write(); os._exit(3)
        time.sleep(0.25)

def timed(fn):
    t = time.perf_counter(); out = fn(); return out, time.perf_counter() - t

def split(ref):
    sheet, cell = ref.rsplit('!', 1)
    col = ''.join(ch for ch in cell if ch.isalpha()); row = int(cell[len(col):])
    n = 0
    for ch in col.upper():
        n = n * 26 + ord(ch) - 64
    if sheet.startswith("'") and sheet.endswith("'"):
        sheet = sheet[1:-1].replace("''", "'")
    return sheet, row, n

def cached_values(xlsm, sheets, sample):
    from openpyxl import load_workbook
    f_wb = load_workbook(xlsm, read_only=True, data_only=False)
    v_wb = load_workbook(xlsm, read_only=True, data_only=True)
    out = {}
    for s in sheets:
        refs = []
        for row in f_wb[s].iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith('='):
                    refs.append((c.row, c.column))
            if len(refs) >= sample:
                break
        refs = refs[:sample]; want = set(refs)
        if not refs:
            continue
        for r, row in enumerate(v_wb[s].iter_rows(max_row=max(r for r, _ in refs)), 1):
            for c, cell in enumerate(row, 1):
                if (r, c) in want:
                    out[(s, r, c)] = cell.value
    f_wb.close()
    v_wb.close()
    return out


def compare_samples(wb, expected, rep):
    parity = {}
    for (s, r, c), exp in expected.items():
        sheet_id = 'sheet-' + hashlib.sha256(s.encode()).hexdigest()[:16]
        got = wb.get_value(s, r, c); p = parity.setdefault(sheet_id, {'checked': 0, 'matched': 0})
        p['checked'] += 1
        if same(exp, got):
            p['matched'] += 1
        else:
            rep.priv['mismatches'].append({'sheet': s, 'row': r, 'col': c, 'formula': wb.get_formula(s, r, c),
                                           'expected': exp, 'got': got})
    rep.pub['parity'] = parity

def pick_inputs(wb, candidates):
    import formualizer
    fan = []
    for ref in candidates:
        try:
            report = wb.dependents(ref, max_results=DEPENDENTS_LIMIT)
            fan.append((len(report.dependents), ref, report.truncation.incomplete))
        except formualizer.InspectionError:
            fan.append((-1, ref, True))
    fan.sort()
    picks = {'narrow': fan[0], 'median': fan[len(fan) // 2], 'worst': fan[-1]} if fan else {}
    return picks

def measure_edits(wb, picks, m):
    edits = {}
    for label, (fanout, ref, truncated) in picks.items():
        s, r, c = split(ref); times = []
        for i in range(m['repeats']):
            v = m['edit_values'][i % len(m['edit_values'])]
            wb.set_value(s, r, c, v)
            _, dt = timed(wb.evaluate_all); times.append(dt * 1000)
        times.sort()
        edits[label] = {'fanout_direct': fanout, 'fanout_truncated': truncated,
                        'p50_ms': round(statistics.median(times), 1),
                        'p95_ms': round(times[min(len(times) - 1, int(0.95 * len(times)))], 1), 'max_ms': round(times[-1], 1)}
    return edits

def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument('--xlsm', required=True); ap.add_argument('--manifest', required=True)
    ap.add_argument('--public-dir', default='results-private/public')
    ap.add_argument('--private-dir', default='results-private')
    ap.add_argument('--max-rss-mb', type=float, default=12000); ap.add_argument('--max-seconds', type=float, default=3600)
    a = ap.parse_args(argv)
    rep = Report(a)
    import formualizer
    from importlib.metadata import version
    m = json.loads(Path(a.manifest).read_text())
    t0 = time.time()
    rep.pub.update(formualizer=version('formualizer'), xlsm_sha256=hashlib.sha256(Path(a.xlsm).read_bytes()).hexdigest())
    threading.Thread(target=watchdog, args=(rep, (a.max_rss_mb, a.max_seconds), t0), daemon=True).start()
    expected = cached_values(a.xlsm, m['parity_sheets'], m['parity_sample'])
    wb, rep.pub['load_s'] = timed(lambda: formualizer.load_workbook(a.xlsm))
    rep.pub['rss_mb_after_load'] = round(rss_mb(), 1)
    _, rep.pub['evaluate_all_s'] = timed(wb.evaluate_all)
    rep.pub['rss_mb_after_eval'] = round(rss_mb(), 1)
    compare_samples(wb, expected, rep)
    picks = pick_inputs(wb, m["candidate_inputs"])
    rep.priv["picked_inputs"] = {k: v[1] for k, v in picks.items()}
    rep.pub["edits"] = measure_edits(wb, picks, m)
    rep.pub['total_s'] = round(time.time() - t0, 1)
    if time.time() - t0 > a.max_seconds:  # the polling watchdog can miss a short run
        rep.pub.update(status='aborted', aborted_reason=f'time > {a.max_seconds} s'); rep.write(); sys.exit(3)
    rep.pub['status'] = 'ok'; rep.write()
    print(json.dumps(rep.pub, indent=2))

if __name__ == '__main__':
    main()
