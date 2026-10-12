"""Synthetic dashboard workbook fixture with pivot and slicer OOXML."""

from pathlib import Path
import shutil
import tempfile
import zipfile
import xml.etree.ElementTree as ET

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Protection


MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PACKAGE_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
CONTENT_TYPES = "http://schemas.openxmlformats.org/package/2006/content-types"
X14 = "http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"
XDR = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
DRAWING = "http://schemas.openxmlformats.org/drawingml/2006/main"
XDR14 = "http://schemas.microsoft.com/office/drawing/2010/slicer"

PIVOT_REL = f"{REL}/pivotTable"
CACHE_REL = f"{REL}/pivotCacheDefinition"
DRAWING_REL = f"{REL}/drawing"
SLICER_REL = "http://schemas.microsoft.com/office/2007/relationships/slicer"
SLICER_CACHE_REL = "http://schemas.microsoft.com/office/2007/relationships/slicerCache"

FIELD_NAMES = [
    "Region", "Department", "Product", "Season", "Class", "Units", "Revenue",
    "Margin", "Cost", "Price", "Channel", "Country", "Currency", "Plan Type",
    "Year", "Quarter", "Month", "Owner", "Status",
]
SHARED_COUNTS = [11, 79, 770, 9, 52]
SLICER_FIELDS = [0, 1, 2, 4]
SLICER_COUNTS = [10, 78, 769, 52]


def _source_row(index):
    region = None if index % 11 == 10 else f"Region {index % 10 + 1:02d}"
    department = None if index % 79 == 78 else f"Department {index % 78 + 1:02d}"
    product = None if index % 770 == 769 else f"Product {index % 769 + 1:03d}"
    return [
        region, department, product, f"Season {index % 9 + 1:02d}",
        f"Class {index % 52 + 1:02d}", index % 20 + 1, (index % 200 + 1) * 9.5,
        (index % 150 + 1) * 2.25, (index % 120 + 1) * 4.75,
        "=Dashboard!$H$1*2" if index == 0 else (index % 80 + 1) * 1.5,
        f"Channel {index % 4 + 1}", f"Country {index % 6 + 1}", "USD",
        "Working", 2027, f"Q{index % 4 + 1}", index % 12 + 1,
        f"Owner {index % 8 + 1}", "Active",
    ]


def _write_base(path, row_count):
    workbook = Workbook(write_only=True)
    dashboard = workbook.create_sheet("Dashboard")
    dashboard.sheet_format.defaultRowHeight = 15
    dashboard.protection.sheet = True
    multiplier = WriteOnlyCell(dashboard, value=2)
    multiplier.protection = Protection(locked=False)
    dashboard.append(["Department", "(All)", None, None, None, None, "Multiplier", multiplier])
    dashboard.append(["Product", "(All)"])
    dashboard.append(["Season", "(All)"])
    dashboard.append([])
    dashboard.append(["Row Labels", "Sum of Units", "Sum of Revenue", "Sum of Margin", "Sum of Cost", "Sum of Price"])
    source = workbook.create_sheet("SourceData")
    source.append(FIELD_NAMES)
    for index in range(row_count):
        source.append(_source_row(index))
    workbook.save(path)


def _xml(data):
    return ET.fromstring(data)


def _serialized(root):
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _relationship(root, values):
    identifier, kind, target = values
    ET.SubElement(root, f"{{{PACKAGE_REL}}}Relationship", {
        "Id": identifier, "Type": kind, "Target": target,
    })


def _workbook_xml(data):
    root = _xml(data)
    pivot_caches = ET.SubElement(root, f"{{{MAIN}}}pivotCaches")
    ET.SubElement(pivot_caches, f"{{{MAIN}}}pivotCache", {
        "cacheId": "1", f"{{{REL}}}id": "rIdPivotCache",
    })
    extensions = ET.SubElement(root, f"{{{MAIN}}}extLst")
    extension = ET.SubElement(extensions, f"{{{MAIN}}}ext", {
        "uri": "{BBE1A952-AA13-448e-AADC-164F8A28A991}",
    })
    caches = ET.SubElement(extension, f"{{{X14}}}slicerCaches")
    for index in range(1, 5):
        ET.SubElement(caches, f"{{{X14}}}slicerCache", {
            f"{{{REL}}}id": f"rIdSlicerCache{index}",
        })
    return _serialized(root)


def _workbook_rels(data):
    root = _xml(data)
    _relationship(root, ("rIdPivotCache", CACHE_REL, "pivotCache/pivotCacheDefinition1.xml"))
    for index in range(1, 5):
        _relationship(root, (f"rIdSlicerCache{index}", SLICER_CACHE_REL,
                             f"slicerCaches/slicerCache{index}.xml"))
    return _serialized(root)


def _dashboard_sheet(data):
    root = _xml(data)
    parts = ET.SubElement(root, f"{{{MAIN}}}pivotTableParts", {"count": "1"})
    ET.SubElement(parts, f"{{{MAIN}}}pivotTablePart", {f"{{{REL}}}id": "rIdPivot"})
    ET.SubElement(root, f"{{{MAIN}}}drawing", {f"{{{REL}}}id": "rIdDrawing"})
    extensions = ET.SubElement(root, f"{{{MAIN}}}extLst")
    extension = ET.SubElement(extensions, f"{{{MAIN}}}ext", {
        "uri": "{A8765BA9-456A-4DAB-B4F3-ACF838C121DE}",
    })
    slicer_list = ET.SubElement(extension, f"{{{X14}}}slicerList")
    ET.SubElement(slicer_list, f"{{{X14}}}slicer", {f"{{{REL}}}id": "rIdSlicers"})
    return _serialized(root)


def _sheet_relationships():
    root = ET.Element(f"{{{PACKAGE_REL}}}Relationships")
    _relationship(root, ("rIdPivot", PIVOT_REL, "../pivotTables/pivotTable1.xml"))
    _relationship(root, ("rIdDrawing", DRAWING_REL, "../drawings/drawing1.xml"))
    _relationship(root, ("rIdSlicers", SLICER_REL, "../slicers/slicer1.xml"))
    return _serialized(root)


def _styles(data):
    root = _xml(data)
    num_formats = root.find(f"{{{MAIN}}}numFmts")
    if num_formats is None:
        num_formats = ET.Element(f"{{{MAIN}}}numFmts")
        root.insert(0, num_formats)
    ET.SubElement(num_formats, f"{{{MAIN}}}numFmt", {
        "numFmtId": "165", "formatCode": "$#,##0.00",
    })
    num_formats.set("count", str(len(num_formats)))
    extensions = ET.SubElement(root, f"{{{MAIN}}}extLst")
    extension = ET.SubElement(extensions, f"{{{MAIN}}}ext", {
        "uri": "{EB79DEF2-80B8-43e5-95BD-54CBDDF9020C}",
    })
    ET.SubElement(extension, f"{{{X14}}}slicerStyles", {
        "defaultSlicerStyle": "SlicerStyleLight1",
    })
    return _serialized(root)


def _content_types(data):
    root = _xml(data)
    overrides = {
        "/xl/pivotTables/pivotTable1.xml": "application/vnd.openxmlformats-officedocument.spreadsheetml.pivotTable+xml",
        "/xl/pivotCache/pivotCacheDefinition1.xml": "application/vnd.openxmlformats-officedocument.spreadsheetml.pivotCacheDefinition+xml",
        "/xl/drawings/drawing1.xml": "application/vnd.openxmlformats-officedocument.drawing+xml",
        "/xl/slicers/slicer1.xml": "application/vnd.ms-excel.slicer+xml",
    }
    overrides.update({
        f"/xl/slicerCaches/slicerCache{index}.xml": "application/vnd.ms-excel.slicerCache+xml"
        for index in range(1, 5)
    })
    for part, content_type in overrides.items():
        ET.SubElement(root, f"{{{CONTENT_TYPES}}}Override", {
            "PartName": part, "ContentType": content_type,
        })
    return _serialized(root)


def _shared_items(field_index):
    count = SHARED_COUNTS[field_index]
    text_count = count - 1 if field_index < 3 else count
    prefix = FIELD_NAMES[field_index]
    items = "".join(f'<s v="{prefix} {index + 1:0{3 if field_index == 2 else 2}d}"/>'
                    for index in range(text_count))
    return items + ("<m/>" if field_index < 3 else "")


def _cache_fields():
    fields = []
    for index, name in enumerate(FIELD_NAMES):
        attributes = f' name="{name}" databaseField="1"'
        if index == 6:
            attributes += ' numFmtId="165"'
        if index < len(SHARED_COUNTS):
            count = SHARED_COUNTS[index]
            blank = ' containsBlank="1"' if index < 3 else ""
            shared = f'<sharedItems count="{count}" containsString="1"{blank}>{_shared_items(index)}</sharedItems>'
        else:
            shared = '<sharedItems containsNumber="1"/>' if 5 <= index <= 9 else ""
        fields.append(f"<cacheField{attributes}>{shared}</cacheField>")
    return "".join(fields)


def _pivot_cache(row_count, external=False):
    source = ('<cacheSource type="external" connectionId="1"/>' if external else
              f'<cacheSource type="worksheet"><worksheetSource ref="A1:S{row_count + 1}" sheet="SourceData"/></cacheSource>')
    return f'''<?xml version="1.0" encoding="utf-8"?>
<pivotCacheDefinition xmlns="{MAIN}" xmlns:x14="{X14}" saveData="0" refreshOnLoad="1" recordCount="{row_count}">
  {source}
  <cacheFields count="19">{_cache_fields()}</cacheFields>
  <extLst><ext uri="{{725AE2AE-9491-48be-B2B4-4EB974FC3084}}"><x14:pivotCacheDefinition pivotCacheId="1"/></ext></extLst>
</pivotCacheDefinition>'''.encode()


def _pivot_field(index):
    if index == 0:
        items = "".join(
            f'<item x="{item}"{" h=" + chr(34) + "1" + chr(34) if item == 3 else ""}/>'
            for item in (2, 0, 1, *range(3, 11))
        ) + '<item t="default"/>'
        return (f'<pivotField axis="axisRow" name="Region" showAll="0" outline="1" '
                f'subtotalTop="1" defaultSubtotal="0" sumSubtotal="1" sortType="descending" '
                f'rankBy="6" dataSourceSort="1"><items count="12">{items}</items></pivotField>')
    if 1 <= index <= 3:
        count = SHARED_COUNTS[index]
        items = "".join(f'<item x="{item}"/>' for item in range(count)) + '<item t="default"/>'
        return (f'<pivotField axis="axisPage" name="{FIELD_NAMES[index]}" showAll="1">'
                f'<items count="{count + 1}">{items}</items></pivotField>')
    if 5 <= index <= 9:
        return '<pivotField dataField="1" showAll="0"/>'
    return '<pivotField showAll="1"/>'


def _pivot_table():
    fields = "".join(_pivot_field(index) for index in range(19))
    data_fields = "".join(
        f'<dataField name="Sum of {FIELD_NAMES[field]}" fld="{field}" subtotal="sum"'
        f'{" numFmtId=" + chr(34) + "165" + chr(34) if field == 6 else ""}/>'
        for field in range(5, 10)
    )
    row_items = '<i><x v="2"/></i><i><x v="0"/></i><i><x v="1"/></i><i t="grand"/>'
    column_items = "".join(
        f'<i i="{index}"><x v="{index}"/></i>' for index in (1, 0, 2, 3, 4)
    )
    return f'''<?xml version="1.0" encoding="utf-8"?>
<pivotTableDefinition xmlns="{MAIN}" name="DashboardPivot" cacheId="1" dataCaption="Values"
 rowHeaderCaption="Row Labels" colHeaderCaption="Column Labels" grandTotalCaption="Grand Total"
 rowGrandTotals="1" colGrandTotals="1" outline="1" outlineData="1">
 <location ref="A5:F14" firstHeaderRow="1" firstDataRow="2" firstDataCol="1"/>
 <pivotFields count="19">{fields}</pivotFields>
 <rowFields count="1"><field x="0"/></rowFields>
 <rowItems count="4">{row_items}</rowItems>
 <colFields count="1"><field x="-2"/></colFields>
 <colItems count="5">{column_items}</colItems>
 <pageFields count="3"><pageField fld="1" hier="-1"/><pageField fld="2" hier="-1"/><pageField fld="3" hier="-1"/></pageFields>
 <dataFields count="5">{data_fields}</dataFields>
 <pivotTableStyleInfo name="PivotStyleLight16" showRowHeaders="1" showColHeaders="1"
  showRowStripes="0" showColStripes="0" showLastColumn="0"/>
</pivotTableDefinition>'''.encode()


def _pivot_relationships():
    root = ET.Element(f"{{{PACKAGE_REL}}}Relationships")
    _relationship(root, ("rIdCache", CACHE_REL, "../pivotCache/pivotCacheDefinition1.xml"))
    return _serialized(root)


def _slicer_cache(index):
    field = SLICER_FIELDS[index]
    count = SLICER_COUNTS[index]
    items = "".join(
        f'<x14:i x="{item}" s="{int(item == 0 if index == 0 else True)}"/>'
        for item in range(count)
    )
    name = f"Slicer_{FIELD_NAMES[field].replace(' ', '_')}"
    return f'''<?xml version="1.0" encoding="utf-8"?>
<x14:slicerCacheDefinition xmlns:x14="{X14}" name="{name}" sourceName="{FIELD_NAMES[field]}">
 <x14:pivotTables><x14:pivotTable tabId="1" name="DashboardPivot"/></x14:pivotTables>
 <x14:data><x14:tabular pivotCacheId="1" sortOrder="ascending" customListSort="1"
  showMissing="1" crossFilter="showItemsWithDataAtTop"><x14:items count="{count}">{items}</x14:items></x14:tabular></x14:data>
</x14:slicerCacheDefinition>'''.encode()


def _slicers():
    slicers = []
    for field in SLICER_FIELDS:
        name = f"Slicer_{FIELD_NAMES[field].replace(' ', '_')}"
        slicers.append(
            f'<x14:slicer name="{name}" cache="{name}" caption="{FIELD_NAMES[field]}" '
            'startItem="0" columnCount="1" showCaption="1" style="SlicerStyleLight1" '
            'lockedPosition="0" rowHeight="241300"/>'
        )
    return (f'<?xml version="1.0" encoding="utf-8"?><x14:slicers xmlns:x14="{X14}">'
            f'{"".join(slicers)}</x14:slicers>').encode()


def _marker(tag, column, row, column_offset=0):
    return (f'<xdr:{tag}><xdr:col>{column}</xdr:col><xdr:colOff>{column_offset}</xdr:colOff>'
            f'<xdr:row>{row}</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:{tag}>')


def _drawing():
    anchors = []
    frames = [((2, 0), (3, 300000)), ((3, 300000), (4, 600000)),
              ((4, 600000), (6, 300000)), ((6, 300000), (7, 600000))]
    for index, field in enumerate(SLICER_FIELDS):
        name = f"Slicer_{FIELD_NAMES[field].replace(' ', '_')}"
        start, end = frames[index]
        anchors.append(
            f'<xdr:twoCellAnchor>{_marker("from", start[0], 16, start[1])}'
            f'{_marker("to", end[0], 29, end[1])}'
            f'<xdr:graphicFrame><xdr:nvGraphicFramePr><xdr:cNvPr id="{index + 2}" name="{name}"/>'
            '<xdr:cNvGraphicFramePr/></xdr:nvGraphicFramePr><xdr:xfrm/>'
            f'<a:graphic><a:graphicData uri="{XDR14}"><xdr14:slicer name="{name}"/>'
            '</a:graphicData></a:graphic></xdr:graphicFrame><xdr:clientData/></xdr:twoCellAnchor>'
        )
    return (f'<?xml version="1.0" encoding="utf-8"?><xdr:wsDr xmlns:xdr="{XDR}" '
            f'xmlns:a="{DRAWING}" xmlns:xdr14="{XDR14}">{"".join(anchors)}</xdr:wsDr>').encode()


def _new_parts(row_count, external=False):
    parts = {
        "xl/worksheets/_rels/sheet1.xml.rels": _sheet_relationships(),
        "xl/pivotTables/pivotTable1.xml": _pivot_table(),
        "xl/pivotTables/_rels/pivotTable1.xml.rels": _pivot_relationships(),
        "xl/pivotCache/pivotCacheDefinition1.xml": _pivot_cache(row_count, external),
        "xl/slicers/slicer1.xml": _slicers(),
        "xl/drawings/drawing1.xml": _drawing(),
    }
    parts.update({
        f"xl/slicerCaches/slicerCache{index + 1}.xml": _slicer_cache(index)
        for index in range(4)
    })
    return parts


def _inject_dashboard_parts(path, row_count, external=False):
    replacements = {
        "[Content_Types].xml": _content_types,
        "xl/workbook.xml": _workbook_xml,
        "xl/_rels/workbook.xml.rels": _workbook_rels,
        "xl/worksheets/sheet1.xml": _dashboard_sheet,
        "xl/styles.xml": _styles,
    }
    new_parts = _new_parts(row_count, external)
    handle = tempfile.NamedTemporaryFile(dir=path.parent, suffix=".xlsm", delete=False)
    temporary = Path(handle.name)
    handle.close()
    try:
        with zipfile.ZipFile(path) as source, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as target:
            for info in source.infolist():
                data = source.read(info.filename)
                target.writestr(info, replacements.get(info.filename, lambda value: value)(data))
            for name, data in new_parts.items():
                target.writestr(name, data)
        shutil.move(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def make_dashboard(path, row_count=98_700, external=False):
    """Write a synthetic 19-field dashboard workbook with no cached source records.

    ``external=True`` makes the pivot cache a Power Query connection (the real template's shape).
    """
    if not isinstance(row_count, int) or isinstance(row_count, bool) or row_count < 1:
        raise ValueError("row_count must be a positive integer")
    path = Path(path)
    _write_base(path, row_count)
    _inject_dashboard_parts(path, row_count, external)
    return {"source": {"sheet": "SourceData", "ref": f"A1:S{row_count + 1}"}}
