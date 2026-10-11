# ops/fixtures.py
"""Deterministic synthetic workbooks. No company data. Used by CI and by the Univer gate."""
import xml.etree.ElementTree as ET
import zipfile
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Protection
from openpyxl.workbook.defined_name import DefinedName

CF_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"

def make_semantic(path):
    wb = Workbook()
    inp = wb.active
    inp.title = 'Inputs'
    inp.protection.sheet = True
    rows = [('Cost', 10), ('Retail', 20), ('Blank', None), ('EmptyText', ''), ('Rate', '1.234,5')]
    inputs = []
    for i, (label, val) in enumerate(rows, start=2):
        inp.cell(i, 1, label)
        c = inp.cell(i, 2, val)
        c.protection = Protection(locked=False)
        inputs.append(f'Inputs!B{i}')
    lk = wb.create_sheet('Lookup')
    for i in range(1, 11):
        lk.cell(i, 1, f'K{i}')
        lk.cell(i, 2, i * 1.5)
        lk.cell(i, 3, 'A' if i % 2 else 'B')
    calc = wb.create_sheet('Calc')
    formulas = {
        'B2': '=Inputs!B2*2',
        'B3': '=IFERROR(1-Inputs!B2/Inputs!B3,"")',
        'B4': '=IF(Inputs!B4="","blank","set")',
        'B5': '=IF(ISBLANK(Inputs!B4),1,0)',
        'B6': '=IFERROR(1/0,"err")',
        'B7': '=_xlfn.XLOOKUP("K3",Lookup!A1:A10,Lookup!B1:B10,"none")',
        'B8': '=_xlfn.LET(_xlpm.x,Inputs!B2,_xlpm.x+1)',
        'B9': '=_xlfn.SWITCH(Inputs!B2,10,"ten",20,"twenty","other")',
        'B10': '=_xlfn.IFS(Inputs!B2>15,"hi",TRUE,"lo")',
        'B11': '=_xlfn.NUMBERVALUE(Inputs!B6,",",".")',
        'B12': '=SUMIFS(Lookup!B1:B10,Lookup!C1:C10,"A")',
        'B13': '=_xlfn.SUM(_xlfn._xlws.FILTER(Lookup!B1:B10,Lookup!C1:C10="B"))',
        'B14': '=ROWS(_xlfn.UNIQUE(Lookup!C1:C10))',
        'B15': '=SUM(_xlfn.SEQUENCE(3))',
        'B16': '=VLOOKUP("K5",Lookup!A1:C10,2,FALSE)',
        'B17': '=SUBSTITUTE(TRIM("  a  b "),"a","x")',
    }
    for ref, f in formulas.items():
        calc[ref] = f
    wb.save(path)
    return {'inputs': inputs, 'formulas': {f'Calc!{k}': v for k, v in formulas.items()}}

def make_scale(path, rows=4088, cols=312, lookup_rows=100000):
    wb = Workbook(write_only=True)
    lk = wb.create_sheet('Lookup')
    for r in range(1, lookup_rows + 1):
        lk.append([f'K{r}', r * 0.5, f'G{r % 97}'] + [r % 13] * 14)
    grid = wb.create_sheet('Grid')
    for r in range(1, rows + 1):
        row = [f'K{(r * 7) % lookup_rows + 1}', r % 50 + 1]
        for c in range(3, cols + 1):
            row.append(f'=IFERROR(VLOOKUP($A{r},Lookup!$A:$C,2,FALSE)*$B{r}+{c},0)' if c % 3 else f'=SUBSTITUTE(TRIM($A{r}),"K","")')
        grid.append(row)
    wb.save(path)


def make_conditional_formatting(path):
    """Create a synthetic workbook, then inject the tested OOXML CF rules."""
    book = Workbook()
    inputs = book.active
    inputs.title = "Inputs"
    inputs.protection.sheet = True
    values = {
        "A1": 10, "B2": 12, "B3": 4, "C2": "match literal",
        "D2": "duplicate", "D1026": "duplicate", "E2": "prefix-value",
        "G2": "ERR", "H2": 3, "I2": 1, "M2": 1, "P2": "Case",
        "P1026": "case", "Q2": 1, "Q1026": "=1.0", "R2": 1,
        "S2": 1, "T2": 1, "U2": 1, "V257": 0, "W2": 1,
        "AB2": 0, "AB5": 1, "AC6": 0, "AD2": 1, "AE2": 1, "AF2": 1,
    }
    for address, value in values.items():
        inputs[address] = value
    inputs["B2"].protection = Protection(locked=False)
    for address in ("F2", "L2", "N3", "O3"):
        inputs[address] = "#DIV/0!"
        inputs[address].data_type = "e"
    inputs["F2"].font = Font(bold=True)
    inputs["F2"].fill = PatternFill("solid", fgColor="FF112233")
    inputs["AE2"].fill = PatternFill("solid", fgColor="FF445566")
    for address, value in (("J2", "=literal"), ("K2", "#DIV/0!")):
        inputs[address] = value
        inputs[address].data_type = "s"
    calc = book.create_sheet("Calc")
    calc["B2"] = "=Inputs!B2*2"
    calc["B3"] = "=Inputs!B3*2"
    calc["C2"] = "target"
    quoted = book.create_sheet("O'Brien")
    quoted["A1"] = 7
    quoted["B1"] = "target"
    scratch_row = book.create_sheet("Scratch Row")
    scratch_row["J1"], scratch_row["A2"] = "x", "target"
    scratch_column = book.create_sheet("Scratch Column")
    scratch_column["AF1"], scratch_column["A2"] = 1, "target"
    book.create_sheet("Plain")
    book.defined_names.add(DefinedName("GlobalThreshold", attr_text="'Inputs'!$A$1"))
    book.defined_names.add(DefinedName("SimpleLimit", attr_text="10"))
    book.defined_names.add(DefinedName("Threshold", attr_text="99"))
    inputs.defined_names.add(DefinedName("LocalThreshold", attr_text="'Inputs'!$A$1"))
    inputs.defined_names.add(DefinedName("Threshold", attr_text="'Inputs'!$A$1"))
    calc.defined_names.add(DefinedName("Threshold", attr_text="'Calc'!$B$2"))
    book.save(path)
    _inject_conditional_formats(path)


def _inject_conditional_formats(path):
    with zipfile.ZipFile(path, "r") as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    members["xl/styles.xml"] = _with_dxfs(members["xl/styles.xml"])
    members["xl/worksheets/sheet1.xml"] = _with_rules(
        members["xl/worksheets/sheet1.xml"],
    )
    extra_rules = {
        2: (("C2", 'type="expression" dxfId="0" priority="1"><formula>Threshold=24</formula>'),),
        3: (("B1", 'type="expression" dxfId="2" priority="1"><formula>Threshold=99</formula>'),),
        4: (("A2", 'type="expression" dxfId="0" priority="1"><formula>COUNTIF($1:$1,&quot;x&quot;)=1</formula>'),),
        5: (("A2", 'type="expression" dxfId="1" priority="1"><formula>COUNT(A:AF)=1</formula>'),),
    }
    for sheet, rules in extra_rules.items():
        name = f"xl/worksheets/sheet{sheet}.xml"
        members[name] = _with_rules(members[name], rules)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in members.items():
            archive.writestr(name, data)


def _with_dxfs(data):
    root = ET.fromstring(data)
    styles = (
        '<fill><patternFill patternType="solid"><fgColor rgb="FF00FFFF"/><bgColor rgb="FFFF0000"/></patternFill></fill>',
        '<font><color rgb="FF00AA00"/></font>',
        '<font><b val="1"/></font>',
        '<border><bottom style="thin"><color rgb="FF0000FF"/></bottom></border>',
        '<numFmt numFmtId="164" formatCode="0.000"/>',
        '<font><b val="0"/></font><fill><patternFill patternType="none"/></fill>',
        '<fill><patternFill patternType="solid"><fgColor rgb="FF111111"/></patternFill></fill>',
        '<font><b val="1"/></font>',
        '<font><color theme="1"/></font>',
        '<fill><patternFill patternType="solid"><fgColor indexed="10"/></patternFill></fill>',
        '<border><top style="thin"><color rgb="FF000000" tint="0.5"/></top></border>',
        '<fill><patternFill patternType="solid"/></fill>',
    )
    dxfs = ET.Element(f"{{{CF_NS}}}dxfs", {"count": str(len(styles))})
    for style in styles:
        dxfs.append(ET.fromstring(f'<dxf xmlns="{CF_NS}">{style}</dxf>'))
    root.append(dxfs)
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _with_rules(data, rules=None):
    root = ET.fromstring(data)
    rules = rules or (
        ("A2:A3", 'type="expression" dxfId="0" priority="1"><formula>AND(Calc!B2&gt;$A$1,MOD(B2,2)=0)</formula>'),
        ("B2", 'type="cellIs" operator="greaterThan" dxfId="1" priority="2"><formula>10</formula>'),
        ("C2", 'type="containsText" operator="containsText" text="match" dxfId="2" priority="3"><formula>NOT(ISERROR(SEARCH(&quot;match&quot;,C2)))</formula>'),
        ("D2:D1026", 'type="duplicateValues" dxfId="3" priority="4">'),
        ("E2", 'type="beginsWith" operator="beginsWith" text="prefix" dxfId="4" priority="5"><formula>LEFT(E2,LEN(&quot;prefix&quot;))=&quot;prefix&quot;</formula>'),
        ("F2:G2", 'type="containsErrors" dxfId="5" priority="6"><formula>ISERROR(F2)</formula>'),
        ("H2", 'type="cellIs" operator="greaterThan" dxfId="6" priority="7" stopIfTrue="1"><formula>1</formula>'),
        ("H2", 'type="cellIs" operator="greaterThan" dxfId="7" priority="8"><formula>0</formula>'),
        ("I2", 'type="expression" dxfId="0" priority="9"><formula>AND(ROW()=2,COLUMN()=9)</formula>'),
        ("J2", 'type="expression" dxfId="1" priority="10"><formula>EXACT(J2,&quot;=literal&quot;)</formula>'),
        ("K2", 'type="expression" dxfId="2" priority="11"><formula>EXACT(K2,&quot;#DIV/0!&quot;)</formula>'),
        ("L2", 'type="expression" dxfId="3" priority="12"><formula>ISERROR(L2)</formula>'),
        ("M2", 'type="expression" dxfId="4" priority="13"><formula>AND(GlobalThreshold=10,LocalThreshold=10,SimpleLimit=10,Threshold=10)</formula>'),
        ("N2", 'type="containsText" text="none" dxfId="0" priority="14">'),
        ("N3", 'type="containsText" text="div" dxfId="0" priority="15">'),
        ("O2", 'type="beginsWith" text="none" dxfId="0" priority="16">'),
        ("O3", 'type="beginsWith" text="{" dxfId="0" priority="17">'),
        ("P2:P1026", 'type="duplicateValues" dxfId="0" priority="18">'),
        ("Q2:Q1026", 'type="duplicateValues" dxfId="1" priority="19">'),
        ("R2", 'type="expression" dxfId="2" priority="20"><formula>COUNTIF(P:P,&quot;case&quot;)=2</formula>'),
        ("S2", 'type="cellIs" operator="greaterThan" dxfId="8" priority="21"><formula>0</formula>'),
        ("T2", 'type="cellIs" operator="greaterThan" dxfId="9" priority="22"><formula>0</formula>'),
        ("U2", 'type="cellIs" operator="greaterThan" dxfId="10" priority="23"><formula>0</formula>'),
        ("V2:V257", 'type="expression" dxfId="0" priority="24"><formula>ROW()&gt;0</formula>'),
        ("V2:V257", 'type="expression" dxfId="1" priority="25"><formula>COLUMN()=22</formula>'),
        ("W2", 'type="expression" dxfId="2" priority="26"><formula>COUNTIF(X:X,5)=1</formula>'),
        ("AA2:AA3 AC5:AC6", 'type="expression" dxfId="0" priority="27"><formula>$AB2&gt;0</formula>'),
        ("AD2", 'type="expression" dxfId="2" priority="28"><formula>\'O\'\'Brien\'!A1=7</formula>'),
        ("AE2", 'type="cellIs" operator="greaterThan" dxfId="11" priority="29"><formula>0</formula>'),
        ("XFD:XFD", 'type="expression" dxfId="0" priority="30"><formula>TRUE</formula>'),
        ("AF2000:AF2001", 'type="expression" dxfId="0" priority="31"><formula>TRUE</formula>'),
    )
    for sqref, body in rules:
        conditional = ET.Element(f"{{{CF_NS}}}conditionalFormatting", {"sqref": sqref})
        conditional.append(ET.fromstring(f'<cfRule xmlns="{CF_NS}" {body}</cfRule>'))
        root.append(conditional)
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)
